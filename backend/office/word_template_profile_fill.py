"""Fill Word template profiles (M1a): content controls, bookmarks, placeholders."""

from __future__ import annotations

import uuid
from pathlib import Path
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
from jinja2.sandbox import SandboxedEnvironment

from .errors import (
    OfficeFileNotFoundError,
    OfficeParseError,
    OfficePathError,
    OfficeTemplateFillError,
)
from .path_safety import resolve_within
from .word_template import _validate_template_safety
from .word_template_profile import (
    _PLACEHOLDER_RE,
    _sdt_name,
    _walk_bookmarks,
    analyze_word_template_profile,
)


def _value_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):  # noqa: UP038 — py38 isinstance
        return ", ".join(str(item) for item in value)
    return str(value)


def _check_output_name(name: str) -> str:
    value = (name or "").strip()
    if not value:
        raise OfficeTemplateFillError("output_filename 不能为空")
    if Path(value).is_absolute() or "/" in value or "\\" in value:
        raise OfficeTemplateFillError("output_filename 不能包含路径分隔符")
    if value in (".", ".."):
        raise OfficeTemplateFillError("output_filename 非法")
    return value


def _set_sdt_text(sdt: Any, value: str) -> bool:
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return False
    texts = list(content.iter(qn("w:t")))
    if not texts:
        paragraph = content.find(qn("w:p"))
        if paragraph is None:
            paragraph = OxmlElement("w:p")
            content.append(paragraph)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        run.append(text)
        paragraph.append(run)
        texts = [text]
    lines = value.split("\n")
    first = texts[0]
    first.text = lines[0]
    first.set(qn("xml:space"), "preserve")
    for extra in texts[1:]:
        extra.text = ""
    run = first.getparent()
    for line in lines[1:]:
        run.append(OxmlElement("w:br"))
        text = OxmlElement("w:t")
        text.text = line
        text.set(qn("xml:space"), "preserve")
        run.append(text)
    return True


def _replace_text_range(paragraph: Any, start: int, end: int, value: str) -> bool:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(value)
        return True
    offsets: List[Tuple[int, int, Any]] = []
    position = 0
    for run in runs:
        text = run.text or ""
        offsets.append((position, position + len(text), run))
        position += len(text)
    affected = [
        (run_start, run_end, run)
        for run_start, run_end, run in offsets
        if (run_end > start and run_start < end)
        or (start == end and run_start <= start <= run_end)
    ]
    if not affected:
        runs[-1].text = (runs[-1].text or "") + value
        return True
    first_start, _first_end, first_run = affected[0]
    last_start, _last_end, last_run = affected[-1]
    if first_run is last_run:
        text = first_run.text or ""
        local_start = max(0, start - first_start)
        local_end = min(len(text), end - first_start)
        first_run.text = text[:local_start] + value + text[local_end:]
        return True
    first_text = first_run.text or ""
    first_run.text = first_text[: max(0, start - first_start)] + value
    last_text = last_run.text or ""
    last_run.text = last_text[max(0, end - last_start) :]
    for _run_start, _run_end, run in affected[1:-1]:
        run.text = ""
    return True


def _bookmark_ranges(document: Any) -> Dict[str, Tuple[int, int, int, int]]:
    open_map: Dict[str, Tuple[str, int, int]] = {}
    ranges: Dict[str, Tuple[int, int, int, int]] = {}
    for index, paragraph in enumerate(document.paragraphs):
        _, starts, ends = _walk_bookmarks(paragraph._p)
        for name, ident, position in starts:
            open_map[ident] = (name, index, position)
        for ident, position in ends:
            record = open_map.pop(ident, None)
            if record is None:
                continue
            name, start_index, start_pos = record
            ranges.setdefault(name, (start_index, start_pos, index, position))
    return ranges


def _fill_bookmarks(document: Any, data: Dict[str, Any]) -> int:
    filled = 0
    ranges = _bookmark_ranges(document)
    paragraphs = list(document.paragraphs)
    for name, value in data.items():
        record = ranges.get(name)
        if record is None:
            continue
        start_index, start_pos, end_index, end_pos = record
        if start_index != end_index:
            raise OfficeTemplateFillError(f"书签 {name} 跨段落，暂不支持填充")
        _replace_text_range(
            paragraphs[start_index], start_pos, end_pos, _value_text(value)
        )
        filled += 1
    return filled


def _iter_paragraphs(document: Any) -> List[Any]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        for story in (section.header, section.footer):
            try:
                paragraphs.extend(story.paragraphs)
                for table in story.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            paragraphs.extend(cell.paragraphs)
            except (AttributeError, ValueError):
                continue
    return paragraphs


def _fill_story_content_controls(document: Any, data: Dict[str, Any]) -> None:
    for section in document.sections:
        for story in (section.header, section.footer):
            try:
                elements = story._element.iter(qn("w:sdt"))
            except (AttributeError, ValueError):
                continue
            for sdt in elements:
                name = _sdt_name(sdt)
                if name and name in data:
                    _set_sdt_text(sdt, _value_text(data[name]))


def _fill_placeholders(document: Any, data: Dict[str, Any]) -> int:
    filled = 0
    for paragraph in _iter_paragraphs(document):
        for _ in range(200):
            match = None
            matched_name = None
            for candidate in _PLACEHOLDER_RE.finditer(paragraph.text):
                raw = candidate.group(0)
                if raw.startswith("{{"):
                    name = raw[2:-2].split("|", 1)[0].strip()
                elif raw.startswith(("<<", "[[")):
                    name = raw[2:-2].strip()
                else:
                    name = raw[1:-1].strip()
                if name in data:
                    match = candidate
                    matched_name = name
                    break
            if match is None or matched_name is None:
                break
            _replace_text_range(
                paragraph,
                match.start(),
                match.end(),
                _value_text(data[matched_name]),
            )
            filled += 1
    return filled


def fill_word_template_profile(
    template_path: Path,
    data: Dict[str, Any],
    output_filename: str,
    *,
    workspace_path: str,
) -> Dict[str, Any]:
    """Fill M1a template slots into a new .docx, preserving non-slot formatting."""
    root = Path(workspace_path).resolve()
    if not root.is_dir():
        raise OfficeParseError("工作区目录不存在")
    template = Path(template_path)
    if not template.is_file():
        raise OfficeFileNotFoundError(template)
    resolved = resolve_within(root, template)
    profile = analyze_word_template_profile(resolved, workspace_path=str(root))
    if not isinstance(data, dict):
        raise OfficeTemplateFillError("data 必须是对象")
    missing = sorted(
        {
            slot["name"]
            for slot in profile["slots"]
            if slot["required"] and slot["name"] not in data
        }
    )
    if missing:
        raise OfficeTemplateFillError("缺少必填槽位: " + ", ".join(missing))

    try:
        document = Document(str(resolved))
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to parse DOCX template: {exc}") from exc

    filled_names = set()
    for sdt in document.element.body.iter(qn("w:sdt")):
        name = _sdt_name(sdt)
        if name and name in data and _set_sdt_text(sdt, _value_text(data[name])):
            filled_names.add(name)
    _fill_story_content_controls(document, data)
    _fill_bookmarks(document, data)
    _fill_placeholders(document, data)
    for slot in profile["slots"]:
        if slot["name"] in data:
            filled_names.add(slot["name"])

    output_name = _check_output_name(output_filename)
    try:
        target = resolve_within(root, resolved.parent / output_name)
    except OfficePathError as exc:
        raise OfficeTemplateFillError("输出路径不在工作区内") from exc
    if target == resolved:
        raise OfficeTemplateFillError("输出文件不能覆盖模板")
    if target.exists():
        raise OfficeTemplateFillError("输出文件已存在，拒绝覆盖")

    has_control_flow = bool(profile["summary"]["has_jinja_control"])
    work_path = target.with_name(f".{target.stem}.{uuid.uuid4().hex}.m1b.docx")
    try:
        document.save(str(work_path))
        if has_control_flow:
            _validate_template_safety(work_path)
            template = DocxTemplate(str(work_path))
            template.render(
                data, jinja_env=SandboxedEnvironment(autoescape=False)
            )
            template.save(str(target))
        else:
            work_path.replace(target)
    except OSError as exc:
        raise OfficeTemplateFillError(f"写入输出文件失败: {exc}") from exc
    finally:
        try:
            if work_path.exists():
                work_path.unlink()
        except OSError:
            pass

    residual: List[str] = []
    warnings: List[Dict[str, Any]] = list(profile["warnings"])
    if has_control_flow:
        warnings.append(
            {
                "code": "jinja_control_rendered",
                "message": "Jinja 控制流已由 docxtpl 渲染",
            }
        )
    try:
        output_profile = analyze_word_template_profile(target, workspace_path=str(root))
        residual = sorted(
            {
                slot["name"]
                for slot in output_profile["slots"]
                if slot["kind"] in ("placeholder", "jinja")
            }
        )
        if residual:
            warnings.append(
                {
                    "code": "residual_placeholders",
                    "message": "输出仍包含占位符: " + ", ".join(residual[:10]),
                }
            )
    except OfficeParseError as exc:
        warnings.append(
            {
                "code": "output_parse_warning",
                "message": f"输出文件已写入，但模板解析失败: {exc}",
            }
        )

    unfilled = sorted(
        {slot["name"] for slot in profile["slots"] if slot["name"] not in data}
    )
    return {
        "ok": True,
        "template_path": profile["source_path"],
        "output_path": target.relative_to(root).as_posix(),
        "filename": output_name,
        "filled_count": len(filled_names),
        "unfilled_slots": unfilled,
        "residual_placeholders": residual,
        "jinja_rendered": has_control_flow,
        "warnings": warnings,
    }


__all__ = ["fill_word_template_profile"]

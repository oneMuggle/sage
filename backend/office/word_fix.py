"""Deterministic Word fixes (M2b) for audit rule packs.

Only mechanically safe properties are changed: page setup, style fonts and
paragraph formats, page-number fields, and excess empty paragraphs. The
original file is never modified unless ``overwrite=True``. After saving,
the same rule pack is re-audited and returned as ``remaining``.
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .errors import OfficeFileNotFoundError, OfficeParseError, OfficePathError
from .path_safety import resolve_within
from .word_audit import audit_word_document, validate_rulepack
from .word_rulepack_models import (
    WordAuditFixResult,
    WordFixAction,
    WordRulePack,
    WordRuleSpec,
)

_PAGE_SIZES = {
    "A4": (Cm(21.0), Cm(29.7)),
    "LETTER": (Inches(8.5), Inches(11)),
}
_MARGIN_ATTRS = (
    ("top", "top_margin"),
    ("bottom", "bottom_margin"),
    ("left", "left_margin"),
    ("right", "right_margin"),
)
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
}
_SAVE_RETRY_DELAYS = (0.0, 0.1, 0.25, 0.5, 1.0)


class _FixContext:
    """Mutable python-docx document plus helpers shared by fixers."""

    def __init__(self, document: Any) -> None:
        self.document = document

    @property
    def paragraphs(self) -> List[Any]:
        return list(self.document.paragraphs)


def _find_style(document: Any, style_name: str) -> Any:
    target = style_name.strip().lower()
    for candidate in document.styles:
        name = str(getattr(candidate, "name", "") or "").strip().lower()
        style_id = str(getattr(candidate, "style_id", "") or "").strip().lower()
        if target in (name, style_id):
            return candidate
    return None


def _differs_cm(value: Any, expected: float, tolerance: float = 0.001) -> bool:
    if value is None:
        return True
    try:
        actual = float(value.cm)
    except AttributeError:
        return True
    return abs(actual - expected) > tolerance


def _differs_pt(value: Any, expected: float, tolerance: float = 0.01) -> bool:
    if value is None:
        return True
    try:
        actual = float(value.pt)
    except AttributeError:
        return True
    return abs(actual - expected) > tolerance


def _section_has_page_field(section: Any) -> bool:
    try:
        elements = (section.header._element, section.footer._element)
    except (AttributeError, ValueError):
        return False
    for element in elements:
        for instr in element.iter(qn("w:instrText")):
            if "PAGE" in (instr.text or "").upper():
                return True
        for simple in element.iter(qn("w:fldSimple")):
            if "PAGE" in (simple.get(qn("w:instr")) or "").upper():
                return True
    return False


def _append_page_field(section: Any) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _set_east_asia(style: Any, name: str) -> bool:
    element = getattr(style, "element", None)
    if element is None:
        return False
    rfonts = element.get_or_add_rPr().get_or_add_rFonts()
    if (rfonts.get(qn("w:eastAsia")) or "") == name:
        return False
    rfonts.set(qn("w:eastAsia"), name)
    return True


def _fix_page_size(ctx: _FixContext, rule: WordRuleSpec) -> int:
    expected = str(rule.params.get("size", "A4")).strip().upper()
    if expected not in _PAGE_SIZES:
        raise OfficeParseError(f"规则 {rule.id} 不支持页面尺寸 {expected}")
    width, height = _PAGE_SIZES[expected]
    changed = 0
    for section in ctx.document.sections:
        landscape = (
            section.page_width is not None
            and section.page_height is not None
            and section.page_width > section.page_height
        )
        target_width, target_height = (height, width) if landscape else (width, height)
        if _differs_cm(section.page_width, target_width.cm) or _differs_cm(
            section.page_height, target_height.cm
        ):
            section.page_width = target_width
            section.page_height = target_height
            changed += 1
    return changed


def _fix_page_orientation(ctx: _FixContext, rule: WordRuleSpec) -> int:
    expected = str(rule.params.get("orientation", "portrait")).strip().lower()
    if expected not in ("portrait", "landscape"):
        raise OfficeParseError(f"规则 {rule.id} 不支持页面方向 {expected}")
    expected_landscape = expected == "landscape"
    expected_enum = WD_ORIENT.LANDSCAPE if expected_landscape else WD_ORIENT.PORTRAIT
    changed = 0
    for section in ctx.document.sections:
        current_landscape = (
            section.page_width is not None
            and section.page_height is not None
            and section.page_width > section.page_height
        )
        swapped = current_landscape != expected_landscape
        if swapped:
            section.page_width, section.page_height = (
                section.page_height,
                section.page_width,
            )
        orientation_changed = section.orientation != expected_enum
        if orientation_changed:
            section.orientation = expected_enum
        if swapped or orientation_changed:
            changed += 1
    return changed


def _fix_page_margins(ctx: _FixContext, rule: WordRuleSpec) -> int:
    expected_margins = {
        side: rule.params.get(f"{side}_cm") for side, _attr in _MARGIN_ATTRS
    }
    if all(value is None for value in expected_margins.values()):
        raise OfficeParseError(f"规则 {rule.id} 至少需要一项页边距参数")
    changed = 0
    for section in ctx.document.sections:
        for side, attr in _MARGIN_ATTRS:
            expected = expected_margins[side]
            if expected is None:
                continue
            current = getattr(section, attr)
            if _differs_cm(current, float(expected)):
                setattr(section, attr, Cm(float(expected)))
                changed += 1
    return changed


def _fix_page_number(ctx: _FixContext, rule: WordRuleSpec) -> int:
    required = bool(rule.params.get("required", True))
    expected_format = rule.params.get("format")
    changed = 0
    for section in ctx.document.sections:
        if required and not _section_has_page_field(section):
            _append_page_field(section)
            changed += 1
        if expected_format:
            pgnum = section._sectPr.find(qn("w:pgNumType"))
            if pgnum is None:
                pgnum = OxmlElement("w:pgNumType")
                section._sectPr.append(pgnum)
            expected_value = str(expected_format)
            if (pgnum.get(qn("w:fmt")) or "") != expected_value:
                pgnum.set(qn("w:fmt"), expected_value)
                changed += 1
    return changed


def _fix_style_font(ctx: _FixContext, rule: WordRuleSpec) -> int:
    style_name = str(rule.params.get("style", "Normal"))
    style = _find_style(ctx.document, style_name)
    if style is None:
        return 0
    font = style.font
    changed = 0
    if "font_ascii" in rule.params:
        expected = str(rule.params["font_ascii"])
        if str(font.name or "") != expected:
            font.name = expected
            changed += 1
    if "font_east_asia" in rule.params and _set_east_asia(
        style, str(rule.params["font_east_asia"])
    ):
        changed += 1
    if "size_pt" in rule.params:
        expected_size = float(rule.params["size_pt"])
        tolerance = float(rule.params.get("tolerance_pt", 0.01))
        if _differs_pt(font.size, expected_size, tolerance):
            font.size = Pt(expected_size)
            changed += 1
    for param_key, attr in (
        ("bold", "bold"),
        ("italic", "italic"),
        ("underline", "underline"),
    ):
        if param_key not in rule.params:
            continue
        expected_bool = bool(rule.params[param_key])
        if bool(getattr(font, attr)) != expected_bool:
            setattr(font, attr, expected_bool)
            changed += 1
    if "color" in rule.params:
        expected_color = str(rule.params["color"]).lstrip("#").upper()
        try:
            current_color = str(font.color.rgb) if font.color and font.color.rgb else ""
        except AttributeError:
            current_color = ""
        if current_color.upper() != expected_color:
            font.color.rgb = RGBColor.from_string(expected_color)
            changed += 1
    return changed


def _fix_style_paragraph(ctx: _FixContext, rule: WordRuleSpec) -> int:
    style_name = str(rule.params.get("style", "Normal"))
    style = _find_style(ctx.document, style_name)
    if style is None:
        return 0
    fmt = style.paragraph_format
    changed = 0
    if "line_spacing" in rule.params:
        expected_spacing = float(rule.params["line_spacing"])
        current = fmt.line_spacing
        current_number = (
            float(current)
            if isinstance(current, (int, float))  # noqa: UP038 — py38 isinstance
            and not isinstance(current, bool)
            else None
        )
        if current_number is None or abs(current_number - expected_spacing) > 0.001:
            fmt.line_spacing = expected_spacing
            changed += 1
    if "line_spacing_pt" in rule.params:
        expected_pt = float(rule.params["line_spacing_pt"])
        current = fmt.line_spacing
        current_pt = float(current.pt) if hasattr(current, "pt") else None
        if current_pt is None or abs(current_pt - expected_pt) > 0.1:
            fmt.line_spacing = Pt(expected_pt)
            changed += 1
    if "first_line_indent_cm" in rule.params:
        expected_indent = float(rule.params["first_line_indent_cm"])
        if _differs_cm(fmt.first_line_indent, expected_indent):
            fmt.first_line_indent = Cm(expected_indent)
            changed += 1
    if "alignment" in rule.params:
        expected_alignment = str(rule.params["alignment"]).strip().lower()
        if expected_alignment not in _ALIGNMENTS:
            raise OfficeParseError(
                f"规则 {rule.id} 不支持对齐方式 {expected_alignment}"
            )
        alignment = fmt.alignment
        current_alignment = alignment.name.lower() if alignment is not None else "left"
        if current_alignment != expected_alignment:
            fmt.alignment = _ALIGNMENTS[expected_alignment]
            changed += 1
    return changed


def _fix_empty_paragraphs(ctx: _FixContext, rule: WordRuleSpec) -> int:
    max_consecutive = int(rule.params.get("max_consecutive", 2))
    if max_consecutive < 0:
        raise OfficeParseError(f"规则 {rule.id} 参数 max_consecutive 不能为负")
    to_delete: List[Any] = []
    run: List[Any] = []

    def _flush() -> None:
        if len(run) > max_consecutive:
            to_delete.extend(run[max_consecutive:])
        run.clear()

    for paragraph in ctx.paragraphs:
        if str(getattr(paragraph, "text", "") or "").strip() == "":
            run.append(paragraph)
        else:
            _flush()
    _flush()
    removed = 0
    for paragraph in to_delete:
        element = paragraph._element
        parent = element.getparent()
        if parent is None:
            continue
        parent.remove(element)
        removed += 1
    return removed


_FIXERS: Dict[str, Callable[[_FixContext, WordRuleSpec], int]] = {
    "page.size": _fix_page_size,
    "page.orientation": _fix_page_orientation,
    "page.margins": _fix_page_margins,
    "page.page_number": _fix_page_number,
    "style.font": _fix_style_font,
    "style.paragraph": _fix_style_paragraph,
    "content.empty_paragraphs": _fix_empty_paragraphs,
}


def supported_fix_kinds() -> List[str]:
    """Return every rule kind with a deterministic fixer."""
    return sorted(_FIXERS)


def _check_output_name(name: str) -> str:
    value = (name or "").strip()
    if not value:
        raise OfficeParseError("output_filename 不能为空")
    if Path(value).is_absolute() or "/" in value or "\\" in value:
        raise OfficeParseError("output_filename 不能包含路径分隔符")
    if value in (".", ".."):
        raise OfficeParseError("output_filename 非法")
    if not value.lower().endswith(".docx"):
        raise OfficeParseError("output_filename 必须以 .docx 结尾")
    return value


def _rel(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root).as_posix()
    except ValueError:
        return path.name


def _save_and_replace(document: Any, temporary: Path, target: Path) -> None:
    try:
        document.save(str(temporary))
        last_exc: Optional[Exception] = None
        for delay in _SAVE_RETRY_DELAYS:
            if delay:
                time.sleep(delay)
            try:
                temporary.replace(target)
                return
            except PermissionError as exc:
                last_exc = exc
        if last_exc is not None:
            raise OfficeParseError(f"原子替换失败: {last_exc}") from last_exc
    finally:
        try:
            if temporary.exists():
                temporary.unlink()
        except OSError:
            pass


def fix_word_document(
    file_path: Path,
    rulepack: WordRulePack,
    output_filename: Optional[str] = None,
    *,
    overwrite: bool = False,
    workspace_path: Optional[str] = None,
) -> WordAuditFixResult:
    """Apply deterministic fixes and re-audit the result (never silent)."""
    path = Path(file_path)
    if not path.exists():
        raise OfficeFileNotFoundError(path)
    if not path.is_file():
        raise OfficeParseError("DOCX 路径不是文件")
    if path.suffix.lower() != ".docx":
        raise OfficeParseError("仅支持 .docx 文档修复")
    validate_rulepack(rulepack)
    if overwrite and output_filename:
        raise OfficeParseError("overwrite=true 时不能指定 output_filename")
    root = Path(workspace_path).resolve() if workspace_path else path.parent.resolve()
    if not root.is_dir():
        raise OfficeParseError("工作区目录不存在")

    before = audit_word_document(path, rulepack, workspace_path=workspace_path)
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 — map all parse failures
        raise OfficeParseError(f"Failed to parse DOCX: {exc}") from exc

    ctx = _FixContext(document)
    actions: List[WordFixAction] = []
    warnings: List[Dict[str, str]] = []
    fixed_count = 0
    for rule in rulepack.rules:
        if not rule.fixable:
            continue
        fixer = _FIXERS.get(rule.kind)
        if fixer is None:
            warnings.append(
                {
                    "code": "fix_unsupported_kind",
                    "message": f"规则 {rule.id} 标记可修复，但修复器未实现（{rule.kind}）",
                }
            )
            continue
        count = fixer(ctx, rule)
        if count:
            actions.append(
                WordFixAction(
                    rule_id=rule.id,
                    kind=rule.kind,
                    fixed_count=count,
                    message=rule.title,
                )
            )
            fixed_count += count

    overwritten = False
    if overwrite:
        target = path
        if fixed_count:
            temporary = path.with_name(path.name + ".m2b-tmp")
            _save_and_replace(document, temporary, target)
            overwritten = True
    else:
        name = _check_output_name(output_filename or f"{path.stem}-fixed.docx")
        try:
            target = resolve_within(root, path.parent / name)
        except OfficePathError as exc:
            raise OfficeParseError("输出路径不在工作区内") from exc
        if target == path:
            raise OfficeParseError("输出文件不能覆盖源文件；如需覆盖请使用 overwrite=true")
        if target.exists():
            raise OfficeParseError("输出文件已存在，拒绝覆盖")
        try:
            document.save(str(target))
        except OSError as exc:
            raise OfficeParseError(f"写入输出文件失败: {exc}") from exc

    remaining = audit_word_document(target, rulepack, workspace_path=workspace_path)
    result_warnings = [dict(item) for item in remaining.warnings]
    result_warnings.extend(warnings)
    return WordAuditFixResult(
        ok=remaining.ok,
        source_path=before.source_path,
        output_path=_rel(target, root),
        overwritten=overwritten,
        rule_pack_id=rulepack.id,
        actions=actions,
        fixed_count=fixed_count,
        before=before.summary,
        remaining=remaining,
        warnings=result_warnings,
    )


__all__ = ["fix_word_document", "supported_fix_kinds"]

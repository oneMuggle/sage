"""Read-only Word document profile (M0): page setup, styles, runs, anchors."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn

from backend.office.errors import OfficeFileNotFoundError, OfficeParseError

from .style_resolver import StyleResolver

_HEADING_RE = re.compile(r"^heading\s+(\d+)$", re.I)


def _hash(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:16]


def _heading_level(style_name: str) -> int:
    match = _HEADING_RE.match(style_name.strip())
    return int(match.group(1)) if match else 0


def _rel(path: Path, workspace: Optional[str]) -> str:
    if not workspace:
        return path.name
    try:
        return path.resolve().relative_to(Path(workspace).resolve()).as_posix()
    except ValueError:
        return path.name


def _pt(value: Any) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(float(value.pt), 2)
    except AttributeError:
        return None


def _enum_name(value: Any) -> Optional[str]:
    if value is None:
        return None
    name = getattr(value, "name", None)
    return str(name or value).lower()


def _anchor(
    *,
    part: str = "word/document.xml",
    paragraph_index: Optional[int] = None,
    run_index: Optional[int] = None,
    char_start: int = 0,
    char_end: int = 0,
    text_hash: str = "",
    table_index: Optional[int] = None,
    row: Optional[int] = None,
    col: Optional[int] = None,
) -> Dict[str, Any]:
    return {
        "part": part,
        "paragraph_index": paragraph_index,
        "run_index": run_index,
        "char_start": char_start,
        "char_end": char_end,
        "text_hash": text_hash,
        "table_index": table_index,
        "row": row,
        "col": col,
    }


def _numbering_map(document: Any) -> Dict[str, Dict[str, Dict[str, Optional[str]]]]:
    try:
        numbering = document.part.numbering_part.element
    except (AttributeError, KeyError):
        return {}
    abstracts: Dict[str, Dict[str, Dict[str, Optional[str]]]] = {}
    for abstract in numbering.findall(qn("w:abstractNum")):
        abstract_id = abstract.get(qn("w:abstractNumId"))
        if abstract_id is None:
            continue
        levels: Dict[str, Dict[str, Optional[str]]] = {}
        for level in abstract.findall(qn("w:lvl")):
            ilvl = level.get(qn("w:ilvl")) or "0"
            fmt = level.find(qn("w:numFmt"))
            text = level.find(qn("w:lvlText"))
            levels[ilvl] = {
                "format": fmt.get(qn("w:val")) if fmt is not None else None,
                "text": text.get(qn("w:val")) if text is not None else None,
            }
        abstracts[abstract_id] = levels
    mapping: Dict[str, Dict[str, Dict[str, Optional[str]]]] = {}
    for num in numbering.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        ref = num.find(qn("w:abstractNumId"))
        if num_id is None or ref is None:
            continue
        abstract_id = ref.get(qn("w:val"))
        mapping[num_id] = abstracts.get(str(abstract_id), {})
    return mapping


def _paragraph_numbering(
    paragraph: Any, mapping: Dict[str, Dict[str, Dict[str, Optional[str]]]]
) -> Optional[Dict[str, Any]]:
    ppr = paragraph._p.pPr
    num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
    if num_pr is None:
        return None
    num_id_element = num_pr.find(qn("w:numId"))
    level_element = num_pr.find(qn("w:ilvl"))
    if num_id_element is None:
        return None
    num_id = num_id_element.get(qn("w:val")) or ""
    level = (level_element.get(qn("w:val")) if level_element is not None else "0") or "0"
    try:
        level_number = int(level)
    except ValueError:
        level_number = 0
    definition = mapping.get(num_id, {}).get(str(level_number), {})
    return {
        "num_id": num_id,
        "level": level_number,
        "format": definition.get("format"),
        "text": definition.get("text"),
    }


def _run_profiles(
    paragraph: Any,
    index: int,
    resolver: StyleResolver,
    style_id: Optional[str],
) -> List[Dict[str, Any]]:
    runs: List[Dict[str, Any]] = []
    offset = 0
    text_hash = _hash(paragraph.text)
    for run_index, run in enumerate(paragraph.runs):
        text = run.text or ""
        fmt = resolver.resolve_run(run, style_id)
        start = offset
        end = offset + len(text)
        offset = end
        runs.append(
            {
                "index": run_index,
                "text": text,
                "style": getattr(getattr(run, "style", None), "name", None),
                "bold": bool(fmt.get("bold")),
                "italic": bool(fmt.get("italic")),
                "underline": bool(fmt.get("underline")),
                "font_ascii": fmt.get("font_ascii"),
                "font_east_asia": fmt.get("font_east_asia"),
                "size_pt": fmt.get("size_pt"),
                "color": fmt.get("color"),
                "anchor": _anchor(
                    paragraph_index=index,
                    run_index=run_index,
                    char_start=start,
                    char_end=end,
                    text_hash=text_hash,
                ),
            }
        )
    return runs


def _paragraph_profile(
    paragraph: Any,
    index: int,
    resolver: StyleResolver,
    numbering: Dict[str, Dict[str, Dict[str, Optional[str]]]],
) -> Dict[str, Any]:
    style_id = getattr(getattr(paragraph, "style", None), "style_id", None)
    style_name = (
        getattr(getattr(paragraph, "style", None), "name", None) or style_id or "Normal"
    )
    fmt = paragraph.paragraph_format
    line_spacing = fmt.line_spacing
    if isinstance(line_spacing, (int, float)):  # noqa: UP038 — py38 isinstance
        spacing_value = float(line_spacing)
    else:
        spacing_value = _pt(line_spacing)
    return {
        "index": index,
        "style": str(style_name),
        "text": paragraph.text,
        "level": _heading_level(str(style_name)),
        "alignment": _enum_name(paragraph.alignment),
        "indent_left_pt": _pt(fmt.left_indent),
        "indent_right_pt": _pt(fmt.right_indent),
        "first_line_indent_pt": _pt(fmt.first_line_indent),
        "space_before_pt": _pt(fmt.space_before),
        "space_after_pt": _pt(fmt.space_after),
        "line_spacing": spacing_value,
        "line_spacing_rule": _enum_name(fmt.line_spacing_rule),
        "keep_with_next": bool(fmt.keep_with_next),
        "page_break_before": bool(fmt.page_break_before),
        "numbering": _paragraph_numbering(paragraph, numbering),
        "runs": _run_profiles(paragraph, index, resolver, style_id),
        "anchor": _anchor(
            paragraph_index=index,
            char_start=0,
            char_end=len(paragraph.text),
            text_hash=_hash(paragraph.text),
        ),
    }


def _walk_bookmarks(
    element: Any, position: int = 0
) -> Tuple[int, List[Tuple[str, str, int]], List[Tuple[str, int]]]:
    """Collect (name, id, pos) starts and (id, pos) ends inside an element."""
    starts: List[Tuple[str, str, int]] = []
    ends: List[Tuple[str, int]] = []
    for child in element.iterchildren():
        tag = child.tag
        if tag == qn("w:bookmarkStart"):
            name = child.get(qn("w:name")) or ""
            ident = child.get(qn("w:id")) or ""
            if name and not name.startswith("_"):
                starts.append((name, ident, position))
        elif tag == qn("w:bookmarkEnd"):
            ends.append((child.get(qn("w:id")) or "", position))
        elif tag == qn("w:r"):
            position += len("".join(t.text or "" for t in child.iter(qn("w:t"))))
        elif tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
            position += 1
        else:
            position, nested_starts, nested_ends = _walk_bookmarks(child, position)
            starts.extend(nested_starts)
            ends.extend(nested_ends)
    return position, starts, ends


def _bookmark_profiles(paragraphs: List[Any]) -> List[Dict[str, Any]]:
    records: Dict[str, Dict[str, Any]] = {}
    for index, paragraph in enumerate(paragraphs):
        _position, starts, ends = _walk_bookmarks(paragraph._p)
        text_hash = _hash(paragraph.text)
        for name, ident, position in starts:
            records[ident] = {
                "name": name,
                "anchor": _anchor(
                    paragraph_index=index,
                    char_start=position,
                    char_end=position,
                    text_hash=text_hash,
                ),
                "end_anchor": None,
            }
        for ident, position in ends:
            record = records.get(ident)
            if record is not None and record["end_anchor"] is None:
                record["end_anchor"] = _anchor(
                    paragraph_index=index,
                    char_start=position,
                    char_end=position,
                    text_hash=text_hash,
                )
    return list(records.values())


def _field_profiles(paragraphs: List[Any]) -> List[Dict[str, Any]]:
    fields: List[Dict[str, Any]] = []
    for index, paragraph in enumerate(paragraphs):
        text_hash = _hash(paragraph.text)
        for element in paragraph._p.iter(qn("w:instrText")):
            instr = (element.text or "").strip()
            if instr:
                fields.append(
                    {
                        "instr": instr,
                        "anchor": _anchor(
                            paragraph_index=index,
                            char_start=0,
                            char_end=len(paragraph.text),
                            text_hash=text_hash,
                        ),
                    }
                )
        for element in paragraph._p.iter(qn("w:fldSimple")):
            instr = (element.get(qn("w:instr")) or "").strip()
            if instr:
                fields.append(
                    {
                        "instr": instr,
                        "anchor": _anchor(
                            paragraph_index=index,
                            char_start=0,
                            char_end=len(paragraph.text),
                            text_hash=text_hash,
                        ),
                    }
                )
    return fields


def _cell_text(cell: Any) -> str:
    paragraphs = []
    for paragraph in cell.iter(qn("w:p")):
        paragraphs.append("".join(t.text or "" for t in paragraph.iter(qn("w:t"))))
    return "\n".join(paragraphs)


def _table_profile(table: Any, index: int) -> Dict[str, Any]:
    cells: List[Dict[str, Any]] = []
    rows = list(table._tbl.findall(qn("w:tr")))
    max_cols = 0
    for row_index, row in enumerate(rows):
        col = 0
        for tc in row.findall(qn("w:tc")):
            tc_pr = tc.tcPr
            grid_span = 1
            v_merge = None
            if tc_pr is not None:
                span_element = tc_pr.find(qn("w:gridSpan"))
                if span_element is not None:
                    try:
                        grid_span = int(span_element.get(qn("w:val")) or "1")
                    except ValueError:
                        grid_span = 1
                merge_element = tc_pr.find(qn("w:vMerge"))
                if merge_element is not None:
                    v_merge = merge_element.get(qn("w:val")) or "continue"
            cells.append(
                {
                    "row": row_index,
                    "col": col,
                    "row_span": 1,
                    "col_span": grid_span,
                    "v_merge": v_merge,
                    "text": _cell_text(tc),
                    "anchor": _anchor(table_index=index, row=row_index, col=col),
                }
            )
            col += grid_span
        max_cols = max(max_cols, col)
    return {
        "index": index,
        "rows": len(rows),
        "cols": max_cols,
        "cells": cells,
        "anchor": _anchor(table_index=index),
    }


def _has_page_field(element: Any) -> bool:
    if element is None:
        return False
    for instr in element.iter(qn("w:instrText")):
        if "PAGE" in (instr.text or "").upper():
            return True
    for simple in element.iter(qn("w:fldSimple")):
        if "PAGE" in (simple.get(qn("w:instr")) or "").upper():
            return True
    return False


def _section_profile(section: Any, index: int) -> Dict[str, Any]:
    try:
        header_text = "\n".join(p.text for p in section.header.paragraphs)
        footer_text = "\n".join(p.text for p in section.footer.paragraphs)
        has_page = _has_page_field(section.header._element) or _has_page_field(
            section.footer._element
        )
    except (AttributeError, ValueError):
        header_text = ""
        footer_text = ""
        has_page = False
    return {
        "index": index,
        "start_type": _enum_name(section.start_type),
        "page_width_pt": _pt(section.page_width),
        "page_height_pt": _pt(section.page_height),
        "orientation": _enum_name(section.orientation),
        "margin_top_pt": _pt(section.top_margin),
        "margin_bottom_pt": _pt(section.bottom_margin),
        "margin_left_pt": _pt(section.left_margin),
        "margin_right_pt": _pt(section.right_margin),
        "header_distance_pt": _pt(section.header_distance),
        "footer_distance_pt": _pt(section.footer_distance),
        "header_text": header_text,
        "footer_text": footer_text,
        "has_page_number_field": has_page,
    }


def _has_theme_fonts(document: Any) -> bool:
    for fonts in document.element.iter(qn("w:rFonts")):
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "csTheme"):
            if fonts.get(qn("w:" + attr)):
                return True
    return False


def read_doc_profile(
    file_path: Path,
    *,
    workspace_path: Optional[str] = None,
) -> Dict[str, Any]:
    """Read a .docx file into a read-only DocumentProfile payload."""
    path = Path(file_path)
    if not path.exists():
        raise OfficeFileNotFoundError(path)
    if not path.is_file():
        raise OfficeParseError("DOCX 路径不是文件")
    if path.suffix.lower() != ".docx":
        raise OfficeParseError("仅支持 .docx 文档画像")
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 — map all parse failures
        raise OfficeParseError(f"Failed to parse DOCX: {exc}") from exc

    resolver = StyleResolver(document)
    numbering = _numbering_map(document)
    paragraphs = list(document.paragraphs)
    paragraph_profiles = [
        _paragraph_profile(paragraph, index, resolver, numbering)
        for index, paragraph in enumerate(paragraphs)
    ]
    tables = [
        _table_profile(table, index) for index, table in enumerate(document.tables)
    ]
    sections = [
        _section_profile(section, index)
        for index, section in enumerate(document.sections)
    ]
    bookmarks = _bookmark_profiles(paragraphs)
    fields = _field_profiles(paragraphs)
    warnings: List[Dict[str, str]] = []
    if _has_theme_fonts(document):
        warnings.append(
            {
                "code": "theme_fonts_unresolved",
                "message": "文档使用主题字体；当前仅解析显式 ascii/eastAsia 字体名",
            }
        )

    return {
        "ok": True,
        "source_path": _rel(path, workspace_path),
        "summary": {
            "paragraphs": len(paragraph_profiles),
            "tables": len(tables),
            "sections": len(sections),
            "styles": len(resolver.styles()),
            "bookmarks": len(bookmarks),
            "fields": len(fields),
        },
        "defaults": dict(resolver.defaults),
        "sections": sections,
        "styles": resolver.styles(),
        "paragraphs": paragraph_profiles,
        "tables": tables,
        "bookmarks": bookmarks,
        "fields": fields,
        "warnings": warnings,
    }


__all__ = ["read_doc_profile"]

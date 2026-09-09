# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Read-only diff preview for office update ops (Office parity Item 2.5).

:func:`preview_update` answers "what WOULD this update do?" without
touching the source document:

1. Copy ``source`` to a temp file next to it, apply the ops to the COPY
   via the matching :mod:`backend.office.edit` updater, delete the copy
   in a ``finally`` block.
2. Read structured content before (source) and after (the modified copy)
   with the standard readers (:func:`word.read_docx` /
   :func:`excel.read_xlsx` / :func:`ppt.read_ppt`) and derive a
   human-readable change list per op.
3. Never hard-fail on summarization: an op we can't describe specifically
   becomes a generic ``{op, summary}`` entry. An op the editor rejects
   makes the whole preview fail with ``ok=False`` + a human-readable
   ``error`` so the UI can show WHY the real update would fail.

Change-entry shapes (all ``DiffPreviewChange``: ``{op, target?, before?,
after?, summary?}``):

    word:
        replace_text      {target=find, before/after=~80-char context snippet}
        append_paragraphs {target="document end", after=joined texts}
        append_table      {target="table[N]", after=header row}
        set_table_cell    {target="table[N].cell(r,c)", before/after}
        delete_paragraph  {target=find, before=matched text, summary=count}
        set_paragraph_style / add_image  {target, summary (batch-2 ops)}
    excel:
        set_cells         one entry per cell: {target="Sheet!A1", before/after}
        append_rows       {target=sheet, after=first row preview, summary=count}
        add_sheet / rename_sheet / delete_sheet  {target=sheet, summary}
        (chart / style / freeze ops from batch 2 fall through to the
        generic {op, summary="applies <op> (key=args…)"} entry)
    ppt:
        replace_text      {target="slide[i]", before/after snippets}
        set_slide_title / set_slide_bullets / set_slide_notes
                          {target="slide[i]", before/after}
        append_slide      {target="slide[N]", after=title+bullets preview}
        delete_slide      {target="slide[i]", before=old title}
        add_picture       {target="slide[i]", summary}

The change list is capped at ``MAX_CHANGES`` (200) entries with a
``truncated`` flag — a preview must stay cheap no matter how big the op
list is.
"""

from __future__ import annotations

import contextlib
import os
import shutil
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pydantic import BaseModel, ConfigDict, Field

from .edit import update_document
from .excel import read_xlsx
from .ppt import read_ppt
from .word import read_docx

__all__ = [
    "DiffPreviewChange",
    "DiffPreviewResult",
    "MAX_CHANGES",
    "OfficeExportPdfRequest",
    "OfficeUpdatePreviewRequest",
    "preview_update",
]

#: Upper bound on the returned change list (spec: 200 entries + truncated flag).
MAX_CHANGES = 200

#: Default snippet length for before/after context strings.
_SNIPPET_LIMIT = 80

#: File extension → backend.office.edit doc_type.
_DOC_TYPES = {
    ".docx": "word",
    ".xlsx": "excel",
    ".pptx": "ppt",
}


# ──────────────────────────────────────────────────────────────────────
# Request / response models (imported into backend.api.office_routes —
# models.py is frozen for this wave, so they live here)
# ──────────────────────────────────────────────────────────────────────


class DiffPreviewChange(BaseModel):
    """One human-readable change entry in a preview."""

    model_config = ConfigDict(extra="forbid")

    op: str = Field(description="Op name, e.g. 'replace_text' / 'set_cells'")
    target: Optional[str] = Field(
        default=None, description="Where the change lands: 'Sheet!A1', 'slide[2]', 'table[0]'…"
    )
    before: Optional[str] = Field(default=None, description="Content before the op (snippet)")
    after: Optional[str] = Field(default=None, description="Content after the op (snippet)")
    summary: Optional[str] = Field(
        default=None, description="One-line description when before/after don't tell the story"
    )


class DiffPreviewResult(BaseModel):
    """Result of POST /office/update/preview."""

    model_config = ConfigDict(extra="forbid")

    ok: bool = Field(description="False when applying the ops to the preview copy failed")
    changes: List[DiffPreviewChange] = Field(default_factory=list)
    truncated: bool = Field(
        default=False, description="True when len(changes) was capped at MAX_CHANGES"
    )
    error: Optional[str] = Field(
        default=None, description="Why the real update would fail (set when ok=False)"
    )


class OfficeUpdatePreviewRequest(BaseModel):
    """POST /office/update/preview — exactly one of file_path / doc_id."""

    model_config = ConfigDict(extra="forbid")

    workspace_path: str = Field(description="Absolute path to the workspace dir")
    file_path: Optional[str] = Field(
        default=None, description="Absolute path to the document inside the workspace"
    )
    doc_id: Optional[str] = Field(
        default=None, description="Managed document id (resolved via office_documents)"
    )
    ops: List[Dict[str, Any]] = Field(
        default_factory=list, description="Same op dicts office_update accepts"
    )


class OfficeExportPdfRequest(BaseModel):
    """POST /office/export-pdf."""

    model_config = ConfigDict(extra="forbid")

    workspace_path: str = Field(description="Absolute path to the workspace dir")
    file_path: str = Field(description="Absolute path to the document inside the workspace")


# ──────────────────────────────────────────────────────────────────────
# Small string helpers
# ──────────────────────────────────────────────────────────────────────


def _clamp(text: Any, limit: int = _SNIPPET_LIMIT) -> Optional[str]:
    """Stringify + ellipsize to ``limit`` chars; None in, None out."""
    if text is None:
        return None
    s = str(text).strip()
    if len(s) <= limit:
        return s
    return s[: limit - 1] + "…"


def _context(text: str, needle: str, limit: int = _SNIPPET_LIMIT) -> str:
    """A ~``limit``-char window of ``text`` around the first ``needle`` hit.

    The needle is guaranteed to appear in full inside the window, so
    ``window.replace(needle, replacement)`` renders the after-snippet.
    """
    pos = text.find(needle)
    if pos < 0:
        return _clamp(text, limit) or ""
    if len(text) <= limit:
        return text
    center = pos + len(needle) // 2
    start = max(0, center - limit // 2)
    end = min(len(text), start + limit)
    start = max(0, end - limit)
    prefix = "…" if start > 0 else ""
    suffix = "…" if end < len(text) else ""
    return f"{prefix}{text[start:end]}{suffix}"


def _short_repr(value: Any, limit: int = 40) -> str:
    """Compact repr of an op argument for generic summaries."""
    text = str(value)
    if len(text) > limit:
        text = text[: limit - 1] + "…"
    return text


def _generic_change(op: Any) -> DiffPreviewChange:
    """Last-resort entry so unknown-but-valid ops never break the preview."""
    if not isinstance(op, dict):
        return DiffPreviewChange(
            op="unknown", summary=f"unrecognized op payload: {_short_repr(op)}"
        )
    name = str(op.get("op") or "unknown")
    args = ", ".join(f"{k}={_short_repr(v)}" for k, v in op.items() if k != "op")
    summary = f"applies {name}" + (f" ({args})" if args else "")
    return DiffPreviewChange(op=name, summary=summary)


def _first_op_error(ops: List[Dict[str, Any]], results: List[Dict[str, Any]]) -> str:
    """Human-readable message for the first failed op (edit.py contract)."""
    for op, res in zip(ops, results):  # noqa: B905 — py3.8 无 strict 参数
        if isinstance(res, dict) and not res.get("ok", True):
            name = res.get("op") or (op.get("op") if isinstance(op, dict) else "unknown")
            return f"op {name} failed: {res.get('error')}"
    return "one or more ops failed"


# ──────────────────────────────────────────────────────────────────────
# Structured before/after readers
# ──────────────────────────────────────────────────────────────────────


def _read_structured(doc_type: str, path: Path):
    """Dispatch to the standard reader (workspace_path unused for previews)."""
    if doc_type == "word":
        return read_docx(path)
    if doc_type == "excel":
        return read_xlsx(path)
    return read_ppt(path)


def _xlsx_cell_map(path: Path) -> Dict[str, Dict[str, str]]:
    """``{sheet: {"A1": "text"}}`` of non-empty cells, formulas as text.

    Why not read_xlsx().sheets.rows: the reader skips fully-empty rows, so
    array indexes drift from real row numbers once a blank row exists.
    set_cells previews need address-accurate before-values (including
    formula strings, hence ``data_only=False``).
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    out: Dict[str, Dict[str, str]] = {}
    for ws in wb.worksheets:
        cells: Dict[str, str] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = str(cell.value)
        out[ws.title] = cells
    return out


# ──────────────────────────────────────────────────────────────────────
# Per-format summarizers (each returns a list — set_cells fans out)
# ──────────────────────────────────────────────────────────────────────


def _word_snippet(before: Any, find: str) -> Optional[str]:
    """First body-paragraph or table-cell text containing ``find``."""
    for para in before.paragraphs:
        if find in para.text:
            return para.text
    for table in before.tables:
        for row in table.rows:
            for cell_text in row:
                if find in cell_text:
                    return cell_text
    return None


def _word_changes(  # noqa: PLR0911 — op 分发表
    op: Dict[str, Any],
    result: Dict[str, Any],
    before: Any,
    after: Any,
) -> List[DiffPreviewChange]:
    name = str(op.get("op"))

    if name == "replace_text":
        find, replace = str(op.get("find", "")), str(op.get("replace", ""))
        snippet = _word_snippet(before, find) or find
        before_snippet = _context(snippet, find)
        after_snippet = before_snippet.replace(find, replace)
        hits = result.get("replacements") if isinstance(result, dict) else None
        summary = (
            f"replaces {hits} occurrence(s) of {find!r}" if hits is not None else None
        )
        return [
            DiffPreviewChange(
                op=name, target=find, before=before_snippet, after=after_snippet, summary=summary
            )
        ]

    if name == "append_paragraphs":
        specs = op.get("paragraphs") or []
        texts = [str(s.get("text", "")) for s in specs if isinstance(s, dict)]
        return [
            DiffPreviewChange(
                op=name,
                target="document end",
                after=_clamp("; ".join(texts)),
                summary=f"appends {len(specs)} paragraph(s) at the end of the document",
            )
        ]

    if name == "append_table":
        headers = op.get("headers") or []
        rows = op.get("rows") or []
        return [
            DiffPreviewChange(
                op=name,
                target=f"table[{len(before.tables)}]",
                after=_clamp(" | ".join(str(h) for h in headers)),
                summary=f"appends a table with {1 + len(rows)} row(s) including header",
            )
        ]

    if name == "set_table_cell":
        ti = int(op.get("table_index", -1))
        ri = int(op.get("row", -1))
        ci = int(op.get("col", -1))
        before_val: Optional[str] = None
        if 0 <= ti < len(before.tables):
            rows = before.tables[ti].rows
            if 0 <= ri < len(rows) and 0 <= ci < len(rows[ri]):
                before_val = _clamp(rows[ri][ci])
        return [
            DiffPreviewChange(
                op=name,
                target=f"table[{ti}].cell({ri},{ci})",
                before=before_val,
                after=_clamp(op.get("text")),
            )
        ]

    if name == "delete_paragraph":
        find = str(op.get("find", ""))
        snippet = _word_snippet(before, find)
        removed = result.get("removed") if isinstance(result, dict) else None
        count = f"{removed} " if removed is not None else ""
        return [
            DiffPreviewChange(
                op=name,
                target=find,
                before=_context(snippet, find) if snippet else None,
                summary=f"deletes {count}paragraph(s) matching {find!r}",
            )
        ]

    if name == "set_paragraph_style":
        target = op.get("match", op.get("index"))
        style_bits = ", ".join(
            f"{k}={_short_repr(v)}" for k, v in op.items() if k not in ("op", "index", "match")
        )
        return [
            DiffPreviewChange(
                op=name,
                target=_short_repr(target) if target is not None else None,
                summary="restyles paragraph" + (f" ({style_bits})" if style_bits else ""),
            )
        ]

    if name == "add_image":
        src = op.get("path") or ("inline base64 image" if op.get("base64") else None)
        delta = None
        if after.images != before.images:
            delta = f"{before.images} → {after.images} image(s)"
        return [
            DiffPreviewChange(
                op=name,
                target=_short_repr(src) if src else None,
                after=delta,
                summary=f"inserts an image ({after.images} in the document after this update)",
            )
        ]

    return []


def _excel_changes(  # noqa: PLR0911 — op 分发表
    op: Dict[str, Any],
    result: Dict[str, Any],
    cell_map: Optional[Dict[str, Dict[str, str]]],
) -> List[DiffPreviewChange]:
    name = str(op.get("op"))

    if name == "set_cells":
        sheet = str(op.get("sheet", ""))
        sheet_cells = (cell_map or {}).get(sheet, {})
        changes = []
        for spec in op.get("cells") or []:
            if not isinstance(spec, dict) or not spec.get("addr"):
                continue
            addr = str(spec["addr"])
            value = spec.get("value")
            changes.append(
                DiffPreviewChange(
                    op=name,
                    target=f"{sheet}!{addr}",
                    before=_clamp(sheet_cells.get(addr)),
                    after=_clamp(value),
                )
            )
        return changes or [_generic_change(op)]

    if name == "append_rows":
        sheet = str(op.get("sheet", ""))
        rows = op.get("rows") or []
        first = rows[0] if rows and isinstance(rows[0], list) else []
        max_row = result.get("max_row") if isinstance(result, dict) else None
        suffix = f", new last row {max_row}" if max_row is not None else ""
        return [
            DiffPreviewChange(
                op=name,
                target=sheet,
                after=_clamp(" | ".join(str(v) for v in first)),
                summary=f"appends {len(rows)} row(s){suffix}",
            )
        ]

    if name == "add_sheet":
        sheet = str(op.get("name", ""))
        headers = op.get("headers") or []
        row_count = len(op.get("rows") or [])
        return [
            DiffPreviewChange(
                op=name,
                target=sheet,
                after=_clamp(" | ".join(str(h) for h in headers)),
                summary=(
                    f"creates sheet {sheet!r} with {len(headers)} header(s) "
                    f"and {row_count} row(s)"
                ),
            )
        ]

    if name == "rename_sheet":
        return [
            DiffPreviewChange(
                op=name,
                target="sheets",
                before=str(op.get("from", "")),
                after=str(op.get("to", "")),
            )
        ]

    if name == "delete_sheet":
        sheet = str(op.get("name", ""))
        return [DiffPreviewChange(op=name, target=sheet, summary=f"deletes sheet {sheet!r}")]

    # Batch-2 style/chart ops (add_chart / set_column_width /
    # set_number_format / set_fill / freeze_panes / …) land here: describe
    # generically by op name + key args so valid-but-unknown ops still show.
    return []


def _ppt_snippet(before: Any, find: str) -> Tuple[Optional[int], Optional[str]]:
    """First slide whose joined text contains ``find`` → (index, text)."""
    for slide in before.slides:
        parts = [slide.title or ""] + list(slide.text_blocks)
        joined = " | ".join(parts)
        if find in joined:
            return slide.index, joined
    return None, None


def _ppt_changes(  # noqa: PLR0911 — op 分发表
    op: Dict[str, Any],
    result: Dict[str, Any],
    before: Any,
    after: Any,
) -> List[DiffPreviewChange]:
    name = str(op.get("op"))

    if name == "replace_text":
        find, replace = str(op.get("find", "")), str(op.get("replace", ""))
        idx, snippet = _ppt_snippet(before, find)
        if snippet is None:
            return [DiffPreviewChange(op=name, target=find, summary=f"replaces {find!r}")]
        before_snippet = _context(snippet, find)
        hits = result.get("replacements") if isinstance(result, dict) else None
        summary = (
            f"replaces {hits} occurrence(s) of {find!r}" if hits is not None else None
        )
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                before=before_snippet,
                after=before_snippet.replace(find, replace),
                summary=summary,
            )
        ]

    if name == "set_slide_title":
        idx = int(op.get("index", -1))
        old = before.slides[idx].title if 0 <= idx < len(before.slides) else None
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                before=_clamp(old),
                after=_clamp(op.get("title")),
            )
        ]

    if name == "set_slide_bullets":
        idx = int(op.get("index", -1))
        bullets = [str(b) for b in op.get("bullets") or []]
        old_blocks = before.slides[idx].text_blocks if 0 <= idx < len(before.slides) else []
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                before=_clamp("; ".join(old_blocks)),
                after=_clamp("; ".join(bullets)),
                summary=f"sets {len(bullets)} bullet(s)",
            )
        ]

    if name == "set_slide_notes":
        idx = int(op.get("index", -1))
        old = before.slides[idx].notes if 0 <= idx < len(before.slides) else None
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                before=_clamp(old),
                after=_clamp(op.get("notes")),
            )
        ]

    if name == "append_slide":
        title = str(op.get("title", ""))
        bullets = [str(b) for b in op.get("bullets") or []]
        idx = result.get("index") if isinstance(result, dict) else None
        if idx is None:
            idx = len(after.slides) - 1
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                after=_clamp("; ".join([title] + bullets)),
                summary=f"appends a slide titled {title!r}",
            )
        ]

    if name == "delete_slide":
        idx = int(op.get("index", -1))
        old_title = before.slides[idx].title if 0 <= idx < len(before.slides) else None
        suffix = f" ({old_title!r})" if old_title else ""
        return [
            DiffPreviewChange(
                op=name,
                target=f"slide[{idx}]",
                before=_clamp(old_title),
                summary=f"deletes slide {idx}{suffix}",
            )
        ]

    if name == "add_picture":
        idx = op.get("index", op.get("slide"))
        target = f"slide[{idx}]" if idx is not None else None
        return [
            DiffPreviewChange(
                op=name,
                target=target,
                summary="inserts a picture" + (f" on slide {idx}" if idx is not None else ""),
            )
        ]

    # Other future batch-2 ops fall through to the generic entry.
    return []


# ──────────────────────────────────────────────────────────────────────
# Public entry point
# ──────────────────────────────────────────────────────────────────────


def preview_update(source: Path, ops: List[Dict[str, Any]]) -> DiffPreviewResult:
    """Dry-run ``ops`` against a throwaway copy of ``source``.

    The source file is never modified: ops are applied to a temp copy in
    the same directory (removed in ``finally``), and the change list is
    derived from structured reads of source vs. modified copy. Invalid ops
    surface as ``DiffPreviewResult(ok=False, error=…)`` instead of raising.
    """
    source = Path(source)
    doc_type = _DOC_TYPES.get(source.suffix.lower())
    if doc_type is None:
        return DiffPreviewResult(
            ok=False, error=f"unsupported extension for preview: {source.suffix!r}"
        )

    try:
        before = _read_structured(doc_type, source)
    except Exception as exc:  # noqa: BLE001 — preview 把一切失败折算成 ok=False
        return DiffPreviewResult(ok=False, error=f"failed to read source document: {exc}")

    fd, tmp_name = tempfile.mkstemp(
        prefix=f".{source.stem}.preview.", suffix=source.suffix, dir=str(source.parent)
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        shutil.copyfile(source, tmp_path)
        saved, per_op_results = update_document(doc_type, tmp_path, ops)
        if not saved:
            return DiffPreviewResult(ok=False, error=_first_op_error(ops, per_op_results))
        after = _read_structured(doc_type, tmp_path)
    except Exception as exc:  # noqa: BLE001 — 同上
        return DiffPreviewResult(ok=False, error=f"failed to apply ops to preview copy: {exc}")
    finally:
        with contextlib.suppress(OSError):
            tmp_path.unlink()

    # Excel set_cells needs address-accurate before-values from the source.
    cell_map: Optional[Dict[str, Dict[str, str]]] = None
    if doc_type == "excel":
        try:
            cell_map = _xlsx_cell_map(source)
        except Exception:  # noqa: BLE001 — before 值尽力而为，失败退化为 None
            cell_map = {}

    changes: List[DiffPreviewChange] = []
    for i, op in enumerate(ops):
        result = per_op_results[i] if i < len(per_op_results) else {}
        if doc_type == "word":
            entries = _word_changes(
                op, result if isinstance(result, dict) else {}, before, after
            )
        elif doc_type == "excel":
            entries = _excel_changes(op, result if isinstance(result, dict) else {}, cell_map)
        else:
            entries = _ppt_changes(
                op, result if isinstance(result, dict) else {}, before, after
            )
        if not entries:
            entries = [_generic_change(op)]
        changes.extend(entries)

    truncated = len(changes) > MAX_CHANGES
    if truncated:
        changes = changes[:MAX_CHANGES]
    return DiffPreviewResult(ok=True, changes=changes, truncated=truncated)

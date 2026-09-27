# ruff: noqa: UP038 — Python 3.8 does not support isinstance union types
"""Isolated formula entry point. Heavy work starts only after parent handshake."""
from __future__ import annotations

import contextlib
import json
import math
import sys
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

# -I omits the script directory and PYTHONPATH; use only this installed tree.
if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from backend.office.worker_process import MAX_INPUT_BYTES, MAX_OUTPUT_BYTES, apply_posix_limits

_MAX_INT_FLOAT = 1e15


def _load_formula_cells(file_path: Path, max_cells: int) -> Dict[str, Dict[str, str]]:
    """openpyxl ``data_only=False`` → ``{sheet标题: {坐标: 公式文本}}``.

    只收集字符串值且以 ``=`` 开头的单元格（ArrayFormula 对象跳过，
    与读取侧 ``_extract_sheet_formulas`` 的口径一致）。
    """
    from openpyxl import load_workbook

    mapping: Dict[str, Dict[str, str]] = {}
    count = 0
    wb = load_workbook(str(file_path), data_only=False, read_only=True, keep_links=False)
    try:
        for ws in wb.worksheets:
            cells: Dict[str, str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        count += 1
                        if count > max_cells:
                            raise ValueError("formula cell limit exceeded")
                        cells[cell.coordinate] = value
            if cells:
                mapping[ws.title] = cells
    finally:
        wb.close()
    return mapping


def _split_solution_key(
    key: str,
    book_name: str,
    sheets_by_upper: Dict[str, str],
) -> Optional[Tuple[str, str]]:
    """把 ``formulas`` 输出键归一化为 ``(sheet标题, 单元格坐标)``.

    键形如 ``"'[book.xlsx]SHEET1'!B4"``：``!`` 前是 ``'<book>]SHEET'``
    引用部（book 为路径 basename、sheet 名大写），后是单元格地址或
    区域。book 与传入文件 basename 忽略大小写比对、sheet 大写名经
    ``sheets_by_upper`` 映射回 openpyxl 原标题；对不上或缺成分 → None。
    """
    ref_part, sep, addr = key.rpartition("!")
    if not sep or ":" in addr:
        return None
    if not (ref_part.startswith("'") and ref_part.endswith("'")):
        return None
    inner = ref_part[1:-1]
    book, bracket, sheet_upper = inner.partition("]")
    if not bracket or not sheet_upper or not book.startswith("["):
        return None
    book = book[1:]  # partition 保留前导 '['
    if book.lower() != book_name.lower():
        return None
    sheet_title = sheets_by_upper.get(sheet_upper.upper())
    if sheet_title is None:
        return None
    return sheet_title, addr


def _scalar_value(raw: Any) -> Any:
    """``formulas`` 的 1x1 数组包装 → JSON/字符串友好的原生标量。

    ``array([[7.0]])`` → ``7``；字符串/布尔原样；提取失败退回原对象
    （显示层 ``_cell_value_to_str`` 还有一层 str 兜底）。
    """
    value = getattr(raw, "value", raw)
    with contextlib.suppress(Exception):  # 非二维下标形态一律退回原值
        value = value[0, 0]
    to_item = getattr(value, "item", None)
    if callable(to_item):
        with contextlib.suppress(Exception):  # 标量化失败不致命
            value = to_item()
    # formulas 全程按 float 计算：整数值转 int，让展示为 "30" 而非 "30.0"
    if (
        isinstance(value, float)
        and not isinstance(value, bool)
        and value.is_integer()
        and abs(value) < _MAX_INT_FLOAT
    ):
        return int(value)
    return value


def _calculate(file_path: Path) -> Dict[str, Any]:
    """在 worker 进程里跑 formulas 的 load + calculate。

    返回 ``formulas`` 原生解算 dict（键形如 ``"'[book.xlsx]SHEET1'!B4"``，
    值为 1x1 数组包装），由调用方按 openpyxl 侧的公式清单归一化。
    ``import formulas`` 放在子进程内：模块加载本身就有可观的启动开销，
    与编译/计算一并计入同一份超时预算。
    """
    import formulas

    model = formulas.ExcelModel().loads(str(file_path)).finish()
    return model.calculate()


def calculate_values(file_path: Path, max_cells: int) -> Dict[str, Dict[str, Any]]:
    formula_cells = _load_formula_cells(file_path, max_cells)
    if not formula_cells:
        return {}
    solution = _calculate(file_path)
    sheets_by_upper = {title.upper(): title for title in formula_cells}
    evaluated: Dict[str, Dict[str, Any]] = {}
    for key, raw in solution.items():
        parsed = _split_solution_key(str(key), file_path.name, sheets_by_upper)
        if parsed is None:
            continue
        sheet_title, addr = parsed
        if addr not in formula_cells[sheet_title]:
            continue
        value = _scalar_value(raw)
        # Unsupported/nonfinite values remain unresolved, never fabricated zero.
        if not isinstance(value, (str, int, float, bool)) or (
            isinstance(value, float) and not math.isfinite(value)
        ):
            continue
        evaluated.setdefault(sheet_title, {})[addr] = value
    return evaluated


def main() -> int:
    if sys.stdin.buffer.readline(4) != b"GO\n":
        return 1
    try:
        source, destination, max_cells, cpu_seconds = sys.argv[1:]
        apply_posix_limits(int(cpu_seconds))
        path = Path(source)
        if not path.is_file() or path.stat().st_size > MAX_INPUT_BYTES:
            return 1
        values = calculate_values(path, int(max_cells))
        payload = json.dumps({"version": 1, "values": values}, allow_nan=False).encode("utf-8")
        if len(payload) > MAX_OUTPUT_BYTES:
            return 1
        Path(destination).write_bytes(payload)
        return 0
    except Exception:  # noqa: BLE001 — parent observes nonzero exit, not engine objects
        return 1


if __name__ == "__main__":
    sys.exit(main())

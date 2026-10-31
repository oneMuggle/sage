"""Word template slot profile (M1a): content controls, bookmarks, placeholders."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn

from .errors import OfficeFileNotFoundError, OfficeParseError
from .word_template import _validate_docx_zip

_PLACEHOLDER_RE = re.compile(
    r"(\{\{[^{}\n]{1,120}\}\}|<<[^<>\n]{1,120}>>|\[\[[^\[\]\n]{1,120}\]\]|"
    r"【[^【】\n]{1,120}】)"
)
_JINJA_CONTROL_RE = re.compile(r"\{[%#]")


def _hash(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:16]


def _rel(path: Path, workspace: Optional[str]) -> str:
    if not workspace:
        return path.name
    try:
        return path.resolve().relative_to(Path(workspace).resolve()).as_posix()
    except ValueError:
        return path.name


def _anchor(
    *,
    part: str = "word/document.xml",
    paragraph_index: Optional[int] = None,
    table_index: Optional[int] = None,
    row: Optional[int] = None,
    col: Optional[int] = None,
    char_start: int = 0,
    char_end: int = 0,
    sdt_index: Optional[int] = None,
    tag: Optional[str] = None,
    bookmark: Optional[str] = None,
    section: Optional[int] = None,
    story: Optional[str] = None,
) -> Dict[str, Any]:
    payload: Dict[str, Any] = {
        "part": part,
        "paragraph_index": paragraph_index,
        "table_index": table_index,
        "row": row,
        "col": col,
        "char_start": char_start,
        "char_end": char_end,
        "text_hash": "",
    }
    if sdt_index is not None:
        payload["sdt_index"] = sdt_index
    if tag:
        payload["tag"] = tag
    if bookmark:
        payload["bookmark"] = bookmark
    if section is not None:
        payload["section"] = section
    if story:
        payload["story"] = story
    return payload


def _placeholder_slots(
    text: str,
    *,
    paragraph_index: Optional[int] = None,
    table_index: Optional[int] = None,
    row: Optional[int] = None,
    col: Optional[int] = None,
    part: str = "word/document.xml",
    section: Optional[int] = None,
    story: Optional[str] = None,
) -> List[Dict[str, Any]]:
    slots: List[Dict[str, Any]] = []
    for match in _PLACEHOLDER_RE.finditer(text):
        raw = match.group(0)
        if raw.startswith("{{"):
            kind = "jinja"
            name = raw[2:-2].split("|", 1)[0].strip()
        elif raw.startswith(("<<", "[[")):
            kind = "placeholder"
            name = raw[2:-2].strip()
        else:
            kind = "placeholder"
            name = raw[1:-1].strip()
        if not name:
            continue
        slots.append(
            {
                "name": name,
                "kind": kind,
                "type": "text",
                "required": True,
                "description": f"模板占位符：{raw[:80]}",
                "current_value": raw,
                "confidence": 0.9,
                "anchor": _anchor(
                    part=part,
                    paragraph_index=paragraph_index,
                    table_index=table_index,
                    row=row,
                    col=col,
                    char_start=match.start(),
                    char_end=match.end(),
                    section=section,
                    story=story,
                ),
            }
        )
    return slots


def _sdt_name(sdt: Any) -> Optional[str]:
    props = sdt.find(qn("w:sdtPr"))
    if props is None:
        return None
    for tag_name in ("w:tag", "w:alias"):
        element = props.find(qn(tag_name))
        if element is not None:
            value = (element.get(qn("w:val")) or "").strip()
            if value:
                return value
    return None


def _sdt_text(sdt: Any) -> str:
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return ""
    return "".join(t.text or "" for t in content.iter(qn("w:t")))


def _walk_bookmarks(
    element: Any, position: int = 0
) -> Tuple[int, List[Tuple[str, str, int]], List[Tuple[str, int]]]:
    starts: List[Tuple[str, str, int]] = []
    ends: List[Tuple[str, int]] = []
    for child in element.iterchildren():
        if child.tag == qn("w:bookmarkStart"):
            name = child.get(qn("w:name")) or ""
            ident = child.get(qn("w:id")) or ""
            if name and not name.startswith("_"):
                starts.append((name, ident, position))
        elif child.tag == qn("w:bookmarkEnd"):
            ends.append((child.get(qn("w:id")) or "", position))
        elif child.tag == qn("w:r"):
            position += len("".join(t.text or "" for t in child.iter(qn("w:t"))))
        elif child.tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
            position += 1
        else:
            position, nested_starts, nested_ends = _walk_bookmarks(child, position)
            starts.extend(nested_starts)
            ends.extend(nested_ends)
    return position, starts, ends


def _bookmark_slots(
    document: Any,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
    slots: List[Dict[str, Any]] = []
    warnings: List[Dict[str, str]] = []
    open_map: Dict[str, Dict[str, Any]] = {}
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        _, starts, ends = _walk_bookmarks(paragraph._p)
        for name, ident, position in starts:
            open_map[ident] = {
                "name": name,
                "paragraph": index,
                "start": position,
                "text": paragraph.text,
            }
        for ident, position in ends:
            record = open_map.pop(ident, None)
            if record is None:
                continue
            if record["paragraph"] == index:
                current = record["text"][record["start"] : position]
                slots.append(
                    {
                        "name": record["name"],
                        "kind": "bookmark",
                        "type": "text",
                        "required": False,
                        "description": f"书签：{record['name']}",
                        "current_value": current,
                        "confidence": 0.7,
                        "anchor": _anchor(
                            paragraph_index=index,
                            char_start=record["start"],
                            char_end=position,
                            bookmark=record["name"],
                        ),
                    }
                )
            else:
                warnings.append(
                    {
                        "code": "bookmark_cross_paragraph",
                        "message": f"书签 {record['name']} 跨段落，M1a 暂不填充",
                    }
                )
    return slots, warnings


def _sdt_slot(
    sdt: Any,
    sdt_index: int,
    *,
    part: str = "word/document.xml",
    section: Optional[int] = None,
    story: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    name = _sdt_name(sdt)
    if not name:
        return None
    text = _sdt_text(sdt)
    return {
        "name": name,
        "kind": "content_control",
        "type": "rich_text" if "\n" in text else "text",
        "required": True,
        "description": f"内容控件：{name}",
        "current_value": text,
        "confidence": 0.95,
        "anchor": _anchor(
            part=part,
            sdt_index=sdt_index,
            tag=name,
            section=section,
            story=story,
        ),
    }


def _story_slots(
    document: Any,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]], List[str]]:
    slots: List[Dict[str, Any]] = []
    warnings: List[Dict[str, str]] = []
    texts: List[str] = []
    for section_index, section in enumerate(document.sections):
        for story_name, story in (
            ("header", section.header),
            ("footer", section.footer),
        ):
            try:
                part = str(story.part.partname).lstrip("/")
                paragraphs = list(story.paragraphs)
                tables = list(story.tables)
            except (AttributeError, KeyError, ValueError):
                warnings.append(
                    {
                        "code": "story_unavailable",
                        "message": f"无法读取第 {section_index + 1} 节 {story_name}",
                    }
                )
                continue
            for paragraph_index, paragraph in enumerate(paragraphs):
                texts.append(paragraph.text)
                slots.extend(
                    _placeholder_slots(
                        paragraph.text,
                        paragraph_index=paragraph_index,
                        part=part,
                        section=section_index,
                        story=story_name,
                    )
                )
                for sdt_index, sdt in enumerate(paragraph._p.iter(qn("w:sdt"))):
                    slot = _sdt_slot(
                        sdt,
                        sdt_index,
                        part=part,
                        section=section_index,
                        story=story_name,
                    )
                    if slot is not None:
                        slots.append(slot)
            for table_index, table in enumerate(tables):
                for row_index, row in enumerate(table.rows):
                    for col_index, cell in enumerate(row.cells):
                        for paragraph in cell.paragraphs:
                            texts.append(paragraph.text)
                            slots.extend(
                                _placeholder_slots(
                                    paragraph.text,
                                    table_index=table_index,
                                    row=row_index,
                                    col=col_index,
                                    part=part,
                                    section=section_index,
                                    story=story_name,
                                )
                            )
    return slots, warnings, texts


def _analysis(document: Any, source_path: str) -> Dict[str, Any]:
    slots: List[Dict[str, Any]] = []
    warnings: List[Dict[str, str]] = []
    for index, paragraph in enumerate(document.paragraphs):
        slots.extend(_placeholder_slots(paragraph.text, paragraph_index=index))
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    slots.extend(
                        _placeholder_slots(
                            paragraph.text,
                            table_index=table_index,
                            row=row_index,
                            col=col_index,
                        )
                    )
    sdt_index = 0
    for sdt in document.element.body.iter(qn("w:sdt")):
        slot = _sdt_slot(sdt, sdt_index)
        if slot is not None:
            slots.append(slot)
            sdt_index += 1
    story_slots, story_warnings, story_texts = _story_slots(document)
    slots.extend(story_slots)
    warnings.extend(story_warnings)
    bookmark_slots, bookmark_warnings = _bookmark_slots(document)
    slots.extend(bookmark_slots)
    warnings.extend(bookmark_warnings)

    all_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs] + story_texts
    )
    has_control_flow = bool(_JINJA_CONTROL_RE.search(all_text))
    if has_control_flow:
        warnings.append(
            {
                "code": "jinja_control_flow",
                "message": "模板包含 Jinja 控制流；M1a 填充不支持，请使用现有 /office/word/template/fill",
            }
        )
    if not slots:
        warnings.append({"code": "no_slots", "message": "未发现可填充槽位"})
    counts: Dict[str, int] = {}
    for slot in slots:
        counts[slot["kind"]] = counts.get(slot["kind"], 0) + 1
    return {
        "ok": True,
        "kind": "word-template-profile",
        "source_path": source_path,
        "slots": slots,
        "warnings": warnings,
        "summary": {
            "slots": len(slots),
            "kinds": counts,
            "has_jinja_control": has_control_flow,
        },
    }


def analyze_word_template_profile(
    file_path: Path,
    *,
    workspace_path: Optional[str] = None,
) -> Dict[str, Any]:
    """Analyze a .docx template into a unified slot profile (M1a)."""
    path = Path(file_path)
    if not path.exists():
        raise OfficeFileNotFoundError(path)
    if not path.is_file():
        raise OfficeParseError("模板路径不是文件")
    if path.suffix.lower() != ".docx":
        raise OfficeParseError("仅支持 .docx 模板")
    _validate_docx_zip(path)
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to parse DOCX template: {exc}") from exc
    return _analysis(document, _rel(path, workspace_path))


__all__ = ["analyze_word_template_profile"]

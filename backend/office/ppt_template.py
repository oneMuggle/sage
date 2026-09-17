"""PPT 模板占位符分析与填充（office-p5a）。

与 word_template（docxtpl 模式）对偶的 pptx 模板能力：

- **分析**：枚举模板的母版版式（layout 名 + 占位符 idx/类型），供生成端
  的 ``slides[].layout`` 字段按名引用；
- **填充**：以模板副本为基础，按 ``{1 起页号, 占位符 idx, 文本}`` 逐条
  写入占位符文本框（``text_frame.text`` 赋值替换全部内容），另存为
  新文件 —— 模板原件绝不修改。

安全与失败契约：
- 源/目标都在 workspace 围栏内（``resolve_within``）；
- 目标文件已存在 → 拒绝（生成类端点同口径）；
- 页号/占位符 idx 不存在 → op 级失败（all-or-nothing）；
- Python 3.8 兼容（typing.* generics），win7 cherry-pick 友好。
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pydantic import BaseModel, ConfigDict, Field

from .errors import OfficeError, OfficeParseError
from .path_safety import resolve_within
from .storage import validate_workspace

logger = logging.getLogger(__name__)

__all__ = [
    "PptTemplateAnalyzeRequest",
    "PptTemplateAnalyzeResult",
    "PptTemplateLayoutInfo",
    "PptTemplatePlaceholderInfo",
    "PptTemplateFillRequest",
    "PptTemplateFillResult",
    "analyze_ppt_template",
    "fill_ppt_template",
]


class PptTemplatePlaceholderInfo(BaseModel):
    """模板版式里的一个占位符。"""

    model_config = ConfigDict(extra="forbid")

    idx: int = Field(description="占位符 idx（fill 端按此定位）")
    type: str = Field(description="占位符类型名（如 TITLE / BODY / PICTURE）")
    is_title: bool = Field(description="是否标题占位符")


class PptTemplateLayoutInfo(BaseModel):
    """模板里的一个版式。"""

    model_config = ConfigDict(extra="forbid")

    name: str = Field(description="版式名（generate 的 slides[].layout 按名引用）")
    placeholders: List[PptTemplatePlaceholderInfo] = Field(default_factory=list)


class PptTemplateAnalyzeRequest(BaseModel):
    """POST /office/ppt/analyze-template 请求体。"""

    model_config = ConfigDict(extra="forbid")

    workspace_path: str
    file_path: str


class PptTemplateAnalyzeResult(BaseModel):
    """模板分析结果：版式清单。"""

    model_config = ConfigDict(extra="forbid")

    ok: bool
    layouts: List[PptTemplateLayoutInfo] = Field(default_factory=list)
    error: Optional[str] = None


class PptTemplateFillItem(BaseModel):
    """一条占位符填充：1 起页号 + 占位符 idx + 文本。"""

    model_config = ConfigDict(extra="forbid")

    slide_number: int = Field(ge=1, le=1000, description="1 起页号")
    placeholder_idx: int = Field(ge=0, description="占位符 idx")
    text: str = Field(min_length=1, max_length=5000)


class PptTemplateFillRequest(BaseModel):
    """POST /office/ppt/fill-template 请求体。"""

    model_config = ConfigDict(extra="forbid")

    workspace_path: str
    file_path: str
    output_filename: str = Field(
        min_length=1,
        max_length=200,
        description="输出文件名（自动补 .pptx 扩展名；须与模板同目录产出）",
    )
    fills: List[PptTemplateFillItem] = Field(min_length=1, max_length=200)


class PptTemplateFillResult(BaseModel):
    """填充结果：成功携带输出路径。"""

    model_config = ConfigDict(extra="forbid")

    ok: bool
    output_path: Optional[str] = None
    filename: Optional[str] = None
    filled_count: int = 0
    error: Optional[str] = None


def _load_prs(file_path: Path):
    """打开 pptx（python-pptx 懒加载），失败折算 OfficeParseError。"""
    try:
        from pptx import Presentation

        return Presentation(str(file_path))
    except Exception as exc:  # noqa: BLE001 — 统一折算解析错误
        raise OfficeParseError(
            f"Failed to parse PPTX: {exc}", file_path=file_path
        ) from exc


def analyze_ppt_template(request: PptTemplateAnalyzeRequest) -> PptTemplateAnalyzeResult:
    """枚举模板的母版版式与占位符（workspace 围栏内）。"""
    try:
        workspace = validate_workspace(Path(request.workspace_path))
        source = resolve_within(workspace, Path(request.file_path))
        if not source.is_file():
            return PptTemplateAnalyzeResult(ok=False, error="源路径不是常规文件")
        prs = _load_prs(source)
    except OfficeError as exc:
        return PptTemplateAnalyzeResult(ok=False, error=str(exc))
    except Exception as exc:  # noqa: BLE001 — 解析失败降级
        return PptTemplateAnalyzeResult(ok=False, error=f"模板解析失败: {exc}")

    layouts: List[PptTemplateLayoutInfo] = []
    for layout in prs.slide_layouts:
        phs: List[PptTemplatePlaceholderInfo] = []
        for ph in layout.placeholders:
            phs.append(
                PptTemplatePlaceholderInfo(
                    idx=int(ph.placeholder_format.idx),
                    type=str(ph.placeholder_format.type),
                    is_title=ph.placeholder_format.idx == 0,
                )
            )
        layouts.append(PptTemplateLayoutInfo(name=layout.name or "", placeholders=phs))
    return PptTemplateAnalyzeResult(ok=True, layouts=layouts)


def fill_ppt_template(request: PptTemplateFillRequest) -> PptTemplateFillResult:  # noqa: PLR0912, PLR0911 — 失败契约的逐条早退
    """以模板副本为基础按占位符填充，另存为新文件（模板原件不动）。"""
    try:
        return _fill_inner(request)
    except OfficeError as exc:
        return PptTemplateFillResult(ok=False, error=str(exc))
    except Exception as exc:  # noqa: BLE001 — 契约：绝不向路由层抛异常
        logger.exception("ppt template fill crashed unexpectedly")
        return PptTemplateFillResult(
            ok=False, error=f"模板填充失败：内部错误（{type(exc).__name__}，详见日志）"
        )


def _fill_inner(request: PptTemplateFillRequest) -> PptTemplateFillResult:  # noqa: PLR0911 — 失败如约的逐条早退
    workspace = validate_workspace(Path(request.workspace_path))
    source = resolve_within(workspace, Path(request.file_path))
    if not source.is_file():
        return PptTemplateFillResult(ok=False, error="源路径不是常规文件")

    filename = request.output_filename.strip()
    if not filename or "/" in filename or "\\" in filename or ".." in filename:
        return PptTemplateFillResult(ok=False, error="非法输出文件名")
    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    target = resolve_within(workspace, source.parent / filename)
    if target.exists():
        return PptTemplateFillResult(ok=False, error=f"输出文件已存在: {filename}")

    prs = _load_prs(source)
    slides = prs.slides
    applied = 0
    for item in request.fills:
        if not (1 <= item.slide_number <= len(slides)):
            return PptTemplateFillResult(
                ok=False,
                error=f"slide_number 超出范围: {item.slide_number}（共 {len(slides)} 页）",
            )
        slide = slides[item.slide_number - 1]
        matched: Optional[Tuple[Any, Dict[str, Any]]] = None
        for ph in slide.placeholders:
            if int(ph.placeholder_format.idx) == item.placeholder_idx:
                matched = (ph, {"idx": item.placeholder_idx})
                break
        if matched is None:
            return PptTemplateFillResult(
                ok=False,
                error=f"第 {item.slide_number} 页不存在占位符 idx={item.placeholder_idx}",
            )
        matched[0].text_frame.text = item.text
        applied += 1

    if applied != len(request.fills):
        return PptTemplateFillResult(ok=False, error="存在未应用的填充项")
    try:
        prs.save(str(target))
    except Exception as exc:  # noqa: BLE001 — 保存失败折算
        return PptTemplateFillResult(ok=False, error=f"保存失败: {exc}")

    return PptTemplateFillResult(
        ok=True,
        output_path=str(target),
        filename=target.name,
        filled_count=applied,
    )

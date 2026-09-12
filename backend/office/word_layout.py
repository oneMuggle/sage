"""Word 版式引擎（Round 7 FormatSpec —— "版式即配置"）。

把 :class:`~backend.office.models.WordFormatSpec` 描述的版式要求以确定性
代码注入 python-docx 文档对象（styles.xml 样式定义 + 节属性 + 页眉页脚），
使期刊论文/项目文档/公文类硬性格式不再依赖 LLM 在 prompt 里"口头约定"。

设计约定：
- 模型（pydantic，零 docx 依赖）定义在 ``models.py``，本模块只含 docx 侧
  应用逻辑 —— 维持 ``scripts/verify-office-paths.py`` canary 的
  "models 仅依赖 pydantic" 前提；
- ``apply_format_spec`` 对 ``spec`` 的每个子项独立判空：None 项完全不触碰
  文档，因此不传 ``format_spec`` 时 ``generate_docx`` 行为与历史版本一致；
- 样式补丁只动字号/加粗/颜色/间距/对齐，字体（rFonts）仍由
  ``word.set_doc_default_font`` 统一负责，避免两条路径互相覆盖；
- 颜色写入时同步清除 ``w:themeColor`` 等 theme 属性 —— 与
  ``_patch_style_rfonts`` 清 theme 引用同理，theme 属性优先级高于显式值，
  不清掉某些渲染端会忽略显式颜色。

Word 兼容性说明：页码使用 ``w:fldSimple w:instr="PAGE"`` 域，Word / WPS /
LibreOffice 打开即渲染（fldSimple 自带占位 run，无需打开后手动更新域）。
"""

from __future__ import annotations

from typing import Any, Dict, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

from .models import (
    WordBodyStyleSpec,
    WordFormatSpec,
    WordHeaderFooterSpec,
    WordHeadingStyleSpec,
    WordPageSetupSpec,
)

#: format_spec 的 align 取值 → python-docx 段落对齐枚举
_ALIGN_MAP: Dict[str, int] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

#: headings 键（h1/h2/h3）→ Word 内置样式名
_HEADING_STYLE_NAMES: Dict[str, str] = {
    "h1": "Heading 1",
    "h2": "Heading 2",
    "h3": "Heading 3",
}

#: 颜色元素上需要清除的 theme 属性（优先级高于显式 w:val）
_COLOR_THEME_ATTRS = ("themeColor", "themeShade", "themeTint")


def _style_or_none(doc: Document, style_name: str):
    """按名取样式；默认模板缺该样式时静默跳过（与 _STYLES_TO_PATCH 容错一致）。"""
    try:
        return doc.styles[style_name]
    except KeyError:
        return None


def _clear_color_theme(style) -> None:
    """清除样式颜色上的 theme 属性，保证显式 RGB 生效。"""
    rPr = style.element.get_or_add_rPr()
    color = rPr.find(qn("w:color"))
    if color is None:
        return
    for attr in _COLOR_THEME_ATTRS:
        qualified = qn(f"w:{attr}")
        if qualified in color.attrib:
            del color.attrib[qualified]


def _hex_to_rgb(color: str) -> RGBColor:
    """"#2F5496" / "2F5496" → RGBColor（pattern 校验已在模型层完成）。"""
    return RGBColor.from_string(color.lstrip("#").upper())


def _apply_page_setup(doc: Document, page: WordPageSetupSpec) -> None:
    section = doc.sections[0]
    width = None
    height = None
    if page.size == "A4":
        width, height = Mm(210), Mm(297)
    elif page.size == "letter":
        width, height = Inches(8.5), Inches(11)
    if width is not None and height is not None:
        if page.orientation == "landscape":
            width, height = height, width
        section.page_width = width
        section.page_height = height
    if page.orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
    elif page.orientation == "portrait":
        section.orientation = WD_ORIENT.PORTRAIT
    if page.margins_cm is not None:
        margins = page.margins_cm
        if margins.top is not None:
            section.top_margin = Cm(margins.top)
        if margins.bottom is not None:
            section.bottom_margin = Cm(margins.bottom)
        if margins.left is not None:
            section.left_margin = Cm(margins.left)
        if margins.right is not None:
            section.right_margin = Cm(margins.right)


def _apply_body_style(doc: Document, body: WordBodyStyleSpec) -> None:
    style = _style_or_none(doc, "Normal")
    if style is None:
        return
    if body.font_size_pt is not None:
        style.font.size = Pt(body.font_size_pt)
    paragraph_format = style.paragraph_format
    if body.line_spacing is not None:
        paragraph_format.line_spacing = body.line_spacing
    if body.first_line_indent_cm is not None:
        paragraph_format.first_line_indent = Cm(body.first_line_indent_cm)
    if body.space_after_pt is not None:
        paragraph_format.space_after = Pt(body.space_after_pt)
    if body.align is not None:
        paragraph_format.alignment = _ALIGN_MAP[body.align]


def _apply_heading_style(doc: Document, style_name: str, spec: WordHeadingStyleSpec) -> None:
    style = _style_or_none(doc, style_name)
    if style is None:
        return
    if spec.font_size_pt is not None:
        style.font.size = Pt(spec.font_size_pt)
    if spec.bold is not None:
        style.font.bold = spec.bold
    if spec.color is not None:
        style.font.color.rgb = _hex_to_rgb(spec.color)
        _clear_color_theme(style)
    paragraph_format = style.paragraph_format
    if spec.space_before_pt is not None:
        paragraph_format.space_before = Pt(spec.space_before_pt)
    if spec.space_after_pt is not None:
        paragraph_format.space_after = Pt(spec.space_after_pt)
    if spec.align is not None:
        paragraph_format.alignment = _ALIGN_MAP[spec.align]


def _use_header_part(section, part_attr: str):
    """取页眉/页脚部件并断开链接（对首节是显式创建空部件），返回其首段。"""
    part = getattr(section, part_attr)
    if part.is_linked_to_previous:
        part.is_linked_to_previous = False
    return part.paragraphs[0]


def _apply_header(doc: Document, header: WordHeaderFooterSpec) -> None:
    paragraph = _use_header_part(doc.sections[0], "header")
    if header.text is not None:
        paragraph.text = header.text
    if header.align is not None:
        paragraph.alignment = _ALIGN_MAP[header.align]


def _apply_footer(doc: Document, footer: WordHeaderFooterSpec) -> None:
    paragraph = _use_header_part(doc.sections[0], "footer")
    if footer.text is not None:
        paragraph.text = footer.text
    if footer.align is not None:
        paragraph.alignment = _ALIGN_MAP[footer.align]
    if footer.page_number:
        # w:fldSimple 携带占位 run，Word/WPS/LibreOffice 打开即显示页码，
        # 无需"打开后更新域"步骤（对比 fldChar 复杂域）。
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        fld.append(run)
        paragraph._p.append(fld)


def apply_format_spec(doc: Document, spec: Optional[WordFormatSpec]) -> None:
    """把版式规范应用到文档（样式定义 + 节属性 + 页眉页脚）。

    在 ``generate_docx`` 中于 ``set_doc_default_font`` 之后、写正文之前调用，
    使后续 add_heading/add_paragraph 直接继承补丁后的样式。
    """
    if spec is None:
        return
    if spec.page is not None:
        _apply_page_setup(doc, spec.page)
    if spec.body is not None:
        _apply_body_style(doc, spec.body)
    if spec.title is not None:
        _apply_heading_style(doc, "Title", spec.title)
    if spec.headings:
        for key, heading_spec in spec.headings.items():
            style_name = _HEADING_STYLE_NAMES.get(key)
            if style_name is not None:
                _apply_heading_style(doc, style_name, heading_spec)
    if spec.header is not None:
        _apply_header(doc, spec.header)
    if spec.footer is not None:
        _apply_footer(doc, spec.footer)


# ──────────────────────────────────────────────────────────────────────
# 表格排版（Round 8）：三线表 / 表头跨页重复 / 固定列宽 / 合并单元格
#
# 风格说明：与 Round 7 样式补丁同理，直接操作 oxml——python-docx 的
# Table 对象不暴露 borders/tblHeader 等接口。
# 边框粗细单位：w:sz 为 1/8 磅（1.5pt = 12，0.75pt = 6）。
# ──────────────────────────────────────────────────────────────────────


def _border_element(tag: str, *, val: str, sz: Optional[str] = None) -> Any:
    """构造单个 w:top/bottom/left/right/insideH/insideV 边框元素。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), sz or "0")
    el.set(qn("w:space"), "0")
    if val == "single":
        el.set(qn("w:color"), "000000")
    return el


def apply_three_line_table(table: Any) -> None:
    """学术三线表：顶/底线 1.5pt，表头下边线 0.75pt，其余无框线。

    覆盖顺序：先清既有 tblBorders（幂等），再写表级边框，最后给表头行
    每个单元格写 tcBorders bottom（单元格边框优先于表级 insideH 生效）。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    borders.append(_border_element("top", val="single", sz="12"))
    borders.append(_border_element("bottom", val="single", sz="12"))
    for edge in ("left", "right", "insideH", "insideV"):
        borders.append(_border_element(edge, val="none"))
    tbl_pr.append(borders)

    if not table.rows:
        return
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        tc_borders.append(_border_element("bottom", val="single", sz="6"))
        tc_pr.append(tc_borders)


def enable_header_repeat(table: Any) -> None:
    """表头跨页重复：首行 trPr 追加 w:tblHeader（Word 标准排版标记）。"""
    from docx.oxml import OxmlElement

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def set_fixed_column_widths(table: Any, widths_cm: Any) -> None:
    """固定布局 + 逐行设置列宽（厘米）。

    Word 对列宽的解析以 tcW 为准（tblLayout fixed 关闭自动伸缩），
    python-docx 的 ``cell.width`` 写的正是 tcW，故逐行全量设置。
    """
    from docx.shared import Cm

    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_cm):
                cell.width = Cm(widths_cm[idx])


def apply_cell_merges(table: Any, merges: Any, *, n_rows: int, n_cols: int) -> None:
    """应用合并区域；越界抛 OfficeGenerateError（进生成失败通路而非静默）。"""
    from .errors import OfficeGenerateError

    for merge in merges:
        if (
            merge.min_row > merge.max_row
            or merge.min_col > merge.max_col
            or merge.max_row >= n_rows
            or merge.max_col >= n_cols
        ):
            raise OfficeGenerateError(
                f"cell merge out of range: rows {merge.min_row}-{merge.max_row}, "
                f"cols {merge.min_col}-{merge.max_col} (table {n_rows}x{n_cols})"
            )
        anchor = table.cell(merge.min_row, merge.min_col)
        other = table.cell(merge.max_row, merge.max_col)
        anchor.merge(other)


def add_caption(doc: Document, text: str, *, kind: str, number: int) -> None:
    """追加居中题注段落（"图N　caption" / "表N　caption"，9pt 小五）。

    ``kind``: "figure"（图，题注在图下方）/ "table"（表，题注在表上方）——
    位置由调用方的插入顺序决定，本函数只负责段落样式与编号文本。
    """
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{'图' if kind == 'figure' else '表'}{number}　{text}")
    run.font.size = Pt(9)


def heading_number_prefix(counters: Any, level: int) -> str:
    """推进 h1/h2/h3 计数器并返回 "N" / "N.M" / "N.M.K" 前缀。

    ``counters`` 为长度 3 的可变列表；高级别出现时重置下级计数
    （h2 变化 → h3 归零），与常规文档编号规则一致。
    """
    counters[level - 1] += 1
    for idx in range(level, 3):
        counters[idx] = 0
    return ".".join(str(counters[i]) for i in range(level))


def insert_toc_field(doc: Document, toc: Any) -> None:
    """在文档当前末尾（生成流程中即标题之后）插入目录标题 + TOC 域 + 分页。

    - 目录标题用加粗居中普通段落（非 Heading 样式）——避免被目录域
      自我收录、也不参与多级标题编号检查（与参考文献节标题同理）；
    - TOC 域为 w:fldSimple，instr 形如 'TOC 反斜杠o "1-3" 反斜杠h
      反斜杠z 反斜杠u'，占位 run 提示用户在 Word/WPS 中更新域生成目录。
    """
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(toc.heading_text))
    run.bold = True
    run.font.size = Pt(16)

    first, last = toc.level_range()
    if first == 1:
        instr = 'TOC \\o "1-' + str(last) + '" \\h \\z \\u'
    else:
        instr = 'TOC \\o "' + str(first) + "-" + str(last) + '" \\h \\z \\u'

    toc_paragraph = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    placeholder = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = str(toc.placeholder_text)
    placeholder.append(text)
    fld.append(placeholder)
    toc_paragraph._p.append(fld)

    doc.add_page_break()

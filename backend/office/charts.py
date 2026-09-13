# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Chart & image pipeline for Office generation（批次 2.1）.

Three responsibilities, all with lazy third-party imports so the office
package keeps importing cleanly when optional deps are absent:

1. :func:`render_chart_png` — matplotlib rendering of a
   :class:`~backend.office.models.ChartSpec` to a ``<uuid>.png`` file.
   matplotlib is a main-channel-only dependency (not bundled for Win7),
   so the import sits inside the function and an ImportError is mapped to
   :class:`OfficeGenerateError` with a user-readable message.
2. :func:`build_openpyxl_chart` — Excel 原生图表（LineChart/BarChart/
   PieChart）构建，供 ``generate_xlsx``（生成期挂载）与 ``edit.add_chart``
   （编辑期挂载）复用，避免两份 openpyxl 校验逻辑。
3. :func:`decode_image_base64` / :func:`resolve_image_payload` — 图片素材
   解析（base64 data URI 或文件路径 → 字节，≤10MB），供 word.py / ppt.py /
   edit.py 复用。

图表中文字体：matplotlib 默认 DejaVu Sans 无 CJK 字形，中文标题会变
「豆腐块」；这里把 ``font.sans-serif`` 依次指到微软雅黑 / 黑体 /
Noto Sans CJK，全部缺失时回退 DejaVu（行为与 pdf.py 的 CJK 兜底同思路）。
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

from .errors import OfficeGenerateError
from .models import ChartSpec

logger = logging.getLogger(__name__)

#: 嵌入图片字节上限（与 word_template._MAX_IMAGE_BYTES 保持一致）。
MAX_IMAGE_BYTES = 10 * 1024 * 1024
#: base64 解码前的字符数快速上限：4/3 膨胀 + 余量，超限直接拒绝，
#: 避免为明显超大的 payload 付出解码内存。
_MAX_IMAGE_B64_CHARS = (MAX_IMAGE_BYTES // 3 + 1) * 4 + 16

#: 数值类型元组常量：py38 兼容 + 绕开 UP038（同 excel._NUMERIC_TYPES 惯例）。
_NUMERIC_TYPES = (int, float)

#: 中文字体候选（按优先级）；前三个覆盖 Windows / 常见 Linux 发行版。
_CJK_FONT_CANDIDATES = ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans")

#: 固定画布：6x4 英寸 @150dpi → 900x600 像素，固定尺寸保证生成物可预期。
_CHART_SIZE_INCHES = (6.0, 4.0)
_CHART_DPI = 150

#: render 支持的图表类型（与 models.ChartSpec.type 的 Literal 对齐；
#: render 对 model_construct 绕过校验的直接调用仍做运行时守卫）。
_CHART_TYPES = ("line", "bar", "hbar", "pie")

#: Excel 原生图表类型 → openpyxl Chart 类（不支持 'hbar'，Excel 原生
#: 横向条形图后续需要时再以 BarChart(type="bar") 扩展）。
_EXCEL_CHART_MAKERS = ("line", "bar", "pie")


# ──────────────────────────────────────────────────────────────────────
# matplotlib 渲染
# ──────────────────────────────────────────────────────────────────────


def _numeric(values: List[Any]) -> bool:
    """True 当列表全为数值（bool 不算，matplotlib 也不该把它当 0/1 轴）。"""
    return all(isinstance(v, _NUMERIC_TYPES) and not isinstance(v, bool) for v in values)


def _category_labels(spec: ChartSpec) -> List[str]:
    """类别标签：显式 labels > 首系列 x（转文本）> 序号 1..n。"""
    if spec.labels:
        return [str(v) for v in spec.labels]
    first = spec.series[0]
    if first.x:
        return [str(v) for v in first.x]
    return [str(i + 1) for i in range(len(first.y))]


def _validate_lengths(spec: ChartSpec, labels: List[str]) -> None:
    """x/y 与类别标签长度一致性校验，避免 matplotlib 半路抛难懂的错误。"""
    if spec.labels:
        for series in spec.series:
            if len(series.y) != len(labels):
                raise OfficeGenerateError(
                    f"图表数据长度不一致: labels={len(labels)}, "
                    f"系列 {series.name!r} y={len(series.y)}"
                )
    for series in spec.series:
        if series.x and len(series.x) != len(series.y):
            raise OfficeGenerateError(
                f"图表数据长度不一致: 系列 {series.name!r} x={len(series.x)}, y={len(series.y)}"
            )


def _apply_cjk_fonts(plt: Any) -> None:
    """让中文标题/标签不出现「豆腐块」（axes.unicode_minus 修负号显示）。"""
    plt.rcParams["font.sans-serif"] = list(_CJK_FONT_CANDIDATES)
    plt.rcParams["axes.unicode_minus"] = False


def render_chart_png(spec: ChartSpec, output_dir: Path) -> Path:
    """Render ``spec`` to a ``<uuid>.png`` under ``output_dir``; return the path.

    matplotlib 懒加载：本机未安装时抛
    :class:`OfficeGenerateError`（"图表生成需要 matplotlib（本机未安装）"），
    而不是裸 ImportError。Agg 后端 + 固定尺寸（6x4in @150dpi）+
    ``tight_layout``，输出文件名用 uuid 避免并发覆盖。

    Raises:
        OfficeGenerateError: matplotlib 缺失、图表类型不支持或数据长度不一致。
    """
    chart_type = str(getattr(spec, "type", "") or "")
    if chart_type not in _CHART_TYPES:
        raise OfficeGenerateError(f"不支持的图表类型: {chart_type!r}（支持 {_CHART_TYPES}）")

    try:
        import matplotlib
    except ImportError as exc:
        raise OfficeGenerateError("图表生成需要 matplotlib（本机未安装）") from exc

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    _apply_cjk_fonts(plt)
    labels = _category_labels(spec)
    _validate_lengths(spec, labels)

    fig, ax = plt.subplots(figsize=_CHART_SIZE_INCHES, dpi=_CHART_DPI)
    try:
        if chart_type == "pie":
            # 饼图只取第一条系列（多条系列无几何意义）；无 labels 时回退序号。
            if len(spec.series) > 1:
                logger.warning("pie 图表仅使用第一条系列（共 %s 条）", len(spec.series))
            first = spec.series[0]
            ax.pie(first.y, labels=labels, autopct="%1.1f%%")
            ax.axis("equal")
        elif chart_type == "line":
            _plot_line(ax, spec, labels)
        elif chart_type == "hbar":
            _plot_hbar(ax, spec, labels)
        else:  # "bar"
            _plot_bar(ax, spec, labels)

        if spec.title:
            ax.set_title(spec.title)
        fig.tight_layout()

        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        out_path = output_dir / f"{uuid.uuid4().hex}.png"
        fig.savefig(str(out_path), dpi=_CHART_DPI, format="png")
    finally:
        plt.close(fig)
    return out_path


def _plot_line(ax: Any, spec: ChartSpec, labels: List[str]) -> None:
    """折线图：x 全数值时按真值画，否则按类别位置 + tick 标签。"""
    positions = list(range(len(labels)))
    use_numeric_x = not spec.labels and all(
        series.x and _numeric(series.x) for series in spec.series
    )
    for series in spec.series:
        xs = series.x if use_numeric_x and series.x else positions
        ax.plot(xs, series.y, marker="o", label=series.name)
    if not use_numeric_x:
        ax.set_xticks(positions)
        ax.set_xticklabels(labels)
    if len(spec.series) > 1:
        ax.legend(loc="best")
    ax.grid(True, axis="y", alpha=0.3)


def _plot_bar(ax: Any, spec: ChartSpec, labels: List[str]) -> None:
    """纵向柱状图：多系列按类别分组并排（宽度均分）。"""
    n = len(spec.series)
    width = 0.8 / n
    positions = list(range(len(labels)))
    for i, series in enumerate(spec.series):
        ax.bar([p + i * width for p in positions], series.y, width, label=series.name)
    ax.set_xticks([p + width * (n - 1) / 2 for p in positions])
    ax.set_xticklabels(labels)
    if n > 1:
        ax.legend(loc="best")


def _plot_hbar(ax: Any, spec: ChartSpec, labels: List[str]) -> None:
    """横向条形图（类别在 y 轴，首个类别在顶部，符合阅读顺序）。"""
    n = len(spec.series)
    height = 0.8 / n
    positions = list(range(len(labels)))
    for i, series in enumerate(spec.series):
        ax.barh([p + i * height for p in positions], series.y, height, label=series.name)
    ax.set_yticks([p + height * (n - 1) / 2 for p in positions])
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    if n > 1:
        ax.legend(loc="best")


# ──────────────────────────────────────────────────────────────────────
# Excel 原生图表（openpyxl）
# ──────────────────────────────────────────────────────────────────────


def _validated_range(data_ref: Any) -> Dict[str, int]:
    """校验 {min_col, min_row, max_col, max_row}（编辑 op 传 plain dict，
    生成路径传 ExcelCellRange 模型 —— 两者都接受）。

    Returns a plain dict of ints suitable for ``openpyxl.chart.Reference(ws, **...)``.
    Raises ValueError on any shape violation.
    """
    if data_ref is None:
        raise ValueError("data_ref_required")
    if not isinstance(data_ref, dict):
        data_ref = {
            key: getattr(data_ref, key, None)
            for key in ("min_col", "min_row", "max_col", "max_row")
        }
    out: Dict[str, int] = {}
    for key in ("min_col", "min_row", "max_col", "max_row"):
        value = data_ref.get(key)
        if not isinstance(value, int) or isinstance(value, bool) or value < 1:
            raise ValueError(f"invalid_data_ref: {key} 必须为 >=1 的整数")
        out[key] = value
    if out["max_col"] < out["min_col"] or out["max_row"] < out["min_row"]:
        raise ValueError("invalid_data_ref: max_* 不得小于 min_*")
    return out


def build_openpyxl_chart(ws: Any, spec: Any) -> Any:
    """按 :class:`ExcelChartSpec`（或同形 dict/model）在 ``ws`` 上构建原生图表。

    只构建并 ``add_chart``，不保存工作簿 —— 调用方（generate_xlsx /
    edit.update_xlsx）各自掌控保存时机。openpyxl 懒加载。

    Raises:
        ValueError: 图表类型不支持、anchor/data_ref 非法或 sheet 上无法定位。
    """
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    # dict（编辑 op）与 ExcelChartSpec 统一按属性/键双通道取值。
    def _get(name: str, default: Any = None) -> Any:
        if isinstance(spec, dict):
            return spec.get(name, default)
        return getattr(spec, name, default)

    chart_type = _get("type")
    if chart_type not in _EXCEL_CHART_MAKERS:
        raise ValueError(f"unsupported_chart_type: {chart_type!r}（支持 {_EXCEL_CHART_MAKERS}）")
    data_ref = _validated_range(_get("data_ref"))
    anchor = _get("anchor")
    if not isinstance(anchor, str) or not anchor.strip():
        raise ValueError("anchor_required")

    maker = {"line": LineChart, "bar": BarChart, "pie": PieChart}[chart_type]
    chart = maker()
    chart.add_data(
        Reference(ws, **data_ref),
        titles_from_data=bool(_get("titles_from_data", False)),
        from_rows=bool(_get("from_rows", False)),
    )
    categories_ref = _get("categories_ref")
    if categories_ref is not None:
        chart.set_categories(Reference(ws, **_validated_range(categories_ref)))
    title = _get("title")
    if title:
        chart.title = str(title)
    ws.add_chart(chart, str(anchor).strip())
    return chart


# ──────────────────────────────────────────────────────────────────────
# 图片素材解析
# ──────────────────────────────────────────────────────────────────────


def decode_image_base64(encoded: str) -> bytes:
    """解码 base64 图片（接受裸 base64 或 ``data:image/...`` data URI）。

    超 :data:`MAX_IMAGE_BYTES` 抛 ValueError（含编码前字符数快速预检，
    避免为超大 payload 白白付出解码内存）。
    """
    if not isinstance(encoded, str) or not encoded.strip():
        raise ValueError("image_source_required")
    text = encoded.strip()
    if text.startswith("data:"):
        _, _, text = text.partition(",")
        if not text:
            raise ValueError("invalid_image_data: data URI 缺少 payload")
    if len(text) > _MAX_IMAGE_B64_CHARS:
        raise ValueError(
            f"image_too_large: base64 超过 {MAX_IMAGE_BYTES} 字节上限"
        )
    try:
        payload = base64.b64decode(text, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError(f"invalid_image_data: base64 解码失败（{exc}）") from exc
    if len(payload) > MAX_IMAGE_BYTES:
        raise ValueError(f"image_too_large: 图片超过 {MAX_IMAGE_BYTES} 字节上限")
    return payload


def resolve_image_payload(source: str, *, search_dirs: Optional[List[Path]] = None) -> bytes:
    """把 ``source``（data URI base64 或图片路径）解析为图片字节。

    路径解析顺序：绝对路径直接使用；相对路径依次尝试 ``search_dirs``
    （调用方按需传工作区根 / 文档所在目录）。命中即读，超限拒绝；
    全部未命中抛 ``ValueError("image_file_not_found: ...")``。

    Round 22：命中后接入懒加载压缩管线——>8MB 的图片在 Pillow 可用时
    自动降采样压缩到阈值内；未安装 Pillow 时原样返回（行为零变化）。
    """
    if not isinstance(source, str) or not source.strip():
        raise ValueError("image_source_required")
    text = source.strip()
    if text.startswith("data:"):
        from .image_optimize import optimize_image_bytes

        return optimize_image_bytes(decode_image_base64(text))

    candidates: List[Path] = []
    raw = Path(text)
    if raw.is_absolute():
        candidates.append(raw)
    else:
        for directory in search_dirs or ():
            candidates.append(Path(directory) / raw)
    for candidate in candidates:
        if candidate.is_file():
            payload = candidate.read_bytes()
            if len(payload) > MAX_IMAGE_BYTES:
                raise ValueError(f"image_too_large: 图片超过 {MAX_IMAGE_BYTES} 字节上限")
            from .image_optimize import optimize_image_bytes

            return optimize_image_bytes(payload)
    raise ValueError(f"image_file_not_found: {text}")


def image_bytes_to_stream(payload: bytes) -> io.BytesIO:
    """包一层 BytesIO（python-docx / python-pptx 的 add_picture 都吃流）。"""
    return io.BytesIO(payload)


__all__ = [
    "MAX_IMAGE_BYTES",
    "build_openpyxl_chart",
    "decode_image_base64",
    "image_bytes_to_stream",
    "render_chart_png",
    "resolve_image_payload",
]

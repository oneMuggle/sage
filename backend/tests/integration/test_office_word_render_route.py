"""Integration test for the rendered audit route (M5a)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict

import pytest
from docx import Document

import backend.api.office_word_audit_routes as routes
from backend.api.office_word_audit_routes import (
    WordAuditRenderRequest,
    render_audit_endpoint,
)
from backend.office.errors import OfficePathError, OfficeSizeLimitError
from backend.office.word_rulepack_models import WordRenderAuditResult, WordRenderIssue


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def test_render_endpoint_passes_options_to_engine(
    monkeypatch: pytest.MonkeyPatch, workspace: Path
) -> None:
    document = Document()
    document.add_paragraph("正文")
    path = workspace / "doc.docx"
    document.save(str(path))
    captured: Dict[str, Any] = {}

    def fake_audit(*args: Any, **kwargs: Any) -> WordRenderAuditResult:
        captured["args"] = args
        captured["kwargs"] = kwargs
        return WordRenderAuditResult(
            ok=True,
            source_path="doc.docx",
            rendered=True,
            method="libreoffice",
            page_count=1,
            issues=[
                WordRenderIssue(
                    page_index=0,
                    category="page_number",
                    severity="warning",
                    message="第 1 页底部未检测到页码",
                )
            ],
            warnings=[],
        )

    monkeypatch.setattr(routes, "audit_word_render", fake_audit)

    result = render_audit_endpoint(
        WordAuditRenderRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            expected_size="A4",
            margins_cm={"top": 3.7},
            allowed_fonts=["仿宋_GB2312"],
            require_page_number=True,
        )
    )

    assert result.rendered is True
    assert result.page_count == 1
    options = captured["kwargs"]["options"]
    assert options["expected_size"] == "A4"
    assert options["margins_cm"] == {"top": 3.7}
    assert options["allowed_fonts"] == ["仿宋_GB2312"]
    assert options["require_page_number"] is True


def test_render_endpoint_rejects_size_and_escape(workspace: Path) -> None:
    document = Document()
    document.add_paragraph("正文")
    path = workspace / "doc.docx"
    document.save(str(path))

    with pytest.raises(OfficeSizeLimitError):
        render_audit_endpoint(
            WordAuditRenderRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                max_size_bytes=1024,
            )
        )
    with pytest.raises(OfficePathError):
        render_audit_endpoint(
            WordAuditRenderRequest(
                workspace_path=str(workspace),
                file_path="../outside.docx",
            )
        )

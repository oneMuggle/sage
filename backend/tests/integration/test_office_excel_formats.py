"""Integration tests for Excel sheet format enhancements (Round 14).

Covers: header style (bold/fill/center), freeze header, autofit columns
(respecting explicit column_widths), per-column number formats (unknown
headers ignored, formula cells skipped), the office_create tool path, and
backward compatibility for payloads without the new fields.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, sheets: list) -> Path:
    req = OfficeExcelGenerateRequest(
        workspace_path="",
        filename="fmt.xlsx",
        sheets=sheets,
    )
    return generate_xlsx(req, output_dir=str(tmp_path))


def _sheet(tmp_path: Path, name: str = "Sheet1"):
    wb = load_workbook(str(tmp_path / "fmt.xlsx"))
    return wb[name]


# ──────────────────────────────────────────────────────────────────────
# Backward compatibility
# ──────────────────────────────────────────────────────────────────────


def test_legacy_payload_unchanged(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{"name": "S", "headers": ["A", "B"], "rows": [["1", "2"]]}],
    )
    ws = _sheet(tmp_path, "S")
    # pandas to_excel 表头默认加粗+居中（既有基线）；Round 14 未启用时
    # 不加填充/冻结/数字格式
    assert ws["A1"].fill.fill_type != "solid"
    assert ws.freeze_panes is None
    assert ws["A2"].number_format == "General"


# ──────────────────────────────────────────────────────────────────────
# Header style / freeze
# ──────────────────────────────────────────────────────────────────────


def test_header_style_applied_to_first_row_only(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["名称", "金额"],
            "rows": [["x", "100"]],
            "header_style": True,
            "freeze_header": True,
        }],
    )
    ws = _sheet(tmp_path, "S")
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.start_color.rgb == "00D9D9D9"
    # 数据行不受影响（无填充）
    assert ws["A2"].fill.fill_type != "solid"
    assert ws.freeze_panes == "A2"


# ──────────────────────────────────────────────────────────────────────
# Autofit + explicit widths precedence
# ──────────────────────────────────────────────────────────────────────


def test_autofit_widens_columns(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["很短", "一个很长的中文表头列"],
            "rows": [["短", "这是一段相当长的中文内容用于撑宽列"]],
            "autofit_columns": True,
        }],
    )
    ws = _sheet(tmp_path, "S")
    wide = ws.column_dimensions["B"].width
    assert wide is not None
    assert wide > 10
    # autofit 对所有有内容的列生效：A 列也被设置，但按其内容明显更窄
    narrow = ws.column_dimensions["A"].width
    assert narrow is not None
    assert narrow < wide


def test_explicit_column_widths_win_over_autofit(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["长表头内容撑宽", "B"],
            "rows": [["很长的内容AAAA", "1"]],
            "column_widths": [8.0],
            "autofit_columns": True,
        }],
    )
    ws = _sheet(tmp_path, "S")
    assert ws.column_dimensions["A"].width == 8.0  # 显式值不被 autofit 覆盖


# ──────────────────────────────────────────────────────────────────────
# Number formats
# ──────────────────────────────────────────────────────────────────────


def test_number_formats_by_header_name(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["名称", "金额", "占比", "不存在的列"],
            "rows": [["甲", "1234.5", "0.25", "x"]],
            "number_formats": {
                "金额": "#,##0.00",
                "占比": "0.0%",
                "幽灵列": "0.00",
            },
        }],
    )
    ws = _sheet(tmp_path, "S")
    assert ws["B2"].number_format == "#,##0.00"
    assert ws["C2"].number_format == "0.0%"
    # 未知列名的映射不落到任何列；表头行不受数字格式影响
    assert ws["A2"].number_format == "General"
    assert ws["D2"].number_format == "General"
    assert ws["B1"].number_format == "General"


def test_formula_cells_skip_number_format(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["金额"],
            "rows": [["=SUM(1,2)"], ["10"]],
            "number_formats": {"金额": "#,##0.00"},
        }],
    )
    ws = _sheet(tmp_path, "S")
    # 公式单元格保留公式语义（data_type=f），不套数字格式
    assert ws["A2"].data_type == "f"
    assert ws["A2"].value == "=SUM(1,2)"
    assert ws["A3"].number_format == "#,##0.00"


# ──────────────────────────────────────────────────────────────────────
# Tool path
# ──────────────────────────────────────────────────────────────────────


def test_office_create_tool_excel_formats(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "预算",
                "headers": ["科目", "金额"],
                "rows": [["差旅", "1200"]],
                "header_style": True,
                "freeze_header": True,
                "number_formats": {"金额": "#,##0.00"},
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["预算"]
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws["B2"].number_format == "#,##0.00"

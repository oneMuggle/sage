"""journal generator (structured-fill) 集成测试。

TDD: 先写失败测试 → generator.py 实现。
"""
from pathlib import Path

import pytest
from docx import Document

from backend.office.errors import OfficePathError
from backend.office.journal.errors import (
    JournalContentShapeError,
    JournalGenerationError,
)
from backend.office.journal.generator import generate_structured
from backend.office.journal.models import JournalContent
from backend.office.journal.parser import parse_journal_spec
from backend.office.journal.persistence import list_generations

FIXTURE_DIR = (
    Path(__file__).resolve().parent.parent.parent.parent
    / "fixtures"
    / "journal"
)


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    (tmp_path / "office" / "journal" / "specs").mkdir(parents=True)
    (tmp_path / "office" / "journal" / "cache").mkdir(parents=True)
    (tmp_path / "office" / "journal" / "generated").mkdir(parents=True)
    return tmp_path


# --- basic happy path ---


def test_generate_structured_returns_record(workspace):
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试论文",
        abstract="测试摘要内容。",
        sections={"keywords": "测试；关键词"},
    )
    rec = generate_structured(spec, content, workspace, output_filename="test.docx")
    assert rec.mode == "structured_fill"
    assert rec.spec_id == spec.spec_id
    out = Path(rec.output_path)
    assert out.exists()
    assert out.stat().st_size > 0
    assert out.name == "test.docx"


def test_generate_structured_records_generation_in_db(workspace):
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试", abstract="摘要", sections={"keywords": "关键词"}
    )
    generate_structured(spec, content, workspace, "t.docx")
    rows = list_generations(workspace, spec_id=spec.spec_id)
    assert len(rows) == 1


# --- content validation ---


def test_generate_structured_content_missing_abstract_raises(workspace):
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(title="测试", abstract="", sections={"keywords": "k"})
    with pytest.raises(JournalContentShapeError):
        generate_structured(spec, content, workspace, "t.docx")


def test_generate_structured_content_missing_keywords_raises(workspace):
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试", abstract="摘要内容", sections={"keywords": ""}
    )
    with pytest.raises(JournalContentShapeError):
        generate_structured(spec, content, workspace, "t.docx")


# --- path safety ---


def test_generate_structured_path_traversal_raises(workspace):
    """output_filename with '..' must NOT escape generated/ directory."""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试", abstract="摘要", sections={"keywords": "关键词"}
    )
    with pytest.raises(OfficePathError, match="not within base workspace"):
        generate_structured(
            spec, content, workspace, output_filename="../../../etc/cron.d/evil.docx"
        )


def test_generate_structured_duplicate_filename_raises(workspace):
    """Writing to the same output_filename twice must fail (no silent overwrite)."""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试", abstract="摘要", sections={"keywords": "关键词"}
    )
    generate_structured(spec, content, workspace, "dup.docx")
    with pytest.raises(JournalGenerationError):
        generate_structured(spec, content, workspace, "dup.docx")


# --- docx content verification ---


def test_generate_structured_output_is_valid_docx(workspace):
    """Output file must be a valid .docx that python-docx can re-open."""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="验证标题",
        abstract="验证摘要文字。",
        sections={"keywords": "验证；关键词"},
    )
    rec = generate_structured(spec, content, workspace, "verify.docx")
    # Re-open and verify structure
    doc = Document(rec.output_path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("验证标题" in t for t in texts)
    assert any("验证摘要文字" in t for t in texts)
    assert any("验证；关键词" in t for t in texts)


# --- Round 21: structured references → formatted bibliography ---


def test_generate_structured_structured_references_formatted(workspace):
    """structured_references 非空 → 文末参考文献用 [N] + GB/T 格式化。"""
    from backend.office.models import ReferenceSpec

    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="引用测试",
        abstract="摘要",
        sections={"keywords": "引用"},
        structured_references=[
            ReferenceSpec(
                key="li2020",
                ref_type="journal",
                title="分布式系统综述",
                authors=["李四", "王五"],
                year="2020",
                source="计算机学报",
                volume="43",
                issue="2",
                pages="10-20",
            )
        ],
    )
    rec = generate_structured(spec, content, workspace, output_filename="ref.docx")
    doc = Document(Path(rec.output_path))
    texts = [p.text for p in doc.paragraphs]
    assert "[1] 李四, 王五. 分布式系统综述[J]. 计算机学报, 2020, 43(2): 10-20." in texts


def test_generate_structured_plain_references_unchanged(workspace):
    """references 纯文本（无 structured_references）→ 原样写入（零变化）。"""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content = JournalContent(
        title="测试",
        abstract="摘要",
        sections={"keywords": "k"},
        references=["原始文献一", "原始文献二"],
    )
    rec = generate_structured(spec, content, workspace, output_filename="plain.docx")
    doc = Document(Path(rec.output_path))
    joined = chr(10).join(p.text for p in doc.paragraphs)
    assert "原始文献一" in joined
    assert "原始文献二" in joined
    assert "[1]" not in joined  # 纯文本模式无编号

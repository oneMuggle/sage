"""journal generator (LLM generate_article) 集成测试。

覆盖：
- 单轮通过：LLM 返回合法 content → 1 次 LLM 调用；
- 两轮自纠：round 0 content 缺 keywords → round 1 修正 → 2 次 LLM 调用。
"""
from pathlib import Path

import pytest
from docx import Document

from backend.office.journal.generator import generate_article
from backend.office.journal.parser import parse_journal_spec

FIXTURE_DIR = (
    Path(__file__).resolve().parent.parent.parent.parent
    / "fixtures"
    / "journal"
)


class MockLLMProxy:
    """返回结构化 content JSON 的 mock LLM proxy。"""

    def __init__(self, responses):
        self._responses = list(responses)
        self.calls = []

    async def generate(self, *, system_prompt, user_prompt, output_schema=None, **kwargs):
        self.calls.append({"system": system_prompt, "user": user_prompt})
        return self._responses.pop(0)


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    (tmp_path / "office" / "journal" / "specs").mkdir(parents=True)
    (tmp_path / "office" / "journal" / "cache").mkdir(parents=True)
    (tmp_path / "office" / "journal" / "generated").mkdir(parents=True)
    return tmp_path


@pytest.mark.asynci()o()
async def test_generate_article_one_round_passes(workspace):
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content_round1 = {
        "title": "测试论文",
        "abstract": "测试摘要内容。",
        "sections": {"keywords": "测试；关键词"},
        "references": [],
        "citations": [],
    }
    mock = MockLLMProxy([content_round1])
    rec = await generate_article(
        spec,
        user_request="写一个测试论文",
        llm_proxy=mock,
        workspace=workspace,
        output_filename="llm.docx",
        max_rounds=2,
    )
    assert rec.mode == "llm_generate"
    assert len(mock.calls) == 1
    assert Path(rec.output_path).exists()


@pytest.mark.asynci()o()
async def test_generate_article_two_rounds_with_self_correction(workspace):
    """Round 1 缺 '关键词' → validator 报 warning；Round 2 补全。"""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content_round1 = {
        "title": "论文",
        "abstract": "摘要。",
        "sections": {"keywords": ""},  # 空 → validator 报 warning
        "references": [],
        "citations": [],
    }
    content_round2 = {
        "title": "论文",
        "abstract": "摘要。",
        "sections": {"keywords": "已修正"},
        "references": [],
        "citations": [],
    }
    mock = MockLLMProxy([content_round1, content_round2])
    rec = await generate_article(
        spec,
        user_request="论文",
        llm_proxy=mock,
        workspace=workspace,
        output_filename="llm.docx",
        max_rounds=2,
    )
    assert rec.mode == "llm_generate"
    assert len(mock.calls) == 2  # 跑了两轮


@pytest.mark.asynci()o()
async def test_generate_article_structured_references_formatted(workspace):
    """Round 25：LLM 返回 structured_references → 参考文献 [N] 格式化。"""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content_round1 = {
        "title": "引用论文",
        "abstract": "摘要内容",
        "sections": {"keywords": "引用；测试"},
        "references": [],
        "citations": [],
        "structured_references": [
            {
                "key": "wang2021",
                "ref_type": "journal",
                "title": "大模型对齐研究",
                "authors": ["王五", "赵六"],
                "year": "2021",
                "source": "人工智能学报",
                "volume": "44",
                "issue": "3",
                "pages": "55-66",
            }
        ],
    }
    mock = MockLLMProxy([content_round1])
    rec = await generate_article(
        spec,
        user_request="写一篇带引用的测试论文",
        llm_proxy=mock,
        workspace=workspace,
        output_filename="cited.docx",
        max_rounds=2,
    )
    assert rec.mode == "llm_generate"
    from docx import Document

    doc = Document(Path(rec.output_path))
    joined = chr(10).join(p.text for p in doc.paragraphs)
    assert "[1] 王五, 赵六. 大模型对齐研究[J]. 人工智能学报, 2021, 44(3): 55-66." in joined


@pytest.mark.asynci()o()
async def test_generate_article_sanitizes_structured_references(workspace):
    """Round 30：次品条目剔除（缺 title/ref_type 非法），合法条目保留；
    key 重复自动补唯一后缀。"""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content_round1 = {
        "title": "清洗测试",
        "abstract": "摘要",
        "sections": {"keywords": "清洗"},
        "structured_references": [
            {  # 合法
                "key": "good1",
                "ref_type": "journal",
                "title": "合法条目",
                "authors": ["甲"],
                "year": "2020",
            },
            {  # 次品：缺 title
                "key": "bad1",
                "ref_type": "journal",
            },
            {  # 合法（key 重复 → 自动改 key）
                "key": "good1",
                "ref_type": "book",
                "title": "重复 key 的书",
                "authors": ["乙"],
                "year": "2001",
            },
        ],
    }
    mock = MockLLMProxy([content_round1])
    rec = await generate_article(
        spec,
        user_request="清洗测试",
        llm_proxy=mock,
        workspace=workspace,
        output_filename="clean.docx",
        max_rounds=2,
    )
    from docx import Document

    doc = Document(Path(rec.output_path))
    joined = chr(10).join(p.text for p in doc.paragraphs)
    # 合法条目保留并格式化（GB/T 类型码 + 编号）
    assert "[1] 甲. 合法条目[J]." in joined
    # 次品（缺 title）不出现
    assert "bad1" not in joined


@pytest.mark.asynci()o()
async def test_generate_article_all_bad_refs_falls_back(workspace):
    """全为次品条目 → structured_references 清空，纯文本回退，生成不失败。"""
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    content_round1 = {
        "title": "回退测试",
        "abstract": "摘要",
        "sections": {"keywords": "k"},
        "references": ["手工文献一条"],
        "structured_references": [
            {"key": "bad", "ref_type": "journal"},  # 缺 title
        ],
    }
    mock = MockLLMProxy([content_round1])
    rec = await generate_article(
        spec,
        user_request="回退测试",
        llm_proxy=mock,
        workspace=workspace,
        output_filename="fallback.docx",
        max_rounds=2,
    )
    assert Path(rec.output_path).exists()
    doc = Document(Path(rec.output_path))
    joined = chr(10).join(p.text for p in doc.paragraphs)
    assert "手工文献一条" in joined  # 纯文本 references 通路兜底

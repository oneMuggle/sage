"""Integration tests for Excel print setup (Round 23).

Covers: landscape orientation, fit-to-width scaling, print area, the
no-setup zero-change baseline, and the office_create tool path.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, sheets: list, name: str = "ps.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(workspace_path="", filename=name, sheets=sheets)
    return generate_xlsx(req, output_dir=str(tmp_path))


def test_print_setup_full(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A", "B", "C", "D", "E", "F"],
            "rows": [["1", "2", "3", "4", "5", "6"]],
            "print_setup": {
                "orientation": "landscape",
                "fit_to_width": 1,
                "print_area": "A1:F40",
            },
        }],
    )
    ws = load_workbook(str(tmp_path / "ps.xlsx"))["S"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert str(ws.print_area) == "'S'!$A$1:$F$40" or "A1:F40" in str(ws.print_area)


def test_no_print_setup_zero_change(tmp_path: Path) -> None:
    _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
    ws = load_workbook(str(tmp_path / "ps.xlsx"))["S"]
    assert ws.page_setup.orientation in (None, "default", "portrait")
    assert ws.print_area is None or str(ws.print_area) == ""


def test_partial_print_setup_only_sets_given_fields(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {"orientation": "landscape"},
        }],
    )
    ws = load_workbook(str(tmp_path / "ps.xlsx"))["S"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth is None


def test_office_create_tool_print_setup(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "报表",
                "headers": ["科目", "金额"],
                "rows": [["差旅", "1200"]],
                "print_setup": {"orientation": "landscape", "fit_to_width": 1},
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["报表"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1

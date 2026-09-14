"""Integration tests for Excel print margins (Round 31).

Covers: cm→inch conversion for print margins, partial margins (only
given edges set), zero change when print_setup.margins_cm is absent,
and the office_create tool path.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool

_IN = 2.54  # 1 inch in cm


def _generate(tmp_path: Path, sheets: list, name: str = "pm.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(workspace_path="", filename=name, sheets=sheets)
    return generate_xlsx(req, output_dir=str(tmp_path))


def test_print_margins_converted_to_inches(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {"margins_cm": {"top": 2.54, "left": 1.27}},
        }],
    )
    ws = load_workbook(str(tmp_path / "pm.xlsx"))["S"]
    # 2.54cm = 1 inch, 1.27cm = 0.5 inch
    assert ws.page_margins.top == pytest.approx(1.0)
    assert ws.page_margins.left == pytest.approx(0.5)
    # 未给的边保持 openpyxl 默认（0.75 inch）
    assert ws.page_margins.right == pytest.approx(0.75)


def test_print_margins_partial(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {"margins_cm": {"bottom": 5.08}},
        }],
    )
    ws = load_workbook(str(tmp_path / "pm.xlsx"))["S"]
    assert ws.page_margins.bottom == pytest.approx(2.0)
    assert ws.page_margins.top == pytest.approx(1.0)  # openpyxl 默认


def test_no_print_margins_zero_change(tmp_path: Path) -> None:
    _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
    ws = load_workbook(str(tmp_path / "pm.xlsx"))["S"]
    assert ws.page_margins.top == pytest.approx(1.0)  # openpyxl 默认


def test_office_create_tool_print_margins(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "报表",
                "headers": ["科目"],
                "rows": [["差旅"]],
                "print_setup": {"margins_cm": {"top": 2.0}},
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["报表"]
    assert ws.page_margins.top == pytest.approx(2.0 / 2.54)

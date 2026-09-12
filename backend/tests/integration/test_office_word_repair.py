"""Integration tests for the Word format auto-repair (Round 12).

Strategy: generate a docx violating a known FormatSpec, repair it, and
assert the repaired_rules report plus a passing re-lint on the output.
Covers: style/page repairs (via word_layout applier), numbering and
caption renumbering (XML-level edits on the broken doc), citation/coverage
remaining unrepairable, new-file vs overwrite semantics, the REST
endpoint, the WRITE_LOCAL agent tool, and the tool_names/profiles drift
gates.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import OfficeWordGenerateRequest, WordFormatSpec
from backend.office.word import generate_docx
from backend.office.word_lint import lint_docx
from backend.office.word_repair import repair_docx

_SPEC = {
    "page": {"size": "A4", "margins_cm": {"top": 2.5, "bottom": 2.5}},
    "body": {"font_size_pt": 12, "line_spacing": 1.5},
    "headings": {"h1": {"font_size_pt": 16, "bold": True}},
    "header": {"text": "内部资料"},
    "footer": {"page_number": True},
    "numbering": True,
}


def _generate(tmp_path: Path, name: str, spec: dict, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="测试文档",
        format_spec=spec,
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


def _break_page_fields(path: Path) -> None:
    """把一页合规文档破坏成多类违规：删页码域、清页眉、改错边距与字号。"""
    from docx.shared import Cm

    doc = Document(str(path))
    section = doc.sections[0]
    section.top_margin = Cm(1.0)
    section.header.paragraphs[0].text = ""
    for p in section.footer.paragraphs:
        for fld in p._p.findall(".//" + qn("w:fldSimple")):
            fld.getparent().remove(fld)
    doc.styles["Normal"].font.size = None
    doc.save(str(path))


# ──────────────────────────────────────────────────────────────────────
# Style / page repairs
# ──────────────────────────────────────────────────────────────────────


def test_style_and_page_repairs_to_clean(tmp_path: Path) -> None:
    path = _generate(tmp_path, "doc.docx", _SPEC, paragraphs=[{"text": "正文"}])
    _break_page_fields(path)
    before = lint_docx(path, WordFormatSpec(**_SPEC))
    assert not before.ok

    result = repair_docx(path, WordFormatSpec(**_SPEC))
    # 默认写新文件，原文件保持破损
    assert result.output_path.endswith("doc-repaired.docx")
    assert Path(result.output_path).exists()
    assert lint_docx(path, WordFormatSpec(**_SPEC)).error_count == before.error_count

    assert result.ok, [i.message for i in result.remaining.issues]
    repaired = set(result.repaired_rules)
    assert {"page/margins", "header/text", "footer/page_number", "body/font_size"} <= repaired
    assert result.remaining.issue_count == 0
    # 复检是在修复产物上做的
    assert result.remaining.checked_rules


def test_overwrite_repairs_in_place(tmp_path: Path) -> None:
    path = _generate(tmp_path, "doc.docx", _SPEC, paragraphs=[{"text": "正文"}])
    _break_page_fields(path)
    result = repair_docx(path, WordFormatSpec(**_SPEC), overwrite=True)
    assert result.overwrite is True
    assert Path(result.output_path) == path
    assert not path.with_name("doc-repaired.docx").exists()
    assert lint_docx(path, WordFormatSpec(**_SPEC)).ok


def test_conforming_document_repair_is_noop(tmp_path: Path) -> None:
    path = _generate(tmp_path, "ok.docx", _SPEC, paragraphs=[{"text": "正文"}])
    result = repair_docx(path, WordFormatSpec(**_SPEC))
    assert result.repaired_rules == []
    assert result.ok


# ──────────────────────────────────────────────────────────────────────
# Numbering / caption renumbering
# ──────────────────────────────────────────────────────────────────────


def test_heading_numbering_repair(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "num.docx",
        {"numbering": True},
        paragraphs=[
            {"text": "第一章", "heading": "h1"},
            {"text": "第二章", "heading": "h1"},
            {"text": "小节", "heading": "h2"},
        ],
    )
    # 生成后手工破坏第二个 h1 的编号（生成器会重排生成期文本，
    # 破坏须发生在生成之后——真实用户改标题的场景）
    doc = Document(str(path))
    doc.paragraphs[2].text = "9.9 破坏的标题"
    doc.save(str(path))
    assert not lint_docx(path, WordFormatSpec(**{"numbering": True})).ok

    result = repair_docx(path, WordFormatSpec(**{"numbering": True}))
    assert "numbering/sequence" in result.repaired_rules
    texts = [p.text for p in Document(str(result.output_path)).paragraphs]
    assert "1 第一章" in texts
    assert "2 破坏的标题" in texts  # 旧前缀剥离后重排
    assert "2.1 小节" in texts
    assert result.ok


def test_bibliography_heading_not_renumbered(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "bib.docx",
        {"numbering": True},
        paragraphs=[{"text": "引言", "heading": "h1"}],
        references=[{"key": "k1", "title": "文献一", "authors": ["甲"]}],
        citation_style="gbt7714",
    )
    result = repair_docx(path, WordFormatSpec(**{"numbering": True}))
    texts = [p.text for p in Document(str(result.output_path)).paragraphs]
    assert "参考文献" in texts  # 文献节标题保持无编号
    assert "1 引言" in texts


def test_caption_sequence_repair(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "cap.docx",
        {},
        paragraphs=[
            {"text": "图1　合法"},
            {"text": "图5　跳号"},
            {"text": "表2　错号表题"},
            {"text": "表3　顺延表题"},
        ],
    )
    result = repair_docx(path, WordFormatSpec(**{"page": {"size": "A4"}}))
    assert "caption/sequence" in result.repaired_rules
    texts = [p.text for p in Document(str(result.output_path)).paragraphs]
    assert "图1　合法" in texts
    assert "图2　跳号" in texts
    assert "表1　错号表题" in texts
    assert "表2　顺延表题" in texts
    assert result.ok


# ──────────────────────────────────────────────────────────────────────
# Semantic issues remain unrepairable
# ──────────────────────────────────────────────────────────────────────


def test_citation_coverage_remains_in_remaining(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "cite.docx",
        {},
        paragraphs=[{"text": "讨论 [1] 与 [3]。"}],
    )
    result = repair_docx(path, WordFormatSpec())
    assert result.repaired_rules == []
    coverage = [
        i for i in result.remaining.issues if i.rule_id == "citation/coverage"
    ]
    assert coverage
    assert coverage[0].severity == "warning"


# ──────────────────────────────────────────────────────────────────────
# REST endpoint + agent tool + drift gates
# ──────────────────────────────────────────────────────────────────────


def test_repair_endpoint(tmp_path: Path) -> None:
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from backend.api.office_routes import repair_word_endpoint, router
    from backend.office.errors import OfficePathError
    from backend.office.models import WordRepairRequest

    ws = tmp_path / "ws"
    ws.mkdir()
    doc_path = ws / "doc.docx"
    _generate(tmp_path, "doc.docx", {}, paragraphs=[{"text": "图3　跳号"}])
    doc_path.write_bytes((tmp_path / "doc.docx").read_bytes())

    app = FastAPI()
    app.include_router(router)
    client = TestClient(app, headers={"Authorization": "Bearer test-local-auth-token"})
    resp = client.post(
        "/office/word/repair",
        json={
            "workspace_path": str(ws),
            "file_path": str(doc_path),
            "format_spec": {"page": {"size": "A4"}},
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["ok"] is True
    assert "caption/sequence" in body["repaired_rules"]
    assert body["output_path"].endswith("doc-repaired.docx")

    # 工作区围栏：越界 file_path 拒绝（裸 app 无全局错误信封，直接调函数）
    with pytest.raises(OfficePathError):
        repair_word_endpoint(
            WordRepairRequest(
                workspace_path=str(ws),
                file_path=str(tmp_path / "doc.docx"),
                format_spec=WordFormatSpec(),
            )
        )


def test_office_repair_tool_roundtrip(tmp_path: Path) -> None:
    from backend.tools.context import ToolExecutionContext, reset_tool_context, set_tool_context
    from backend.tools.office_repair_tool import OfficeRepairWordTool

    tool = OfficeRepairWordTool()
    assert tool.schema.name == "office_repair_word"
    path = _generate(
        tmp_path,
        "doc.docx",
        {"numbering": True},
        paragraphs=[{"text": "第一章", "heading": "h1"}],
    )
    doc = Document(str(path))
    doc.paragraphs[1].text = "9 破坏的标题"
    doc.save(str(path))

    # requires_tool_context=True：无上下文 fail-closed
    assert not tool.execute(
        file_path=str(path), format_spec={"numbering": True}
    ).success

    token = set_tool_context(
        ToolExecutionContext(
            session_id="sess-repair",
            stream_id="stream-repair",
            binding_generation=1,
            office_doc_scope=frozenset(),
        )
    )
    try:
        result = tool.execute(
            file_path=str(path),
            format_spec={"numbering": True},
        )
        assert result.success, result.error
        assert "numbering/sequence" in result.content["repaired_rules"]
        assert result.content["ok"] is True
        # 缺参数 fail-fast
        assert not tool.execute(file_path=str(path), format_spec={}).success
    finally:
        reset_tool_context(token)


def test_repair_tool_registered_and_known() -> None:
    """防漂移三件套：注册面 / tool_names / writer profile 三处一致。"""
    from backend.domain.tool_names import ALL_BUILTIN_TOOL_NAMES
    from backend.domain.tool_policy import ToolPolicy
    from backend.tools import ToolRegistry, register_all_tools

    registry = ToolRegistry()
    register_all_tools(registry, policy=ToolPolicy())
    assert "office_repair_word" in set(registry.list_names())
    assert "office_repair_word" in set(ALL_BUILTIN_TOOL_NAMES)

    from backend.agents.profiles import create_default_agents

    writer = next(p for p in create_default_agents() if p.id == "writer")
    assert "office_repair_word" in writer.tools

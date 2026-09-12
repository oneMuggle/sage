"""Integration tests for Excel data validation dropdowns (Round 18).

Covers: dropdown rule written via generate_xlsx and read back through
ws.data_validations (type/formula1/sqref/prompt), over-255-char option
lists skipped without blocking, model-layer rejection of empty options,
zero change when no rules, and the office_create tool path.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, sheets: list, name: str = "dv.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(workspace_path="", filename=name, sheets=sheets)
    return generate_xlsx(req, output_dir=str(tmp_path))


def _dvs(ws):
    return [(dv.type, dv.formula1, str(dv.sqref)) for dv in ws.data_validations.dataValidation]


def test_dropdown_rule_written(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["任务", "状态"],
            "rows": [["a", "未开始"]],
            "data_validations": [{
                "range": "B2:B100",
                "options": ["未开始", "进行中", "已完成"],
                "prompt_title": "选择状态",
                "prompt": "从下拉列表选择",
            }],
        }],
    )
    ws = load_workbook(str(tmp_path / "dv.xlsx"))["S"]
    dv_list = ws.data_validations.dataValidation
    assert len(dv_list) == 1
    dv = dv_list[0]
    assert dv.type == "list"
    assert dv.formula1 == '"未开始,进行中,已完成"'
    assert str(dv.sqref) == "B2:B100"
    assert dv.allow_blank is True
    assert dv.promptTitle == "选择状态"


def test_allow_blank_false(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["状态"],
            "data_validations": [{
                "range": "A2:A50",
                "options": ["是", "否"],
                "allow_blank": False,
            }],
        }],
    )
    ws = load_workbook(str(tmp_path / "dv.xlsx"))["S"]
    assert ws.data_validations.dataValidation[0].allow_blank is False


def test_overlong_options_skipped_without_blocking(tmp_path: Path) -> None:
    long_opts = ["选项" + "x" * 50 for _ in range(6)]  # 总长 > 255
    path = _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "data_validations": [
                {"range": "A2:A50", "options": long_opts},
                {"range": "B2:B50", "options": ["短", "选项"]},
            ],
        }],
    )
    ws = load_workbook(str(path))["S"]
    dv_list = ws.data_validations.dataValidation
    # 超长条跳过，短条正常写入
    assert len(dv_list) == 1
    assert str(dv_list[0].sqref) == "B2:B50"


def test_empty_options_rejected_at_model() -> None:
    with pytest.raises(ValidationError):
        OfficeExcelGenerateRequest(
            workspace_path="",
            filename="x.xlsx",
            sheets=[{
                "name": "S",
                "data_validations": [{"range": "A2:A50", "options": []}],
            }],
        )


def test_no_validations_zero_change(tmp_path: Path) -> None:
    path = _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
    ws = load_workbook(str(path))["S"]
    assert len(ws.data_validations.dataValidation) == 0


def test_office_create_tool_datavalidation(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "台账",
                "headers": ["任务", "状态"],
                "rows": [["a", "未开始"]],
                "data_validations": [{
                    "range": "B2:B100",
                    "options": ["未开始", "进行中", "已完成"],
                }],
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["台账"]
    assert ws.data_validations.dataValidation[0].formula1 == '"未开始,进行中,已完成"'

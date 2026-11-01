"""Integration tests for Word rule extraction routes (M3a)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm, Pt

from backend.api.office_word_audit_routes import (
    WordRuleExtractRequest,
    WordRuleExtractSaveRequest,
    extract_rulepack_endpoint,
    list_rulepacks_endpoint,
    save_extracted_rulepack_endpoint,
)
from backend.office.errors import OfficeParseError, OfficePathError, OfficeSizeLimitError
from backend.office.word_rulepack_models import WordRulePack, WordRuleSpec


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _save_spec(workspace: Path, name: str = "spec.docx") -> Path:
    document = Document()
    for text in [
        "页面使用 A4 纸张，页边距上 3.7 厘米、下 3.5 厘米、左 2.8 厘米、右 2.6 厘米。",
        "正文采用仿宋_GB2312 三号字。",
        "页脚应包含阿拉伯数字页码。",
    ]:
        document.add_paragraph(text)
    path = workspace / name
    document.save(str(path))
    return path


def _save_template(workspace: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    normal = document.styles["Normal"]
    normal.font.size = Pt(16)
    document.add_paragraph("正文")
    path = workspace / "template.docx"
    document.save(str(path))
    return path


def test_extract_save_and_list(workspace: Path) -> None:
    path = _save_spec(workspace)

    result = extract_rulepack_endpoint(
        WordRuleExtractRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            pack_id="journal-spec",
            pack_name="期刊规范",
        )
    )

    assert result.ok is True
    assert result.file_type == "spec"
    assert result.pack.id == "journal-spec"
    assert result.pack.rules
    assert all(rule.source.strip() for rule in result.pack.rules)

    saved = save_extracted_rulepack_endpoint(
        WordRuleExtractSaveRequest(workspace_path=str(workspace), pack=result.pack)
    )
    assert saved.origin == "workspace"
    assert saved.pack.id == "journal-spec"

    listing = list_rulepacks_endpoint(workspace_path=str(workspace))
    assert any(
        pack.id == "journal-spec" and pack.origin == "workspace"
        for pack in listing.packs
    )


def test_extract_template_endpoint(workspace: Path) -> None:
    path = _save_template(workspace)

    result = extract_rulepack_endpoint(
        WordRuleExtractRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            file_type="template",
            pack_id="template-derived",
        )
    )

    assert result.ok is True
    assert result.file_type == "template"
    kinds = {rule.kind for rule in result.pack.rules}
    assert {"page.size", "style.font"} <= kinds


def test_save_rejects_rules_without_source(workspace: Path) -> None:
    pack = WordRulePack(
        id="no-source",
        name="No source",
        rules=[
            WordRuleSpec(
                id="bad",
                title="bad",
                kind="content.placeholder_residue",
                params={},
                source="",
            )
        ],
    )

    with pytest.raises(OfficeParseError):
        save_extracted_rulepack_endpoint(
            WordRuleExtractSaveRequest(workspace_path=str(workspace), pack=pack)
        )


def test_extract_rejects_size_escape_and_bad_template_type(workspace: Path) -> None:
    path = _save_spec(workspace)

    with pytest.raises(OfficeSizeLimitError):
        extract_rulepack_endpoint(
            WordRuleExtractRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                max_size_bytes=1024,
            )
        )
    with pytest.raises(OfficePathError):
        extract_rulepack_endpoint(
            WordRuleExtractRequest(
                workspace_path=str(workspace),
                file_path="../outside.docx",
            )
        )

    text_path = workspace / "spec.txt"
    text_path.write_text("A4", encoding="utf-8")
    with pytest.raises(OfficeParseError):
        extract_rulepack_endpoint(
            WordRuleExtractRequest(
                workspace_path=str(workspace),
                file_path=str(text_path),
                file_type="template",
            )
        )

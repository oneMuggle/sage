# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Integration tests for the batch-2 office preview/export routes (Item 2.5).

Covers the HTTP endpoint functions end-to-end (route functions called
directly, mirroring test_office_doc_actions_routes.py):
- ``POST /office/update/preview`` via ``doc_id`` AND via ``file_path``:
  returns DiffPreviewResult, leaves the managed file untouched;
- preview with bad ops → ``ok=False`` + error (UI shows why it would fail);
- unknown doc_id → OfficeFileNotFoundError (404 via exception handler);
- neither file_path nor doc_id → OfficePathError (400);
- ``POST /office/export-pdf``: passes the validated file through to
  ``backend.office.export_pdf.export_to_pdf`` and returns its result
  verbatim. Patch point: the ``export_to_pdf`` attribute on the
  ``backend.office.export_pdf`` module object (the route imports the
  module lazily and looks the attribute up at call time).
"""

from __future__ import annotations

import sys
import types
from pathlib import Path

import pytest
from docx import Document

from backend.api.office_routes import export_pdf_endpoint, preview_update_endpoint
from backend.data.database import get_database
from backend.office.diff_preview import (
    OfficeExportPdfRequest,
    OfficeUpdatePreviewRequest,
)
from backend.office.errors import OfficeFileNotFoundError, OfficePathError
from backend.office.models import (
    OfficeDocStatus,
    OfficeDocType,
    OfficeDocumentMetadata,
    OfficeDocumentSummary,
)
from backend.office.storage import generate_document_dir

pytestmark = pytest.mark.integration


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _seed_docx(
    workspace: Path, *, paragraphs: tuple = ("Hello world", "Second paragraph")
) -> tuple:
    """Create a managed .docx + its DB row; return (doc_id, managed_path)."""
    doc_id = "doc-1"
    directory = generate_document_dir(workspace, OfficeDocType.WORD, doc_id)
    managed = directory / "report.docx"
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(managed))
    now = 1_700_000_000_000
    save_summary = OfficeDocumentSummary(
        id=doc_id,
        workspace_path=str(workspace.resolve()),
        doc_type=OfficeDocType.WORD,
        original_filename=None,
        generated_filename="report.docx",
        status=OfficeDocStatus.PARSED,
        created_at=now,
        updated_at=now,
        metadata=OfficeDocumentMetadata(file_size_bytes=managed.stat().st_size),
    )
    from backend.office.storage import save_document

    save_document(get_database().get_connection(), save_summary)
    return doc_id, managed


def _export_pdf_module(monkeypatch):
    """Return ``backend.office.export_pdf``, stubbing it if needed.

    The real module is built by a parallel batch-2 agent and may not exist
    when these tests run; the route imports it lazily, so a module injected
    into sys.modules (+ the package attribute) satisfies the call path.
    """
    import backend.office

    try:
        from backend.office import export_pdf
    except ImportError:
        export_pdf = types.ModuleType("backend.office.export_pdf")
        monkeypatch.setitem(sys.modules, "backend.office.export_pdf", export_pdf)
        monkeypatch.setattr(backend.office, "export_pdf", export_pdf, raising=False)
    return export_pdf


# ──────────────────────────────────────────────────────────────────────
# POST /office/update/preview
# ──────────────────────────────────────────────────────────────────────


def test_preview_via_doc_id_returns_changes_and_leaves_file_untouched(workspace: Path):
    doc_id, managed = _seed_docx(workspace)
    result = preview_update_endpoint(
        OfficeUpdatePreviewRequest(
            workspace_path=str(workspace),
            doc_id=doc_id,
            ops=[{"op": "replace_text", "find": "Hello", "replace": "Goodbye"}],
        )
    )
    assert result.ok is True
    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.op == "replace_text"
    assert "Hello" in (change.before or "")
    assert "Goodbye" in (change.after or "")
    # The managed file on disk still says "Hello world".
    assert Document(str(managed)).paragraphs[0].text == "Hello world"


def test_preview_via_file_path(workspace: Path):
    _, managed = _seed_docx(workspace)
    result = preview_update_endpoint(
        OfficeUpdatePreviewRequest(
            workspace_path=str(workspace),
            file_path=str(managed),
            ops=[{"op": "replace_text", "find": "Second", "replace": "Final"}],
        )
    )
    assert result.ok is True
    assert "Final" in (result.changes[0].after or "")
    assert "Second" in Document(str(managed)).paragraphs[1].text


def test_preview_with_bad_ops_returns_ok_false(workspace: Path):
    doc_id, _ = _seed_docx(workspace)
    result = preview_update_endpoint(
        OfficeUpdatePreviewRequest(
            workspace_path=str(workspace),
            doc_id=doc_id,
            ops=[{"op": "replace_text", "find": "no-such-text"}],
        )
    )
    assert result.ok is False
    assert result.error is not None
    assert "text_not_found" in result.error
    assert result.changes == []


def test_preview_unknown_doc_id_raises_404_error(workspace: Path):
    with pytest.raises(OfficeFileNotFoundError):
        preview_update_endpoint(
            OfficeUpdatePreviewRequest(
                workspace_path=str(workspace),
                doc_id="ghost-id",
                ops=[],
            )
        )


def test_preview_requires_file_path_or_doc_id(workspace: Path):
    with pytest.raises(OfficePathError):
        preview_update_endpoint(
            OfficeUpdatePreviewRequest(workspace_path=str(workspace), ops=[])
        )


# ──────────────────────────────────────────────────────────────────────
# POST /office/export-pdf
# ──────────────────────────────────────────────────────────────────────


def test_export_pdf_passes_validated_path_through(monkeypatch, workspace: Path):
    export_pdf = _export_pdf_module(monkeypatch)

    calls = {}
    sentinel = object()

    def fake_export_to_pdf(source, ws):
        calls["source"] = source
        calls["workspace"] = ws
        return sentinel

    monkeypatch.setattr(export_pdf, "export_to_pdf", fake_export_to_pdf, raising=False)

    _, managed = _seed_docx(workspace)
    result = export_pdf_endpoint(
        OfficeExportPdfRequest(workspace_path=str(workspace), file_path=str(managed))
    )

    # result passed through verbatim; args are the workspace-validated
    # file path and the resolved workspace.
    assert result is sentinel
    assert calls["source"] == managed
    assert calls["workspace"] == workspace.resolve()


def test_export_pdf_rejects_paths_outside_workspace(workspace: Path):
    with pytest.raises(OfficePathError):
        export_pdf_endpoint(
            OfficeExportPdfRequest(
                workspace_path=str(workspace), file_path="C:/elsewhere/secret.docx"
            )
        )

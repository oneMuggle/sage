"""Integration tests for Excel print header/footer (Round 32).

Covers: print_header/print_footer written via generate_xlsx and read
back through ws.oddHeader/oddFooter center text, partial sets (only
header or only footer), zero change when absent, and the office_create
tool path.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, sheets: list, name: str = "hf.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(workspace_path="", filename=name, sheets=sheets)
    return generate_xlsx(req, output_dir=str(tmp_path))


def test_print_header_and_footer_written(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {
                "print_header": "公司机密",
                "print_footer": "第 &P 页",
            },
        }],
    )
    ws = load_workbook(str(tmp_path / "hf.xlsx"))["S"]
    assert ws.oddHeader.center.text == "公司机密"
    assert "第" in ws.oddFooter.center.text
    assert "&P" in ws.oddFooter.center.text


def test_print_header_only(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {"print_header": "仅页眉"},
        }],
    )
    ws = load_workbook(str(tmp_path / "hf.xlsx"))["S"]
    assert ws.oddHeader.center.text == "仅页眉"
    assert ws.oddFooter.center.text is None or ws.oddFooter.center.text == ""


def test_print_footer_only(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "print_setup": {"print_footer": "第 &P 页"},
        }],
    )
    ws = load_workbook(str(tmp_path / "hf.xlsx"))["S"]
    assert ws.oddHeader.center.text is None or ws.oddHeader.center.text == ""
    assert "&P" in ws.oddFooter.center.text


def test_no_print_hf_zero_change(tmp_path: Path) -> None:
    path = _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
    ws = load_workbook(str(path))["S"]
    assert ws.oddHeader.center.text is None or ws.oddHeader.center.text == ""
    assert ws.oddFooter.center.text is None or ws.oddFooter.center.text == ""


def test_office_create_tool_print_hf(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "报表",
                "headers": ["科目"],
                "rows": [["差旅"]],
                "print_setup": {
                    "print_header": "报表页眉",
                    "print_footer": "第 &P 页",
                },
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["报表"]
    assert ws.oddHeader.center.text == "报表页眉"
    assert "&P" in ws.oddFooter.center.text

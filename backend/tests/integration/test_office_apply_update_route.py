"""Integration tests for POST /office/doc/{doc_id}/update (round-2 Item R1).

Covers the new apply endpoint through the route function (same style as
test_office_doc_actions_routes.py — direct endpoint fn calls over the shared
temp DB from the autouse ``setup_test_db`` fixture):
- apply replace_text → file content changed, DB row status=edited with
  refreshed size/updated_at, archived state untouched, self_check present
  with ``paragraph_count``.
- apply a rejected op → OfficeOpRejectedError (an OfficeEditError, 422-mapped
  via the existing office_error_to_http_status branches), file + DB row
  unchanged.
- the pre-edit snapshot lands in ``<managed_dir>/.snapshots/`` BEFORE the
  apply (same bytes as the pre-edit file) and shows up in the snapshots
  endpoint.
- unknown doc id → OfficeFileNotFoundError (404-mapped).
- word add_comment op passes through the route unchanged (arbitrary op
  passthrough) and round-trips through the comments reader.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from backend.api.office_routes import (
    list_snapshots_endpoint,
    update_document_endpoint,
)
from backend.data.database import get_database
from backend.office.apply_update import (
    OfficeDocUpdateRequest,
    OfficeOpRejectedError,
)
from backend.office.errors import (
    OfficeEditError,
    OfficeFileNotFoundError,
    office_error_to_http_status,
)
from backend.office.models import (
    OfficeDocStatus,
    OfficeDocType,
    OfficeDocumentMetadata,
    OfficeDocumentSummary,
)
from backend.office.storage import (
    document_path,
    generate_document_dir,
    get_document,
    save_document,
)

pytestmark = pytest.mark.integration


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _seed_docx(workspace: Path) -> str:
    """Create a managed .docx with real paragraphs + its DB row; return id."""
    from docx import Document

    doc_id = "doc-1"
    directory = generate_document_dir(workspace, OfficeDocType.WORD, doc_id)
    managed = directory / "report.docx"
    document = Document()
    document.add_paragraph("Hello world")
    document.add_paragraph("Second paragraph")
    document.save(str(managed))
    now = 1_700_000_000_000
    save_document(
        get_database().get_connection(),
        OfficeDocumentSummary(
            id=doc_id,
            workspace_path=str(workspace.resolve()),
            doc_type=OfficeDocType.WORD,
            original_filename=None,
            generated_filename="report.docx",
            status=OfficeDocStatus.PARSED,
            created_at=now,
            updated_at=now,
            metadata=OfficeDocumentMetadata(file_size_bytes=managed.stat().st_size),
        ),
    )
    return doc_id


def _row(conn, doc_id: str):
    return conn.execute(
        """
        SELECT status, archived_at, updated_at, metadata
        FROM office_documents WHERE id = ?
        """,
        (doc_id,),
    ).fetchone()


# ──────────────────────────────────────────────────────────────────────
# happy path: replace_text
# ──────────────────────────────────────────────────────────────────────


def test_apply_replace_text_updates_file_db_and_self_check(workspace: Path):
    doc_id = _seed_docx(workspace)
    conn = get_database().get_connection()
    before = get_document(conn, doc_id)
    assert before is not None
    managed = document_path(before)

    result = update_document_endpoint(
        doc_id,
        OfficeDocUpdateRequest(
            ops=[{"op": "replace_text", "find": "Hello", "replace": "Bonjour"}]
        ),
    )

    assert result.ok is True

    # 文件内容真的改了
    from docx import Document

    parsed = Document(str(managed))
    texts = [p.text for p in parsed.paragraphs]
    assert "Bonjour world" in texts
    assert not any("Hello" in t for t in texts)

    # DB 行：status=edited、archived 未动、updated_at/size 刷新
    row = _row(conn, doc_id)
    assert row["status"] == "edited"
    assert row["archived_at"] is None
    assert row["updated_at"] > before.updated_at
    assert json.loads(row["metadata"])["file_size_bytes"] == managed.stat().st_size

    # 响应里的 summary 与持久化行一致
    assert result.summary.id == doc_id
    assert result.summary.status is OfficeDocStatus.EDITED
    assert result.summary.archived_at is None
    assert result.summary.metadata.file_size_bytes == managed.stat().st_size

    # self_check 回读：word → paragraph_count
    assert result.self_check["ok"] is True
    assert result.self_check["summary"]["paragraph_count"] >= 2

    # per-op 结果随响应返回
    assert result.results[0]["ok"] is True
    assert result.results[0]["op"] == "replace_text"


# ──────────────────────────────────────────────────────────────────────
# rejected op: 422-shaped error, file + row untouched
# ──────────────────────────────────────────────────────────────────────


def test_apply_rejected_op_raises_422_error_and_leaves_file_untouched(
    workspace: Path,
):
    doc_id = _seed_docx(workspace)
    conn = get_database().get_connection()
    summary = get_document(conn, doc_id)
    assert summary is not None
    managed = document_path(summary)
    original_bytes = managed.read_bytes()

    with pytest.raises(OfficeOpRejectedError) as excinfo:
        update_document_endpoint(
            doc_id,
            OfficeDocUpdateRequest(
                ops=[{"op": "replace_text", "find": "does-not-exist", "replace": "x"}]
            ),
        )

    # 是编辑层错误，且经既有映射折算为 422（OfficeContentShapeError 分支）
    assert isinstance(excinfo.value, OfficeEditError)
    assert office_error_to_http_status(excinfo.value) == 422
    # per-op 失败信息在 message 里
    assert "replace_text" in excinfo.value.message
    assert "text_not_found" in excinfo.value.message

    # 文件原封不动
    assert managed.read_bytes() == original_bytes
    # DB 行未变（仍是 parsed，updated_at 未动）
    row = _row(conn, doc_id)
    assert row["status"] == "parsed"
    assert row["updated_at"] == summary.updated_at


# ──────────────────────────────────────────────────────────────────────
# pre-edit snapshot before apply
# ──────────────────────────────────────────────────────────────────────


def test_apply_creates_pre_edit_snapshot_before_apply(workspace: Path):
    doc_id = _seed_docx(workspace)
    conn = get_database().get_connection()
    summary = get_document(conn, doc_id)
    assert summary is not None
    managed = document_path(summary)
    original_bytes = managed.read_bytes()
    snapshots_dir = managed.parent / ".snapshots"
    assert not snapshots_dir.exists()

    result = update_document_endpoint(
        doc_id,
        OfficeDocUpdateRequest(
            ops=[{"op": "replace_text", "find": "Hello", "replace": "Bonjour"}]
        ),
    )
    assert result.ok is True

    # 快照目录在 apply 前创建，留存的是 pre-edit 字节
    assert snapshots_dir.is_dir()
    snapshots = [p for p in snapshots_dir.iterdir() if p.is_file()]
    assert len(snapshots) == 1
    assert snapshots[0].read_bytes() == original_bytes
    assert snapshots[0].name.endswith("report.docx")

    # snapshots endpoint 也能看到它
    listed = list_snapshots_endpoint(doc_id)
    assert listed.total == 1


# ──────────────────────────────────────────────────────────────────────
# unknown doc id
# ──────────────────────────────────────────────────────────────────────


def test_apply_unknown_doc_raises_404_error(workspace: Path):
    with pytest.raises(OfficeFileNotFoundError):
        update_document_endpoint("ghost-id", OfficeDocUpdateRequest(ops=[]))


# ──────────────────────────────────────────────────────────────────────
# arbitrary op passthrough: word add_comment
# ──────────────────────────────────────────────────────────────────────


def test_apply_add_comment_op_passthrough(workspace: Path):
    doc_id = _seed_docx(workspace)
    conn = get_database().get_connection()

    result = update_document_endpoint(
        doc_id,
        OfficeDocUpdateRequest(
            ops=[
                {
                    "op": "add_comment",
                    "find": "Hello world",
                    "comment": "please review",
                    "author": "tester",
                }
            ]
        ),
    )

    assert result.ok is True
    assert result.results[0]["ok"] is True
    assert result.results[0]["comment_id"] == "0"
    assert result.summary.status is OfficeDocStatus.EDITED
    assert result.self_check["ok"] is True

    # 批注经 comments reader 回读可验证（证明 op 真的写进了文件）
    from backend.office.word import read_docx_comments

    summary = get_document(conn, doc_id)
    assert summary is not None
    comments = read_docx_comments(document_path(summary)).comments
    assert len(comments) == 1
    assert comments[0].text == "please review"
    assert comments[0].author == "tester"

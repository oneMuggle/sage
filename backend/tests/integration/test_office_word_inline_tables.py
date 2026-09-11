"""Integration tests for Word inline images / captions / table layout (Round 8).

Covers the P1 content-element batch:

1. inline image placement — the picture lands in the body element right
   after the anchor paragraph; ``after_paragraph`` out of range falls back
   to trailing append (legacy behavior byte-compatible for caption-less
   trailing images);
2. captions — figures and tables number independently ("图1"/"表1"),
   centered 9pt, and caption-less elements consume no numbers;
3. three-line table borders (tblBorders sz=12 top/bottom, insideV none,
   header bottom tcBorders sz=6), header row repeat (w:tblHeader),
   fixed column widths, cell merges (incl. out-of-range rejection);
4. multi-level heading numbering "1 / 1.1 / 1.1.1" with lower-level reset,
   and zero change when ``numbering`` is false;
5. managed-path passthrough regression: ``_coerce_word_request`` no longer
   drops ``images``;
6. backward compatibility — payloads without the new fields behave as in
   Round 7.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import ValidationError

from backend.office.errors import OfficeGenerateError
from backend.office.models import (
    OfficeWordGenerateRequest,
    WordFormatSpec,
    WordImageSpec,
    WordTableSpec,
)
from backend.office.tool_service import _coerce_word_request
from backend.office.word import generate_docx
from backend.office.word_layout import heading_number_prefix

# 1x1 red PNG（charts.resolve_image_payload 接受 data URI，与生产路径一致）
_TINY_PNG = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _generate(tmp_path: Path, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="r8.docx",
        title="测试文档",
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


# ──────────────────────────────────────────────────────────────────────
# Inline images + figure captions
# ──────────────────────────────────────────────────────────────────────


def test_inline_image_lands_after_anchor_paragraph(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "第一段"}, {"text": "第二段"}],
        images=[{"source": _TINY_PNG, "after_paragraph": 0, "caption": "架构图"}],
    )
    doc = Document(str(path))
    # body 布局: title / p1 / [图片段] / [题注段] / p2 / sectPr
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "测试文档"
    assert texts[1] == "第一段"
    # 图片段：含 w:drawing
    pic_paragraph = doc.paragraphs[2]
    assert pic_paragraph._p.findall(".//" + qn("w:drawing"))
    # 题注段紧跟图片段
    assert texts[3].startswith("图1")
    assert "架构图" in texts[3]
    assert texts[4] == "第二段"


def test_figure_and_table_numbers_are_independent(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文"}],
        tables=[{"headers": ["A"], "rows": [["1"]], "caption": "汇总表"}],
        images=[{"source": _TINY_PNG, "caption": "示意图"}],
    )
    doc = Document(str(path))
    caption_texts = [p.text for p in doc.paragraphs if p.text.startswith(("图", "表"))]
    assert caption_texts == ["表1　汇总表", "图1　示意图"]


def test_caption_style_centered_9pt(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文"}],
        images=[{"source": _TINY_PNG, "caption": "Logo"}],
    )
    doc = Document(str(path))
    caption = next(p for p in doc.paragraphs if p.text.startswith("图1"))
    run = caption.runs[0]
    assert run.font.size == Pt(9)


def test_captionless_image_consumes_no_number(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文"}],
        images=[
            {"source": _TINY_PNG},
            {"source": _TINY_PNG, "caption": "有题注的图"},
        ],
    )
    doc = Document(str(path))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("图")]
    assert captions == ["图1　有题注的图"]
    assert len(doc.inline_shapes) == 2


def test_after_paragraph_out_of_range_clamps_to_trailing(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "唯一一段"}],
        images=[{"source": _TINY_PNG, "after_paragraph": 99, "caption": "尾部图"}],
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts[-1] == "图1　尾部图"  # 钳到文末：题注是最后一个段落


# ──────────────────────────────────────────────────────────────────────
# Table layout
# ──────────────────────────────────────────────────────────────────────


def test_three_line_table_borders(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        tables=[{"headers": ["列A", "列B"], "rows": [["1", "2"]], "style": "three_line"}],
    )
    doc = Document(str(path))
    table = doc.tables[0]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    top = borders.find(qn("w:top"))
    assert top.get(qn("w:val")) == "single"
    assert top.get(qn("w:sz")) == "12"  # 1.5pt = 12 × 1/8pt
    inside_v = borders.find(qn("w:insideV"))
    assert inside_v.get(qn("w:val")) == "none"
    # 表头单元格底线 0.75pt = sz 6
    header_cell = table.rows[0].cells[0]
    tc_borders = header_cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    assert tc_borders is not None
    assert tc_borders.find(qn("w:bottom")).get(qn("w:sz")) == "6"


def test_grid_style_sets_table_grid(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        tables=[{"headers": ["A"], "rows": [["1"]], "style": "grid"}],
    )
    table = Document(str(path)).tables[0]
    assert table.style.name == "Table Grid"


def test_table_caption_above_and_numbered(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        tables=[
            {"headers": ["A"], "rows": []},
            {"headers": ["B"], "rows": [], "caption": "对比表"},
        ],
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "表1　对比表" in texts  # 无题注的第一张表不占号
    # 题注在表前：题注段索引早于第二张表的 XML 位置
    caption_index = texts.index("表1　对比表")
    assert doc.tables[1]._tbl in list(doc.element.body)[caption_index + 1 :]


def test_header_repeat_and_column_widths(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        tables=[
            {
                "headers": ["名称", "数量"],
                "rows": [["x", "1"]],
                "header_repeat": True,
                "column_widths_cm": [4.0, 2.5],
            }
        ],
    )
    doc = Document(str(path))
    table = doc.tables[0]
    tr_pr = table.rows[0]._tr.trPr
    assert tr_pr is not None
    assert tr_pr.find(qn("w:tblHeader")) is not None
    assert table.rows[0].cells[0].width.cm == pytest.approx(4.0, abs=0.01)
    assert table.rows[0].cells[1].width.cm == pytest.approx(2.5, abs=0.01)


def test_cell_merge(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        tables=[
            {
                "headers": ["H0", "H1", "H2"],
                "rows": [["a", "", "c"], ["d", "e", "f"]],
                "merges": [{"min_row": 1, "max_row": 2, "min_col": 0, "max_col": 0}],
            }
        ],
    )
    table = Document(str(path)).tables[0]
    # 合并后 (1,0) 与 (2,0) 是同一个 tc
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc
    # 非合并区域不受影响
    assert table.cell(2, 2).text == "f"


def test_column_widths_length_mismatch_rejected(tmp_path: Path) -> None:
    with pytest.raises(OfficeGenerateError, match="不一致"):
        _generate(
            tmp_path,
            tables=[
                {"headers": ["A", "B"], "rows": [], "column_widths_cm": [3.0]}
            ],
        )


def test_merge_out_of_range_rejected(tmp_path: Path) -> None:
    with pytest.raises(OfficeGenerateError, match="out of range"):
        _generate(
            tmp_path,
            tables=[
                {
                    "headers": ["A"],
                    "rows": [["x"]],
                    "merges": [{"min_row": 0, "max_row": 9, "min_col": 0, "max_col": 0}],
                }
            ],
        )


# ──────────────────────────────────────────────────────────────────────
# Heading numbering
# ──────────────────────────────────────────────────────────────────────


def test_heading_numbering_prefixes_with_reset(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[
            {"text": "引言", "heading": "h1"},
            {"text": "背景", "heading": "h2"},
            {"text": "现状", "heading": "h3"},
            {"text": "问题", "heading": "h3"},
            {"text": "方案", "heading": "h2"},
            {"text": "总结", "heading": "h1"},
        ],
        format_spec={"numbering": True},
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts[1] == "1 引言"
    assert texts[2] == "1.1 背景"
    assert texts[3] == "1.1.1 现状"
    assert texts[4] == "1.1.2 问题"
    assert texts[5] == "1.2 方案"  # h2 变化 → h3 重置
    assert texts[6] == "2 总结"


def test_numbering_false_keeps_text_verbatim(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "引言", "heading": "h1"}],
    )
    assert Document(str(path)).paragraphs[1].text == "引言"


def test_heading_number_prefix_unit() -> None:
    counters = [0, 0, 0]
    assert heading_number_prefix(counters, 1) == "1"
    assert heading_number_prefix(counters, 2) == "1.1"
    assert heading_number_prefix(counters, 3) == "1.1.1"
    assert heading_number_prefix(counters, 3) == "1.1.2"
    assert heading_number_prefix(counters, 2) == "1.2"
    assert heading_number_prefix(counters, 3) == "1.2.1"


# ──────────────────────────────────────────────────────────────────────
# Managed-path passthrough regression + validation + backward compat
# ──────────────────────────────────────────────────────────────────────


def test_coerce_word_request_passes_images_through() -> None:
    req = _coerce_word_request(
        "docid",
        "report.docx",
        {
            "title": "T",
            "paragraphs": [{"text": "正文"}],
            "images": [{"source": _TINY_PNG, "caption": "图", "after_paragraph": 0}],
        },
        "",
    )
    assert req.images is not None
    assert len(req.images) == 1
    assert req.images[0].caption == "图"


def test_word_image_spec_validation() -> None:
    with pytest.raises(ValidationError):
        WordImageSpec(source=_TINY_PNG, after_paragraph=-1)
    with pytest.raises(ValidationError):
        WordImageSpec(source=_TINY_PNG, caption="x" * 201)


def test_table_spec_validation() -> None:
    with pytest.raises(ValidationError):
        WordTableSpec(headers=["A"], style="dashed")  # type: ignore[arg-type]
    with pytest.raises(ValidationError):
        WordTableSpec(headers=["A"], merges=[{"min_row": 0}])


def test_legacy_payload_backward_compatible(tmp_path: Path) -> None:
    """旧 payload（无任何新字段）生成结果与 Round 7 行为一致。"""
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文", "heading": "h1"}],
        tables=[{"headers": ["A"], "rows": [["1"]]}],
        images=[{"source": _TINY_PNG}],
        format_spec={"page": {"margins_cm": {"top": 2.5}}},
    )
    doc = Document(str(path))
    assert doc.paragraphs[1].text == "正文"
    assert doc.paragraphs[1].style.name == "Heading 1"
    assert doc.tables[0].style.name != "Table Grid"  # 未指定 style 不触碰
    # 未指定 column_widths_cm 不改写 tcW（python-docx 默认 6 英寸）
    assert doc.tables[0].rows[0].cells[0].width.inches == pytest.approx(6.0)
    assert doc.sections[0].top_margin.cm == pytest.approx(2.5, abs=0.01)
    assert len(doc.inline_shapes) == 1


def test_format_spec_numbering_model_default() -> None:
    assert WordFormatSpec().numbering is False

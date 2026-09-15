"""Integration tests for the Word FormatSpec layout engine (Round 7).

Covers "版式即配置" wiring end to end:

1. backward compatibility — omitting ``format_spec`` leaves the generated
   document identical to the historical behavior (default template margins);
2. page setup — margins / orientation / paper size land on section 0;
3. body + heading styles — Normal / Title / Heading 1-3 style definitions
   are patched (font size, line spacing, first-line indent, spacing, color);
4. header text + footer PAGE field (``w:fldSimple``) presence;
5. pydantic validation of the spec models (extra="forbid", numeric bounds);
6. agent-tool path — ``OfficeCreateTool.execute`` accepts ``format_spec``
   inside ``content`` and the generated file carries the layout.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from pydantic import ValidationError

from backend.office.models import OfficeWordGenerateRequest
from backend.office.word import generate_docx
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="spec.docx",
        title="测试文档",
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


# ──────────────────────────────────────────────────────────────────────
# Backward compatibility
# ──────────────────────────────────────────────────────────────────────


def test_no_format_spec_keeps_default_layout(tmp_path: Path) -> None:
    path = _generate(tmp_path, paragraphs=[{"text": "正文"}])
    doc = Document(str(path))
    # python-docx default template: 1-inch margins, portrait letter-ish page.
    section = doc.sections[0]
    assert section.top_margin.inches == pytest.approx(1.0)
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert section.header.is_linked_to_previous
    assert section.footer.is_linked_to_previous


def test_format_spec_none_explicitly_is_backward_compatible(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文"}],
        format_spec=None,
    )
    assert path.exists()
    assert Document(str(path)).sections[0].top_margin.inches == pytest.approx(1.0)


# ──────────────────────────────────────────────────────────────────────
# Page setup
# ──────────────────────────────────────────────────────────────────────


def test_page_margins_applied(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        format_spec={
            "page": {
                "margins_cm": {"top": 2.5, "bottom": 2.5, "left": 3.0, "right": 3.0}
            }
        },
    )
    section = Document(str(path)).sections[0]
    assert section.top_margin.cm == pytest.approx(2.5, abs=0.01)
    assert section.bottom_margin.cm == pytest.approx(2.5, abs=0.01)
    assert section.left_margin.cm == pytest.approx(3.0, abs=0.01)
    assert section.right_margin.cm == pytest.approx(3.0, abs=0.01)


def test_a4_landscape_applied(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        format_spec={"page": {"size": "A4", "orientation": "landscape"}},
    )
    section = Document(str(path)).sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.mm == pytest.approx(297, abs=0.05)
    assert section.page_height.mm == pytest.approx(210, abs=0.05)
    assert section.page_width > section.page_height


# ──────────────────────────────────────────────────────────────────────
# Body / heading / title styles
# ──────────────────────────────────────────────────────────────────────


def test_body_style_applied(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "正文段落"}],
        format_spec={
            "body": {
                "font_size_pt": 12,
                "line_spacing": 1.5,
                "first_line_indent_cm": 0.74,
                "space_after_pt": 6,
                "align": "justify",
            }
        },
    )
    normal = Document(str(path)).styles["Normal"]
    assert normal.font.size.pt == pytest.approx(12)
    assert normal.paragraph_format.line_spacing == pytest.approx(1.5)
    assert normal.paragraph_format.first_line_indent.cm == pytest.approx(0.74, abs=0.01)
    assert normal.paragraph_format.space_after.pt == pytest.approx(6)


def test_heading_and_title_styles_applied(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "第一章", "heading": "h1"}],
        format_spec={
            "title": {"font_size_pt": 22, "bold": True, "align": "center"},
            "headings": {
                "h1": {"font_size_pt": 16, "bold": True, "color": "#2F5496"},
                "h2": {"font_size_pt": 14},
            },
        },
    )
    doc = Document(str(path))
    title = doc.styles["Title"]
    assert title.font.size.pt == pytest.approx(22)
    assert title.font.bold is True

    h1 = doc.styles["Heading 1"]
    assert h1.font.size.pt == pytest.approx(16)
    assert h1.font.bold is True
    assert str(h1.font.color.rgb) == "2F5496"
    # theme 属性已清除，显式颜色不会被 theme 覆盖
    color_el = h1.element.get_or_add_rPr().find(qn("w:color"))
    assert color_el is not None
    assert qn("w:themeColor") not in color_el.attrib

    assert doc.styles["Heading 2"].font.size.pt == pytest.approx(14)
    # 正文里 h1 段落实际使用补丁后的样式
    assert doc.paragraphs[1].style.name == "Heading 1"


# ──────────────────────────────────────────────────────────────────────
# Header / footer
# ──────────────────────────────────────────────────────────────────────


def test_header_text_applied(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        format_spec={"header": {"text": "内部资料 · 请勿外传", "align": "center"}},
    )
    header = Document(str(path)).sections[0].header
    assert not header.is_linked_to_previous
    assert header.paragraphs[0].text == "内部资料 · 请勿外传"


def test_footer_page_number_field(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        format_spec={"footer": {"page_number": True, "align": "center"}},
    )
    footer = Document(str(path)).sections[0].footer
    assert not footer.is_linked_to_previous
    flds = footer.paragraphs[0]._p.findall(qn("w:fldSimple"))
    assert len(flds) == 1
    assert flds[0].get(qn("w:instr")) == "PAGE"


# ──────────────────────────────────────────────────────────────────────
# Validation
# ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "payload",
    [
        {"unknown_key": 1},
        {"body": {"font_size_pt": 0}},
        {"body": {"font_size_pt": 999}},
        {"page": {"margins_cm": {"top": -1}}},
        {"headings": {"h6": {"font_size_pt": 12}}},  # Round 20 起 h4 已合法
        {"title": {"color": "red"}},
    ],
)
def test_format_spec_validation_rejects(payload: dict) -> None:
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="",
            filename="spec.docx",
            title="T",
            format_spec=payload,
        )


def test_format_spec_dict_coerced_to_model() -> None:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="spec.docx",
        title="T",
        format_spec={"page": {"margins_cm": {"top": 2.0}}},
    )
    assert req.format_spec is not None
    assert req.format_spec.page.margins_cm.top == pytest.approx(2.0)


# ──────────────────────────────────────────────────────────────────────
# Agent tool path
# ──────────────────────────────────────────────────────────────────────


def test_office_create_tool_accepts_format_spec(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="word",
        output_dir=str(tmp_path),
        filename="layout.docx",
        content={
            "title": "项目报告",
            "paragraphs": [{"text": "第一章", "heading": "h1"}, {"text": "正文"}],
            "format_spec": {
                "page": {"size": "A4", "margins_cm": {"top": 2.54}},
                "body": {"font_size_pt": 12, "line_spacing": 1.5},
                "headings": {"h1": {"font_size_pt": 15, "bold": True}},
                "footer": {"page_number": True},
            },
        },
    )
    assert result.success, getattr(result, "error", None)

    doc = Document(str(tmp_path / "layout.docx"))
    assert doc.sections[0].top_margin.cm == pytest.approx(2.54, abs=0.01)
    assert doc.styles["Heading 1"].font.size.pt == pytest.approx(15)
    flds = doc.sections[0].footer.paragraphs[0]._p.findall(qn("w:fldSimple"))
    assert len(flds) == 1

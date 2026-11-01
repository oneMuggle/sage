"""Integration tests for the Word audit/fix route (M2b)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from backend.api.office_word_audit_routes import (
    WordAuditFixRequest,
    fix_word_endpoint,
)
from backend.office.errors import (
    OfficeFileNotFoundError,
    OfficeParseError,
    OfficePathError,
    OfficeSizeLimitError,
)


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _save_doc(workspace: Path, name: str = "doc.docx") -> Path:
    document = Document()
    document.add_paragraph("正文")
    path = workspace / name
    document.save(str(path))
    return path


def test_fix_endpoint_creates_compliant_copy(workspace: Path) -> None:
    path = _save_doc(workspace)

    result = fix_word_endpoint(
        WordAuditFixRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            rulepack_id="gb-t-9704-general",
            output_filename="fixed.docx",
        )
    )

    assert result.ok is True
    assert result.fixed_count > 0
    assert result.overwritten is False
    assert result.remaining.ok is True
    assert (workspace / "fixed.docx").is_file()
    original = Document(str(path))
    assert round(original.sections[0].page_width.mm) == 216


def test_fix_endpoint_overwrite_is_idempotent(workspace: Path) -> None:
    path = _save_doc(workspace)

    first = fix_word_endpoint(
        WordAuditFixRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            rulepack_id="gb-t-9704-general",
            overwrite=True,
        )
    )
    assert first.overwritten is True
    assert first.fixed_count > 0

    second = fix_word_endpoint(
        WordAuditFixRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            rulepack_id="gb-t-9704-general",
            overwrite=True,
        )
    )
    assert second.fixed_count == 0
    assert second.overwritten is False
    assert second.remaining.ok is True


def test_fix_endpoint_rejects_unsafe_requests(workspace: Path) -> None:
    path = _save_doc(workspace)

    with pytest.raises(OfficeFileNotFoundError):
        fix_word_endpoint(
            WordAuditFixRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                rulepack_id="missing-pack",
            )
        )
    with pytest.raises(OfficeSizeLimitError):
        fix_word_endpoint(
            WordAuditFixRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                max_size_bytes=1024,
            )
        )
    with pytest.raises(OfficePathError):
        fix_word_endpoint(
            WordAuditFixRequest(
                workspace_path=str(workspace),
                file_path="../outside.docx",
            )
        )
    with pytest.raises(OfficeParseError):
        fix_word_endpoint(
            WordAuditFixRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                output_filename="fixed.docx",
                overwrite=True,
            )
        )

    fix_word_endpoint(
        WordAuditFixRequest(
            workspace_path=str(workspace),
            file_path=str(path),
            output_filename="fixed.docx",
        )
    )
    with pytest.raises(OfficeParseError):
        fix_word_endpoint(
            WordAuditFixRequest(
                workspace_path=str(workspace),
                file_path=str(path),
                output_filename="fixed.docx",
            )
        )

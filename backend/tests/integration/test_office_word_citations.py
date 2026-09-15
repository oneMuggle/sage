"""Integration tests for Word citations wiring (Round 9).

Covers the generate_docx citation pipeline end to end (in-text superscript
markers, first-appearance numbering, auto bibliography section, strict
consistency errors), the REST parse-bibtex endpoint, the agent tool, and
backward compatibility for payloads without references/citations.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from pydantic import ValidationError

from backend.office.errors import OfficeGenerateError
from backend.office.models import OfficeWordGenerateRequest
from backend.office.word import generate_docx
from backend.tools.office_bibtex_tool import OfficeBibTexTool


def _generate(tmp_path: Path, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="cites.docx",
        title="测试论文",
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


_REFS = [
    {
        "key": "zhang2023",
        "ref_type": "journal",
        "title": "量子计算综述",
        "authors": ["李四", "王五"],
        "year": "2023",
        "source": "计算机学报",
        "volume": "46",
        "issue": "5",
        "pages": "100-110",
    },
    {
        "key": "li2016",
        "ref_type": "book",
        "title": "Deep Learning",
        "authors": ["Li Y"],
        "year": "2016",
        "address": "Cambridge",
        "publisher": "MIT Press",
        "language": "en",
    },
]


def _cited_request() -> dict:
    return {
        "paragraphs": [
            {"text": "引言", "heading": "h1"},
            {"text": "文献一。", "citations": ["zhang2023", "li2016"]},
            {"text": "文献二。", "citations": ["zhang2023"]},
        ],
        "references": _REFS,
    }


# ──────────────────────────────────────────────────────────────────────
# Generation: in-text markers + numbering + bibliography
# ──────────────────────────────────────────────────────────────────────


def test_in_text_superscript_and_first_appearance_ordering(tmp_path: Path) -> None:
    path = _generate(tmp_path, **_cited_request())
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    # 编号 = 首现顺序：zhang2023 → 1, li2016 → 2；连续编号合并 [1-2]
    assert texts[2] == "文献一。[1-2]"
    assert texts[3] == "文献二。[1]"
    sup_runs = [r.text for r in doc.paragraphs[2].runs if r.font.superscript]
    assert sup_runs == ["[1-2]"]


def test_bibliography_section_generated_sorted_with_hanging_indent(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=_cited_request()["paragraphs"],
        references=_REFS,
        format_spec={"bibliography": {"heading_text": "参考文献", "font_size_pt": 9}},
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "参考文献" in texts
    bib_entries = [t for t in texts if t.startswith("[")]
    # 条目按引用编号排序（zhang2023=1 在前），带 "[N] " 前缀
    assert bib_entries[0].startswith("[1] 李四, 王五. 量子计算综述[J].")
    assert bib_entries[1].startswith("[2] Li Y. Deep Learning[M].")
    entry = next(p for p in doc.paragraphs if p.text.startswith("[1] "))
    assert entry.runs[0].font.size == Pt(9)
    # 悬挂缩进：左缩进 0.74cm、首行 -0.74cm（默认值）
    assert entry.paragraph_format.left_indent.cm == pytest.approx(0.74, abs=0.01)
    assert entry.paragraph_format.first_line_indent.cm == pytest.approx(-0.74, abs=0.01)


def test_bibliography_heading_not_numbered_by_format_spec(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=_cited_request()["paragraphs"],
        references=_REFS,
        format_spec={"numbering": True},
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "1 引言" in texts  # 正式标题有编号
    assert "参考文献" in texts  # 文献节标题不参与多级编号


def test_apa_style_rendering(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        paragraphs=[{"text": "引用。", "citations": ["zhang2023"]}],
        references=_REFS[:1],
        citation_style="apa",
    )
    doc = Document(str(path))
    entry = next(p for p in doc.paragraphs if p.text.startswith("[1] "))
    assert "(2023)." in entry.text
    assert "[J]." not in entry.text  # APA 无 GB/T 类型码


# ──────────────────────────────────────────────────────────────────────
# Strict consistency errors
# ──────────────────────────────────────────────────────────────────────


def test_citation_to_unknown_key_rejected(tmp_path: Path) -> None:
    with pytest.raises(OfficeGenerateError, match="未定义"):
        _generate(
            tmp_path,
            paragraphs=[{"text": "正文", "citations": ["ghost"]}],
            references=_REFS,
        )


def test_uncited_reference_rejected_when_citations_exist(tmp_path: Path) -> None:
    extra = dict(_REFS[1])
    extra["key"] = "unused1"
    with pytest.raises(OfficeGenerateError, match="未被引用"):
        _generate(
            tmp_path,
            paragraphs=[{"text": "正文", "citations": ["zhang2023"]}],
            references=[_REFS[0], extra],
        )


def test_references_without_citations_render_full_bibliography(tmp_path: Path) -> None:
    """不带 citations 时全部条目入表（"只要文献表"生成需求）。"""
    path = _generate(tmp_path, paragraphs=[{"text": "正文"}], references=_REFS)
    doc = Document(str(path))
    entries = [p.text for p in doc.paragraphs if p.text.startswith("[")]
    assert len(entries) == 2


def test_citations_on_heading_rejected(tmp_path: Path) -> None:
    with pytest.raises(OfficeGenerateError, match="标题段落不支持"):
        _generate(
            tmp_path,
            paragraphs=[{"text": "标题", "heading": "h1", "citations": ["x"]}],
            references=_REFS[:1],
        )


# ──────────────────────────────────────────────────────────────────────
# Validation
# ──────────────────────────────────────────────────────────────────────


def test_reference_spec_validation() -> None:
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="", filename="x.docx", title="T",
            references=[{"key": "k"}],  # 缺 title
        )
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="", filename="x.docx", title="T",
            citation_style="chicago",  # type: ignore[arg-type]
        )
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="", filename="x.docx", title="T",
            references=[{"key": "k", "title": "T", "ref_type": "blog"}],  # type: ignore[arg-type]
        )


# ──────────────────────────────────────────────────────────────────────
# REST endpoint + agent tool + backward compat
# ──────────────────────────────────────────────────────────────────────


def test_parse_bibtex_endpoint() -> None:
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from backend.api.office_routes import router

    app = FastAPI()
    app.include_router(router)
    client = TestClient(app, headers={"Authorization": "Bearer test-local-auth-token"})
    resp = client.post(
        "/office/word/parse-bibtex",
        json={"text": "@article{k1, title={T}, author={A and B}, year={2020}}"},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["count"] == 1
    assert body["references"][0]["key"] == "k1"
    assert body["references"][0]["authors"] == ["A", "B"]


def test_office_bibtex_tool_roundtrip() -> None:
    tool = OfficeBibTexTool()
    schema = tool.schema
    assert schema.name == "office_parse_bibtex"
    result = tool.execute(text="@book{b1, title={Book}, publisher={P}, year={2001}}")
    assert result.success
    assert result.content["count"] == 1
    assert result.content["references"][0]["ref_type"] == "book"
    # 空输入 fail-fast
    assert not tool.execute(text="   ").success


def test_managed_path_passthrough_references() -> None:
    from backend.office.tool_service import _coerce_word_request

    req = _coerce_word_request(
        "docid",
        "paper.docx",
        {
            "title": "T",
            "paragraphs": [{"text": "正文", "citations": ["k1"]}],
            "references": _REFS[:1],
            "citation_style": "apa",
        },
        "",
    )
    assert len(req.references) == 1
    assert req.citation_style == "apa"


def test_legacy_payload_unchanged_without_citations(tmp_path: Path) -> None:
    path = _generate(tmp_path, paragraphs=[{"text": "正文"}])
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts[-1] == "正文"  # 没有多出参考文献节
    assert not any(t.startswith("[") for t in texts)

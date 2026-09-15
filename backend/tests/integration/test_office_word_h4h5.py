"""Integration tests for Word h4/h5 heading support (Round 20).

Covers the full pipeline for the extended heading levels: generation
(Heading 4/5 styles), multi-level numbering to 5 levels (skipped levels
omit zero segments), format_spec headings style override for h4/h5,
linter numbering detection at depth 4, and model validation.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from backend.office.models import OfficeWordGenerateRequest, WordFormatSpec
from backend.office.word import generate_docx
from backend.office.word_lint import lint_docx


def _generate(tmp_path: Path, name: str, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="测试文档",
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_h4_h5_generation_and_numbering(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "h45.docx",
        paragraphs=[
            {"text": "章", "heading": "h1"},
            {"text": "节", "heading": "h2"},
            {"text": "小节", "heading": "h3"},
            {"text": "四级", "heading": "h4"},
            {"text": "五级", "heading": "h5"},
        ],
        format_spec={"numbering": True},
    )
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts[1] == "1 章"
    assert texts[2] == "1.1 节"
    assert texts[3] == "1.1.1 小节"
    assert texts[4] == "1.1.1.1 四级"
    assert texts[5] == "1.1.1.1.1 五级"
    assert doc.paragraphs[4].style.name == "Heading 4"
    assert doc.paragraphs[5].style.name == "Heading 5"


def test_skip_level_omits_zero_segments(tmp_path: Path) -> None:
    """h1 后直接 h4：省略中间 0 段 → "1.1"（Word 常见编号习惯）。"""
    path = _generate(
        tmp_path,
        "skip.docx",
        paragraphs=[
            {"text": "一", "heading": "h1"},
            {"text": "甲", "heading": "h4"},
            {"text": "二", "heading": "h1"},
            {"text": "乙", "heading": "h4"},
        ],
        format_spec={"numbering": True},
    )
    texts = [p.text for p in Document(str(path)).paragraphs]
    assert texts[1] == "1 一"
    assert texts[2] == "1.1 甲"
    assert texts[3] == "2 二"
    assert texts[4] == "2.1 乙"


def test_h4_style_override_via_format_spec(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "style.docx",
        paragraphs=[{"text": "四级标题", "heading": "h4"}],
        format_spec={
            "headings": {
                "h4": {"font_size_pt": 13, "bold": False, "color": "333333"}
            },
        },
    )
    h4 = Document(str(path)).styles["Heading 4"]
    assert h4.font.size.pt == pytest.approx(13)
    assert h4.font.bold is False


def test_h5_style_override_via_format_spec(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "style5.docx",
        paragraphs=[{"text": "五级标题", "heading": "h5"}],
        format_spec={
            "headings": {"h5": {"font_size_pt": 12, "bold": True}},
        },
    )
    h5 = Document(str(path)).styles["Heading 5"]
    assert h5.font.size.pt == pytest.approx(12)
    assert h5.font.bold is True


def test_linter_numbering_detects_broken_h4(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "broken.docx",
        paragraphs=[
            {"text": "章", "heading": "h1"},
            {"text": "节", "heading": "h2"},
            {"text": "小节", "heading": "h3"},
            {"text": "四级", "heading": "h4"},
        ],
        format_spec={"numbering": True},
    )
    # 手工破坏 h4 编号（模拟用户改动）
    doc = Document(str(path))
    doc.paragraphs[4].text = "9.9 破坏"
    doc.save(str(path))
    result = lint_docx(path, WordFormatSpec(**{"numbering": True}))
    issues = [i for i in result.issues if i.rule_id == "numbering/sequence"]
    assert issues
    assert "1.1.1.1" in issues[0].message


def test_model_rejects_heading_beyond_h5() -> None:
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="",
            filename="x.docx",
            title="T",
            paragraphs=[{"text": "正文", "heading": "h6"}],
        )

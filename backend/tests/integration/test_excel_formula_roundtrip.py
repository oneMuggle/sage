"""Excel formula closed-loop integration tests (Item 1.4).

端到端锁定公式闭环：
    generate_xlsx（'=' 字符串 → 真公式）
    → update_xlsx（编辑不改写、不丢公式）
    → read_xlsx(include_formulas=True)（公式文本 + 缺缓存值提示）
    → OfficeToolService.read(formula_mode=True)（工具调用同一视图）

与既有 test_excel_roundtrip.py 的 caveat 锁定测试互补：那边锁定
"编辑后缓存计算值会丢"，这里锁定 "公式文本不丢、公式视图可读回"。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.office.edit import update_xlsx
from backend.office.excel import generate_xlsx, read_xlsx
from backend.office.models import ExcelSheetSpec, OfficeExcelGenerateRequest

_FORMULA_CACHE_NOTE = "公式计算值需在 Excel 中打开后生效"


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    """Per-test scratch workspace matching ``managed_document_path`` layout."""
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _formula_req(workspace: Path, filename: str) -> OfficeExcelGenerateRequest:
    return OfficeExcelGenerateRequest(
        workspace_path=str(workspace),
        filename=filename,
        sheets=[
            ExcelSheetSpec(
                name="Calc",
                headers=["Item", "Amount"],
                rows=[["A", "10"], ["B", "20"], ["Sum", "=SUM(B2:B3)"]],
            ),
        ],
    )


# ─────────────────────────────────────────────────────────────────────────
# generate → edit → read（文件级闭环）
# ─────────────────────────────────────────────────────────────────────────


def test_roundtrip_generate_edit_read_formulas(workspace: Path) -> None:
    path = generate_xlsx(_formula_req(workspace, "formula-loop"))

    # 生成即公式（不是文本）
    wb = load_workbook(str(path), data_only=False)
    assert wb["Calc"]["B4"].value == "=SUM(B2:B3)"
    assert wb["Calc"]["B4"].data_type == "f"
    wb.close()

    # 编辑：改数值 + 追加公式行 —— 原公式保留、新公式生效
    saved, results = update_xlsx(
        path,
        [
            {
                "op": "set_cells",
                "sheet": "Calc",
                "cells": [{"addr": "B2", "value": "15"}],
            },
            {"op": "append_rows", "sheet": "Calc", "rows": [["Double", "=B3*2"]]},
        ],
    )
    assert saved
    assert all(r["ok"] for r in results)

    wb = load_workbook(str(path), data_only=False)
    try:
        ws = wb["Calc"]
        assert ws["B4"].value == "=SUM(B2:B3)"
        assert ws["B4"].data_type == "f"
        assert ws["B5"].value == "=B3*2"
        assert ws["B5"].data_type == "f"
        # 数值 coerce 语义不受公式改动影响（"15" → 15）
        assert ws["B2"].value == 15
    finally:
        wb.close()

    # 读取：公式视图列出公式 + 缺缓存值提示
    result = read_xlsx(path, include_formulas=True)
    sheet = result.sheets[0]
    assert sheet.formulas == ["B4=SUM(B2:B3)", "B5=B3*2"]
    assert sheet.note == _FORMULA_CACHE_NOTE
    assert sheet.rows[-1] == ["Double", ""]


def test_roundtrip_default_read_hides_formulas(workspace: Path) -> None:
    """默认读取（不带公式视图）输出形状与旧行为一致。"""
    path = generate_xlsx(_formula_req(workspace, "formula-default"))
    result = read_xlsx(path)
    sheet = result.sheets[0]
    assert sheet.formulas is None
    assert sheet.note is None


# ─────────────────────────────────────────────────────────────────────────
# OfficeToolService 闭环（create → read，工具调用同一管道）
# ─────────────────────────────────────────────────────────────────────────


def _seed_session(conn, session_id: str) -> None:
    conn.execute(
        "INSERT INTO sessions (id, title, created_at, updated_at) VALUES (?, ?, ?, ?)",
        (session_id, "t", 1, 1),
    )
    conn.commit()


def test_roundtrip_service_read_formula_mode(tmp_path: Path, workspace: Path) -> None:
    """OfficeToolService.create → read(formula_mode=True) 暴露同一公式视图。"""
    from backend.data.database import Database
    from backend.office.session_workspace import bind_session_workspace
    from backend.office.tool_service import OfficeToolService

    db = Database(db_path=str(tmp_path / "t.db"))
    db.init_db()
    conn = db.get_connection()
    _seed_session(conn, "sess-1")
    binding = bind_session_workspace(conn, "sess-1", str(workspace), now_ms=1)

    service = OfficeToolService()
    created = service.create(
        conn,
        "sess-1",
        binding.generation,
        doc_type="excel",
        filename="formula.xlsx",
        content={
            "sheets": [
                {
                    "name": "Calc",
                    "headers": ["Item", "Amount"],
                    "rows": [["A", "10"], ["B", "20"], ["Sum", "=SUM(B2:B3)"]],
                }
            ]
        },
    )
    assert created["success"], created
    doc_id = created["content"]["document_id"]

    # 默认读取：公式视图关闭（formulas 为 None）
    plain = service.read(conn, "sess-1", binding.generation, doc_id, section="all")
    assert plain["success"]
    sheet_plain = plain["content"]["sheets"][0]
    assert sheet_plain["formulas"] is None
    assert sheet_plain["note"] is None

    # formula_mode=True：公式文本 + 缺缓存值提示
    result = service.read(
        conn,
        "sess-1",
        binding.generation,
        doc_id,
        section="all",
        formula_mode=True,
    )
    assert result["success"]
    sheet = result["content"]["sheets"][0]
    assert sheet["formulas"] == ["B4=SUM(B2:B3)"]
    assert sheet["note"] == _FORMULA_CACHE_NOTE

    # section="summary" 语义不变：忽略 formula_mode，只返回 summary
    summary = service.read(
        conn,
        "sess-1",
        binding.generation,
        doc_id,
        section="summary",
        formula_mode=True,
    )
    assert summary["success"]
    assert set(summary["content"].keys()) == {"summary"}

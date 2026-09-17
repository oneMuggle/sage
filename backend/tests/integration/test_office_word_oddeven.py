"""Integration tests for Word odd/even page headers (Round 34).

Covers: settings.odd_and_even_pages_header_footer global flag, even-page
header/footer content written independently, zero change when
odd_even_pages is false, and model fields.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from backend.office.models import OfficeWordGenerateRequest, WordFormatSpec
from backend.office.word import generate_docx


def _generate(tmp_path: Path, name: str, spec: dict, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="测试文档",
        format_spec=spec,
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_odd_even_written(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "oe.docx",
        {
            "odd_even_pages": True,
            "even_page_header": {"text": "偶数页页眉"},
            "even_page_footer": {"text": "偶数页页脚"},
        },
        paragraphs=[{"text": "正文"}],
    )
    doc = Document(str(path))
    assert doc.settings.odd_and_even_pages_header_footer is True
    section = doc.sections[0]
    assert section.even_page_header.paragraphs[0].text == "偶数页页眉"
    assert section.even_page_footer.paragraphs[0].text == "偶数页页脚"


def test_odd_even_disabled_zero_change(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "off.docx",
        {
            "odd_even_pages": False,
            "even_page_header": {"text": "不应出现"},
        },
        paragraphs=[{"text": "正文"}],
    )
    doc = Document(str(path))
    assert doc.settings.odd_and_even_pages_header_footer is False
    section = doc.sections[0]
    assert section.even_page_header.is_linked_to_previous
    assert section.even_page_header.paragraphs[0].text != "不应出现"


def test_model_odd_even_fields() -> None:
    spec = WordFormatSpec(odd_even_pages=True)
    assert spec.odd_even_pages is True
    assert spec.even_page_header is None


def test_model_even_page_spec_accepted() -> None:
    from backend.office.models import WordHeaderFooterSpec

    spec = WordFormatSpec(
        odd_even_pages=True,
        even_page_header=WordHeaderFooterSpec(text="偶"),
    )
    assert spec.even_page_header is not None
    assert spec.even_page_header.text == "偶"

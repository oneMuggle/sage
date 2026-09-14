"""Integration tests for Word section breaks / landscape pages (Round 26).

Covers: multi-section documents with per-section page setup (landscape
orientation, margins), section break placement at start_paragraph, zero
change without section_breaks, and model validation (negative
start_paragraph rejected).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from pydantic import ValidationError

from backend.office.models import OfficeWordGenerateRequest
from backend.office.word import generate_docx


def _generate(tmp_path: Path, name: str, **kwargs) -> Path:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="测试文档",
        **kwargs,
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_landscape_section_for_wide_table(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "sec.docx",
        paragraphs=[
            {"text": "第一章", "heading": "h1"},
            {"text": "纵向内容"},
            {"text": "横向附表页", "heading": "h1"},
            {"text": "横向内容"},
        ],
        format_spec={
            "section_breaks": [
                {
                    "start_paragraph": 2,
                    "page_setup": {
                        "orientation": "landscape",
                        "margins_cm": {"top": 1.5},
                    },
                }
            ]
        },
    )
    doc = Document(str(path))
    assert len(doc.sections) == 2
    first, second = doc.sections[0], doc.sections[1]
    assert first.orientation == WD_ORIENT.PORTRAIT
    assert second.orientation == WD_ORIENT.LANDSCAPE
    assert second.page_width > second.page_height
    assert second.top_margin.cm == pytest.approx(1.5, abs=0.01)

    # 第一节边距未被波及
    assert first.top_margin.cm == pytest.approx(2.54, abs=0.01)
    # 内容归属：纵向节含前两段，横向节含后两段
    texts = [p.text for p in doc.paragraphs]
    assert texts[1] == "第一章"
    assert texts[2] == "纵向内容"
    # add_section 插入承载分节符的空段落，其后的段落属横向节
    assert texts[3] == ""
    assert texts[4] == "横向附表页"
    assert texts[5] == "横向内容"


def test_two_section_breaks_three_sections(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        "multi.docx",
        paragraphs=[
            {"text": "开篇"},
            {"text": "宽表页"},
            {"text": "收尾"},
        ],
        format_spec={
            "section_breaks": [
                {"start_paragraph": 1, "page_setup": {"orientation": "landscape"}},
                {
                    "start_paragraph": 2,
                    "page_setup": {"orientation": "portrait"},
                },
            ]
        },
    )
    doc = Document(str(path))
    assert len(doc.sections) == 3
    assert doc.sections[0].orientation == WD_ORIENT.PORTRAIT
    assert doc.sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert doc.sections[2].orientation == WD_ORIENT.PORTRAIT


def test_no_section_breaks_zero_change(tmp_path: Path) -> None:
    path = _generate(
        tmp_path, "plain.docx", paragraphs=[{"text": "正文"}]
    )
    doc = Document(str(path))
    assert len(doc.sections) == 1


def test_negative_start_paragraph_rejected() -> None:
    with pytest.raises(ValidationError):
        OfficeWordGenerateRequest(
            workspace_path="",
            filename="x.docx",
            title="T",
            format_spec={
                "section_breaks": [
                    {"start_paragraph": -1, "page_setup": {"orientation": "landscape"}}
                ]
            },
        )

"""Integration tests for Excel conditional formatting (Round 17).

Covers data_bar / color_scale / duplicate rules written via
generate_xlsx and read back through ws.conditional_formatting; illegal
ranges skipped without blocking; zero change when no rules; and the
office_create tool path.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest
from backend.tools.office_create_tool import OfficeCreateTool


def _generate(tmp_path: Path, sheets: list, name: str = "cf.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(
        workspace_path="",
        filename=name,
        sheets=sheets,
    )
    return generate_xlsx(req, output_dir=str(tmp_path))


def _cf_rules(ws):
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            out.append((str(cf.sqref), rule.type))
    return out


# ──────────────────────────────────────────────────────────────────────
# Rules
# ──────────────────────────────────────────────────────────────────────


def test_data_bar_rule(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["项目", "金额"],
            "rows": [["a", "10"], ["b", "50"]],
            "conditional_formats": [{"rule_type": "data_bar", "range": "B2:B3"}],
        }],
    )
    ws = load_workbook(str(tmp_path / "cf.xlsx"))["S"]
    rules = _cf_rules(ws)
    assert ("B2:B3", "dataBar") in rules


def test_color_scale_rule(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["项", "值"],
            "rows": [["a", "1"], ["b", "9"]],
            "conditional_formats": [{
                "rule_type": "color_scale",
                "range": "B2:B3",
                "min_color": "F8696B",
                "max_color": "63BE7B",
            }],
        }],
    )
    ws = load_workbook(str(tmp_path / "cf.xlsx"))["S"]
    rules = _cf_rules(ws)
    assert ("B2:B3", "colorScale") in rules


def test_duplicate_rule(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["名"],
            "rows": [["x"], ["x"], ["y"]],
            "conditional_formats": [{
                "rule_type": "duplicate",
                "range": "A2:A4",
                "fill_color": "FFFF00",
            }],
        }],
    )
    ws = load_workbook(str(tmp_path / "cf.xlsx"))["S"]
    rules = _cf_rules(ws)
    assert ("A2:A4", "expression") in rules  # FormulaRule 序列化为 expression


# ──────────────────────────────────────────────────────────────────────
# Robustness
# ──────────────────────────────────────────────────────────────────────


def test_illegal_range_rejected_at_model(tmp_path: Path) -> None:
    """非法 range 在模型层即拒绝（pattern），非法语义由生成器兜底。"""
    with pytest.raises(ValidationError):
        _generate(
            tmp_path,
            [{
                "name": "S",
                "headers": ["A"],
                "rows": [["1"]],
                "conditional_formats": [
                    {"rule_type": "data_bar", "range": "不是范围"},
                ],
            }],
        )


def test_openpyxl_level_failure_skipped_without_blocking(tmp_path: Path) -> None:
    """range 过 pattern 但超出 openpyxl 最大列（XFD）→ 单条跳过，生成成功。"""
    path = _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["A"],
            "rows": [["1"]],
            "conditional_formats": [
                {"rule_type": "data_bar", "range": "XFE1:XFF1"},
                {"rule_type": "data_bar", "range": "B2:B3"},
            ],
        }],
    )
    ws = load_workbook(str(path))["S"]
    assert ("B2:B3", "dataBar") in _cf_rules(ws)


def test_no_rules_zero_change(tmp_path: Path) -> None:
    path = _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
    ws = load_workbook(str(path))["S"]
    assert len(list(ws.conditional_formatting)) == 0


# ──────────────────────────────────────────────────────────────────────
# Tool path
# ──────────────────────────────────────────────────────────────────────


def test_office_create_tool_conditional_formats(tmp_path: Path) -> None:
    tool = OfficeCreateTool()
    result = tool.execute(
        doc_type="excel",
        output_dir=str(tmp_path),
        filename="tool.xlsx",
        content={
            "sheets": [{
                "name": "预算",
                "headers": ["科目", "金额"],
                "rows": [["差旅", "1200"], ["餐饮", "80"]],
                "conditional_formats": [
                    {"rule_type": "data_bar", "range": "B2:B3"}
                ],
            }],
        },
    )
    assert result.success, getattr(result, "error", None)
    ws = load_workbook(str(tmp_path / "tool.xlsx"))["预算"]
    assert ("B2:B3", "dataBar") in _cf_rules(ws)


# ──────────────────────────────────────────────────────────────────────
# Icon set (Round 19)
# ──────────────────────────────────────────────────────────────────────


def test_icon_set_rule(tmp_path: Path) -> None:
    _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["任务", "进度"],
            "rows": [["a", "10"], ["b", "90"]],
            "conditional_formats": [{
                "rule_type": "icon_set",
                "range": "B2:B3",
                "icon_style": "3TrafficLights1",
            }],
        }],
    )
    ws = load_workbook(str(tmp_path / "cf.xlsx"))["S"]
    rules = _cf_rules(ws)
    assert ("B2:B3", "iconSet") in rules


def test_icon_set_invalid_style_rejected_at_model() -> None:
    from pydantic import ValidationError

    from backend.office.models import ExcelConditionalFormatSpec

    with pytest.raises(ValidationError):
        ExcelConditionalFormatSpec(
            rule_type="icon_set",
            range="B2:B3",
            icon_style="7Smileys",
        )


def test_icon_set_default_style(tmp_path: Path) -> None:
    path = _generate(
        tmp_path,
        [{
            "name": "S",
            "headers": ["进度"],
            "rows": [["50"]],
            "conditional_formats": [{"rule_type": "icon_set", "range": "B2:B2"}],
        }],
    )
    ws = load_workbook(str(path))["S"]
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            assert rule.iconSet.iconSet == "3Arrows"

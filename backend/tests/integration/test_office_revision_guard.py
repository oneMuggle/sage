"""Integration tests for the F1 revision guard on the office write path.

The analysis report (docs/mcp-office-capability-analysis-20260926.md, F1)
reproduced this hole: preview ``Data!A1 10→20`` → an external process writes
``999`` → the apply still succeeds and lands ``20``, silently discarding the
external write. These tests pin the fixed contract:

* preview stamps the source revision it actually read;
* an apply carrying that revision is refused with 409 once the file moved
  on, and the file stays byte-for-byte identical;
* an apply whose revision still matches succeeds and reports the new one;
* a retry with the same idempotency key replays instead of re-applying;
* omitting ``expected_revision`` keeps the pre-existing behaviour, so
  existing callers are not broken by the new guard.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from backend.api.office_routes import (
    get_document_revision_endpoint,
    preview_update_endpoint,
    update_document_endpoint,
)
from backend.data.database import get_database
from backend.office.apply_update import OfficeDocUpdateRequest
from backend.office.diff_preview import OfficeUpdatePreviewRequest
from backend.office.errors import (
    OfficeRevisionConflictError,
    office_error_to_http_status,
)
from backend.office.models import (
    OfficeDocStatus,
    OfficeDocType,
    OfficeDocumentMetadata,
    OfficeDocumentSummary,
)
from backend.office.revision import _reset_for_tests
from backend.office.storage import (
    document_path,
    generate_document_dir,
    get_document,
    save_document,
)

pytestmark = pytest.mark.integration


@pytest.fixture(autouse=True)
def _clean_revision_state():
    _reset_for_tests()
    yield
    _reset_for_tests()


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


def _seed_xlsx(workspace: Path, doc_id: str = "xl-1") -> str:
    """Managed .xlsx with ``Data!A1 = 10`` + its DB row; returns the doc id."""
    from openpyxl import Workbook

    directory = generate_document_dir(workspace, OfficeDocType.EXCEL, doc_id)
    managed = directory / "book.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet["A1"] = 10
    book.save(str(managed))
    now = 1_700_000_000_000
    save_document(
        get_database().get_connection(),
        OfficeDocumentSummary(
            id=doc_id,
            workspace_path=str(workspace.resolve()),
            doc_type=OfficeDocType.EXCEL,
            original_filename=None,
            generated_filename="book.xlsx",
            status=OfficeDocStatus.PARSED,
            created_at=now,
            updated_at=now,
            metadata=OfficeDocumentMetadata(file_size_bytes=managed.stat().st_size),
        ),
    )
    return doc_id


def _seed_docx(workspace: Path, doc_id: str = "wd-1") -> str:
    from docx import Document

    directory = generate_document_dir(workspace, OfficeDocType.WORD, doc_id)
    managed = directory / "report.docx"
    document = Document()
    document.add_paragraph("Hello world")
    document.save(str(managed))
    now = 1_700_000_000_000
    save_document(
        get_database().get_connection(),
        OfficeDocumentSummary(
            id=doc_id,
            workspace_path=str(workspace.resolve()),
            doc_type=OfficeDocType.WORD,
            original_filename=None,
            generated_filename="report.docx",
            status=OfficeDocStatus.PARSED,
            created_at=now,
            updated_at=now,
            metadata=OfficeDocumentMetadata(file_size_bytes=managed.stat().st_size),
        ),
    )
    return doc_id


def _managed(doc_id: str) -> Path:
    doc = get_document(get_database().get_connection(), doc_id)
    assert doc is not None
    return document_path(doc)


def _sha(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _set_a1(path: Path, value: int) -> None:
    """Simulate an external editor writing the file behind our back."""
    from openpyxl import load_workbook

    book = load_workbook(str(path))
    book["Data"]["A1"] = value
    book.save(str(path))


def _preview(workspace: Path, doc_id: str, ops):
    return preview_update_endpoint(
        OfficeUpdatePreviewRequest(
            workspace_path=str(workspace.resolve()), doc_id=doc_id, ops=ops
        )
    )


_SET_A1_TO_20 = [
    {"op": "set_cells", "sheet": "Data", "cells": [{"addr": "A1", "value": 20}]}
]


# ──────────────────────────────────────────────────────────────────────
# preview stamps the version it read
# ──────────────────────────────────────────────────────────────────────


def test_preview_returns_the_source_revision_it_read(workspace: Path):
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)

    preview = _preview(workspace, doc_id, _SET_A1_TO_20)

    assert preview.ok is True
    assert preview.source_revision == _sha(managed)
    assert preview.ops_hash and preview.ops_hash.startswith("ops:")
    assert preview.preview_id and preview.preview_id.startswith("pv_")
    # A dry run must not touch the source.
    assert _sha(managed) == preview.source_revision


def test_revision_endpoint_matches_the_file_on_disk(workspace: Path):
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)

    payload = get_document_revision_endpoint(doc_id)

    assert payload["doc_id"] == doc_id
    assert payload["revision"] == _sha(managed)
    assert payload["size_bytes"] == managed.stat().st_size


# ──────────────────────────────────────────────────────────────────────
# the F1 reproduction: stale apply must be refused
# ──────────────────────────────────────────────────────────────────────


def test_external_write_after_preview_makes_the_apply_conflict(workspace: Path):
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)
    preview = _preview(workspace, doc_id, _SET_A1_TO_20)

    # External editor writes 999 between preview and apply.
    _set_a1(managed, 999)
    external_bytes = managed.read_bytes()

    with pytest.raises(OfficeRevisionConflictError) as excinfo:
        update_document_endpoint(
            doc_id,
            OfficeDocUpdateRequest(
                ops=_SET_A1_TO_20, expected_revision=preview.source_revision
            ),
        )

    assert office_error_to_http_status(excinfo.value) == 409
    assert excinfo.value.expected == preview.source_revision
    assert excinfo.value.actual == _sha(managed)
    # Nothing was written: the external 999 survives byte-for-byte.
    assert managed.read_bytes() == external_bytes


def test_matching_revision_applies_and_reports_the_new_one(workspace: Path):
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)
    preview = _preview(workspace, doc_id, _SET_A1_TO_20)

    result = update_document_endpoint(
        doc_id,
        OfficeDocUpdateRequest(
            ops=_SET_A1_TO_20, expected_revision=preview.source_revision
        ),
    )

    from openpyxl import load_workbook

    assert result.ok is True
    assert result.previous_revision == preview.source_revision
    assert result.revision == _sha(managed)
    assert result.revision != result.previous_revision
    assert result.idempotent_replay is False
    assert load_workbook(str(managed))["Data"]["A1"].value == 20


def test_second_apply_with_the_now_stale_revision_conflicts(workspace: Path):
    """A double-submit without an idempotency key is a stale write, not a retry."""
    doc_id = _seed_xlsx(workspace)
    preview = _preview(workspace, doc_id, _SET_A1_TO_20)
    update_document_endpoint(
        doc_id,
        OfficeDocUpdateRequest(
            ops=_SET_A1_TO_20, expected_revision=preview.source_revision
        ),
    )

    with pytest.raises(OfficeRevisionConflictError):
        update_document_endpoint(
            doc_id,
            OfficeDocUpdateRequest(
                ops=_SET_A1_TO_20, expected_revision=preview.source_revision
            ),
        )


# ──────────────────────────────────────────────────────────────────────
# idempotency
# ──────────────────────────────────────────────────────────────────────


def test_retry_with_the_same_idempotency_key_replays_once(workspace: Path):
    """An append op retried with the same key must not append twice."""
    from docx import Document

    doc_id = _seed_docx(workspace)
    managed = _managed(doc_id)
    ops = [{"op": "append_paragraphs", "paragraphs": [{"text": "Appended once"}]}]

    first = update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=ops, idempotency_key="retry-1")
    )
    after_first = managed.read_bytes()

    second = update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=ops, idempotency_key="retry-1")
    )

    assert first.idempotent_replay is False
    assert second.idempotent_replay is True
    assert second.revision == first.revision
    assert managed.read_bytes() == after_first
    texts = [p.text for p in Document(str(managed)).paragraphs]
    assert texts.count("Appended once") == 1


def test_a_different_key_applies_again(workspace: Path):
    from docx import Document

    doc_id = _seed_docx(workspace)
    managed = _managed(doc_id)
    ops = [{"op": "append_paragraphs", "paragraphs": [{"text": "Appended once"}]}]

    update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=ops, idempotency_key="retry-1")
    )
    second = update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=ops, idempotency_key="retry-2")
    )

    assert second.idempotent_replay is False
    texts = [p.text for p in Document(str(managed)).paragraphs]
    assert texts.count("Appended once") == 2


def test_replay_is_dropped_once_the_file_changes_again(workspace: Path):
    """A remembered outcome only replays while its revision is still current."""
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)

    update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=_SET_A1_TO_20, idempotency_key="k")
    )
    _set_a1(managed, 999)

    replayed = update_document_endpoint(
        doc_id, OfficeDocUpdateRequest(ops=_SET_A1_TO_20, idempotency_key="k")
    )

    from openpyxl import load_workbook

    assert replayed.idempotent_replay is False
    assert load_workbook(str(managed))["Data"]["A1"].value == 20


# ──────────────────────────────────────────────────────────────────────
# backwards compatibility
# ──────────────────────────────────────────────────────────────────────


def test_apply_without_expected_revision_keeps_legacy_behaviour(workspace: Path):
    """Existing callers (no version fields) still apply — and now learn the revision."""
    doc_id = _seed_xlsx(workspace)
    managed = _managed(doc_id)
    _preview(workspace, doc_id, _SET_A1_TO_20)
    _set_a1(managed, 999)

    result = update_document_endpoint(doc_id, OfficeDocUpdateRequest(ops=_SET_A1_TO_20))

    from openpyxl import load_workbook

    assert result.ok is True
    assert result.revision == _sha(managed)
    assert load_workbook(str(managed))["Data"]["A1"].value == 20

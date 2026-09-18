"""Integration tests for Round 31 configurable tunables.

Covers: Excel freeze_panes parameterization (B2/C2 notations, precedence
over freeze_header, zero change when absent) and the Pillow threshold
environment variable override (SAGE_IMAGE_OPTIMIZE_THRESHOLD_BYTES,
0 = disable optimization).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import OfficeExcelGenerateRequest


def _generate(tmp_path: Path, sheets: list, name: str = "fp.xlsx") -> Path:
    req = OfficeExcelGenerateRequest(workspace_path="", filename=name, sheets=sheets)
    return generate_xlsx(req, output_dir=str(tmp_path))


class TestFreezePanes:
    def test_b2_freezes_row_and_column(self, tmp_path: Path) -> None:
        _generate(
            tmp_path,
            [{
                "name": "S",
                "headers": ["A", "B"],
                "rows": [["1", "2"]],
                "freeze_panes": "B2",
            }],
        )
        ws = load_workbook(str(tmp_path / "fp.xlsx"))["S"]
        assert ws.freeze_panes == "B2"

    def test_c2_notation(self, tmp_path: Path) -> None:
        _generate(
            tmp_path,
            [{
                "name": "S",
                "headers": ["A", "B", "C"],
                "rows": [["1", "2", "3"]],
                "freeze_panes": "C2",
            }],
        )
        ws = load_workbook(str(tmp_path / "fp.xlsx"))["S"]
        assert ws.freeze_panes == "C2"

    def test_freeze_panes_wins_over_freeze_header(self, tmp_path: Path) -> None:
        _generate(
            tmp_path,
            [{
                "name": "S",
                "headers": ["A", "B"],
                "rows": [["1", "2"]],
                "freeze_header": True,
                "freeze_panes": "B2",
            }],
        )
        ws = load_workbook(str(tmp_path / "fp.xlsx"))["S"]
        # freeze_panes 优先（B2 而非 A2）
        assert ws.freeze_panes == "B2"

    def test_freeze_header_still_works(self, tmp_path: Path) -> None:
        _generate(
            tmp_path,
            [{
                "name": "S",
                "headers": ["A", "B"],
                "rows": [["1", "2"]],
                "freeze_header": True,
            }],
        )
        ws = load_workbook(str(tmp_path / "fp.xlsx"))["S"]
        assert ws.freeze_panes == "A2"

    def test_neither_set_zero_change(self, tmp_path: Path) -> None:
        _generate(tmp_path, [{"name": "S", "headers": ["A"], "rows": [["1"]]}])
        ws = load_workbook(str(tmp_path / "fp.xlsx"))["S"]
        assert ws.freeze_panes is None

    def test_invalid_panes_rejected(self) -> None:
        from pydantic import ValidationError

        from backend.office.models import ExcelSheetSpec

        with pytest.raises(ValidationError):
            ExcelSheetSpec(name="S", headers=["A"], freeze_panes="not-a-cell")


class TestThresholdEnv:
    def test_env_override_disables(self, monkeypatch) -> None:
        import importlib

        from backend.office import image_optimize

        monkeypatch.setenv("SAGE_IMAGE_OPTIMIZE_THRESHOLD_BYTES", "0")
        importlib.reload(image_optimize)
        data = b"\xff\xd8\xff" + b"x" * (
            image_optimize.DEFAULT_OPTIMIZE_THRESHOLD_BYTES + 1
        )
        result = image_optimize.optimize_image_bytes(data)
        assert result == data  # 阈值 0 = 禁用优化，原样返回
        monkeypatch.delenv("SAGE_IMAGE_OPTIMIZE_THRESHOLD_BYTES")
        importlib.reload(image_optimize)

    def test_default_threshold_is_8mb(self, monkeypatch) -> None:
        monkeypatch.delenv("SAGE_IMAGE_OPTIMIZE_THRESHOLD_BYTES", raising=False)
        import importlib

        from backend.office import image_optimize

        importlib.reload(image_optimize)
        assert image_optimize.optimize_threshold_bytes() == 8 * 1024 * 1024

"""Office 图表/图片 round-trip integration tests（批次 2.1 + 2.3）.

端到端锁定：
    generate_xlsx（charts + column_widths）→ read_xlsx（内容完好）
                            → load_workbook（原生图表可读回）
    update_xlsx（add_chart 编辑期挂载）→ read_xlsx（内容完好）
    generate_docx（images base64/工作区路径）→ read_docx（inline images >= 1）
    generate_ppt（slide.image）→ read_ppt（image_count >= 1）

与 test_excel_roundtrip / test_images_styles 的分工：那边锁定编辑 op 的
单点行为，这里锁定「生成 → 落盘 → 读回」整条链路不被图表/图片破坏。
"""

from __future__ import annotations

import base64
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.office.edit import update_xlsx
from backend.office.errors import OfficeGenerateError
from backend.office.excel import generate_xlsx, read_xlsx
from backend.office.models import (
    ChartSeriesSpec,
    ChartSpec,
    ExcelCellRange,
    ExcelChartSpec,
    ExcelSheetSpec,
    ImageSourceSpec,
    OfficeExcelGenerateRequest,
    OfficePptGenerateRequest,
    OfficeWordGenerateRequest,
    PptSlideSpec,
    WordParagraphSpec,
)
from backend.office.ppt import generate_ppt, read_ppt
from backend.office.word import generate_docx, read_docx

pytestmark = pytest.mark.integration

_MINIMAL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP8z8Dw"
    "HwAFBQIAX8v0gQAAAABJRU5ErkJggg=="
)
_MINIMAL_PNG_BYTES = base64.b64decode(_MINIMAL_PNG_B64, validate=True)
_DATA_URI = "data:image/png;base64," + _MINIMAL_PNG_B64


# ──────────────────────────────────────────────────────────────────────
# Excel：generate with charts → read
# ──────────────────────────────────────────────────────────────────────


def _chart_req(tmp_path: Path, filename: str = "报表.xlsx") -> OfficeExcelGenerateRequest:
    return OfficeExcelGenerateRequest(
        workspace_path=str(tmp_path),
        filename=filename,
        sheets=[
            ExcelSheetSpec(
                name="销售",
                headers=["月份", "销售额", "成本"],
                rows=[
                    ["一月", "100", "60"],
                    ["二月", "180", "90"],
                    ["三月", "150", "70"],
                ],
                column_widths=[14, 20, 16],
            ),
            ExcelSheetSpec(name="汇总", headers=["指标"], rows=[["合计"]]),
        ],
        charts=[
            ExcelChartSpec(
                type="line",
                anchor="F2",
                data_ref=ExcelCellRange(min_col=2, min_row=1, max_col=3, max_row=4),
                titles_from_data=True,
                categories_ref=ExcelCellRange(min_col=1, min_row=2, max_col=1, max_row=4),
                title="销售额 vs 成本",
            ),
            ExcelChartSpec(
                sheet="汇总",
                type="pie",
                anchor="A5",
                data_ref=ExcelCellRange(min_col=2, min_row=1, max_col=2, max_row=4),
                titles_from_data=True,
            ),
        ],
    )


def test_generate_xlsx_with_charts_reads_back(tmp_path: Path) -> None:
    req = _chart_req(tmp_path)
    out = generate_xlsx(req)
    # read_xlsx：内容完好（图表/列宽不影响网格读取）
    result = read_xlsx(out, workspace_path=str(tmp_path))
    assert [s.name for s in result.sheets] == ["销售", "汇总"]
    assert result.sheets[0].rows[0] == ["月份", "销售额", "成本"]
    assert result.sheets[0].rows[3] == ["三月", "150", "70"]
    # 原生图表落盘且可读回：'销售' 1 张 line、'汇总' 1 张 pie
    wb = load_workbook(str(out))
    assert len(wb["销售"]._charts) == 1
    assert "LineChart" in type(wb["销售"]._charts[0]).__name__
    assert len(wb["汇总"]._charts) == 1
    # 列宽 generation 侧生效
    assert wb["销售"].column_dimensions["A"].width == 14


def test_generate_xlsx_chart_on_missing_sheet_fails(tmp_path: Path) -> None:
    req = OfficeExcelGenerateRequest(
        workspace_path=str(tmp_path),
        filename="bad.xlsx",
        sheets=[ExcelSheetSpec(name="数据", headers=["A"], rows=[["1"]])],
        charts=[
            ExcelChartSpec(
                sheet="不存在",
                type="bar",
                anchor="C1",
                data_ref=ExcelCellRange(min_col=1, min_row=1, max_col=1, max_row=1),
            )
        ],
    )
    with pytest.raises(OfficeGenerateError):
        generate_xlsx(req)


def test_update_xlsx_add_chart_roundtrip(tmp_path: Path) -> None:
    req = OfficeExcelGenerateRequest(
        workspace_path=str(tmp_path),
        filename="plain.xlsx",
        sheets=[
            ExcelSheetSpec(
                name="数据",
                headers=["类别", "数量"],
                rows=[["甲", "10"], ["乙", "20"]],
            )
        ],
    )
    out = generate_xlsx(req)
    saved, results = update_xlsx(
        out,
        [
            {
                "op": "add_chart",
                "sheet": "数据",
                "type": "pie",
                "anchor": "D2",
                "data_ref": {"min_col": 2, "min_row": 1, "max_col": 2, "max_row": 3},
                "titles_from_data": True,
                "categories_ref": {"min_col": 1, "min_row": 2, "max_col": 1, "max_row": 3},
            }
        ],
    )
    assert saved
    assert results[0]["ok"]
    result = read_xlsx(out, workspace_path=str(tmp_path))
    assert result.sheets[0].rows[1] == ["甲", "10"]
    assert len(load_workbook(str(out))["数据"]._charts) == 1


# ──────────────────────────────────────────────────────────────────────
# Word / PPT：image round-trip
# ──────────────────────────────────────────────────────────────────────


def test_generate_docx_with_image_roundtrip(tmp_path: Path) -> None:
    req = OfficeWordGenerateRequest(
        workspace_path=str(tmp_path),
        filename="带图.docx",
        title="带图文档",
        paragraphs=[WordParagraphSpec(text="正文一段")],
        images=[
            ImageSourceSpec(source=_DATA_URI, width_inches=2.5),
        ],
    )
    out = generate_docx(req, output_dir=str(tmp_path))
    result = read_docx(out, workspace_path=str(tmp_path))
    assert result.images >= 1
    assert result.paragraphs[0].text == "带图文档"
    assert "正文一段" in [p.text for p in result.paragraphs]


def test_generate_docx_with_workspace_image_roundtrip(tmp_path: Path) -> None:
    (tmp_path / "logo.png").write_bytes(_MINIMAL_PNG_BYTES)
    req = OfficeWordGenerateRequest(
        workspace_path=str(tmp_path),
        filename="带图2.docx",
        title="工作区图片",
        paragraphs=[],
        images=[ImageSourceSpec(source="logo.png")],
    )
    out = generate_docx(req)
    assert read_docx(out, workspace_path=str(tmp_path)).images == 1


def test_generate_ppt_with_image_roundtrip(tmp_path: Path) -> None:
    req = OfficePptGenerateRequest(
        workspace_path=str(tmp_path),
        filename="deck.pptx",
        slides=[
            PptSlideSpec(
                title="图表页",
                bullets=["要点"],
                image=ImageSourceSpec(source=_DATA_URI, width_inches=4.0),
            )
        ],
    )
    out = generate_ppt(req, output_dir=str(tmp_path))
    result = read_ppt(out, workspace_path=str(tmp_path))
    assert result.slides[0].image_count == 1
    assert result.slides[0].title == "图表页"


# ──────────────────────────────────────────────────────────────────────
# matplotlib 渲染图表 → 嵌入 Word（2.1 全链路，缺 matplotlib 跳过）
# ──────────────────────────────────────────────────────────────────────


def test_render_chart_png_embeds_into_docx(tmp_path: Path) -> None:
    pytest.importorskip("matplotlib", reason="matplotlib 未安装")
    from backend.office.charts import render_chart_png

    png_path = render_chart_png(
        ChartSpec(
            type="bar",
            title="季度销量",
            labels=["Q1", "Q2", "Q3"],
            series=[ChartSeriesSpec(name="销量", y=[10, 20, 30])],
        ),
        tmp_path,
    )
    req = OfficeWordGenerateRequest(
        workspace_path=str(tmp_path),
        filename="图表报告.docx",
        title="图表报告",
        paragraphs=[WordParagraphSpec(text="看下图")],
        images=[ImageSourceSpec(source=str(png_path), width_inches=5.0)],
    )
    out = generate_docx(req, output_dir=str(tmp_path))
    assert read_docx(out, workspace_path=str(tmp_path)).images == 1

"""Integration tests for the template library routes (batch 3 — Item 3.2).

Exercises the new HTTP endpoints through the route functions (same style as
test_office_doc_actions_routes.py):
- ``GET /office/templates``: builtin registry + workspace templates scan.
- ``POST /office/templates/instantiate``: fills a builtin or workspace
  template into the managed Word layout and persists a ``generated`` row so
  the document appears in ``GET /documents``.
- Unknown template ids raise ``OfficeFileNotFoundError`` (mapped to 404 by
  the registered exception handler).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from backend.api.office_routes import (
    instantiate_template_endpoint,
    list_documents_endpoint,
    list_templates_endpoint,
)
from backend.data.database import get_database
from backend.office.errors import OfficeFileNotFoundError
from backend.office.models import OfficeDocType, OfficeTemplateInstantiateRequest

pytestmark = pytest.mark.integration


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


WEEKLY_DATA = {
    "title": "后端组周报",
    "author": "李四",
    "week_range": "2026-09-01 ~ 2026-09-07",
    "highlights": "模板库上线",
    "next_plan": "前端模板选择器",
    "risks": "无",
}


def test_list_templates_endpoint_returns_builtins(workspace: Path):
    response = list_templates_endpoint(workspace_path=str(workspace))
    builtins = [t for t in response.templates if t.source == "builtin"]
    assert len(builtins) == 6
    assert {t.id for t in builtins} >= {"weekly_report", "meeting_minutes", "resume"}


def test_list_templates_endpoint_scans_workspace_templates(workspace: Path):
    templates_dir = workspace / "office" / "templates"
    templates_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("项目：{{project_name}}")
    doc.save(str(templates_dir / "my-template.docx"))

    response = list_templates_endpoint(workspace_path=str(workspace))
    workspace_entries = [t for t in response.templates if t.source == "workspace"]
    assert len(workspace_entries) == 1
    assert workspace_entries[0].filename == "my-template.docx"
    assert {p.name for p in workspace_entries[0].placeholders} == {"project_name"}


def test_instantiate_builtin_template_persists_row_in_documents(workspace: Path):
    req = OfficeTemplateInstantiateRequest(
        workspace_path=str(workspace),
        template_id="weekly_report",
        filename="周报-w37.docx",
        data=WEEKLY_DATA,
    )
    result = instantiate_template_endpoint(req)

    output = Path(result.output_path)
    assert output.is_file()
    assert output.parent.parent == workspace / "office" / "word"
    assert result.filled_count == 6

    listing = list_documents_endpoint(workspace_path=str(workspace))
    assert listing.total == 1
    row = listing.documents[0]
    assert row.doc_type == OfficeDocType.WORD
    assert row.status.value == "generated"
    assert row.generated_filename == "周报-w37.docx"
    assert row.id == output.parent.name

    # DB 行确实落库（存储层可查）
    saved = get_database().get_connection().execute(
        "SELECT generated_filename, status FROM office_documents WHERE id = ?",
        (row.id,),
    ).fetchone()
    assert saved is not None
    assert saved["generated_filename"] == "周报-w37.docx"
    assert saved["status"] == "generated"

    # 填充内容落到最终文档
    text = "\n".join(p.text for p in Document(str(output)).paragraphs)
    assert "后端组周报" in text
    assert "模板库上线" in text


def test_instantiate_workspace_template_persists_row_in_documents(workspace: Path):
    templates_dir = workspace / "office" / "templates"
    templates_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("项目：{{project_name}}")
    doc.save(str(templates_dir / "kickoff.docx"))

    req = OfficeTemplateInstantiateRequest(
        workspace_path=str(workspace),
        workspace_template="kickoff.docx",
        filename="kickoff-filled.docx",
        data={"project_name": "启明星"},
    )
    result = instantiate_template_endpoint(req)

    assert Path(result.output_path).is_file()
    text = "\n".join(p.text for p in Document(result.output_path).paragraphs)
    assert "项目：启明星" in text

    listing = list_documents_endpoint(workspace_path=str(workspace))
    assert listing.total == 1
    assert listing.documents[0].generated_filename == "kickoff-filled.docx"


def test_instantiate_unknown_template_id_raises_404_error(workspace: Path):
    req = OfficeTemplateInstantiateRequest(
        workspace_path=str(workspace),
        template_id="ghost_template",
        filename="x.docx",
        data={},
    )
    with pytest.raises(OfficeFileNotFoundError):
        instantiate_template_endpoint(req)

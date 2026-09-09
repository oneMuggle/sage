"""1.5 digest 升级单元测试: 真实小文件 → 新 digest 形状 (不再有损)。

覆盖 (item 1.5):
- _digest_word: heading 层级 / 全段文本 (不再砍首句) / 列表 / GFM 表格
- _digest_excel: 表头 + 数值列统计 + 自适应行数 (小表 >5 行全量)
- _digest_pdf: .pdf 进 OFFICE_EXTS, digest 含页数 + 页文本
- budget: 超预算时出现 `[…已截断，共 N …]` marker (heading/表头仍保留)
- 边界: 非法 ext 仍拒绝; >50MB (mock size 上限, 不真造大文件) 拒绝

docx/xlsx 用 python-docx / openpyxl 现场生成, pdf 用
backend.office.pdf.generate_pdf (纯 ASCII 内容, 不依赖 CJK 字体)。
"""

from __future__ import annotations

import pytest

from backend.chat import attachment_resolver
from backend.chat.attachment_resolver import (
    MAX_ATTACHMENT_DIGEST_BYTES,
    OFFICE_EXTS,
    _digest_excel,
    _digest_pdf,
    _digest_word,
    extract_mentions,
    process,
)

# ─── fixtures: 现场生成真实小文件 ────────────────────────────────


@pytest.fixture()
def word_workspace(tmp_path):
    """带结构化 docx 的工作区: heading + 多句段落 + 两种列表 + 表格。"""
    from docx import Document

    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Intro paragraph before the details.")
    doc.add_heading("Details", level=2)
    doc.add_paragraph(
        "Revenue grew in every region. The growth accelerated in Q4. "
        "Management expects the trend to continue."
    )
    doc.add_paragraph("First bullet item", style="List Bullet")
    doc.add_paragraph("Second bullet item", style="List Bullet")
    doc.add_paragraph("First numbered step", style="List Number")
    doc.add_paragraph("Second numbered step", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "EMEA"
    table.cell(1, 1).text = "1200"
    doc.save(str(tmp_path / "report.docx"))
    return tmp_path


@pytest.fixture()
def excel_workspace(tmp_path):
    """带 xlsx 的工作区: Sales 12 行数值数据 + Notes 纯文本表。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Qty", "Price"])
    for i in range(1, 13):  # 12 行数据 → 验证不再硬编码 5 行
        ws.append([f"Item{i:02d}", i, i * 2])
    notes = wb.create_sheet("Notes")
    notes.append(["Comment"])
    notes.append(["plain text only"])
    wb.save(str(tmp_path / "book.xlsx"))
    return tmp_path


@pytest.fixture()
def pdf_workspace(tmp_path):
    """带 2 页 ASCII pdf 的工作区 (走 backend.office.pdf 的生成纯函数)。"""
    from backend.office.models import PdfGenerateRequest, PdfPageSpec
    from backend.office.pdf import generate_pdf

    req = PdfGenerateRequest(
        workspace_path=str(tmp_path),
        filename="handbook.pdf",
        pages=[
            PdfPageSpec(title="Overview", paragraphs=["Alpha page plain content."]),
            PdfPageSpec(title="Appendix", paragraphs=["Beta page plain content."]),
        ],
    )
    generate_pdf(req)
    return tmp_path


# ─── _digest_word ────────────────────────────────────────────────


def test_digest_word_keeps_full_paragraph_text(word_workspace) -> None:
    """多句段落不再砍到首句 — 第 2/3 句必须完整出现在 digest 里。"""
    out = _digest_word(str(word_workspace / "report.docx"), str(word_workspace))
    assert "Revenue grew in every region." in out
    assert "The growth accelerated in Q4. Management expects the trend to continue." in out


def test_digest_word_heading_levels_lists(word_workspace) -> None:
    out = _digest_word(str(word_workspace / "report.docx"), str(word_workspace))
    assert "# Quarterly Report" in out
    assert "## Details" in out
    assert "- First bullet item" in out
    assert "- Second bullet item" in out
    assert "1. First numbered step" in out
    assert "2. Second numbered step" in out


def test_digest_word_table_as_gfm_markdown(word_workspace) -> None:
    out = _digest_word(str(word_workspace / "report.docx"), str(word_workspace))
    assert "| Region | Revenue |" in out
    assert "| --- | --- |" in out
    assert "| EMEA | 1200 |" in out


def test_digest_word_truncation_marker_under_budget(
    word_workspace, monkeypatch
) -> None:
    """收紧 budget (mock, 不造超大文档): 正文被截断并带 marker, heading 仍全保留。"""
    monkeypatch.setattr(attachment_resolver, "MAX_ATTACHMENT_DIGEST_BYTES", 150)
    out = _digest_word(str(word_workspace / "report.docx"), str(word_workspace))
    assert "[…已截断，共 5 段]" in out
    assert "# Quarterly Report" in out
    assert "## Details" in out
    # 表格不受 budget 影响, 全量保留
    assert "| EMEA | 1200 |" in out


def test_digest_word_within_budget_no_marker(word_workspace) -> None:
    """默认 8KB 预算下小文档不应出现截断 marker。"""
    out = _digest_word(str(word_workspace / "report.docx"), str(word_workspace))
    assert "已截断" not in out
    assert len(out.encode("utf-8")) <= MAX_ATTACHMENT_DIGEST_BYTES + 200


# ─── _digest_excel ───────────────────────────────────────────────


def test_digest_excel_header_stats_and_adaptive_rows(excel_workspace) -> None:
    out = _digest_excel(str(excel_workspace / "book.xlsx"), str(excel_workspace))
    assert "sheets: Sales, Notes" in out
    assert "--- Sales (" in out  # 行列数标注
    assert "Product\tQty\tPrice" in out  # 表头行
    # 数值列统计: Qty/Price 是数值列, Product 不是
    assert "stat Qty: count=12, non_null=12, min=1, max=12, mean=6.5" in out
    assert "stat Price: count=12, non_null=12, min=2, max=24, mean=13" in out
    # 自适应行数: 12 行小表应全量输出 (旧实现只会给 header + 前 4 行数据)
    assert "Item05\t5\t10" in out
    assert "Item12\t12\t24" in out


def test_digest_excel_text_column_gets_no_stats(excel_workspace) -> None:
    out = _digest_excel(str(excel_workspace / "book.xlsx"), str(excel_workspace))
    assert "stat Product:" not in out
    assert "stat Comment:" not in out


def test_digest_excel_truncation_marker_under_budget(
    excel_workspace, monkeypatch
) -> None:
    """收紧 budget: 数据行按预算装载, 装不下的行数显式标注。"""
    monkeypatch.setattr(attachment_resolver, "MAX_ATTACHMENT_DIGEST_BYTES", 200)
    out = _digest_excel(str(excel_workspace / "book.xlsx"), str(excel_workspace))
    assert "[…已截断，共" in out
    # 表头 + 统计等固定段仍保留
    assert "--- Sales (" in out
    assert "Product\tQty\tPrice" in out
    assert "stat Qty:" in out


# ─── _digest_pdf / .pdf 接入 ─────────────────────────────────────


def test_pdf_ext_accepted() -> None:
    assert ".pdf" in OFFICE_EXTS
    mentions = extract_mentions("看 @handbook.pdf")
    assert len(mentions) == 1
    assert mentions[0].kind == "office-pdf"


def test_digest_pdf_contains_page_count_and_text(pdf_workspace) -> None:
    out = _digest_pdf(str(pdf_workspace / "handbook.pdf"), str(pdf_workspace))
    assert out.startswith("pages: 2")
    assert "--- page 1 ---" in out
    assert "Alpha page plain content." in out
    assert "--- page 2 ---" in out
    assert "Beta page plain content." in out


def test_process_injects_pdf_attachment_block(pdf_workspace) -> None:
    """端到端: .pdf mention 走原 <attachments> 注入契约, 块结构不变。"""
    out = process("总结 @handbook.pdf", workspace=str(pdf_workspace))
    assert out.startswith("<attachments>")
    assert out.endswith("</attachments>")
    assert "=== handbook.pdf ===" in out
    assert "Alpha page plain content." in out


# ─── 边界 ────────────────────────────────────────────────────────


def test_unsupported_ext_still_rejected(tmp_path) -> None:
    mentions = extract_mentions("@notes.txt")
    assert len(mentions) == 1
    assert mentions[0].kind is None
    assert process("see @notes.txt", workspace=str(tmp_path)) == ""


def test_oversized_file_rejected(word_workspace, monkeypatch) -> None:
    """mock size 上限 (收紧常数等价触发同一分支), 不真造 50MB 文件。"""
    monkeypatch.setattr(attachment_resolver, "MAX_ATTACHMENT_FILE_SIZE_BYTES", 10)
    out = process("看 @report.docx", workspace=str(word_workspace))
    assert out == ""

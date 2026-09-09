# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Round-2 tests: office_update ``dry_run`` (R4) + archive/restore
``self_check`` readback (R7) + word digest 批注概况 (R3-digest).

Covers:

- R4 schema: ``dry_run`` boolean parameter exposed
- R4 file_path mode: dry_run returns ``{dry_run, changes, truncated}``
  and the source file's bytes/mtime are untouched
- R4 doc_id mode: preview against the managed doc; file bytes and DB
  status (PARSED, not EDITED) untouched
- R4 invalid ops → ``success=False`` with ``preview_failed: …``
- R4 regression: normal path (no dry_run) still applies ops + self_check
- R7: archive → ``self_check.summary.archived_count`` + document
  fingerprint; restore → ``live_count``; readback failure degrades to
  ``{ok: False, error}`` without failing the tool result
- R3-digest: commented docx digest ends with ``comments: N 条`` + up to
  3 ``[author: anchor → text]`` lines; uncommented doc → no line
  (skipped while the parallel ``read_docx`` comments merge hasn't landed)

Fixtures mirror ``test_office_archive_tool.py`` / ``test_office_update_tool.py``.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional
from unittest.mock import patch

import pytest

from backend.data.database import Database
from backend.domain.tool_policy import ToolPolicy
from backend.office.edit import update_document as apply_update
from backend.office.models import (
    OfficeDocStatus,
    OfficeDocType,
    OfficeDocumentMetadata,
    OfficeDocumentSummary,
)
from backend.office.session_workspace import bind_session_workspace
from backend.office.storage import save_document
from backend.tools.context import ToolExecutionContext, reset_tool_context, set_tool_context
from backend.tools.office_archive_tool import OfficeArchiveTool
from backend.tools.office_restore_tool import OfficeRestoreTool
from backend.tools.office_update_tool import OfficeUpdateTool

pytestmark = pytest.mark.unit


# ── Helpers (mirror test_office_archive_tool.py / test_office_update_tool.py) ──


def _update_tool(**policy_kwargs) -> OfficeUpdateTool:
    return OfficeUpdateTool(policy=ToolPolicy(**policy_kwargs))


def _archive_tool() -> OfficeArchiveTool:
    return OfficeArchiveTool(policy=ToolPolicy())


def _restore_tool() -> OfficeRestoreTool:
    return OfficeRestoreTool(policy=ToolPolicy())


def _make_doc(
    *,
    doc_id: str,
    workspace_path: str,
    archived_at: Optional[int] = None,
    doc_type: OfficeDocType = OfficeDocType.WORD,
    original_filename: str = "上传.docx",
) -> OfficeDocumentSummary:
    return OfficeDocumentSummary(
        id=doc_id,
        workspace_path=workspace_path,
        doc_type=doc_type,
        original_filename=original_filename,
        generated_filename=f"{doc_id}.docx",
        status=OfficeDocStatus.PARSED,
        created_at=1_700_000_000_000,
        updated_at=1_700_000_000_000,
        metadata=OfficeDocumentMetadata(file_size_bytes=1024),
        archived_at=archived_at,
    )


def _seed_session(conn, session_id: str) -> None:
    conn.execute(
        "INSERT INTO sessions (id, title, created_at, updated_at) VALUES (?, ?, ?, ?)",
        (session_id, "t", 1, 1),
    )
    conn.commit()


def _ctx(session_id: str, binding_generation: int = 1) -> ToolExecutionContext:
    return ToolExecutionContext(
        session_id=session_id,
        stream_id="stream-x",
        binding_generation=binding_generation,
        office_doc_scope=frozenset(),
    )


def _make_docx(path: Path, *paragraphs: str) -> Path:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _word_fingerprint(path: Path):
    """mtime_ns + size — dry_run 断言「文件一个字节都没动」的最低成本指纹。"""
    stat = path.stat()
    return stat.st_mtime_ns, stat.st_size


# ── R4: schema ────────────────────────────────────────────────────────


def test_schema_exposes_dry_run_boolean():
    props = _update_tool().schema.parameters["properties"]
    assert props["dry_run"]["type"] == "boolean"
    assert "预览" in props["dry_run"]["description"]


# ── R4: file_path 模式 ────────────────────────────────────────────────


def test_dry_run_by_path_returns_changes_without_touching_file(tmp_path: Path):
    target = _make_docx(tmp_path / "doc.docx", "旧文本第一段", "旧文本第二段")
    before = _word_fingerprint(target)

    result = _update_tool().execute(
        file_path=str(target),
        ops=[{"op": "replace_text", "find": "旧文本", "replace": "新文本"}],
        dry_run=True,
    )

    assert result.success is True
    content = result.content
    assert content["dry_run"] is True
    assert content["truncated"] is False
    changes = content["changes"]
    assert isinstance(changes, list)
    assert changes
    assert changes[0]["op"] == "replace_text"
    assert "旧文本" in (changes[0]["before"] or "")
    assert "新文本" in (changes[0]["after"] or "")
    # 源文件零写入
    assert _word_fingerprint(target) == before
    from backend.office.word import read_docx

    assert any("旧文本" in p.text for p in read_docx(target, workspace_path="").paragraphs)


def test_dry_run_by_path_invalid_ops_fail_without_writing(tmp_path: Path):
    """op 本身合法但目标文本不存在 → preview 失败，error 带原因，文件不动。"""
    target = _make_docx(tmp_path / "doc.docx", "真实内容")
    before = _word_fingerprint(target)

    result = _update_tool().execute(
        file_path=str(target),
        ops=[{"op": "replace_text", "find": "不存在的文本", "replace": "x"}],
        dry_run=True,
    )

    assert result.success is False
    assert result.error.startswith("preview_failed")
    assert _word_fingerprint(target) == before


def test_dry_run_by_path_still_enforces_location_guards(tmp_path: Path):
    result = _update_tool().execute(
        file_path="relative.docx",
        ops=[{"op": "replace_text", "find": "a", "replace": "b"}],
        dry_run=True,
    )
    assert result.success is False
    assert result.error.startswith("file_path_absolute_required")

    missing = _update_tool().execute(
        file_path=str(tmp_path / "ghost.docx"),
        ops=[{"op": "replace_text", "find": "a", "replace": "b"}],
        dry_run=True,
    )
    assert missing.success is False
    assert missing.error == "file_not_found"


def test_dry_run_without_target_reports_missing_mode():
    result = _update_tool().execute(
        ops=[{"op": "replace_text", "find": "a", "replace": "b"}], dry_run=True
    )
    assert result.success is False
    assert result.error == "doc_id_or_file_path_required"


# ── R4: doc_id 模式 ───────────────────────────────────────────────────


def _seed_managed_docx(tmp_path: Path, doc_id: str = "doc-a", *paragraphs: str):
    """绑定工作区 + 受管 docx 行；返回 (db, conn, binding, doc_file)。"""
    db = Database(db_path=str(tmp_path / "t.db"))
    db.init_db()
    conn = db.get_connection()
    _seed_session(conn, "sess-1")
    work = tmp_path / "work"
    work.mkdir()
    binding = bind_session_workspace(conn, "sess-1", str(work), now_ms=1)
    doc_dir = work / "office" / "word" / doc_id
    doc_dir.mkdir(parents=True)
    doc_file = _make_docx(doc_dir / f"{doc_id}.docx", *paragraphs)
    save_document(conn, _make_doc(doc_id=doc_id, workspace_path=binding.workspace_path))
    return db, conn, binding, doc_file


def test_dry_run_by_doc_id_previews_managed_doc_without_side_effects(tmp_path: Path):
    db, conn, binding, doc_file = _seed_managed_docx(
        tmp_path, "doc-a", "旧文本第一段", "旧文本第二段"
    )
    before = _word_fingerprint(doc_file)

    with patch("backend.tools.office_update_tool.get_database", return_value=db):
        token = set_tool_context(_ctx("sess-1", binding.generation))
        try:
            result = _update_tool().execute(
                doc_id="doc-a",
                ops=[{"op": "replace_text", "find": "旧文本", "replace": "新文本"}],
                dry_run=True,
            )
        finally:
            reset_tool_context(token)

    assert result.success is True
    assert result.content["dry_run"] is True
    assert result.content["changes"][0]["op"] == "replace_text"
    # 文件字节未动 + DB 状态未翻转（保持 PARSED，不是 EDITED）
    assert _word_fingerprint(doc_file) == before
    row = conn.execute("SELECT status FROM office_documents WHERE id = 'doc-a'").fetchone()
    assert row["status"] == OfficeDocStatus.PARSED.value


def test_dry_run_by_doc_id_without_context_fails_closed():
    result = _update_tool().execute(
        doc_id="doc-a",
        ops=[{"op": "replace_text", "find": "a", "replace": "b"}],
        dry_run=True,
    )
    assert result.success is False
    assert result.error == "missing_tool_context"


def test_dry_run_by_doc_id_unknown_doc_is_not_found(tmp_path: Path):
    db, _conn, binding, _doc_file = _seed_managed_docx(tmp_path, "内容")
    with patch("backend.tools.office_update_tool.get_database", return_value=db):
        token = set_tool_context(_ctx("sess-1", binding.generation))
        try:
            result = _update_tool().execute(
                doc_id="ghost",
                ops=[{"op": "replace_text", "find": "a", "replace": "b"}],
                dry_run=True,
            )
        finally:
            reset_tool_context(token)
    assert result.success is False
    assert result.error == "document_not_found"


def test_dry_run_by_doc_id_archived_doc_is_not_found(tmp_path: Path):
    """dry_run 解析走「隐藏归档行」的 lookup，与正式 update 同语义。"""
    db, _conn, binding, _doc_file = _seed_managed_docx(tmp_path, "内容")
    # seed an archived doc row directly
    conn = db.get_connection()
    save_document(
        conn,
        _make_doc(
            doc_id="doc-arc",
            workspace_path=binding.workspace_path,
            archived_at=1_700_000_000_001,
        ),
    )
    with patch("backend.tools.office_update_tool.get_database", return_value=db):
        token = set_tool_context(_ctx("sess-1", binding.generation))
        try:
            result = _update_tool().execute(
                doc_id="doc-arc",
                ops=[{"op": "replace_text", "find": "a", "replace": "b"}],
                dry_run=True,
            )
        finally:
            reset_tool_context(token)
    assert result.success is False
    assert result.error == "document_not_found"


# ── R4 regression: 正式路径不受影响 ───────────────────────────────────


def test_normal_path_still_applies_and_self_checks(tmp_path: Path):
    """不传 dry_run（默认 False）→ 原有 apply + self_check 行为不变。"""
    target = _make_docx(tmp_path / "doc.docx", "旧文本")
    result = _update_tool().execute(
        file_path=str(target),
        ops=[{"op": "replace_text", "find": "旧文本", "replace": "新文本"}],
    )
    assert result.success is True
    assert "dry_run" not in result.content
    assert result.content["results"][0]["replacements"] == 1
    assert result.content["self_check"]["ok"] is True
    from backend.office.word import read_docx

    assert [p.text for p in read_docx(target, workspace_path="").paragraphs] == ["新文本"]


def test_explicit_dry_run_false_behaves_like_normal_path(tmp_path: Path):
    target = _make_docx(tmp_path / "doc.docx", "旧文本")
    result = _update_tool().execute(
        file_path=str(target),
        ops=[{"op": "replace_text", "find": "旧文本", "replace": "新文本"}],
        dry_run=False,
    )
    assert result.success is True
    assert "dry_run" not in result.content


# ── R7: archive / restore self_check ─────────────────────────────────


def _seed_two_docs(tmp_path: Path):
    """工作区 + 两个受管 docx（doc-a / doc-b）；返回 (db, conn, binding)。"""
    db = Database(db_path=str(tmp_path / "t.db"))
    db.init_db()
    conn = db.get_connection()
    _seed_session(conn, "sess-1")
    work = tmp_path / "work"
    work.mkdir()
    binding = bind_session_workspace(conn, "sess-1", str(work), now_ms=1)
    for doc_id in ("doc-a", "doc-b"):
        doc_dir = work / "office" / "word" / doc_id
        doc_dir.mkdir(parents=True)
        (doc_dir / f"{doc_id}.docx").write_bytes(b"placeholder")
        save_document(
            conn,
            _make_doc(doc_id=doc_id, workspace_path=binding.workspace_path),
        )
    return db, conn, binding


def test_archive_attaches_archived_count_self_check(tmp_path: Path):
    db, conn, binding = _seed_two_docs(tmp_path)

    with patch("backend.tools.office_archive_tool.get_database", return_value=db):
        token = set_tool_context(_ctx("sess-1", binding.generation))
        try:
            result = _archive_tool().execute(doc_id="doc-a")
        finally:
            reset_tool_context(token)

    assert result.success is True
    self_check = result.content["self_check"]
    assert self_check["ok"] is True
    assert self_check["summary"]["archived_count"] == 1
    assert self_check["summary"]["document"]["filename"] == "上传.docx"
    assert self_check["summary"]["document"]["doc_type"] == "word"
    assert len(str(self_check)) < 1024  # keep-tiny 契约


def test_restore_attaches_live_count_self_check(tmp_path: Path):
    db, conn, binding = _seed_two_docs(tmp_path)
    ctx = _ctx("sess-1", binding.generation)

    with patch("backend.tools.office_archive_tool.get_database", return_value=db), patch(
        "backend.tools.office_restore_tool.get_database", return_value=db
    ):
        token = set_tool_context(ctx)
        try:
            archived = _archive_tool().execute(doc_id="doc-a")
            restored = _restore_tool().execute(doc_id="doc-a")
        finally:
            reset_tool_context(token)

    assert archived.success is True
    assert restored.success is True
    self_check = restored.content["self_check"]
    assert self_check["ok"] is True
    # doc-a 已还原 → 工作区两份文档全部存活
    assert self_check["summary"]["live_count"] == 2
    assert self_check["summary"]["document"]["filename"] == "上传.docx"
    assert self_check["summary"]["document"]["doc_type"] == "word"


def test_self_check_readback_failure_degrades_without_failing_tool(tmp_path: Path):
    """计数查询炸掉 → {ok: False, error} 占位，工具主结果仍 success。"""
    db, _conn, binding = _seed_two_docs(tmp_path)

    def _boom(*_args, **_kwargs):
        raise RuntimeError("db gone")

    with patch("backend.tools.office_archive_tool.list_documents", side_effect=_boom), patch(
        "backend.tools.office_archive_tool.get_database", return_value=db
    ):
        token = set_tool_context(_ctx("sess-1", binding.generation))
        try:
            result = _archive_tool().execute(doc_id="doc-a")
        finally:
            reset_tool_context(token)

    assert result.success is True
    assert result.content["document_id"] == "doc-a"
    self_check = result.content["self_check"]
    assert self_check["ok"] is False
    assert "self_check_failed" in self_check["error"]


# ── R3-digest: word digest 批注概况 ──────────────────────────────────


def test_digest_word_appends_comment_overview(tmp_path: Path):
    """edit.add_comment 造出带批注 docx → digest 末尾有 `comments: 1 条`。"""
    from backend.chat.attachment_resolver import _digest_word
    from backend.office.word import read_docx

    target = _make_docx(tmp_path / "reviewed.docx", "This sentence needs review.")
    saved, results = apply_update(
        "word",
        target,
        [
            {
                "op": "add_comment",
                "find": "This sentence needs review.",
                "comment": "Please polish the wording here.",
                "author": "李四",
            }
        ],
    )
    assert saved is True
    assert results[0]["ok"] is True

    # 并行批次（models.py/word.py 归他人）未把 comments 并入 read_docx
    # 时，digest 只会静默省略概况行 —— 本用例此时跳过而非失败。
    parsed = read_docx(target, workspace_path="")
    if not list(getattr(parsed, "comments", None) or []):
        pytest.skip("read_docx does not expose comments field yet (parallel batch)")

    out = _digest_word(str(target), str(tmp_path))
    assert "comments: 1 条" in out
    assert "[李四: This sentence needs review. → Please polish the wording here.]" in out


def test_digest_word_uncommented_doc_has_no_comment_line(tmp_path: Path):
    from backend.chat.attachment_resolver import _digest_word

    target = _make_docx(tmp_path / "plain.docx", "普通段落，没有批注。")
    out = _digest_word(str(target), str(tmp_path))
    assert "comments:" not in out


def test_digest_word_caps_comment_overview_at_three_and_clamps_text(tmp_path: Path):
    """4 条批注 → `comments: 4 条` 但明细只列前 3 条；超长正文截到 ≤60 字。"""
    from backend.chat.attachment_resolver import _digest_word
    from backend.office.word import read_docx

    target = _make_docx(tmp_path / "many.docx", "Anchor paragraph body text.")
    ops = [
        {
            "op": "add_comment",
            "find": "Anchor paragraph body text.",
            "comment": f"批注{chr(0x4E00 + i)}" * 30,  # 90 字，远超 60
            "author": f"作者{i}",
        }
        for i in range(4)
    ]
    saved, results = apply_update("word", target, ops)
    assert saved is True
    assert all(r["ok"] for r in results)

    parsed = read_docx(target, workspace_path="")
    if len(list(getattr(parsed, "comments", None) or [])) < 4:
        pytest.skip("read_docx does not expose comments field yet (parallel batch)")

    out = _digest_word(str(target), str(tmp_path))
    assert "comments: 4 条" in out
    detail_lines = [ln for ln in out.splitlines() if ln.startswith("[作者")]
    assert len(detail_lines) == 3  # 只列前 3 条
    clamped = detail_lines[0].rsplit("→", 1)[1].strip().rstrip("]")
    assert len(clamped) == 60  # ≤60 字（含结尾 …）
    assert clamped.endswith("…")

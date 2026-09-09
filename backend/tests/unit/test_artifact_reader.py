# backend/tests/unit/test_artifact_reader.py
from unittest.mock import patch

from backend.data import artifact_reader, artifact_repo


def test_read_text_markdown(tmp_path):
    f = tmp_path / "doc.md"
    f.write_text("# Hello\n\nWorld", encoding="utf-8")
    aid = artifact_repo.record_artifact("sess_001", str(f), "doc.md", "markdown", 14)

    result = artifact_reader.read_text(aid)

    assert result["ok"] is True
    assert result["kind"] == "markdown"
    assert result["content"] == "# Hello\n\nWorld"
    assert result["truncated"] is False


def test_read_text_truncates_long_content(tmp_path):
    f = tmp_path / "big.md"
    f.write_text("x" * 600_000, encoding="utf-8")
    aid = artifact_repo.record_artifact("sess_001", str(f), "big.md", "markdown", 600_000)

    result = artifact_reader.read_text(aid)

    assert result["ok"] is True
    assert result["truncated"] is True
    assert len(result["content"]) <= 500_000


def test_read_text_missing_file(tmp_path):
    aid = artifact_repo.record_artifact("sess_001", str(tmp_path / "gone.md"), "gone.md", "markdown", 0)
    result = artifact_reader.read_text(aid)
    assert result["ok"] is False
    assert "not found" in result["error"].lower()


def test_read_text_missing_artifact():
    result = artifact_reader.read_text("nonexistent")
    assert result["ok"] is False


def test_read_image_returns_data_url(tmp_path):
    png_bytes = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d49444154789c6300010000000500010d0a2db40000000049454e44ae426082"
    )
    f = tmp_path / "pixel.png"
    f.write_bytes(png_bytes)
    aid = artifact_repo.record_artifact("sess_001", str(f), "pixel.png", "image", len(png_bytes))

    result = artifact_reader.read_image(aid)

    assert result["ok"] is True
    assert result["kind"] == "image"
    assert result["data_url"].startswith("data:image/png;base64,")


def test_reveal_in_file_manager(tmp_path):
    f = tmp_path / "doc.md"
    f.write_text("test", encoding="utf-8")
    aid = artifact_repo.record_artifact("sess_001", str(f), "doc.md", "markdown", 4)

    with patch("subprocess.run", return_value=None) as mock_run:
        result = artifact_reader.reveal_in_file_manager(aid)

    assert result["ok"] is True
    mock_run.assert_called_once()


def test_read_pdf_returns_data_url(tmp_path):
    # F11 (round4): PDF 走 base64 data URL,前端 iframe 内嵌渲染
    pdf_bytes = b"%PDF-1.4 fake body"
    f = tmp_path / "out.pdf"
    f.write_bytes(pdf_bytes)
    aid = artifact_repo.record_artifact("sess_001", str(f), "out.pdf", "pdf", len(pdf_bytes))

    result = artifact_reader.read_pdf(aid)

    assert result["ok"] is True
    assert result["kind"] == "pdf"
    assert result["data_url"].startswith("data:application/pdf;base64,")


def test_read_pdf_oversize_rejected(tmp_path, monkeypatch):
    pdf_bytes = b"%PDF-1.4 big"
    f = tmp_path / "big.pdf"
    f.write_bytes(pdf_bytes)
    aid = artifact_repo.record_artifact("sess_001", str(f), "big.pdf", "pdf", len(pdf_bytes))

    monkeypatch.setattr(artifact_reader, "MAX_PDF_BYTES", 4)
    result = artifact_reader.read_pdf(aid)

    assert result["ok"] is False
    assert "20MB" in result["error"]


def test_read_office_docx_roundtrip(tmp_path):
    """C-2 (round5 批次 C): docx → 全转义 HTML 预览。"""
    from docx import Document

    f = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("季度报告", level=1)
    doc.add_paragraph("正文 <script>alert(1)</script> 内容")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "100"
    doc.save(str(f))

    aid = artifact_repo.record_artifact("sess_001", str(f), "report.docx", "docx", f.stat().st_size)
    result = artifact_reader.read_office(aid, kind="docx")
    assert result["ok"] is True
    assert result["kind"] == "docx"
    assert "<h2>季度报告</h2>" in result["html"]
    assert "<td>营收</td>" in result["html"]
    # 用户数据必须被转义（script 标签不得原样出现）
    assert "<script>" not in result["html"]
    assert "&lt;script&gt;" in result["html"]


def test_read_office_xlsx_roundtrip_with_row_cap(tmp_path, monkeypatch):
    """C-2: xlsx → 分 sheet HTML 表格,行数按 MAX_SHEET_PREVIEW_ROWS 截断。"""
    from openpyxl import Workbook

    f = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["月份", "金额"])
    for i in range(1, 260):
        ws.append([f"m{i}", i])
    wb.save(str(f))

    aid = artifact_repo.record_artifact("sess_001", str(f), "data.xlsx", "xlsx", f.stat().st_size)
    result = artifact_reader.read_office(aid, kind="xlsx")
    assert result["ok"] is True
    assert "<h3>销售</h3>" in result["html"]
    assert "预览已截断" in result["html"]
    assert result["html"].count("<tr>") == artifact_reader.MAX_SHEET_PREVIEW_ROWS


def test_read_office_pptx_roundtrip(tmp_path):
    """C-2: pptx → slide 标题 + 文本块大纲。"""
    from pptx import Presentation

    f = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "项目计划"
    slide.placeholders[1].text = "第一要点"
    prs.save(str(f))

    aid = artifact_repo.record_artifact("sess_001", str(f), "deck.pptx", "pptx", f.stat().st_size)
    result = artifact_reader.read_office(aid, kind="pptx")
    assert result["ok"] is True
    assert "项目计划" in result["html"]
    assert "第一要点" in result["html"]


def test_read_office_oversize_rejected(tmp_path, monkeypatch):
    """C-2: 超过 20MB 上限拒绝预览。"""
    from docx import Document

    f = tmp_path / "big.docx"
    Document().save(str(f))
    monkeypatch.setattr(artifact_reader, "MAX_OFFICE_BYTES", 10)
    aid = artifact_repo.record_artifact("sess_001", str(f), "big.docx", "docx", f.stat().st_size)
    result = artifact_reader.read_office(aid, kind="docx")
    assert result["ok"] is False
    assert "20MB" in result["error"]


def test_read_office_invalid_file_degrades(tmp_path):
    """C-2: 非 office 内容解析失败 → ok=False 不抛异常。"""
    f = tmp_path / "fake.docx"
    f.write_bytes(b"not a zip archive")
    aid = artifact_repo.record_artifact("sess_001", str(f), "fake.docx", "docx", f.stat().st_size)
    result = artifact_reader.read_office(aid, kind="docx")
    assert result["ok"] is False
    assert "解析失败" in result["error"]

"""Unit tests for Word comment support (批次 3.3, plan §3.3).

python-docx 1.1.2 has no comments API, so both the write side
(edit.add_comment / delete_comment) and the read side
(word.read_docx_comments) work at the OOXML level. These tests verify the
full round-trip through the on-disk zip:

- add_comment creates word/comments.xml (content-type + relationship),
  inserts commentRangeStart/End + commentReference in document.xml, and
  appends the w:comment entry
- read_docx_comments extracts {id, author, date, text, anchor_text}
- anchoring works inside table cells
- delete_comment removes the entry and the document markers
- per-op failure semantics (missing find / unknown comment_id)
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document

from backend.office.edit import update_docx
from backend.office.errors import OfficeFileNotFoundError, OfficeParseError
from backend.office.word import read_docx, read_docx_comments

pytestmark = pytest.mark.unit


# ──────────────────────────────────────────────────────────────────────
# Fixtures / helpers
# ──────────────────────────────────────────────────────────────────────


def _make_docx(path: Path) -> Path:
    """DOCX with body paragraphs (one matches the find needles below)."""
    doc = Document()
    doc.add_heading("评审文档", level=0)
    doc.add_paragraph("This is some text to review.")
    doc.add_paragraph("第二段没有要批注的内容")
    doc.save(str(path))
    return path


def _zip_text(path: Path, member: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(member).decode("utf-8")


# ──────────────────────────────────────────────────────────────────────
# add_comment round-trip
# ──────────────────────────────────────────────────────────────────────


def test_add_comment_round_trip(tmp_path: Path):
    """add_comment → comments.xml part + markers + readable comment."""
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(
        path,
        [
            {
                "op": "add_comment",
                "find": "some text",
                "comment": "请检查这句话",
                "author": "张三",
                "date": "2026-09-09T08:00:00Z",
            }
        ],
    )
    assert saved
    assert results[0]["ok"]
    assert results[0]["comment_id"] == "0"

    # Reopen the saved zip: comments part + document markers all present.
    with zipfile.ZipFile(path) as z:
        assert "word/comments.xml" in z.namelist()
    comments_xml = _zip_text(path, "word/comments.xml")
    assert "请检查这句话" in comments_xml
    assert "张三" in comments_xml
    document_xml = _zip_text(path, "word/document.xml")
    assert "commentRangeStart" in document_xml
    assert "commentRangeEnd" in document_xml
    assert "commentReference" in document_xml

    parsed = read_docx_comments(path)
    assert len(parsed.comments) == 1
    comment = parsed.comments[0]
    assert comment.id == "0"
    assert comment.author == "张三"
    assert comment.date == "2026-09-09T08:00:00Z"
    assert comment.text == "请检查这句话"
    # 整段锚定：锚文本包含被 find 命中的正文
    assert "some text" in comment.anchor_text

    # 既有 read_docx 结果不受批注 part 影响
    doc_result = read_docx(path, workspace_path="")
    assert "This is some text to review." in [p.text for p in doc_result.paragraphs]


def test_add_comment_default_author_and_date(tmp_path: Path):
    """author/date 缺省：author=Sage，date 为当前 UTC 的 ISO 时间；id 递增。"""
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(
        path,
        [
            {"op": "add_comment", "find": "some text", "comment": "第一条"},
            {"op": "add_comment", "find": "第二段", "comment": "第二条"},
        ],
    )
    assert saved
    assert all(r["ok"] for r in results)
    assert [r["comment_id"] for r in results] == ["0", "1"]

    parsed = read_docx_comments(path)
    assert [c.id for c in parsed.comments] == ["0", "1"]
    first = parsed.comments[0]
    assert first.author == "Sage"
    # ISO 8601 UTC（'Z' 结尾，Word 期望的 w:date 形态）
    assert first.date is not None
    assert first.date.endswith("Z")
    assert first.text == "第一条"
    assert "第二段" in parsed.comments[1].anchor_text


def test_add_comment_missing_find_fails_per_op(tmp_path: Path):
    """find 无命中：该 op 失败、不保存、原文件未被触碰。"""
    path = _make_docx(tmp_path / "a.docx")
    before = path.stat().st_mtime_ns
    saved, results = update_docx(
        path, [{"op": "add_comment", "find": "不存在的文本", "comment": "批注"}]
    )
    assert not saved
    assert not results[0]["ok"]
    assert "text_not_found" in results[0]["error"]
    assert path.stat().st_mtime_ns == before
    # 文件未被改动：仍然没有 comments part
    with zipfile.ZipFile(path) as z:
        assert "word/comments.xml" not in z.namelist()


def test_add_comment_requires_find_and_comment(tmp_path: Path):
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(path, [{"op": "add_comment", "find": "x"}])
    assert not saved
    assert "missing_field: comment" in results[0]["error"]
    saved, results = update_docx(path, [{"op": "add_comment", "comment": "y"}])
    assert not saved
    assert "missing_field: find" in results[0]["error"]


def test_add_comment_in_table_cell(tmp_path: Path):
    """批注锚定表格单元格内的文本（range 落在 cell 的段落里）。"""
    path = _make_docx(tmp_path / "a.docx")
    doc = Document(str(path))
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "表格中的分数数据"
    doc.save(str(path))

    saved, results = update_docx(
        path, [{"op": "add_comment", "find": "分数数据", "comment": "核对数值", "author": "李四"}]
    )
    assert saved
    assert results[0]["ok"]

    parsed = read_docx_comments(path)
    assert len(parsed.comments) == 1
    assert parsed.comments[0].author == "李四"
    assert "分数数据" in parsed.comments[0].anchor_text
    # 范围标记位于 document.xml 的表格分支内（仍在 body 下，全局计数为 1 对）
    document_xml = _zip_text(path, "word/document.xml")
    assert document_xml.count('commentRangeStart w:id="0"') == 1
    assert document_xml.count('commentRangeEnd w:id="0"') == 1


# ──────────────────────────────────────────────────────────────────────
# delete_comment
# ──────────────────────────────────────────────────────────────────────


def test_delete_comment_removes_entry_and_ranges(tmp_path: Path):
    path = _make_docx(tmp_path / "a.docx")
    update_docx(path, [{"op": "add_comment", "find": "some text", "comment": "待删除"}])

    saved, results = update_docx(path, [{"op": "delete_comment", "comment_id": "0"}])
    assert saved
    assert results[0]["ok"]
    assert results[0]["comment_id"] == "0"

    document_xml = _zip_text(path, "word/document.xml")
    assert "commentRangeStart" not in document_xml
    assert "commentRangeEnd" not in document_xml
    assert "commentReference" not in document_xml
    # comments.xml part 保留（可为空），但条目已移除
    assert "<w:comment " not in _zip_text(path, "word/comments.xml")
    assert read_docx_comments(path).comments == []


def test_delete_comment_missing_id_fails_per_op(tmp_path: Path):
    path = _make_docx(tmp_path / "a.docx")
    before = path.stat().st_mtime_ns
    saved, results = update_docx(path, [{"op": "delete_comment", "comment_id": "99"}])
    assert not saved
    assert not results[0]["ok"]
    assert "comment_not_found" in results[0]["error"]
    assert path.stat().st_mtime_ns == before


def test_delete_comment_keeps_other_comments(tmp_path: Path):
    """多批注场景：删 0 留 1，剩余批注 id/文本完好。"""
    path = _make_docx(tmp_path / "a.docx")
    update_docx(
        path,
        [
            {"op": "add_comment", "find": "some text", "comment": "第一条"},
            {"op": "add_comment", "find": "第二段", "comment": "第二条"},
        ],
    )
    saved, results = update_docx(path, [{"op": "delete_comment", "comment_id": "0"}])
    assert saved
    assert results[0]["ok"]

    parsed = read_docx_comments(path)
    assert [c.id for c in parsed.comments] == ["1"]
    assert parsed.comments[0].text == "第二条"
    assert "第二段" in parsed.comments[0].anchor_text


# ──────────────────────────────────────────────────────────────────────
# read_docx_comments — file-level behavior
# ──────────────────────────────────────────────────────────────────────


def test_read_docx_comments_without_comments_part(tmp_path: Path):
    """无批注的文档：返回空列表（而不是报错）。"""
    path = _make_docx(tmp_path / "a.docx")
    result = read_docx_comments(path)
    assert result.comments == []


def test_read_docx_comments_missing_file_raises(tmp_path: Path):
    with pytest.raises(OfficeFileNotFoundError):
        read_docx_comments(tmp_path / "missing.docx")


def test_read_docx_comments_corrupt_file_raises(tmp_path: Path):
    garbage = tmp_path / "garbage.docx"
    garbage.write_bytes(b"definitely not a docx file")
    with pytest.raises(OfficeParseError):
        read_docx_comments(garbage)

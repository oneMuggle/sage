"""Unit tests for native tracked revisions (M4b)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.errors import OfficeFileNotFoundError, OfficeParseError
from backend.office.word_revision_writer import apply_word_revisions, build_revision_plan
from backend.office.word_rulepack_models import (
    WordAuditIssue,
    WordRulePack,
    WordRuleSpec,
)


def _doc(path: Path, texts: list[str]) -> Path:
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def _rule(
    rule_id: str,
    kind: str,
    params: dict = None,
    *,
    fixable: bool = True,
) -> WordRuleSpec:
    return WordRuleSpec(
        id=rule_id,
        title=rule_id,
        category="content",
        severity="warning",
        kind=kind,
        params=params or {},
        fixable=fixable,
    )


def _issue(
    rule_id: str,
    *,
    expected: str,
    actual: str,
    fixable: bool = True,
) -> WordAuditIssue:
    return WordAuditIssue(
        rule_id=rule_id,
        rule_pack_id="pack",
        category="content",
        severity="warning",
        message=f"{actual} -> {expected}",
        expected=expected,
        actual=actual,
        fixable=fixable,
    )


def test_tracked_replace_writes_del_ins_and_comment(tmp_path: Path) -> None:
    path = _doc(tmp_path / "doc.docx", ["我们采用 AI 方案。"])

    result = apply_word_revisions(
        path,
        [{"find": "AI", "replace": "人工智能", "comment": "统一术语"}],
        "revised.docx",
        workspace_path=str(tmp_path),
        rule_pack_id="pack",
    )

    assert result.ok is True
    assert result.applied == 1
    assert result.skipped == 0
    assert result.changes[0].applied is True
    fixed = Document(str(tmp_path / "revised.docx"))
    paragraph = fixed.paragraphs[0]
    # python-docx paragraph.text excludes both w:delText and w:ins content;
    # the tracked change must be asserted at the OOXML level.
    assert "AI" not in paragraph.text
    assert "人工智能" not in paragraph.text

    deleted = paragraph._p.findall(qn("w:del"))
    inserted = paragraph._p.findall(qn("w:ins"))
    assert len(deleted) == 1
    assert len(inserted) == 1
    assert deleted[0].get(qn("w:author")) == "Sage"
    assert deleted[0].find(".//" + qn("w:delText")).text == "AI"
    assert inserted[0].find(".//" + qn("w:t")).text == "人工智能"

    comments = [
        part
        for part in fixed.part.package.iter_parts()
        if str(part.partname) == "/word/comments.xml"
    ]
    assert comments
    assert "统一术语" in comments[0].blob.decode("utf-8")


def test_replace_across_multiple_runs(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("我们采用 AI")
    paragraph.add_run(" 方案。")
    path = tmp_path / "multi.docx"
    document.save(str(path))

    result = apply_word_revisions(
        path,
        [{"find": "AI 方案", "replace": "人工智能方案"}],
        "multi-revised.docx",
        workspace_path=str(tmp_path),
        add_comments=False,
    )

    assert result.applied == 1
    fixed = Document(str(tmp_path / "multi-revised.docx"))
    inserted = fixed.paragraphs[0]._p.findall(qn("w:ins"))
    assert len(inserted) == 1
    assert inserted[0].find(".//" + qn("w:t")).text == "人工智能方案"
    deleted_texts = [
        element.text for element in fixed.paragraphs[0]._p.iter(qn("w:delText"))
    ]
    assert "AI" in deleted_texts
    assert " 方案" in deleted_texts


def test_missing_text_is_skipped(tmp_path: Path) -> None:
    path = _doc(tmp_path / "doc.docx", ["正文"])

    result = apply_word_revisions(
        path,
        [{"find": "不存在", "replace": "x"}],
        "revised.docx",
        workspace_path=str(tmp_path),
    )

    assert result.ok is False
    assert result.applied == 0
    assert result.skipped == 1
    assert result.changes[0].reason == "未找到目标文本"
    assert (tmp_path / "revised.docx").is_file()


def test_duplicate_changes_are_skipped(tmp_path: Path) -> None:
    path = _doc(tmp_path / "doc.docx", ["AI 方案"])
    change = {"find": "AI", "replace": "人工智能"}

    result = apply_word_revisions(
        path,
        [change, dict(change)],
        "revised.docx",
        workspace_path=str(tmp_path),
        add_comments=False,
    )

    assert result.applied == 1
    assert result.skipped == 1
    assert result.changes[1].reason == "重复修订"


def test_overwrite_is_safe_and_second_pass_skips(tmp_path: Path) -> None:
    path = _doc(tmp_path / "doc.docx", ["AI 方案"])
    change = {"find": "AI", "replace": "人工智能"}

    first = apply_word_revisions(
        path,
        [change],
        overwrite=True,
        workspace_path=str(tmp_path),
        add_comments=False,
    )
    assert first.overwritten is True
    assert first.applied == 1

    second = apply_word_revisions(
        path,
        [change],
        overwrite=True,
        workspace_path=str(tmp_path),
        add_comments=False,
    )
    assert second.applied == 0
    assert second.overwritten is False
    assert second.skipped == 1


def test_build_revision_plan_from_issues() -> None:
    pack = WordRulePack(
        id="pack",
        name="Pack",
        rules=[
            _rule("content/terminology", "content.terminology", fixable=True),
            _rule("content/date_format", "content.date_format", fixable=True),
            _rule("content/placeholder", "content.placeholder_residue", fixable=False),
        ],
    )
    issues = [
        _issue("content/terminology", expected="人工智能", actual="AI, ML"),
        _issue("content/date_format", expected="YYYY年M月D日", actual="2026-10-31"),
        _issue("content/placeholder", expected="x", actual="{{name}}"),
    ]

    plan = build_revision_plan(issues, pack)

    assert {entry["find"] for entry in plan} == {"AI", "ML", "2026-10-31"}
    date_entry = next(entry for entry in plan if entry["find"] == "2026-10-31")
    assert date_entry["replace"] == "2026年10月31日"
    assert all(entry["rule_id"] != "content/placeholder" for entry in plan)


def test_output_safety_and_missing_file(tmp_path: Path) -> None:
    path = _doc(tmp_path / "doc.docx", ["AI"])
    change = {"find": "AI", "replace": "人工智能"}

    with pytest.raises(OfficeParseError):
        apply_word_revisions(
            path, [change], "sub/revised.docx", workspace_path=str(tmp_path)
        )
    with pytest.raises(OfficeParseError):
        apply_word_revisions(
            path, [change], "revised.txt", workspace_path=str(tmp_path)
        )
    with pytest.raises(OfficeParseError):
        apply_word_revisions(
            path,
            [change],
            "revised.docx",
            overwrite=True,
            workspace_path=str(tmp_path),
        )
    existing = tmp_path / "exists.docx"
    existing.write_bytes(b"")
    with pytest.raises(OfficeParseError):
        apply_word_revisions(
            path, [change], "exists.docx", workspace_path=str(tmp_path)
        )
    with pytest.raises(OfficeFileNotFoundError):
        apply_word_revisions(
            tmp_path / "missing.docx",
            [change],
            workspace_path=str(tmp_path),
        )

"""Unit tests for deterministic Word fixes (M2b)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.errors import OfficeFileNotFoundError, OfficeParseError
from backend.office.word_fix import fix_word_document, supported_fix_kinds
from backend.office.word_rulepack_builtin import builtin_rulepacks
from backend.office.word_rulepack_models import WordRulePack, WordRuleSpec
from backend.office.word_rulepack_store import load_rulepack


def _rule(
    rule_id: str,
    kind: str,
    params: dict = None,
    *,
    category: str = "content",
    severity: str = "error",
    fixable: bool = True,
) -> WordRuleSpec:
    return WordRuleSpec(
        id=rule_id,
        title=rule_id,
        category=category,
        severity=severity,
        kind=kind,
        params=params or {},
        fixable=fixable,
    )


def _pack(*rules: WordRuleSpec, pack_id: str = "fix-pack") -> WordRulePack:
    return WordRulePack(id=pack_id, name="Fix pack", rules=list(rules))


def _builtin(pack_id: str) -> WordRulePack:
    return load_rulepack(None, pack_id)["pack"]


def _default_doc(path: Path) -> Path:
    document = Document()
    document.add_paragraph("正文")
    document.save(str(path))
    return path


def test_supported_fix_kinds_cover_builtin_fixable_rules() -> None:
    kinds = set(supported_fix_kinds())
    assert {"page.size", "page.margins", "style.font", "content.empty_paragraphs"} <= kinds
    for pack in builtin_rulepacks():
        for rule in pack.rules:
            if rule.fixable:
                assert rule.kind in kinds


def test_fix_gb_pack_outputs_compliant_copy(tmp_path: Path) -> None:
    path = _default_doc(tmp_path / "doc.docx")

    result = fix_word_document(
        path,
        _builtin("gb-t-9704-general"),
        "fixed.docx",
        workspace_path=str(tmp_path),
    )

    assert result.ok is True
    assert result.fixed_count > 0
    assert result.overwritten is False
    assert result.output_path == "fixed.docx"
    assert result.remaining.issues == []
    assert {action.rule_id for action in result.actions} >= {
        "page/size",
        "page/margins",
        "page/page_number",
        "style/body_font",
    }

    fixed = Document(str(tmp_path / "fixed.docx"))
    section = fixed.sections[0]
    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297
    assert round(section.top_margin.cm, 1) == 3.7
    assert round(section.bottom_margin.cm, 1) == 3.5
    assert round(section.left_margin.cm, 1) == 2.8
    assert round(section.right_margin.cm, 1) == 2.6
    rfonts = fixed.styles["Normal"].element.rPr.rFonts
    assert rfonts.get(qn("w:eastAsia")) == "仿宋_GB2312"
    assert fixed.styles["Normal"].font.size.pt == pytest.approx(16.0)
    assert "PAGE" in fixed.sections[0].footer._element.xml

    original = Document(str(path))
    assert round(original.sections[0].page_width.mm) == 216


def test_fix_overwrite_is_idempotent(tmp_path: Path) -> None:
    path = _default_doc(tmp_path / "doc.docx")
    pack = _builtin("gb-t-9704-general")

    first = fix_word_document(path, pack, overwrite=True, workspace_path=str(tmp_path))
    assert first.overwritten is True
    assert first.fixed_count > 0
    assert first.remaining.ok is True

    second = fix_word_document(path, pack, overwrite=True, workspace_path=str(tmp_path))
    assert second.fixed_count == 0
    assert second.overwritten is False
    assert second.remaining.ok is True
    assert second.actions == []


def test_fix_empty_paragraphs_removes_excess(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("A")
    for _ in range(4):
        document.add_paragraph("")
    document.add_paragraph("B")
    path = tmp_path / "empty.docx"
    document.save(str(path))

    pack = _pack(
        _rule(
            "content/empty_paragraphs",
            "content.empty_paragraphs",
            {"max_consecutive": 2},
            severity="warning",
        )
    )
    result = fix_word_document(path, pack, "empty-fixed.docx", workspace_path=str(tmp_path))

    assert result.fixed_count == 2
    assert result.remaining.ok is True
    assert not any(
        issue.rule_id == "content/empty_paragraphs" for issue in result.remaining.issues
    )
    fixed = Document(str(tmp_path / "empty-fixed.docx"))
    empty_count = sum(1 for paragraph in fixed.paragraphs if not paragraph.text.strip())
    assert empty_count == 2


def test_fix_style_paragraph_updates_style(tmp_path: Path) -> None:
    path = _default_doc(tmp_path / "style.docx")
    pack = _pack(
        _rule(
            "style/paragraph",
            "style.paragraph",
            {
                "style": "Normal",
                "line_spacing": 2.0,
                "first_line_indent_cm": 0.5,
                "alignment": "center",
            },
            category="style",
        )
    )

    result = fix_word_document(path, pack, "style-fixed.docx", workspace_path=str(tmp_path))

    assert result.fixed_count == 3
    assert result.remaining.ok is True
    fixed = Document(str(tmp_path / "style-fixed.docx"))
    fmt = fixed.styles["Normal"].paragraph_format
    assert fmt.line_spacing == pytest.approx(2.0)
    assert fmt.first_line_indent.cm == pytest.approx(0.5, abs=0.01)
    assert fmt.alignment.name.lower() == "center"


def test_fix_leaves_unfixable_issues_in_remaining(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("残留 {{name}}")
    path = tmp_path / "residue.docx"
    document.save(str(path))
    pack = _pack(
        _rule("content/placeholder_residue", "content.placeholder_residue", fixable=False)
    )

    result = fix_word_document(path, pack, "residue-fixed.docx", workspace_path=str(tmp_path))

    assert result.fixed_count == 0
    assert result.actions == []
    assert result.ok is False
    assert any(
        issue.rule_id == "content/placeholder_residue" for issue in result.remaining.issues
    )
    assert (tmp_path / "residue-fixed.docx").is_file()


def test_fix_unsupported_fixable_kind_warns(tmp_path: Path) -> None:
    path = _default_doc(tmp_path / "unsupported.docx")
    pack = _pack(
        _rule("content/placeholder_residue", "content.placeholder_residue", fixable=True)
    )

    result = fix_word_document(
        path, pack, "unsupported-fixed.docx", workspace_path=str(tmp_path)
    )

    assert result.fixed_count == 0
    assert any(
        warning["code"] == "fix_unsupported_kind" for warning in result.warnings
    )


def test_fix_output_safety(tmp_path: Path) -> None:
    path = _default_doc(tmp_path / "doc.docx")
    pack = _builtin("gb-t-9704-general")

    with pytest.raises(OfficeParseError):
        fix_word_document(path, pack, "sub/fixed.docx", workspace_path=str(tmp_path))
    with pytest.raises(OfficeParseError):
        fix_word_document(path, pack, "fixed.txt", workspace_path=str(tmp_path))
    with pytest.raises(OfficeParseError):
        fix_word_document(path, pack, "doc.docx", workspace_path=str(tmp_path))
    with pytest.raises(OfficeParseError):
        fix_word_document(
            path, pack, "fixed.docx", overwrite=True, workspace_path=str(tmp_path)
        )

    existing = tmp_path / "exists.docx"
    existing.write_bytes(b"")
    with pytest.raises(OfficeParseError):
        fix_word_document(path, pack, "exists.docx", workspace_path=str(tmp_path))


def test_fix_missing_file_rejected(tmp_path: Path) -> None:
    with pytest.raises(OfficeFileNotFoundError):
        fix_word_document(
            tmp_path / "missing.docx",
            _builtin("gb-t-9704-general"),
            workspace_path=str(tmp_path),
        )

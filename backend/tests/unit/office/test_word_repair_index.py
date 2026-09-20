"""Unit tests for Round 48 — repair index 域插入 + SEQ 题注重排兼容。"""

from __future__ import annotations

import base64

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import (
    OfficeWordGenerateRequest,
    WordBodyStyleSpec,
    WordFormatSpec,
    WordImageSpec,
    WordIndexSpec,
    WordParagraphSpec,
)

# captions 规则在"格式化生成语境"（spec 带 page/body/headings）下启用
_LINT_SPEC = WordFormatSpec(body=WordBodyStyleSpec(font_size_pt=12))
from backend.office.word import generate_docx
from backend.office.word_lint import lint_docx
from backend.office.word_repair import repair_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen_with_captions(tmp_path, name, captions):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    images = [
        WordImageSpec(source=str(image_file), caption=caption, after_paragraph=i)
        for i, caption in enumerate(captions)
    ]
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文一段")],
        images=images,
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_repairs_missing_figure_index_field(tmp_path) -> None:
    """spec 声明 figure_index 但文档缺失 → repair 插入 TOF 域。"""
    path = _gen_with_captions(tmp_path, "no-idx.docx", ["架构图"])
    spec = WordFormatSpec(figure_index=WordIndexSpec())
    before = lint_docx(path, spec)
    assert any(i.rule_id == "figure_index/presence" for i in before.issues)

    result = repair_docx(path, spec)
    assert "figure_index/presence" in result.repaired_rules
    assert result.remaining.ok is True

    doc = Document(str(result.output_path))
    instrs = [
        el.text
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:instrText"))
        if el.text and "TOC" in el.text
    ]
    assert any('\\c "图"' in t for t in instrs)
    # 条目来自文档自身 SEQ 题注
    assert any("图1　架构图" in p.text for p in doc.paragraphs)


def test_renumber_captions_preserves_seq_and_bookmark(tmp_path) -> None:
    """SEQ 题注重排只改缓存编号——域与书签结构原样保留。"""
    path = _gen_with_captions(tmp_path, "seq.docx", ["架构图", "部署图"])

    # 人为破坏第一处编号缓存（9 → 触发 caption/sequence）
    doc = Document(str(path))
    caption_paras = [
        p for p in doc.paragraphs if any(
            el.text and "SEQ" in el.text
            for el in p._p.findall(".//" + qn("w:instrText"))
        )
    ]
    assert len(caption_paras) == 2
    first = caption_paras[0]
    for run in first.runs:
        if run.text.strip().isdigit():
            run.text = "9"
            break
    broken = tmp_path / "broken.docx"
    doc.save(str(broken))

    before = lint_docx(broken, _LINT_SPEC)
    assert any(i.rule_id == "caption/sequence" for i in before.issues)

    result = repair_docx(broken, _LINT_SPEC)
    assert "caption/sequence" in result.repaired_rules
    assert result.remaining.ok is True

    # SEQ 域与书签仍在，编号已重排
    fixed = Document(str(result.output_path))
    seq_instrs = [
        el.text
        for p in fixed.paragraphs
        for el in p._p.findall(".//" + qn("w:instrText"))
        if el.text and "SEQ" in el.text
    ]
    assert len(seq_instrs) == 2
    bookmarks = [
        bm.get(qn("w:name"))
        for bm in fixed.element.body.findall(".//" + qn("w:bookmarkStart"))
    ]
    assert "_RefFig1" in bookmarks
    assert "_RefFig2" in bookmarks
    texts = "\n".join(p.text for p in fixed.paragraphs)
    assert "图1　架构图" in texts
    assert "图2　部署图" in texts


def test_duplicate_caption_lint_warning(tmp_path) -> None:
    """重复题注文本 → caption/duplicate warning（不阻断）。"""
    path = _gen_with_captions(tmp_path, "dup.docx", ["架构图", "架构图"])
    result = lint_docx(path, _LINT_SPEC)
    dup = [i for i in result.issues if i.rule_id == "caption/duplicate"]
    assert len(dup) == 1
    assert dup[0].severity == "warning"

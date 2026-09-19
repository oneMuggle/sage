"""Unit tests for Round 60 — residue 补 fn/en + append_paragraphs 交叉引用。"""

from __future__ import annotations

import base64

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.edit import update_document
from backend.office.models import (
    OfficeWordGenerateRequest,
    WordImageSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx
from backend.office.word_lint import lint_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen_with_caption(tmp_path, name):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文一段")],
        images=[WordImageSpec(source=str(image_file), caption="架构图")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]], caption="汇总表")],
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_residue_lint_detects_fn_and_en(tmp_path) -> None:
    """{{fn:}}/{{en:}} 残渍同样触发 cross_ref/residue。"""
    doc = Document()
    doc.add_paragraph("死文本 {{fn:某脚注}} 与 {{en:某尾注}}")
    path = tmp_path / "manual.docx"
    doc.save(str(path))
    from backend.office.models import WordFormatSpec

    result = lint_docx(path, WordFormatSpec())
    rule_ids = {i.rule_id for i in result.issues}
    assert "cross_ref/residue" in rule_ids


def test_append_paragraphs_resolves_cross_ref(tmp_path) -> None:
    """追加段落 {{fig:}}/{{tbl:}} → REF 域（复用生成期书签）。"""
    path = _gen_with_caption(tmp_path, "r.docx")
    saved, results = update_document(
        "word",
        path,
        [{"op": "append_paragraphs", "paragraphs": [
            {"text": "架构见{{fig:架构图}}，数据见{{tbl:汇总表}}。"},
        ]}],
    )
    assert saved is True

    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    last = doc.paragraphs[-1]
    instrs = [
        el.text
        for el in last._p.findall(".//" + qn("w:instrText"))
        if el.text and "REF" in el.text
    ]
    assert " REF _RefFig1 \\h " in instrs
    assert " REF _RefTbl1 \\h " in instrs
    assert last.text == "架构见图1，数据见表1。"


def test_append_paragraphs_unknown_caption_all_or_nothing(tmp_path) -> None:
    """未知题注 → op 失败且不写入（all-or-nothing）。"""
    path = _gen_with_caption(tmp_path, "r2.docx")
    saved, results = update_document(
        "word",
        path,
        [{"op": "append_paragraphs", "paragraphs": [
            {"text": "第一段正常"},
            {"text": "见{{fig:不存在的图}}"},
        ]}],
    )
    assert saved is False

    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "第一段正常" not in texts  # 预校验拒绝，零写入

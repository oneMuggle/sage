"""Unit tests for Round 59 — {{en:}} 内联尾注（Phase C）。"""

from __future__ import annotations

import base64

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import (
    OfficeWordGenerateRequest,
    WordImageSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx, read_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen(tmp_path, name, paragraphs):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=paragraphs,
        images=[WordImageSpec(source=str(image_file), caption="架构图")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_endnote_marker_becomes_reference_and_part(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "en.docx",
        paragraphs=[
            WordParagraphSpec(text="结论见{{en:附注：统计口径说明。}}。"),
        ],
    )
    doc = Document(str(output))

    partnames = [str(p.partname) for p in doc.part.package.iter_parts()]
    assert "/word/endnotes.xml" in partnames

    refs = [
        el
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:endnoteReference"))
    ]
    assert len(refs) == 1
    assert refs[0].get(qn("w:id")) == "1"

    result = read_docx(output, workspace_path="")
    assert result.endnotes == ["附注：统计口径说明。"]
    assert "{{en:" not in "\n".join(p.text for p in doc.paragraphs)

    # 尾注样式注入（与脚注同款幂等）
    style_ids = {
        style.get(qn("w:styleId"))
        for style in doc.styles.element.findall(qn("w:style"))
    }
    assert "EndnoteText" in style_ids
    assert "EndnoteReference" in style_ids


def test_endnote_and_footnote_coexist(tmp_path) -> None:
    """同段脚注+尾注混用：各自独立 part 与编号。"""
    output = _gen(
        tmp_path,
        "mix.docx",
        paragraphs=[
            WordParagraphSpec(text="见{{fn:脚注备注}}与{{en:尾注备注}}。"),
        ],
    )
    doc = Document(str(output))
    partnames = [str(p.partname) for p in doc.part.package.iter_parts()]
    assert "/word/footnotes.xml" in partnames
    assert "/word/endnotes.xml" in partnames

    result = read_docx(output, workspace_path="")
    assert result.footnotes == ["脚注备注"]
    assert result.endnotes == ["尾注备注"]


def test_no_endnote_no_part(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "plain.docx",
        paragraphs=[WordParagraphSpec(text="普通段落")],
    )
    doc = Document(str(output))
    partnames = [str(p.partname) for p in doc.part.package.iter_parts()]
    assert "/word/endnotes.xml" not in partnames
    result = read_docx(output, workspace_path="")
    assert result.endnotes == []

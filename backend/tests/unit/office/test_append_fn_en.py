"""Unit tests for Round 61 — append_paragraphs 支持 {{fn:}}/{{en:}}。"""

from __future__ import annotations

import base64

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.edit import update_document
from backend.office.models import (
    OfficeWordGenerateRequest,
    WordImageSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx, read_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen_with_footnote(tmp_path, name):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文{{fn:已有脚注}}。")],
        images=[WordImageSpec(source=str(image_file), caption="架构图")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_append_fn_continues_numbering(tmp_path) -> None:
    """已有脚注的文档追加 {{fn:}} → part 续接编号（id=2）且回读可见。"""
    path = _gen_with_footnote(tmp_path, "r.docx")
    saved, results = update_document(
        "word",
        path,
        [{"op": "append_paragraphs", "paragraphs": [
            {"text": "补充说明{{fn:追加的脚注}}。"},
        ]}],
    )
    assert saved is True

    result = read_docx(path, workspace_path="")
    assert result.footnotes == ["已有脚注", "追加的脚注"]

    doc = Document(str(path))
    refs = [
        el
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:footnoteReference"))
    ]
    assert [r.get(qn("w:id")) for r in refs] == ["1", "2"]


def test_append_en_creates_endnotes_part(tmp_path) -> None:
    """无尾注文档追加 {{en:}} → 挂载 endnotes part 并写入。"""
    path = _gen_with_footnote(tmp_path, "r2.docx")
    saved, _ = update_document(
        "word",
        path,
        [{"op": "append_paragraphs", "paragraphs": [
            {"text": "补充{{en:尾注备注}}。"},
        ]}],
    )
    assert saved is True

    result = read_docx(path, workspace_path="")
    assert result.endnotes == ["尾注备注"]
    doc = Document(str(path))
    refs = [
        el
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:endnoteReference"))
    ]
    assert len(refs) == 1


def test_append_mixed_fn_en_and_caption(tmp_path) -> None:
    """一段内 fig+fn+en 混用：全部解析，各归其位。"""
    path = _gen_with_footnote(tmp_path, "r3.docx")
    saved, _ = update_document(
        "word",
        path,
        [{"op": "append_paragraphs", "paragraphs": [
            {"text": "架构见{{fig:架构图}}，口径{{fn:口径备注}}，"
                    "延伸{{en:延伸备注}}。"},
        ]}],
    )
    assert saved is True

    result = read_docx(path, workspace_path="")
    assert result.footnotes == ["已有脚注", "口径备注"]
    assert result.endnotes == ["延伸备注"]
    doc = Document(str(path))
    ref_instrs = [
        el.text
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:instrText"))
        if el.text and "REF" in el.text
    ]
    assert " REF _RefFig1 \\h " in ref_instrs

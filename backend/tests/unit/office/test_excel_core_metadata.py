"""Unit tests for Round 50 — Excel core properties（与 Word R49 对称）。"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from backend.office.excel import generate_xlsx
from backend.office.models import ExcelSheetSpec

pytestmark = pytest.mark.unit


def _gen(tmp_path, name, metadata=None):
    from backend.office.models import OfficeExcelGenerateRequest

    req = OfficeExcelGenerateRequest(
        workspace_path="",
        filename=name,
        sheets=[ExcelSheetSpec(name="数据", headers=["A"], rows=[["1"]])],
        metadata=metadata,
    )
    return generate_xlsx(req, output_dir=str(tmp_path))


def test_excel_metadata_written(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "meta.xlsx",
        metadata={
            "author": "张三",
            "subject": "台账",
            "keywords": "预算",
            "comments": "内部",
            "category": "财务",
        },
    )
    props = load_workbook(str(output)).properties
    assert props.creator == "张三"
    assert props.subject == "台账"
    assert props.keywords == "预算"
    assert props.description == "内部"
    assert props.category == "财务"


def test_excel_metadata_absent_defaults(tmp_path) -> None:
    """无 metadata → 不臆造作者（保留 openpyxl 模板默认）。"""
    output = _gen(tmp_path, "plain.xlsx")
    props = load_workbook(str(output)).properties
    assert props.subject in (None, "")

# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Tests for :mod:`backend.office.excel_eval`（Round2 R5 公式本地求值）.

覆盖：

- evaluate_workbook：真实 xlsx（openpyxl 生成、无缓存值）→
  ``{sheet: {"B4": 30}}``；无公式工作簿 → 空 dict。
- 失败路径全部 fail-safe → None：formulas 库缺失（monkeypatch import）、
  超时（monkeypatch _calculate 慢 + 极小超时预算）、求值异常
  （monkeypatch _calculate 抛错）、公式数超上限（monkeypatch 常量）。
- read_xlsx 集成：求值成功 → 条目升级为 ``B4=SUM(B2:B3) → 30 (本地求值)``
  且提示行省略；库缺失 → 原始条目 + 提示行原样（优雅降级）；部分解析
  → 仍有未解析公式时提示行保留。
"""

from __future__ import annotations

import builtins
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest
from openpyxl import Workbook

from backend.office.excel import read_xlsx
from backend.office.excel_eval import evaluate_workbook

pytestmark = pytest.mark.unit

_FORMULA_CACHE_NOTE = "公式计算值需在 Excel 中打开后生效"


def _build_sum_xlsx(path: Path) -> Path:
    """一个 sheet：B2=10, B3=20, B4='=SUM(B2:B3)'（openpyxl 不写缓存值）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = 10
    ws["B3"] = 20
    ws["B4"] = "=SUM(B2:B3)"
    wb.save(str(path))
    return path


def _build_multi_formula_xlsx(path: Path) -> Path:
    """两个公式单元格（B4=SUM、C1=B4*2），便于部分解析 / 上限测试。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["B2"] = 10
    ws["B3"] = 20
    ws["B4"] = "=SUM(B2:B3)"
    ws["C1"] = "=B4*2"
    wb.save(str(path))
    return path


_REAL_IMPORT = builtins.__import__  # monkeypatch 前的原始 import（防自递归）


def _no_formulas_import(name: str, *args: Any, **kwargs: Any):
    """monkeypatch builtins.__import__ 用：让 ``import formulas`` 抛 ImportError。"""
    if name == "formulas" or name.startswith("formulas."):
        raise ImportError("formulas disabled for test")
    return _REAL_IMPORT(name, *args, **kwargs)


# ──────────────────────────────────────────────────────────────────────
# evaluate_workbook
# ──────────────────────────────────────────────────────────────────────


def test_evaluate_workbook_returns_formula_values(fixture_dir: Path) -> None:
    """真实 xlsx 无缓存值 → {sheet: {坐标: 值}}，整数值 float 转 int。"""
    pytest.importorskip("formulas")
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    result = evaluate_workbook(path)
    assert result == {"Sheet1": {"B4": 30}}


def test_evaluate_workbook_only_covers_formula_cells(fixture_dir: Path) -> None:
    """非公式常量单元格（B2/B3）不进结果；范围解算键被跳过。"""
    pytest.importorskip("formulas")
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    result = evaluate_workbook(path)
    assert result is not None
    assert set(result["Sheet1"]) == {"B4"}


def test_evaluate_workbook_no_formulas_returns_empty_dict(fixture_dir: Path) -> None:
    """无公式工作簿 → 空 dict（区别于失败语义的 None）。"""
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = 1
    path = fixture_dir / "plain.xlsx"
    wb.save(str(path))
    assert evaluate_workbook(path) == {}


def test_evaluate_workbook_missing_library_returns_none(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """formulas 缺失 → None（优雅降级），绝不抛 ImportError。"""
    monkeypatch.setattr(builtins, "__import__", _no_formulas_import)
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    assert evaluate_workbook(path) is None


def test_evaluate_workbook_timeout_returns_none(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """求值超过墙钟预算 → None（monkeypatch 慢求值 + 极小超时）。"""
    import time

    def _slow_calculate(file_path: object) -> Dict[str, Any]:
        time.sleep(1.0)
        return {}

    monkeypatch.setattr("backend.office.excel_eval._calculate", _slow_calculate)
    monkeypatch.setattr("backend.office.excel_eval._EVAL_TIMEOUT_SECONDS", 0.05)
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    assert evaluate_workbook(path) is None


def test_evaluate_workbook_eval_failure_returns_none(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """求值线程抛异常 → None，绝不向上传播。"""
    def _boom(file_path: object) -> Dict[str, Any]:
        raise RuntimeError("boom")

    monkeypatch.setattr("backend.office.excel_eval._calculate", _boom)
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    assert evaluate_workbook(path) is None


def test_evaluate_workbook_formula_cap_returns_none(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """公式单元格数超上限 → None（monkeypatch 常量到 1，避免造大簿）。"""
    monkeypatch.setattr("backend.office.excel_eval._MAX_FORMULA_CELLS", 1)
    path = _build_multi_formula_xlsx(fixture_dir / "two.xlsx")  # 2 个公式 > 1
    assert evaluate_workbook(path) is None


def test_evaluate_workbook_corrupt_file_returns_none(fixture_dir: Path) -> None:
    """打不开的文件 → None（公式清单加载失败即降级）。"""
    path = fixture_dir / "corrupt.xlsx"
    path.write_bytes(b"not a zip")
    assert evaluate_workbook(path) is None


# ──────────────────────────────────────────────────────────────────────
# read_xlsx 集成（公式视图升级 / 降级 / 提示行去留）
# ──────────────────────────────────────────────────────────────────────


def test_read_xlsx_upgrades_entries_with_local_eval(fixture_dir: Path) -> None:
    """求值成功：条目升级 '→ 30 (本地求值)'，全部解析 → 提示行省略。"""
    pytest.importorskip("formulas")
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    result = read_xlsx(path, include_formulas=True)
    sheet = result.sheets[0]
    assert sheet.formulas == ["B4=SUM(B2:B3) → 30 (本地求值)"]
    assert sheet.note is None


def test_read_xlsx_graceful_when_formulas_missing(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """formulas 缺失：读取照常成功，原始条目 + 提示行原样保留。"""
    monkeypatch.setattr(builtins, "__import__", _no_formulas_import)
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    result = read_xlsx(path, include_formulas=True)
    sheet = result.sheets[0]
    assert sheet.formulas == ["B4=SUM(B2:B3)"]
    assert sheet.note == _FORMULA_CACHE_NOTE
    # 值网格不受影响：无缓存值 → 公式单元格读为空串
    # （B4 行全空被跳过，末行是 B3=20；网格里不应出现求值结果 30）
    assert sheet.rows[-1] == ["", "20"]
    assert all("30" not in row for row in sheet.rows)


def test_read_xlsx_keeps_note_when_eval_partial(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """部分解析：仍有未解析公式 → 提示行保留，已解析条目照常升级。"""
    path = _build_multi_formula_xlsx(fixture_dir / "multi.xlsx")

    def _partial(file_path: object) -> Optional[Dict[str, Dict[str, Any]]]:
        return {"Calc": {"B4": 30}}  # C1 缺席 → 未解析

    monkeypatch.setattr("backend.office.excel_eval.evaluate_workbook", _partial)
    result = read_xlsx(path, include_formulas=True)
    sheet = result.sheets[0]
    assert sheet.formulas is not None
    assert "B4=SUM(B2:B3) → 30 (本地求值)" in sheet.formulas
    assert "C1=B4*2" in sheet.formulas
    assert sheet.note == _FORMULA_CACHE_NOTE


def test_read_xlsx_without_flag_skips_eval(
    fixture_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """include_formulas=False 永不触发求值（缺 formulas 库也不报错）。"""
    monkeypatch.setattr(builtins, "__import__", _no_formulas_import)

    called: List[str] = []

    def _must_not_run(file_path: object) -> None:
        called.append("eval")

    monkeypatch.setattr("backend.office.excel_eval.evaluate_workbook", _must_not_run)
    path = _build_sum_xlsx(fixture_dir / "sum.xlsx")
    result = read_xlsx(path)
    assert result.sheets[0].formulas is None
    assert result.sheets[0].note is None
    assert called == []

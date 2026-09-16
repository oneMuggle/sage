"""P2-C (office-p2c)：Excel 公式缓存重算（soffice 就地刷新）测试。"""

from __future__ import annotations

import subprocess
from pathlib import Path
from typing import Any

import pytest

from backend.office import excel_recalc
from backend.office.excel import read_xlsx
from backend.office.excel_recalc import ExcelRecalcRequest, refresh_formula_cache

pytestmark = pytest.mark.unit


@pytest.fixture()
def ws(tmp_path: Path) -> Path:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    return workspace


def _staged_xlsx_with_formula(ws: Path) -> Path:
    """受管布局下的 .xlsx：B4 = SUM 公式（无缓存值）。"""
    from openpyxl import Workbook

    managed = ws / "office" / "excel" / "doc-1"
    managed.mkdir(parents=True)
    path = managed / "doc-1.xlsx"
    wb = Workbook()
    sheet = wb.active
    sheet["B2"] = 10
    sheet["B3"] = 20
    sheet["B4"] = "=SUM(B2:B3)"
    wb.save(str(path))
    return path


def _fake_recalc_soffice(monkeypatch: pytest.MonkeyPatch) -> None:
    """fake soffice：把源文件里公式替换为**带缓存值**的等价 xlsx。

    真实 soffice 转换会写缓存值；这里用 openpyxl 直接产出
    「公式 + 模拟缓存」不可行（openpyxl 不写缓存），改为产出
    值已计算的静态 xlsx —— 对调用方语义等价（data_only 读得到值）。
    """

    def fake_run(cmd: list, **kwargs: Any) -> subprocess.CompletedProcess:
        src = Path(cmd[-1])
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        from openpyxl import load_workbook

        wb = load_workbook(str(src))
        sheet = wb.active
        sheet["B4"] = 30
        wb.save(str(outdir / src.name))
        return subprocess.CompletedProcess(cmd, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(excel_recalc, "_locate_soffice", lambda: "/fake/soffice")
    monkeypatch.setattr(excel_recalc.subprocess, "run", fake_run)


def test_recalc_rewrites_formula_with_cached_value(
    ws: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _staged_xlsx_with_formula(ws)
    _fake_recalc_soffice(monkeypatch)

    result = refresh_formula_cache(
        ExcelRecalcRequest(workspace_path=str(ws), file_path=str(path))
    )
    assert result.ok is True
    # data_only=True 现在读得到计算值，不再需要"Excel 打开后生效"提示
    result_read = read_xlsx(path)
    all_cells = [c for r in result_read.sheets[0].rows for c in r]
    assert "30" in all_cells
    assert result_read.sheets[0].note is None


def test_recalc_takes_pre_snapshot(ws: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = _staged_xlsx_with_formula(ws)
    _fake_recalc_soffice(monkeypatch)

    refresh_formula_cache(
        ExcelRecalcRequest(workspace_path=str(ws), file_path=str(path))
    )
    snaps = list((path.parent / ".snapshots").glob("*-doc-1.xlsx"))
    assert len(snaps) == 1
    assert snaps[0].stat().st_size > 0


def test_missing_soffice_returns_guidance(
    ws: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _staged_xlsx_with_formula(ws)
    monkeypatch.setattr(excel_recalc, "_locate_soffice", lambda: None)
    result = refresh_formula_cache(
        ExcelRecalcRequest(workspace_path=str(ws), file_path=str(path))
    )
    assert result.ok is False
    assert "LibreOffice" in (result.error or "")


def test_rejects_non_xlsx_and_escape(ws: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_recalc_soffice(monkeypatch)
    csv_path = ws / "a.csv"
    csv_path.write_text("a\n1\n", encoding="utf-8")
    result = refresh_formula_cache(
        ExcelRecalcRequest(workspace_path=str(ws), file_path=str(csv_path))
    )
    assert result.ok is False

    outside = ws.parent / "out.xlsx"
    outside.write_bytes(b"x")
    result2 = refresh_formula_cache(
        ExcelRecalcRequest(workspace_path=str(ws), file_path=str(outside))
    )
    assert result2.ok is False

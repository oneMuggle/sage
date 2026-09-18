"""Unit tests for Round 42 — SEQ 题注 + 图目录/表目录（TOF 域）。"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import (
    OfficeWordGenerateRequest,
    WordFormatSpec,
    WordImageSpec,
    WordIndexSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx
from backend.office.word_layout import add_caption, insert_tof_field

pytestmark = pytest.mark.unit


def _fld_types(paragraph) -> list:
    return [
        fld.get(qn("w:fldCharType"))
        for fld in paragraph._p.findall(".//" + qn("w:fldChar"))
    ]


def _instr_texts(paragraph) -> list:
    return [
        el.text for el in paragraph._p.findall(".//" + qn("w:instrText"))
    ]


# ── SEQ 题注 ──────────────────────────────────────────────────────────


def test_add_caption_wraps_number_in_seq_field() -> None:
    doc = Document()
    add_caption(doc, "架构图", kind="figure", number=1)

    para = doc.paragraphs[0]
    assert _fld_types(para) == ["begin", "separate", "end"]
    assert _instr_texts(para) == [" SEQ 图 \\* ARABIC "]
    # 回读文本与既有字面形态完全一致（lint caption/sequence 零改动）
    assert para.text == "图1　架构图"


def test_add_caption_table_label() -> None:
    doc = Document()
    add_caption(doc, "季度数据", kind="table", number=3)

    para = doc.paragraphs[0]
    assert _instr_texts(para) == [" SEQ 表 \\* ARABIC "]
    assert para.text == "表3　季度数据"


# ── TOF 域 ────────────────────────────────────────────────────────────


def test_insert_tof_field_with_entries() -> None:
    doc = Document()
    insert_tof_field(doc, WordIndexSpec(), "图", [(1, "架构图"), (2, "部署拓扑")])

    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "图目录"  # 标题段
    instr = _instr_texts(doc.paragraphs[1])
    assert instr == [' TOC \\h \\z \\c "图" ']
    assert _fld_types(doc.paragraphs[1]) == ["begin", "separate"]
    assert "图1　架构图" in texts
    assert "图2　部署拓扑" in texts
    assert _fld_types(doc.paragraphs[-2]) == ["end"]


def test_insert_tof_field_empty_uses_placeholder() -> None:
    doc = Document()
    spec = WordIndexSpec(placeholder_text="（无插图）")
    insert_tof_field(doc, spec, "图", [])

    texts = [p.text for p in doc.paragraphs]
    assert "（无插图）" in texts


# ── 端到端：generate_docx ─────────────────────────────────────────────


_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)
import base64

_TINY_PNG = base64.b64decode(_TINY_PNG_B64)


def _generate(image_source: str, **format_spec_kwargs):
    return OfficeWordGenerateRequest(
        workspace_path="",
        filename="报告.docx",
        title="测试报告",
        paragraphs=[
            WordParagraphSpec(text="第一章开头"),
            WordParagraphSpec(text="第二章开头", heading="h1"),
        ],
        tables=[
            WordTableSpec(
                headers=["A"],
                rows=[["1"]],
                caption="汇总表",
            ),
            WordTableSpec(headers=["B"], rows=[["2"]]),
        ],
        images=[
            WordImageSpec(source=image_source, caption="架构图"),
        ],
        format_spec=WordFormatSpec(
            figure_index=WordIndexSpec(),
            table_index=WordIndexSpec(heading_text="表目录"),
            **format_spec_kwargs,
        ),
    )


def test_generate_docx_with_caption_indexes(tmp_path):
    """图/表目录域 + SEQ 题注端到端：域指令、缓存条目、编号守卫。"""
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(_TINY_PNG)
    req = _generate(str(image_file))
    output = generate_docx(req, output_dir=str(tmp_path))

    doc = Document(str(output))
    instrs = [
        (i, el.text)
        for i, p in enumerate(doc.paragraphs)
        for el in p._p.findall(".//" + qn("w:instrText"))
        if el.text and "TOC" in el.text
    ]
    tof_instrs = [t for _, t in instrs if "\\c" in t]
    assert ' TOC \\h \\z \\c "图" ' in tof_instrs
    assert ' TOC \\h \\z \\c "表" ' in tof_instrs

    all_text = "\n".join(p.text for p in doc.paragraphs)
    # 缓存条目在文档中（页码留待刷新域）
    assert "图1　架构图" in all_text
    assert "表1　汇总表" in all_text
    # 无题注不占号：第二张表无 caption，不出现在目录条目
    assert "表2" not in all_text


def test_lint_ignores_caption_index_cache_lines(tmp_path):
    """TOF 缓存行（图N　标题）不参与 caption/sequence 规则。"""
    from backend.office.word_lint import lint_docx

    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(_TINY_PNG)
    req = _generate(str(image_file))
    output = generate_docx(req, output_dir=str(tmp_path))

    result = lint_docx(output, WordFormatSpec())
    caption_issues = [
        i for i in result.issues if i.rule_id == "caption/sequence"
    ]
    assert caption_issues == []

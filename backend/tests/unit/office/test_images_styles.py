"""Unit tests for 批次 2.1 图片 / 2.3 样式分级（round a）编辑与生成.

Covers:
- Word: 生成期插图（base64 / 工作区路径）、ParagraphSpec 样式字段；
  编辑期 add_image（base64 / 文档目录相对路径 / 失败不落盘）与
  set_paragraph_style（index / match 定位，字号/加粗/颜色/对齐）
- Excel: 编辑期 set_column_width / set_number_format / set_fill /
  freeze_panes / add_chart
- PPT: 生成期 layout 选择（blank 全文本框、title/title_content 走占位符）
  与插图；编辑期 add_picture（index 与 slide 键名都接受）
"""

from __future__ import annotations

import base64
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.office.edit import update_docx, update_pptx, update_xlsx
from backend.office.excel import generate_xlsx, read_xlsx
from backend.office.models import (
    ExcelSheetSpec,
    ImageSourceSpec,
    OfficeExcelGenerateRequest,
    OfficePptGenerateRequest,
    OfficeWordGenerateRequest,
    PptSlideSpec,
    WordParagraphSpec,
)
from backend.office.ppt import generate_ppt, read_ppt
from backend.office.word import generate_docx, read_docx

pytestmark = pytest.mark.unit

#: 1x1 red PNG（与 conftest 的 helper 同源）。
_MINIMAL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP8z8Dw"
    "HwAFBQIAX8v0gQAAAABJRU5ErkJggg=="
)
_MINIMAL_PNG_BYTES = base64.b64decode(_MINIMAL_PNG_B64, validate=True)
_DATA_URI = "data:image/png;base64," + _MINIMAL_PNG_B64


def _make_word_req(tmp_path: Path, **overrides) -> OfficeWordGenerateRequest:
    data = {
        "workspace_path": str(tmp_path),
        "filename": "样式.docx",
        "title": "样式演示",
        "paragraphs": [
            WordParagraphSpec(text="第一段"),
            WordParagraphSpec(text="第二段"),
        ],
    }
    data.update(overrides)
    return OfficeWordGenerateRequest(**data)


# ──────────────────────────────────────────────────────────────────────
# Word 生成：插图 + 段落样式
# ──────────────────────────────────────────────────────────────────────


def test_generate_docx_with_base64_image(tmp_path: Path) -> None:
    req = _make_word_req(
        tmp_path,
        images=[ImageSourceSpec(source=_DATA_URI, width_inches=2.0)],
    )
    out = generate_docx(req, output_dir=str(tmp_path))
    assert read_docx(out, workspace_path="").images == 1


def test_generate_docx_with_workspace_image_path(tmp_path: Path) -> None:
    (tmp_path / "pic.png").write_bytes(_MINIMAL_PNG_BYTES)
    req = _make_word_req(tmp_path, images=[ImageSourceSpec(source="pic.png")])
    out = generate_docx(req)  # 走 workspace 沙箱路径
    assert read_docx(out, workspace_path=str(tmp_path)).images == 1


def test_generate_docx_paragraph_style_fields(tmp_path: Path) -> None:
    req = _make_word_req(
        tmp_path,
        paragraphs=[
            WordParagraphSpec(
                text="红色居中标题段",
                font_size=14,
                bold=True,
                italic=True,
                color="#FF0000",
                align="center",
            ),
        ],
    )
    out = generate_docx(req, output_dir=str(tmp_path))
    doc = Document(str(out))
    para = doc.paragraphs[1]  # 0 是 Title
    run = para.runs[0]
    assert run.font.size.pt == 14
    assert run.font.bold is True
    assert run.font.italic is True
    assert str(run.font.color.rgb) == "FF0000"
    assert para.alignment is not None
    assert "CENTER" in str(para.alignment)


def test_generate_docx_without_style_fields_unchanged(tmp_path: Path) -> None:
    """不传样式字段时保持既有行为（run 无显式 size/bold 覆盖）。"""
    out = generate_docx(_make_word_req(tmp_path), output_dir=str(tmp_path))
    doc = Document(str(out))
    body = doc.paragraphs[1]
    assert body.runs[0].font.size is None
    assert body.runs[0].font.bold is None


# ──────────────────────────────────────────────────────────────────────
# Word 编辑：add_image / set_paragraph_style
# ──────────────────────────────────────────────────────────────────────


def _make_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("会议纪要", level=0)
    doc.add_paragraph("今天天气很好")
    doc.save(str(path))
    return path


def test_edit_docx_add_image_via_base64(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(
        path, [{"op": "add_image", "base64": _DATA_URI, "width_inches": 1.5}]
    )
    assert saved
    assert results[0]["ok"]
    assert read_docx(path, workspace_path="").images == 1


def test_edit_docx_add_image_via_relative_path(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    (tmp_path / "img.png").write_bytes(_MINIMAL_PNG_BYTES)
    saved, results = update_docx(path, [{"op": "add_image", "path": "img.png"}])
    assert saved
    assert results[0]["ok"]
    assert read_docx(path, workspace_path="").images == 1


def test_edit_docx_add_image_missing_source_fails_atomically(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(path, [{"op": "add_image", "path": "ghost.png"}])
    assert not saved
    assert not results[0]["ok"]
    assert "image_file_not_found" in results[0]["error"]
    # 原文件未被破坏
    assert read_docx(path, workspace_path="").images == 0


def test_edit_docx_set_paragraph_style_by_index(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(
        path,
        [
            {
                "op": "set_paragraph_style",
                "index": 1,
                "font_size": 12,
                "bold": True,
                "color": "3366CC",
                "align": "center",
            }
        ],
    )
    assert saved
    assert results[0]["ok"]
    para = Document(str(path)).paragraphs[1]
    assert para.runs[0].font.size.pt == 12
    assert para.runs[0].font.bold is True
    assert str(para.runs[0].font.color.rgb) == "3366CC"
    assert para.alignment is not None
    assert "CENTER" in str(para.alignment)


def test_edit_docx_set_paragraph_style_by_match(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    saved, results = update_docx(
        path, [{"op": "set_paragraph_style", "match": "天气", "italic": True}]
    )
    assert saved
    assert results[0]["ok"]
    assert results[0]["runs"] >= 1
    assert Document(str(path)).paragraphs[1].runs[0].font.italic is True


def test_edit_docx_set_paragraph_style_validation(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "a.docx")
    # index / match 都缺 → 拒绝
    saved, results = update_docx(path, [{"op": "set_paragraph_style", "bold": True}])
    assert not saved
    assert "index_or_match_required" in results[0]["error"]
    # index 越界
    saved, results = update_docx(
        path, [{"op": "set_paragraph_style", "index": 99, "bold": True}]
    )
    assert not saved
    assert "out_of_range" in results[0]["error"]
    # 非法 align
    saved, results = update_docx(
        path, [{"op": "set_paragraph_style", "index": 1, "align": "diagonal"}]
    )
    assert not saved
    assert "invalid_align" in results[0]["error"]
    # 非法 color
    saved, results = update_docx(
        path, [{"op": "set_paragraph_style", "index": 1, "color": "red"}]
    )
    assert not saved
    assert "invalid_color" in results[0]["error"]


# ──────────────────────────────────────────────────────────────────────
# Excel 编辑：列宽 / 数字格式 / 填充 / 冻结 / 原生图表
# ──────────────────────────────────────────────────────────────────────


def _make_xlsx(path: Path) -> Path:
    req = OfficeExcelGenerateRequest(
        workspace_path=str(path.parent),
        filename=path.name,
        sheets=[
            ExcelSheetSpec(
                name="数据",
                headers=["月份", "销售额"],
                rows=[["一月", "100"], ["二月", "200"]],
            )
        ],
    )
    return generate_xlsx(req, output_dir=str(path.parent))


def test_edit_xlsx_set_column_width(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path, [{"op": "set_column_width", "sheet": "数据", "column": "A", "width": 42.5}]
    )
    assert saved
    assert results[0]["ok"]
    ws = load_workbook(str(path))["数据"]
    assert ws.column_dimensions["A"].width == 42.5


def test_edit_xlsx_set_column_width_by_index(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path, [{"op": "set_column_width", "sheet": "数据", "column": 2, "width": 15}]
    )
    assert saved
    assert results[0]["ok"]
    assert results[0]["column"] == "B"
    assert load_workbook(str(path))["数据"].column_dimensions["B"].width == 15


def test_edit_xlsx_set_number_format(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path,
        [
            {
                "op": "set_number_format",
                "sheet": "数据",
                "cells": "B2:B3",
                "format": "0.00%",
            },
            {
                "op": "set_number_format",
                "sheet": "数据",
                "cells": ["A1", "A2"],
                "format": "yyyy-mm-dd",
            },
        ],
    )
    assert saved
    assert results[0]["ok"]
    assert results[0]["cells"] == 2
    assert results[1]["ok"]
    assert results[1]["cells"] == 2
    ws = load_workbook(str(path))["数据"]
    assert ws["B2"].number_format == "0.00%"
    assert ws["B3"].number_format == "0.00%"
    assert ws["A1"].number_format == "yyyy-mm-dd"
    assert ws["A2"].number_format == "yyyy-mm-dd"


def test_edit_xlsx_set_fill(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path, [{"op": "set_fill", "sheet": "数据", "cells": "A1:B1", "color": "#FF0000"}]
    )
    assert saved
    assert results[0]["ok"]
    assert results[0]["color"] == "FFFF0000"
    ws = load_workbook(str(path))["数据"]
    assert str(ws["A1"].fill.start_color.rgb) == "FFFF0000"
    assert ws["B1"].fill.fill_type == "solid"


def test_edit_xlsx_set_fill_invalid_color(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path, [{"op": "set_fill", "sheet": "数据", "cells": "A1", "color": "red"}]
    )
    assert not saved
    assert "invalid_color" in results[0]["error"]


def test_edit_xlsx_freeze_panes(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path, [{"op": "freeze_panes", "sheet": "数据", "cell": "B2"}]
    )
    assert saved
    assert results[0]["ok"]
    assert load_workbook(str(path))["数据"].freeze_panes == "B2"
    # A1 表示取消冻结
    saved, results = update_xlsx(
        path, [{"op": "freeze_panes", "sheet": "数据", "cell": "A1"}]
    )
    assert saved
    assert results[0]["frozen_at"] is None
    assert load_workbook(str(path))["数据"].freeze_panes is None


def test_edit_xlsx_add_chart_op(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path,
        [
            {
                "op": "add_chart",
                "sheet": "数据",
                "type": "bar",
                "anchor": "A10",
                "data_ref": {"min_col": 2, "min_row": 1, "max_col": 2, "max_row": 3},
                "titles_from_data": True,
                "categories_ref": {"min_col": 1, "min_row": 2, "max_col": 1, "max_row": 3},
            }
        ],
    )
    assert saved
    assert results[0]["ok"]
    ws = load_workbook(str(path))["数据"]
    assert len(ws._charts) == 1
    assert "BarChart" in type(ws._charts[0]).__name__


def test_edit_xlsx_add_chart_bad_type(tmp_path: Path) -> None:
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path,
        [
            {
                "op": "add_chart",
                "sheet": "数据",
                "type": "hbar",
                "anchor": "A10",
                "data_ref": {"min_col": 1, "min_row": 1, "max_col": 2, "max_row": 3},
            }
        ],
    )
    assert not saved
    assert "unsupported_chart_type" in results[0]["error"]


# ──────────────────────────────────────────────────────────────────────
# PPT：layout 选择 + 插图 + add_picture
# ──────────────────────────────────────────────────────────────────────


def _make_ppt_req(tmp_path: Path, slides) -> OfficePptGenerateRequest:
    return OfficePptGenerateRequest(
        workspace_path=str(tmp_path),
        filename="deck.pptx",
        slides=slides,
    )


def _gen_ppt(tmp_path: Path, layout, image=None) -> Path:
    spec_kwargs = {"title": "标题页", "bullets": ["要点一", "要点二"]}
    if layout is not None:
        spec_kwargs["layout"] = layout
    if image is not None:
        spec_kwargs["image"] = image
    return generate_ppt(
        _make_ppt_req(tmp_path, [PptSlideSpec(**spec_kwargs)]),
        output_dir=str(tmp_path),
    )


def test_generate_ppt_default_keeps_textbox_geometry(tmp_path: Path) -> None:
    """不传 layout：与既有行为一致 —— 标题/正文均为文本框，无占位符。

    read_ppt 对「无占位符」的幻灯片会把标题框文本一并计入 text_blocks
    （reader 既有行为），这里一并锁定。
    """
    out = _gen_ppt(tmp_path, None)
    result = read_ppt(out, workspace_path="")
    assert result.slides[0].title == "标题页"
    assert result.slides[0].text_blocks == ["标题页", "要点一", "要点二"]
    slide = Presentation(str(out)).slides[0]
    assert not any(shape.is_placeholder for shape in slide.shapes)
    assert all(
        shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        for shape in slide.shapes
        if shape.has_text_frame
    )


def test_generate_ppt_blank_layout_honored(tmp_path: Path) -> None:
    """layout='blank'：显式选 Blank 版式，仍然全文本框（与默认几何一致）。"""
    out = _gen_ppt(tmp_path, "blank")
    slide = Presentation(str(out)).slides[0]
    assert not any(shape.is_placeholder for shape in slide.shapes)
    assert read_ppt(out, workspace_path="").slides[0].title == "标题页"


def test_generate_ppt_title_layout_uses_placeholder(tmp_path: Path) -> None:
    out = _gen_ppt(tmp_path, "title")
    slide = Presentation(str(out)).slides[0]
    assert any(shape.is_placeholder for shape in slide.shapes)
    assert read_ppt(out, workspace_path="").slides[0].title == "标题页"


def test_generate_ppt_title_content_layout_uses_placeholders(tmp_path: Path) -> None:
    out = _gen_ppt(tmp_path, "title_content")
    slide = Presentation(str(out)).slides[0]
    placeholders = [s for s in slide.shapes if s.is_placeholder]
    assert len(placeholders) >= 2
    result = read_ppt(out, workspace_path="")
    assert result.slides[0].title == "标题页"
    assert "要点一" in result.slides[0].text_blocks
    assert "要点二" in result.slides[0].text_blocks


def test_resolve_slide_layout_name_lookup_and_fallback() -> None:
    """版式名查找命中默认模板；未知名字返回 None（回退文本框几何）。"""
    from backend.office.ppt import _resolve_slide_layout

    prs = Presentation()
    assert _resolve_slide_layout(prs, None) is None
    assert _resolve_slide_layout(prs, "title").name == "Title Slide"
    assert _resolve_slide_layout(prs, "title_content").name == "Title and Content"
    assert _resolve_slide_layout(prs, "blank").name == "Blank"
    # 未知名字：不在名字表也超出索引兜底 → None
    assert _resolve_slide_layout(prs, "no-such-layout") is None
    # 原始字符串与 PptLayoutName 枚举都能命中（str-Enum 归一化）
    from backend.office.models import PptLayoutName

    assert _resolve_slide_layout(prs, PptLayoutName.TITLE).name == "Title Slide"


def test_generate_ppt_with_image(tmp_path: Path) -> None:
    out = _gen_ppt(tmp_path, None, image={"source": _DATA_URI, "width_inches": 3.0})
    result = read_ppt(out, workspace_path="")
    assert result.slides[0].image_count == 1


def test_edit_pptx_add_picture(tmp_path: Path) -> None:
    out = _gen_ppt(tmp_path, None)
    saved, results = update_pptx(
        out, [{"op": "add_picture", "index": 0, "base64": _DATA_URI}]
    )
    assert saved
    assert results[0]["ok"]
    assert read_ppt(out, workspace_path="").slides[0].image_count == 1
    # 'slide' 键名同样接受；找不到图片文件时该 op 失败且不落盘
    prs = Presentation(str(out))
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(out))
    saved, results = update_pptx(out, [{"op": "add_picture", "slide": 1, "path": "x"}])
    assert not saved
    assert "image_file_not_found" in results[0]["error"]


def test_read_xlsx_still_works_after_style_ops(tmp_path: Path) -> None:
    """样式编辑后的文件对 read_xlsx 仍是合法 xlsx（round-trip 护栏）。"""
    path = _make_xlsx(tmp_path / "a.xlsx")
    saved, results = update_xlsx(
        path,
        [
            {"op": "set_column_width", "sheet": "数据", "column": "A", "width": 20},
            {"op": "freeze_panes", "sheet": "数据", "cell": "A2"},
        ],
    )
    assert saved
    assert all(r["ok"] for r in results)
    result = read_xlsx(path, workspace_path="")
    assert result.sheets[0].rows[0] == ["月份", "销售额"]

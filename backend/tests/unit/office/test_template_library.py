"""Unit tests for the 中文办公模板库 (batch 3 — Item 3.2).

Covers:
- builtin registry shapes (6 templates, unique ids, non-empty placeholders)
- build() produces valid .docx files (python-docx reopens them)
- scanner extensions: table-cell {{}} → TABLE, ``|``-filter {{}} → RICH_TEXT
- list_templates: builtin + workspace scan (broken docx skipped, not fatal)
- instantiate: builtin end-to-end fill, workspace template fill (incl. a
  ``|``-filter tag whose base name must be addressable in ``data``)
- security: unknown/invalid template id, workspace_template path traversal,
  oversized data values, mutually exclusive selectors
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from backend.office import template_library
from backend.office.errors import (
    OfficeFileNotFoundError,
    OfficePathError,
    OfficeTemplateFillError,
)
from backend.office.models import (
    PlaceholderLocation,
    TemplatePlaceholderType,
)
from backend.office.template_library import (
    BUILTIN_TEMPLATES,
    TEMPLATE_ID_RE,
    instantiate_template,
    list_templates,
)
from backend.office.word_template import analyze_word_template

# ──────────────────────────────────────────────────────────────────────
# builtin registry shapes
# ──────────────────────────────────────────────────────────────────────


def test_builtin_registry_has_six_unique_templates():
    assert len(BUILTIN_TEMPLATES) == 6
    ids = [spec.id for spec in BUILTIN_TEMPLATES]
    assert len(set(ids)) == len(ids)
    for template_id in ids:
        assert TEMPLATE_ID_RE.match(template_id), template_id


def test_builtin_registry_ids_and_names():
    by_id = {spec.id: spec for spec in BUILTIN_TEMPLATES}
    assert set(by_id) == {
        "weekly_report",
        "meeting_minutes",
        "project_plan",
        "daily_report",
        "retrospective",
        "resume",
    }
    assert by_id["weekly_report"].name == "周报"
    assert by_id["meeting_minutes"].name == "会议纪要"
    for spec in BUILTIN_TEMPLATES:
        assert spec.doc_type.value == "word"
        assert spec.description


@pytest.mark.parametrize(
    "spec", BUILTIN_TEMPLATES, ids=[spec.id for spec in BUILTIN_TEMPLATES]
)
def test_builtin_registry_placeholders_non_empty(spec):
    assert spec.placeholders
    names = [placeholder.name for placeholder in spec.placeholders]
    assert len(set(names)) == len(names)
    for placeholder in spec.placeholders:
        assert placeholder.description


# ──────────────────────────────────────────────────────────────────────
# build() produces valid docx
# ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "spec", BUILTIN_TEMPLATES, ids=[spec.id for spec in BUILTIN_TEMPLATES]
)
def test_builtin_build_produces_valid_docx(spec):
    path = spec.build()
    try:
        assert path.is_file()
        doc = Document(str(path))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        for placeholder in spec.placeholders:
            tag = "{{" + placeholder.name + "}}"
            assert tag in all_text or _tag_in_tables(doc, tag), (
                f"{spec.id} missing tag {tag}"
            )
    finally:
        path.unlink(missing_ok=True)


def _tag_in_tables(doc, tag: str) -> bool:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if tag in cell.text:
                    return True
    return False


def test_builtin_meeting_minutes_table_tags_live_in_real_tables():
    spec = {s.id: s for s in BUILTIN_TEMPLATES}["meeting_minutes"]
    path = spec.build()
    try:
        doc = Document(str(path))
        assert len(doc.tables) == 2
        table_text = "\n".join(
            cell.text for t in doc.tables for row in t.rows for cell in row.cells
        )
        assert "{{discussion}}" in table_text
        assert "{{action_items}}" in table_text
    finally:
        path.unlink(missing_ok=True)


# ──────────────────────────────────────────────────────────────────────
# scanner extensions (word_template)
# ──────────────────────────────────────────────────────────────────────


def _analyze_text_template(tmp_path: Path, lines: list) -> object:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    path = tmp_path / "scan.docx"
    doc.save(str(path))
    return analyze_word_template(path, workspace_path=str(tmp_path))


def test_scanner_classifies_filter_tag_as_rich_text(tmp_path: Path):
    result = _analyze_text_template(tmp_path, ["要点：{{highlights | trim}}"])
    assert len(result.placeholders) == 1
    placeholder = result.placeholders[0]
    assert placeholder.type == TemplatePlaceholderType.RICH_TEXT
    # name is stripped to the base variable so the fill context can address it
    assert placeholder.name == "highlights"
    assert placeholder.raw_tag == "{{highlights | trim}}"


def test_scanner_classifies_table_cell_var_as_table(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("正文：{{plain_text}}")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{discussion}}"
    table.cell(0, 1).text = "{{action_items}}"
    path = tmp_path / "table-scan.docx"
    doc.save(str(path))
    result = analyze_word_template(path, workspace_path=str(tmp_path))

    by_name = {p.name: p for p in result.placeholders}
    assert by_name["plain_text"].type == TemplatePlaceholderType.TEXT
    for name in ("discussion", "action_items"):
        assert by_name[name].type == TemplatePlaceholderType.TABLE
        assert by_name[name].location == PlaceholderLocation.TABLE
    # table indices are still recorded for the TABLE classification
    assert by_name["discussion"].table_index == 0


def test_scanner_date_image_keywords_still_win_over_table_context(tmp_path: Path):
    """Pre-existing keyword rules keep precedence so old templates classify unchanged."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{合同日期}}"
    table.cell(0, 1).text = "{{签名图片}}"
    path = tmp_path / "keywords.docx"
    doc.save(str(path))
    result = analyze_word_template(path, workspace_path=str(tmp_path))
    by_name = {p.name: p for p in result.placeholders}
    assert by_name["合同日期"].type == TemplatePlaceholderType.DATE
    assert by_name["签名图片"].type == TemplatePlaceholderType.IMAGE


def test_scanner_filter_in_table_cell_is_rich_text(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{summary | upper}}"
    path = tmp_path / "filter-table.docx"
    doc.save(str(path))
    result = analyze_word_template(path, workspace_path=str(tmp_path))
    assert result.placeholders[0].type == TemplatePlaceholderType.RICH_TEXT


# ──────────────────────────────────────────────────────────────────────
# list_templates
# ──────────────────────────────────────────────────────────────────────


def test_list_templates_without_workspace_returns_builtins():
    response = list_templates(None)
    assert [t.source for t in response.templates].count("builtin") == 6
    assert all(t.filename is None for t in response.templates)
    weekly = next(t for t in response.templates if t.id == "weekly_report")
    assert weekly.name == "周报"
    assert {p.name for p in weekly.placeholders} == {
        "title",
        "author",
        "week_range",
        "highlights",
        "next_plan",
        "risks",
    }


def test_list_templates_workspace_dir_with_broken_docx_skipped(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)

    good = Document()
    good.add_paragraph("标的：{{contract_name}}")
    good.add_paragraph("金额：{{amount}}")
    good.save(str(templates_dir / "contract.docx"))
    (templates_dir / "broken.docx").write_text("not a docx at all")

    response = list_templates(str(tmp_path))
    workspace_entries = [t for t in response.templates if t.source == "workspace"]
    assert len(workspace_entries) == 1
    entry = workspace_entries[0]
    assert entry.filename == "contract.docx"
    assert entry.id.startswith("ws_")
    assert {p.name for p in entry.placeholders} == {"contract_name", "amount"}
    # builtins are still listed alongside
    assert len([t for t in response.templates if t.source == "builtin"]) == 6


def test_list_templates_missing_templates_dir_returns_builtins_only(tmp_path: Path):
    response = list_templates(str(tmp_path))
    assert len(response.templates) == 6
    assert all(t.source == "builtin" for t in response.templates)


# ──────────────────────────────────────────────────────────────────────
# instantiate_template
# ──────────────────────────────────────────────────────────────────────


WEEKLY_DATA = {
    "title": "启明星项目周报",
    "author": "张三",
    "week_range": "2026-09-01 ~ 2026-09-07",
    "highlights": "完成支付链路联调",
    "next_plan": "启动压测",
    "risks": "第三方接口不稳定",
}


def test_instantiate_weekly_report_end_to_end(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path), template_id="weekly_report", filename="本周周报", data=WEEKLY_DATA
    )

    output = Path(result.output_path)
    # managed layout: <workspace>/office/word/<uuid>/<filename>
    assert output.parent.parent == tmp_path / "office" / "word"
    assert output.name == "本周周报.docx"
    assert output.is_file()
    assert result.filename == "本周周报.docx"
    assert result.filled_count == 6
    assert result.unfilled_placeholders == []

    doc = Document(str(output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "启明星项目周报" in text
    assert "作者：张三" in text
    assert "2026-09-01 ~ 2026-09-07" in text
    assert "{{" not in text


def test_instantiate_meeting_minutes_fills_table_cells(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path),
        template_id="meeting_minutes",
        filename="minutes.docx",
        data={
            "meeting_title": "迭代评审会",
            "date": "2026-09-08",
            "attendees": "张三、李四",
            "discussion": "排期确认",
            "action_items": "补充压测报告",
        },
    )
    doc = Document(result.output_path)
    table_text = "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells
    )
    assert "排期确认" in table_text
    assert "补充压测报告" in table_text


def test_instantiate_workspace_template_by_filename(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("标的：{{contract_name}}")
    doc.save(str(templates_dir / "contract.docx"))

    result = instantiate_template(
        str(tmp_path),
        workspace_template="contract.docx",
        filename="filled-contract.docx",
        data={"contract_name": "办公楼租赁"},
    )
    text = "\n".join(p.text for p in Document(result.output_path).paragraphs)
    assert "标的：办公楼租赁" in text
    assert result.unfilled_placeholders == []


def test_instantiate_workspace_template_with_filter_tag_uses_base_name(tmp_path: Path):
    """``{{var | filter}}`` tags are filled by their stripped base variable name."""
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("{{summary | upper}}")
    doc.save(str(templates_dir / "filtered.docx"))

    result = instantiate_template(
        str(tmp_path),
        workspace_template="filtered.docx",
        filename="out.docx",
        data={"summary": "done"},
    )
    assert result.unfilled_placeholders == []
    text = "\n".join(p.text for p in Document(result.output_path).paragraphs)
    assert "DONE" in text


def test_instantiate_unknown_builtin_id_raises(tmp_path: Path):
    with pytest.raises(OfficeFileNotFoundError):
        instantiate_template(
            str(tmp_path), template_id="no_such_template", filename="x.docx", data={}
        )


def test_instantiate_rejects_invalid_template_id(tmp_path: Path):
    for bad_id in ("Weekly-Report", "a/b", "../weekly_report", "a" * 65, ""):
        with pytest.raises(OfficePathError):
            instantiate_template(
                str(tmp_path), template_id=bad_id, filename="x.docx", data={}
            )


def test_instantiate_rejects_workspace_template_traversal(tmp_path: Path):
    outside = tmp_path.parent / "escape.docx"
    Document().save(str(outside))
    with pytest.raises(OfficePathError):
        instantiate_template(
            str(tmp_path), workspace_template="../escape.docx", filename="x.docx", data={}
        )


def test_instantiate_workspace_template_missing_file_raises(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    with pytest.raises(OfficeFileNotFoundError):
        instantiate_template(
            str(tmp_path), workspace_template="ghost.docx", filename="x.docx", data={}
        )


def test_instantiate_rejects_data_value_over_cap(tmp_path: Path):
    big_value = "字" * (template_library.MAX_DATA_VALUE_CHARS + 1)
    with pytest.raises(OfficeTemplateFillError):
        instantiate_template(
            str(tmp_path),
            template_id="weekly_report",
            filename="x.docx",
            data={"title": big_value},
        )


def test_instantiate_requires_template_selector(tmp_path: Path):
    with pytest.raises(OfficePathError):
        instantiate_template(str(tmp_path), filename="x.docx", data={})


def test_instantiate_rejects_conflicting_selectors(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    Document().save(str(templates_dir / "contract.docx"))
    with pytest.raises(OfficePathError):
        instantiate_template(
            str(tmp_path),
            template_id="weekly_report",
            workspace_template="contract.docx",
            filename="x.docx",
            data={},
        )


def test_instantiate_failure_leaves_no_managed_dir(tmp_path: Path, monkeypatch):
    def broken_fill(req):
        raise OfficeTemplateFillError("Template fill failed")

    monkeypatch.setattr(template_library, "fill_word_template", broken_fill)
    with pytest.raises(OfficeTemplateFillError):
        instantiate_template(
            str(tmp_path), template_id="weekly_report", filename="x.docx", data={}
        )
    word_root = tmp_path / "office" / "word"
    assert not word_root.exists() or not any(word_root.iterdir())

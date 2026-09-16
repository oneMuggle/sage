"""Unit tests for word._extract_image_previews (Round C P4 图片缩略).

用 python-docx 真实内嵌一张 PNG 走全链路；Pillow 存在与否两条路径都
覆盖（通过 monkeypatch 模拟 no-Pillow 环境）。
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from backend.office.word import (
    read_docx,
)


def _tiny_png(width: int = 4, height: int = 4) -> bytes:
    """Minimal valid RGB PNG without Pillow (pure struct/zlib)."""
    def chunk(tag: bytes, payload: bytes) -> bytes:
        raw = tag + payload
        return struct.pack(">I", len(payload)) + raw + struct.pack(
            ">I", zlib.crc32(raw) & 0xFFFFFFFF
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    scanlines = b"".join(b"\x00" + b"\x80\x80\x80" * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


def _make_docx_with_image(path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph("图文混排文档")
    doc.add_picture(io.BytesIO(_tiny_png()), width=Inches(1))
    doc.save(str(path))


def test_read_docx_returns_image_previews(tmp_path):
    f = tmp_path / "with-image.docx"
    _make_docx_with_image(f)

    result = read_docx(f)
    assert result.images == 1
    assert len(result.image_previews) == 1
    preview = result.image_previews[0]
    assert preview.index == 0
    assert preview.data_url.startswith("data:image/")
    assert ";base64," in preview.data_url


def test_no_pillow_inlines_small_original(tmp_path, monkeypatch):
    """Pillow 缺失 → ≤150KB 原图直接内联（thumbnail=False）。"""
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "PIL" or name.startswith("PIL."):
            raise ImportError("no PIL in this test")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    f = tmp_path / "no-pillow.docx"
    _make_docx_with_image(f)
    result = read_docx(f)
    assert len(result.image_previews) == 1
    assert result.image_previews[0].thumbnail is False
    assert result.image_previews[0].content_type == "image/png"


def test_no_pillow_skips_oversized_original(tmp_path, monkeypatch):
    """Pillow 缺失且原图 > 上限 → 不产生条目（images 计数仍在）。"""
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "PIL" or name.startswith("PIL."):
            raise ImportError("no PIL in this test")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    monkeypatch.setattr(
        "backend.office.word._IMAGE_PREVIEW_RAW_LIMIT", 10
    )  # 任何真实 PNG 都超限

    f = tmp_path / "big-image.docx"
    _make_docx_with_image(f)
    result = read_docx(f)
    assert result.images == 1
    assert result.image_previews == []


def test_plain_docx_has_no_previews(tmp_path):
    from docx import Document

    f = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("无图")
    doc.save(str(f))
    result = read_docx(f)
    assert result.images == 0
    assert result.image_previews == []

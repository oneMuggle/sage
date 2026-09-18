"""Unit tests for Round 45 — 交叉引用占位符（{{fig:}}/{{tbl:}} → 图N/表N）。"""

from __future__ import annotations

import base64

import pytest
from docx import Document

from backend.office.errors import OfficeGenerateError
from backend.office.models import (
    OfficeWordGenerateRequest,
    WordFormatSpec,
    WordImageSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen(tmp_path, name, *, paragraphs, images=None, tables=None):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=paragraphs,
        images=images
 or [WordImageSpec(source=str(image_file), caption="架构图")],
        tables=tables
 or [WordTableSpec(headers=["A"], rows=[["1"]], caption="汇总表")],
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_cross_ref_placeholder_replaced_with_numbers(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "r.docx",
        paragraphs=[
            WordParagraphSpec(text="系统架构如{{fig:架构图}}所示，数据见{{tbl:汇总表}}。"),
        ],
    )
    doc = Document(str(output))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "如图1所示，数据见表1。" in body
    assert "{{" not in body


def test_cross_ref_placeholder_without_caption_fails_fast(tmp_path) -> None:
    with pytest.raises(OfficeGenerateError, match="cross_ref_not_found"):
        _gen(
            tmp_path,
            "bad.docx",
            paragraphs=[
                WordParagraphSpec(text="见 {{fig:不存在的图}}。"),
            ],
        )


def test_cross_ref_numbering_survives_insertion_order(tmp_path) -> None:
    """题注无题注不占号：第一张图无题注时，第二张图是图1。"""
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    output = _gen(
        tmp_path,
        "r2.docx",
        paragraphs=[WordParagraphSpec(text="见{{fig:部署图}}。")],
        images=[
            WordImageSpec(source=str(image_file)),
            WordImageSpec(source=str(image_file), caption="部署图", after_paragraph=1),
        ],
    )
    doc = Document(str(output))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "见图1。" in body


def test_lint_detects_cross_ref_residue(tmp_path) -> None:
    """残留占位符（手工编辑/外部导入）→ cross_ref/residue error。"""
    from backend.office.word_lint import lint_docx

    doc = Document()
    doc.add_paragraph("正文引用 {{fig:某图}} 是死文本")
    path = tmp_path / "manual.docx"
    doc.save(str(path))

    result = lint_docx(path, WordFormatSpec())
    rule_ids = {i.rule_id for i in result.issues}
    assert "cross_ref/residue" in rule_ids
    assert result.ok is False

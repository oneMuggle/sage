"""Unit tests for Round 53 — 分节页码格式/起始号（w:pgNumType）。"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import (
    OfficeWordGenerateRequest,
    WordFormatSpec,
    WordPageSetupSpec,
    WordParagraphSpec,
    WordSectionBreakSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx
from backend.office.word_lint import lint_docx

pytestmark = pytest.mark.unit


def _pgnum(doc: Document, section_index: int = 0):
    sections = doc.sections
    sect_pr = sections[section_index]._sectPr
    return sect_pr.find(qn("w:pgNumType"))


def test_main_section_pgnum_written(tmp_path) -> None:
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="r.docx",
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
        format_spec=WordFormatSpec(
            page=WordPageSetupSpec(
                page_number_format="lowerRoman", page_number_start=1
            )
        ),
    )
    output = generate_docx(req, output_dir=str(tmp_path))
    pg = _pgnum(Document(str(output)))
    assert pg is not None
    assert pg.get(qn("w:fmt")) == "lowerRoman"
    assert pg.get(qn("w:start")) == "1"


def test_section_break_pgnum_independent(tmp_path) -> None:
    """分节新节独立 fmt/start；主节不受影响。"""
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="r2.docx",
        title="报告",
        paragraphs=[
            WordParagraphSpec(text="前置"),
            WordParagraphSpec(text="正文", heading="h1"),
        ],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
        format_spec=WordFormatSpec(
            page=WordPageSetupSpec(
                page_number_format="lowerRoman", page_number_start=1
            ),
            section_breaks=[
                WordSectionBreakSpec(
                    start_paragraph=1,
                    page_setup=WordPageSetupSpec(
                        orientation="portrait",
                        page_number_format="decimal",
                        page_number_start=5,
                    ),
                )
            ],
        ),
    )
    output = generate_docx(req, output_dir=str(tmp_path))
    doc = Document(str(output))
    first = _pgnum(doc, 0)
    second = _pgnum(doc, 1)
    assert first.get(qn("w:fmt")) == "lowerRoman"
    assert second is not None
    assert second.get(qn("w:fmt")) == "decimal"
    assert second.get(qn("w:start")) == "5"


def test_lint_detects_pgnum_mismatch(tmp_path) -> None:
    """spec 声明 fmt/start 但文档缺失 → page/numbering error。"""
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="plain.docx",
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
    )
    output = generate_docx(req, output_dir=str(tmp_path))
    result = lint_docx(
        output,
        WordFormatSpec(
            page=WordPageSetupSpec(
                page_number_format="lowerRoman", page_number_start=1
            )
        ),
    )
    rule_ids = {i.rule_id for i in result.issues}
    assert "page/numbering" in rule_ids


def test_no_pgnum_spec_leaves_document_untouched(tmp_path) -> None:
    """未声明 pgNum → 不写 pgNumType（既有产物零变化）。"""
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename="none.docx",
        title="报告",
        paragraphs=[WordParagraphSpec(text="正文")],
        tables=[WordTableSpec(headers=["A"], rows=[["1"]])],
    )
    output = generate_docx(req, output_dir=str(tmp_path))
    assert _pgnum(Document(str(output))) is None

"""Unit tests for Round 57 — {{fn:}} 内联脚注（Phase A）。"""

from __future__ import annotations

import base64

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.office.models import (
    OfficeWordGenerateRequest,
    WordImageSpec,
    WordParagraphSpec,
    WordTableSpec,
)
from backend.office.word import generate_docx, read_docx

pytestmark = pytest.mark.unit

_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _gen(tmp_path, name, paragraphs, images=None, tables=None):
    image_file = tmp_path / "tiny.png"
    image_file.write_bytes(base64.b64decode(_TINY_PNG_B64))
    req = OfficeWordGenerateRequest(
        workspace_path="",
        filename=name,
        title="报告",
        paragraphs=paragraphs,
        images=images
        or [WordImageSpec(source=str(image_file), caption="架构图")],
        tables=tables or [WordTableSpec(headers=["A"], rows=[["1"]])],
    )
    return generate_docx(req, output_dir=str(tmp_path))


def test_footnote_marker_becomes_reference_and_part(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "fn.docx",
        paragraphs=[
            WordParagraphSpec(text="营收增长三成{{fn:数据来自内部台账。}}，趋势向好。"),
        ],
    )
    doc = Document(str(output))

    # footnotes part 已挂载
    partnames = [str(p.partname) for p in doc.part.package.iter_parts()]
    assert "/word/footnotes.xml" in partnames

    # 正文段含 footnoteReference run（id=1）
    refs = [
        el
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:footnoteReference"))
    ]
    assert len(refs) == 1
    assert refs[0].get(qn("w:id")) == "1"

    # 回读：脚注清单 + 引用标记不残留
    result = read_docx(output, workspace_path="")
    assert result.footnotes == ["数据来自内部台账。"]
    assert "{{fn:" not in "\n".join(p.text for p in doc.paragraphs)


def test_multiple_footnotes_numbered_in_order(tmp_path) -> None:
    output = _gen(
        tmp_path,
        "fn2.docx",
        paragraphs=[
            WordParagraphSpec(text="第一处{{fn:备注一}}与第二处{{fn:备注二}}。"),
        ],
    )
    result = read_docx(output, workspace_path="")
    assert result.footnotes == ["备注一", "备注二"]
    doc = Document(str(output))
    refs = [
        el
        for p in doc.paragraphs
        for el in p._p.findall(".//" + qn("w:footnoteReference"))
    ]
    assert [r.get(qn("w:id")) for r in refs] == ["1", "2"]


def test_no_footnote_no_part(tmp_path) -> None:
    """无 {{fn:}} → 不挂载 footnotes part（既有产物零变化）。"""
    output = _gen(
        tmp_path,
        "plain.docx",
        paragraphs=[WordParagraphSpec(text="普通段落")],
    )
    doc = Document(str(output))
    partnames = [str(p.partname) for p in doc.part.package.iter_parts()]
    assert "/word/footnotes.xml" not in partnames
    result = read_docx(output, workspace_path="")
    assert result.footnotes == []

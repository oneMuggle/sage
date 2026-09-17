"""Unit tests for backend.office.snapshot_diff (Round B P2 快照可视 diff).

用真实 python-docx / openpyxl / python-pptx 文件走全链路（与
test_artifact_reader 同策略），覆盖三种 doc_type 的增/删/改、
identical 短路、快照 id 穿越拒绝与永不 raise 契约。
"""

from __future__ import annotations

import shutil
import time
from pathlib import Path

import pytest

from backend.office.models import (
    OfficeDocStatus,
    OfficeDocType,
    OfficeDocumentMetadata,
    OfficeDocumentSummary,
)
from backend.office.snapshot_diff import diff_snapshot


def _summary(workspace: Path, doc_type: OfficeDocType, filename: str) -> OfficeDocumentSummary:
    """Build a summary whose managed layout dirs exist on disk."""
    doc_id = f"test-{doc_type.value}-0001"
    managed_dir = workspace / "office" / doc_type.value / doc_id
    managed_dir.mkdir(parents=True, exist_ok=True)
    now = int(time.time() * 1000)
    return OfficeDocumentSummary(
        id=doc_id,
        doc_type=doc_type,
        status=OfficeDocStatus.GENERATED,
        workspace_path=str(workspace),
        generated_filename=filename,
        created_at=now,
        updated_at=now,
        metadata=OfficeDocumentMetadata(file_size_bytes=0),
    )


def _snapshot_current(summary: OfficeDocumentSummary, snapshot_id: str) -> None:
    """Copy the current managed file into .snapshots/<snapshot_id>."""
    from backend.office.storage import document_path

    src = document_path(summary)
    snap_dir = src.parent / ".snapshots"
    snap_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, snap_dir / snapshot_id)


def _write_docx(path: Path, paragraphs) -> None:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


@pytest.fixture()
def ws(tmp_path: Path) -> Path:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    return workspace


def test_word_paragraph_add_change_remove(ws):
    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.WORD, "report.docx")
    target = document_path(summary)
    _write_docx(target, ["第一段", "第二段", "第三段"])
    snap_id = f"{int(time.time() * 1000)}-report.docx"
    _snapshot_current(summary, snap_id)
    # 当前版本：改第二段、删第三段、加第四段
    _write_docx(target, ["第一段", "第二段（改）", "新增第四段"])

    result = diff_snapshot(summary, snap_id)
    assert result.ok is True
    ops = {c.op for c in result.changes}
    assert "changed" in ops
    targets = [c.target for c in result.changes]
    assert any(t and t.startswith("paragraph[") for t in targets)


def test_word_identical_returns_empty(ws):
    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.WORD, "same.docx")
    target = document_path(summary)
    _write_docx(target, ["完全一致的内容"])
    snap_id = f"{int(time.time() * 1000)}-same.docx"
    _snapshot_current(summary, snap_id)

    result = diff_snapshot(summary, snap_id)
    assert result.ok is True
    assert result.changes == []


def test_excel_cell_and_sheet_changes(ws):
    from openpyxl import Workbook

    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.EXCEL, "data.xlsx")
    target = document_path(summary)
    wb = Workbook()
    sh = wb.active
    sh.title = "销售"
    sh.append(["城市", "金额"])
    sh.append(["北京", 100])
    wb.save(str(target))
    snap_id = f"{int(time.time() * 1000)}-data.xlsx"
    _snapshot_current(summary, snap_id)
    # 当前版本：改单元格 + 加 sheet
    wb2 = Workbook()
    sh2 = wb2.active
    sh2.title = "销售"
    sh2.append(["城市", "金额"])
    sh2.append(["北京", 999])
    wb2.create_sheet("汇总")
    wb2.save(str(target))

    result = diff_snapshot(summary, snap_id)
    assert result.ok is True
    assert any(c.op == "added" and c.summary == "新增工作表" for c in result.changes)
    changed = [c for c in result.changes if c.op == "changed"]
    assert any(c.target and "B2" in c.target for c in changed)


def test_ppt_title_and_slide_changes(ws):
    from pptx import Presentation

    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.PPT, "deck.pptx")
    target = document_path(summary)

    def make_deck(titles):
        prs = Presentation()
        for title in titles:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = title
        return prs

    make_deck(["开场", "章节一"]).save(str(target))
    snap_id = f"{int(time.time() * 1000)}-deck.pptx"
    _snapshot_current(summary, snap_id)
    make_deck(["开场（新）", "章节一", "总结"]).save(str(target))

    result = diff_snapshot(summary, snap_id)
    assert result.ok is True
    assert any(
        c.op == "changed" and c.target == "slide[0].title" for c in result.changes
    )
    assert any(c.op == "added" and c.summary == "新增幻灯片" for c in result.changes)


def test_missing_snapshot_folds_to_ok_false(ws):
    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.WORD, "doc.docx")
    _write_docx(document_path(summary), ["内容"])
    result = diff_snapshot(summary, "999999-doc.docx")
    assert result.ok is False
    assert "不存在" in (result.error or "")


def test_path_traversal_snapshot_id_rejected(ws):
    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.WORD, "doc.docx")
    _write_docx(document_path(summary), ["内容"])
    for bad in ("../evil.docx", "a/b.docx", "..\\\\up.docx"):
        result = diff_snapshot(summary, bad)
        assert result.ok is False
        assert "无效" in (result.error or "")


def test_never_raises_on_parse_failure(ws):
    from backend.office.storage import document_path

    summary = _summary(ws, OfficeDocType.WORD, "corrupt.docx")
    target = document_path(summary)
    target.write_bytes(b"this is not a real docx")
    snap_id = f"{int(time.time() * 1000)}-corrupt.docx"
    _snapshot_current(summary, snap_id)

    result = diff_snapshot(summary, snap_id)
    assert result.ok is False
    assert "解析失败" in (result.error or "")

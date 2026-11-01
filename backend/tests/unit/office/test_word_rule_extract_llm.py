"""Unit tests for LLM-assisted rule candidates (M3b)."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest
from docx import Document

from backend.office.word_rule_extract import extract_word_rules_async
from backend.office.word_rule_extract_llm import (
    build_clauses,
    parse_llm_rule_payload,
    propose_llm_candidates,
    validate_llm_params,
)


class FakeClient:
    def __init__(self, response: str) -> None:
        self.response = response
        self.prompts: list[str] = []

    async def complete(self, prompt: str) -> str:
        self.prompts.append(prompt)
        return self.response


class FailingClient:
    async def complete(self, prompt: str) -> str:
        raise RuntimeError("upstream timeout")


def _spec_doc(path: Path, texts: list[str]) -> Path:
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def _paragraphs(texts: list[str]) -> list[dict]:
    return [{"index": index, "text": text} for index, text in enumerate(texts)]


def test_build_clauses_is_bounded_and_stable() -> None:
    clauses = build_clauses(_paragraphs(["第一条", "", "第二条"]))

    assert [clause["clause_id"] for clause in clauses] == ["c1", "c2"]
    assert clauses[0]["paragraph_index"] == 0
    assert clauses[1]["paragraph_index"] == 2


def test_parse_payload_accepts_fences_and_rejects_junk() -> None:
    fenced = '```json\n[{"clause_id": "c1", "kind": "page.size"}]\n```'
    assert parse_llm_rule_payload(fenced) == [
        {"clause_id": "c1", "kind": "page.size"}
    ]
    with pytest.raises(ValueError, match="JSON 数组"):
        parse_llm_rule_payload("not json")
    with pytest.raises(ValueError, match="JSON 数组"):
        parse_llm_rule_payload('{"a": 1}')


def test_validate_params_enforces_allowlists() -> None:
    assert validate_llm_params("page.size", {"size": "a4"}) == {"size": "A4"}
    assert validate_llm_params("page.size", {"size": "A3"}) is None
    assert validate_llm_params("page.margins", {"top_cm": 3.7}) == {"top_cm": 3.7}
    assert validate_llm_params("page.margins", {"top_cm": 100}) is None
    assert validate_llm_params("style.font", {"size_pt": 16}) == {
        "style": "Normal",
        "size_pt": 16.0,
    }
    assert validate_llm_params("style.font", {"size_pt": 100}) is None
    assert validate_llm_params("structure.heading_sequence", {"levels": [0]}) is None
    assert validate_llm_params(
        "structure.required_headings", {"headings": ["摘要", "结论"]}
    ) == {"headings": ["摘要", "结论"], "match": "contains"}
    assert validate_llm_params("evil.kind", {}) is None


def test_propose_candidates_accepts_valid_and_rejects_invalid() -> None:
    response = """```json
[
  {"clause_id":"c1","kind":"structure.required_headings",
   "params":{"headings":["摘要","结论"]},"confidence":0.8,"reason":"条文要求"},
  {"clause_id":"c2","kind":"page.size",
   "params":{"size":"A3"},"confidence":0.9,"reason":"bad size"},
  {"clause_id":"c99","kind":"page.size",
   "params":{"size":"A4"},"confidence":0.9,"reason":"unknown clause"},
  {"clause_id":"c3","kind":"style.font",
   "params":{"font_east_asia":"仿宋","size_pt":16},"confidence":1.5}
]
```"""
    warnings: list[dict] = []
    candidates = asyncio.run(
        propose_llm_candidates(
            "spec.docx",
            _paragraphs(["第一条", "第二条", "第三条"]),
            client=FakeClient(response),
            warnings=warnings,
        )
    )

    assert {candidate.kind for candidate in candidates} == {
        "structure.required_headings",
        "style.font",
    }
    assert all("（LLM）" in candidate.source for candidate in candidates)
    assert max(candidate.confidence for candidate in candidates) <= 0.85
    assert any(warning["code"] == "llm_candidate_rejected" for warning in warnings)
    assert "2 条" in next(
        warning["message"]
        for warning in warnings
        if warning["code"] == "llm_candidate_rejected"
    )


def test_extract_async_merges_llm_candidates(tmp_path: Path) -> None:
    path = _spec_doc(tmp_path / "spec.docx", ["正文采用仿宋_GB2312 三号字。"])
    response = json.dumps(
        [
            {
                "clause_id": "c1",
                "kind": "structure.required_headings",
                "params": {"headings": ["摘要", "结论"]},
                "confidence": 0.8,
                "reason": "条文要求",
            }
        ],
        ensure_ascii=False,
    )
    client = FakeClient(response)

    result = asyncio.run(
        extract_word_rules_async(
            path,
            workspace_path=str(tmp_path),
            use_llm=True,
            llm_client=client,
        )
    )

    assert result.llm_used is True
    kinds = {rule.kind for rule in result.pack.rules}
    assert "style.font" in kinds  # deterministic
    assert "structure.required_headings" in kinds  # LLM
    llm_rule = next(
        rule for rule in result.pack.rules if rule.kind == "structure.required_headings"
    )
    assert "（LLM）" in llm_rule.source
    assert client.prompts
    assert "clause_id" in client.prompts[0]


def test_llm_unavailable_degrades(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from backend.orchestration import llm_factory

    monkeypatch.setattr(llm_factory, "build_llm_client_from_settings", lambda: None)
    path = _spec_doc(tmp_path / "spec.docx", ["正文采用仿宋_GB2312 三号字。"])

    result = asyncio.run(
        extract_word_rules_async(path, workspace_path=str(tmp_path), use_llm=True)
    )

    assert result.llm_used is False
    assert any(warning["code"] == "llm_unavailable" for warning in result.warnings)
    assert any(rule.kind == "style.font" for rule in result.pack.rules)


def test_llm_failure_degrades(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = _spec_doc(tmp_path / "spec.docx", ["正文采用仿宋_GB2312 三号字。"])

    result = asyncio.run(
        extract_word_rules_async(
            path,
            workspace_path=str(tmp_path),
            use_llm=True,
            llm_client=FailingClient(),
        )
    )

    assert result.llm_used is True
    assert any(warning["code"] == "llm_failed" for warning in result.warnings)
    assert any(rule.kind == "style.font" for rule in result.pack.rules)


def test_template_llm_is_unsupported(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("正文")
    path = tmp_path / "template.docx"
    document.save(str(path))
    client = FakeClient("[]")

    result = asyncio.run(
        extract_word_rules_async(
            path,
            file_type="template",
            workspace_path=str(tmp_path),
            use_llm=True,
            llm_client=client,
        )
    )

    assert result.llm_used is False
    assert not client.prompts
    assert any(
        warning["code"] == "llm_template_unsupported" for warning in result.warnings
    )

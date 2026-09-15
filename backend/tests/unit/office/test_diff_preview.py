# ruff: noqa: UP006, UP007, UP035 — release/win7 Python 3.8 兼容，保留 typing 注解
"""Unit tests for backend.office.diff_preview (Item 2.5 update preview).

Covers the contract the preview UI depends on:
- replace_text / set_cells produce human-readable before/after entries
  (including formula-string cells, via a data_only=False cell map);
- invalid ops fail the preview with ok=False + a WHY error instead of
  raising (a preview must show why the update would fail);
- the temp copy is always cleaned up and the source is never modified;
- the change list is capped at MAX_CHANGES with a truncated flag.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from backend.office.diff_preview import MAX_CHANGES, DiffPreviewResult, preview_update

pytestmark = pytest.mark.unit


# ──────────────────────────────────────────────────────────────────────
# Fixtures (real docx / xlsx / pptx files)
# ──────────────────────────────────────────────────────────────────────


@pytest.fixture()
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("The quick brown fox jumps over the lazy dog.")
    doc.add_paragraph("Second paragraph stands alone.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    doc.save(str(path))
    return path


@pytest.fixture()
def xlsx_path(tmp_path: Path) -> Path:
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws["A1"] = 1
    ws["B1"] = 2
    ws["A2"] = 3
    ws["B2"] = "=A1+A2"  # pre-existing formula cell
    wb.save(str(path))
    return path


@pytest.fixture()
def pptx_path(tmp_path: Path) -> Path:
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Quarterly Report"
    prs.save(str(path))
    return path


# ──────────────────────────────────────────────────────────────────────
# Word
# ──────────────────────────────────────────────────────────────────────


def test_replace_text_produces_before_after(docx_path: Path):
    result = preview_update(
        docx_path, [{"op": "replace_text", "find": "brown", "replace": "red"}]
    )
    assert isinstance(result, DiffPreviewResult)
    assert result.ok is True
    assert result.error is None
    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.op == "replace_text"
    assert change.target == "brown"
    assert "brown" in (change.before or "")
    assert "red" in (change.after or "")
    assert "brown" not in (change.after or "")
    # summary carries the replacement count from the editor's per-op result
    assert "1 occurrence" in (change.summary or "")


def test_preview_never_modifies_source_and_cleans_temp(docx_path: Path):
    result = preview_update(
        docx_path, [{"op": "replace_text", "find": "brown", "replace": "red"}]
    )
    assert result.ok is True
    # temp copy removed — the directory holds exactly the source file
    leftovers = [p.name for p in docx_path.parent.iterdir() if p != docx_path]
    assert leftovers == []    # source bytes unchanged: the paragraph still reads "brown"
    assert "The quick brown fox" in Document(str(docx_path)).paragraphs[0].text


def test_set_table_cell_shows_old_and_new(docx_path: Path):
    result = preview_update(
        docx_path,
        [{"op": "set_table_cell", "table_index": 0, "row": 1, "col": 1, "text": "42"}],
    )
    assert result.ok is True
    change = result.changes[0]
    assert change.op == "set_table_cell"
    assert change.target == "table[0].cell(1,1)"
    assert change.before == "1"
    assert change.after == "42"


def test_append_and_delete_paragraph_summaries(docx_path: Path):
    result = preview_update(
        docx_path,
        [
            {"op": "append_paragraphs", "paragraphs": [{"text": "Tail note"}]},
            {"op": "delete_paragraph", "find": "Second paragraph"},
        ],
    )
    assert result.ok is True
    append, delete = result.changes
    assert append.op == "append_paragraphs"
    assert append.target == "document end"
    assert "Tail note" in (append.after or "")
    assert delete.op == "delete_paragraph"
    assert delete.before is not None
    assert "Second paragraph" in delete.before
    assert "deletes 1 paragraph" in (delete.summary or "")


def test_invalid_op_returns_ok_false_with_error(docx_path: Path):
    result = preview_update(docx_path, [{"op": "replace_text", "find": "not-in-doc"}])
    assert result.ok is False
    assert result.error is not None
    assert "text_not_found" in result.error

    result = preview_update(docx_path, [{"op": "frobnicate"}])
    assert result.ok is False
    assert "unsupported_op" in (result.error or "")
    # and the source survived both failed previews
    assert "brown" in Document(str(docx_path)).paragraphs[0].text


def test_unsupported_extension_fails_cleanly(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")
    result = preview_update(path, [])
    assert result.ok is False
    assert "unsupported extension" in (result.error or "")


# ──────────────────────────────────────────────────────────────────────
# Excel
# ──────────────────────────────────────────────────────────────────────


def test_set_cells_per_cell_before_after_including_formulas(xlsx_path: Path):
    result = preview_update(
        xlsx_path,
        [
            {
                "op": "set_cells",
                "sheet": "Sheet",
                "cells": [
                    {"addr": "A1", "value": 10},
                    {"addr": "C1", "value": "=SUM(A1:B1)"},
                    {"addr": "B2", "value": 5},  # overwrites a formula cell
                ],
            }
        ],
    )
    assert result.ok is True
    by_target = {c.target: c for c in result.changes}
    assert by_target["Sheet!A1"].before == "1"
    assert by_target["Sheet!A1"].after == "10"
    # new cell → no before value; formula string passes through verbatim
    assert by_target["Sheet!C1"].before is None
    assert by_target["Sheet!C1"].after == "=SUM(A1:B1)"
    # the before side of a pre-existing formula shows the formula text
    assert by_target["Sheet!B2"].before == "=A1+A2"
    assert by_target["Sheet!B2"].after == "5"


def test_append_rows_and_sheet_ops_summaries(xlsx_path: Path):
    result = preview_update(
        xlsx_path,
        [
            {"op": "append_rows", "sheet": "Sheet", "rows": [[9, 8]]},
            {"op": "rename_sheet", "from": "Sheet", "to": "Data"},
            {"op": "add_sheet", "name": "Summary", "headers": ["k", "v"]},
        ],
    )
    assert result.ok is True
    append, rename, add = result.changes
    assert append.op == "append_rows"
    assert append.target == "Sheet"
    assert append.after == "9 | 8"
    assert rename.before == "Sheet"
    assert rename.after == "Data"
    assert add.target == "Summary"
    assert add.after == "k | v"


def test_changes_capped_at_200_with_truncated_flag(xlsx_path: Path):
    ops = [
        {
            "op": "set_cells",
            "sheet": "Sheet",
            "cells": [{"addr": f"A{i}", "value": i} for i in range(1, 251)],
        }
    ]
    result = preview_update(xlsx_path, ops)
    assert result.ok is True
    assert result.truncated is True
    assert len(result.changes) == MAX_CHANGES
    assert all(c.op == "set_cells" for c in result.changes)


# ──────────────────────────────────────────────────────────────────────
# PPT
# ──────────────────────────────────────────────────────────────────────


def test_ppt_replace_text_targets_slide(pptx_path: Path):
    result = preview_update(
        pptx_path, [{"op": "replace_text", "find": "Quarterly", "replace": "Annual"}]
    )
    assert result.ok is True
    change = result.changes[0]
    assert change.op == "replace_text"
    assert change.target == "slide[0]"
    assert "Quarterly" in (change.before or "")
    assert "Annual" in (change.after or "")


def test_ppt_set_slide_title_and_delete_slide(pptx_path: Path):
    result = preview_update(
        pptx_path,
        [
            {"op": "set_slide_title", "index": 0, "title": "Annual Report"},
            {"op": "delete_slide", "index": 0},
        ],
    )
    assert result.ok is True
    title, delete = result.changes
    assert title.target == "slide[0]"
    assert title.after == "Annual Report"
    assert "Quarterly Report" in (title.before or "")
    assert delete.op == "delete_slide"
    assert "deletes slide 0" in (delete.summary or "")


def test_batch2_style_op_without_dedicated_summary_becomes_generic_entry(
    xlsx_path: Path,
):
    """A valid op without a dedicated summarizer must not break the preview.

    edit.py's batch-2 style ops (set_column_width / add_chart / …) are
    described generically by op name + key args per the Item 2.5 spec.
    """
    result = preview_update(
        xlsx_path,
        [{"op": "set_column_width", "sheet": "Sheet", "column": "A", "width": 30}],
    )
    assert result.ok is True
    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.op == "set_column_width"
    assert "set_column_width" in (change.summary or "")
    assert "width=30" in (change.summary or "")

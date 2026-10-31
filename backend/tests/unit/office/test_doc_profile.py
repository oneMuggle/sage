"""Tests for the read-only Word document profile (M0)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend.office.doc_profile import read_doc_profile
from backend.office.errors import OfficeParseError


def _build_doc(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(0.9)
    section.header.paragraphs[0].text = "Header"

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    custom = document.styles.add_style("BodyCustom", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = normal
    custom.font.name = "Arial"
    custom.font.size = Pt(12)
    custom.font.bold = True

    paragraph = document.add_paragraph("Hello", style="BodyCustom")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.runs[0]
    run.bold = True
    run.italic = True
    run.underline = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    heading = document.add_heading("Title", level=1)
    heading.text = "Title"

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "merged"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"

    marked = document.add_paragraph("Bookmarked")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "42")
    start.set(qn("w:name"), "mark")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "42")
    marked._p.insert(0, start)
    marked._p.append(end)

    field_paragraph = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    field_paragraph._p.append(field)

    document.save(str(path))
    return path


def test_profile_extracts_sections_styles_and_formatting(tmp_path: Path) -> None:
    path = _build_doc(tmp_path / "sample.docx")

    profile = read_doc_profile(path, workspace_path=str(tmp_path))

    assert profile["source_path"] == "sample.docx"
    section = profile["sections"][0]
    assert section["page_width_pt"] == pytest.approx(612.0)
    assert section["page_height_pt"] == pytest.approx(792.0)
    assert section["margin_top_pt"] == pytest.approx(72.0)
    assert section["margin_bottom_pt"] == pytest.approx(57.6)
    assert section["margin_left_pt"] == pytest.approx(79.2)
    assert section["margin_right_pt"] == pytest.approx(64.8)
    assert section["header_text"] == "Header"

    styles = {item["style_id"]: item for item in profile["styles"]}
    assert styles["BodyCustom"]["based_on"] == "Normal"
    assert styles["BodyCustom"]["font_ascii"] == "Arial"
    assert styles["BodyCustom"]["size_pt"] == pytest.approx(12.0)
    assert styles["BodyCustom"]["bold"] is True

    paragraph = profile["paragraphs"][0]
    assert paragraph["style"] == "BodyCustom"
    assert paragraph["alignment"] == "center"
    assert paragraph["indent_left_pt"] == pytest.approx(18.0)
    assert paragraph["space_before_pt"] == pytest.approx(6.0)
    assert paragraph["line_spacing"] == pytest.approx(1.5)
    assert paragraph["keep_with_next"] is True
    run = paragraph["runs"][0]
    assert run["text"] == "Hello"
    assert run["bold"] is True
    assert run["italic"] is True
    assert run["underline"] is True
    assert run["font_ascii"] == "Times New Roman"
    assert run["font_east_asia"] == "宋体"
    assert run["size_pt"] == pytest.approx(14.0)
    assert run["color"] == "112233"
    assert run["anchor"]["char_start"] == 0
    assert run["anchor"]["char_end"] == 5
    assert len(run["anchor"]["text_hash"]) == 16

    title = next(item for item in profile["paragraphs"] if item["text"] == "Title")
    assert title["level"] == 1


def test_profile_extracts_table_merges_bookmarks_and_fields(tmp_path: Path) -> None:
    path = _build_doc(tmp_path / "sample.docx")

    profile = read_doc_profile(path, workspace_path=str(tmp_path))

    table = profile["tables"][0]
    assert table["rows"] == 2
    assert table["cols"] == 2
    first_row = [cell for cell in table["cells"] if cell["row"] == 0]
    assert len(first_row) == 1
    assert first_row[0]["col_span"] == 2

    bookmark = next(item for item in profile["bookmarks"] if item["name"] == "mark")
    assert bookmark["anchor"]["char_start"] == 0
    assert bookmark["end_anchor"]["char_start"] == 10

    assert any("PAGE" in item["instr"].upper() for item in profile["fields"])
    assert profile["summary"]["tables"] == 1
    assert profile["summary"]["bookmarks"] >= 1
    assert profile["summary"]["fields"] >= 1


def test_profile_marks_page_number_field(tmp_path: Path) -> None:
    document = Document()
    footer = document.sections[0].footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    footer._p.append(field)
    path = tmp_path / "page.docx"
    document.save(str(path))

    profile = read_doc_profile(path, workspace_path=str(tmp_path))

    assert profile["sections"][0]["has_page_number_field"] is True


def test_profile_rejects_non_docx(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("not docx", encoding="utf-8")

    with pytest.raises(OfficeParseError):
        read_doc_profile(path, workspace_path=str(tmp_path))

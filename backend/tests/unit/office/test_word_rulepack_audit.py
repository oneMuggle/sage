"""Unit tests for Word rule packs and the deterministic audit engine (M2a)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from backend.office.errors import OfficeFileNotFoundError, OfficeParseError
from backend.office.word_audit import (
    audit_word_document,
    supported_rule_kinds,
    validate_rulepack,
)
from backend.office.word_rulepack_builtin import builtin_rulepacks
from backend.office.word_rulepack_models import WordRulePack, WordRuleSpec
from backend.office.word_rulepack_store import (
    delete_rulepack,
    list_rulepacks,
    load_rulepack,
    save_rulepack,
)


def _rule(
    rule_id: str,
    kind: str,
    params: dict = None,
    *,
    category: str = "content",
    severity: str = "warning",
) -> WordRuleSpec:
    return WordRuleSpec(
        id=rule_id,
        title=rule_id,
        category=category,
        severity=severity,
        kind=kind,
        params=params or {},
        fixable=True,
        fix_hint="fix",
    )


def _pack(*rules: WordRuleSpec, pack_id: str = "test-pack") -> WordRulePack:
    return WordRulePack(id=pack_id, name="Test pack", rules=list(rules))


def _set_east_asia(style: object, name: str) -> None:
    rpr = style.element.get_or_add_rPr()  # type: ignore[attr-defined]
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def _add_page_number(document: Document) -> None:
    footer = document.sections[0].footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)


def _build_clean_doc(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    normal = document.styles["Normal"]
    normal.font.name = "仿宋_GB2312"
    normal.font.size = Pt(16)
    _set_east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.0)
    document.add_paragraph("正文内容")
    _add_page_number(document)
    document.save(str(path))
    return path


def _builtin(pack_id: str) -> WordRulePack:
    return load_rulepack(None, pack_id)["pack"]


def test_builtin_rulepacks_are_valid() -> None:
    packs = builtin_rulepacks()

    assert {pack.id for pack in packs} == {"gb-t-9704-general", "general-quality"}
    kinds = set(supported_rule_kinds())
    for pack in packs:
        validate_rulepack(pack)
        assert pack.rules
        assert {rule.kind for rule in pack.rules} <= kinds


def test_clean_document_passes_gb_pack(tmp_path: Path) -> None:
    path = _build_clean_doc(tmp_path / "clean.docx")

    result = audit_word_document(
        path, _builtin("gb-t-9704-general"), workspace_path=str(tmp_path)
    )

    assert result.ok is True
    assert result.issues == []
    assert result.summary.rules_checked == 5
    assert result.source_path == "clean.docx"


def test_audit_reports_page_style_issues_with_anchor(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("正文")
    path = tmp_path / "default.docx"
    document.save(str(path))

    result = audit_word_document(
        path, _builtin("gb-t-9704-general"), workspace_path=str(tmp_path)
    )

    rule_ids = {issue.rule_id for issue in result.issues}
    assert result.ok is False
    assert {"page/size", "page/margins", "page/page_number", "style/body_font"} <= rule_ids
    font_issue = next(issue for issue in result.issues if issue.rule_id == "style/body_font")
    assert font_issue.scope == "style"
    assert font_issue.anchor is not None
    assert font_issue.anchor.paragraph_index == 0
    assert font_issue.source.startswith("GB/T 9704")


def test_heading_sequence_and_required_headings_are_anchored(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("1 引言", level=1)
    document.add_heading("3 方法", level=1)
    document.add_heading("1.1 背景", level=2)
    document.add_paragraph("正文")
    path = tmp_path / "headings.docx"
    document.save(str(path))

    pack = _pack(
        _rule(
            "structure/heading_sequence",
            "structure.heading_sequence",
            {"levels": [1, 2], "require_prefix": True},
            category="structure",
        ),
        _rule(
            "structure/required_headings",
            "structure.required_headings",
            {"headings": ["结论"]},
            category="structure",
        ),
    )
    result = audit_word_document(path, pack, workspace_path=str(tmp_path))

    sequence = [issue for issue in result.issues if issue.rule_id == "structure/heading_sequence"]
    assert len(sequence) == 2
    assert sequence[0].anchor is not None
    assert sequence[0].anchor.paragraph_index == 1
    assert sequence[0].expected == "2"
    assert sequence[0].actual == "3"
    assert sequence[1].anchor is not None
    assert sequence[1].anchor.paragraph_index == 2

    required = [issue for issue in result.issues if issue.rule_id == "structure/required_headings"]
    assert len(required) == 1
    assert required[0].scope == "document"
    assert required[0].anchor is None


def test_placeholder_residue_and_empty_paragraphs(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("残留 {{name}} 和 <<todo>>")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("正文")
    path = tmp_path / "residue.docx"
    document.save(str(path))

    pack = _pack(
        _rule(
            "content/placeholder_residue",
            "content.placeholder_residue",
            severity="error",
        ),
        _rule(
            "content/empty_paragraphs",
            "content.empty_paragraphs",
            {"max_consecutive": 2},
        ),
    )
    result = audit_word_document(path, pack, workspace_path=str(tmp_path))

    residue = next(
        issue for issue in result.issues if issue.rule_id == "content/placeholder_residue"
    )
    assert residue.severity == "error"
    assert residue.anchor is not None
    assert residue.anchor.paragraph_index == 0
    assert "{{name}}" in residue.actual

    empty = next(
        issue for issue in result.issues if issue.rule_id == "content/empty_paragraphs"
    )
    assert empty.anchor is not None
    assert empty.anchor.paragraph_index == 1
    assert empty.actual == "3 个空段落"
    assert result.ok is False


def test_caption_and_citation_sequences(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("图1 示意图")
    document.add_paragraph("图3 错误编号")
    document.add_paragraph("正文引用[1]与[3]。")
    document.add_paragraph("表1 数据")
    path = tmp_path / "captions.docx"
    document.save(str(path))

    pack = _pack(
        _rule(
            "content/caption_sequence",
            "content.caption_sequence",
            {"labels": ["图", "表"]},
        ),
        _rule("content/citation_coverage", "content.citation_coverage", {}),
    )
    result = audit_word_document(path, pack, workspace_path=str(tmp_path))

    caption = next(
        issue for issue in result.issues if issue.rule_id == "content/caption_sequence"
    )
    assert caption.anchor is not None
    assert caption.anchor.paragraph_index == 1
    assert caption.expected == "2"
    assert caption.actual == "3"

    citation = next(
        issue for issue in result.issues if issue.rule_id == "content/citation_coverage"
    )
    assert citation.anchor is not None
    assert citation.anchor.paragraph_index == 2
    assert "缺失 [2]" in citation.message


def test_style_paragraph_checks_line_spacing_and_indent(tmp_path: Path) -> None:
    path = _build_clean_doc(tmp_path / "style.docx")
    pack = _pack(
        _rule(
            "style/paragraph",
            "style.paragraph",
            {"style": "Normal", "line_spacing": 2.0, "first_line_indent_cm": 0.5},
            category="style",
            severity="error",
        )
    )

    result = audit_word_document(path, pack, workspace_path=str(tmp_path))

    assert len(result.issues) == 2
    assert all(issue.scope == "style" for issue in result.issues)
    assert all(issue.anchor is not None for issue in result.issues)
    assert result.ok is False


def test_unknown_kind_duplicate_ids_and_bad_params_are_rejected(tmp_path: Path) -> None:
    path = _build_clean_doc(tmp_path / "doc.docx")

    with pytest.raises(OfficeParseError):
        validate_rulepack(_pack(_rule("x", "page.unknown", {})))
    with pytest.raises(OfficeParseError):
        validate_rulepack(
            _pack(
                _rule("same", "page.size", {"size": "A4"}),
                _rule("same", "page.size", {"size": "A4"}),
            )
        )
    with pytest.raises(OfficeParseError):
        audit_word_document(
            path,
            _pack(_rule("bad-size", "page.size", {"size": "A3"}, category="page")),
            workspace_path=str(tmp_path),
        )


def test_workspace_rulepack_crud_and_builtin_protection(tmp_path: Path) -> None:
    pack = _pack(
        _rule("content/placeholder_residue", "content.placeholder_residue"),
        pack_id="custom-pack",
    )

    save_rulepack(tmp_path, pack)
    listing = list_rulepacks(tmp_path)
    saved = next(item for item in listing["packs"] if item["id"] == "custom-pack")
    assert saved["origin"] == "workspace"
    assert saved["rule_count"] == 1

    loaded = load_rulepack(tmp_path, "custom-pack")
    assert loaded["origin"] == "workspace"
    assert loaded["pack"].rules[0].kind == "content.placeholder_residue"

    delete_rulepack(tmp_path, "custom-pack")
    assert not (tmp_path / ".sage" / "word-rulepacks" / "custom-pack.json").exists()

    with pytest.raises(OfficeParseError):
        save_rulepack(tmp_path, _builtin("gb-t-9704-general"))
    with pytest.raises(OfficeParseError):
        delete_rulepack(tmp_path, "gb-t-9704-general")
    with pytest.raises(OfficeFileNotFoundError):
        load_rulepack(tmp_path, "missing-pack")

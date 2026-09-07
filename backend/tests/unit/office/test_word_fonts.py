"""字体设置测试。"""
from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from backend.office.word import (
    DEFAULT_ASCII_FONT,
    DEFAULT_EA_FONT,
    set_doc_default_font,
)


def test_default_constants() -> None:
    """常量必须是 Times New Roman + 宋体。"""
    assert DEFAULT_ASCII_FONT == "Times New Roman"
    assert DEFAULT_EA_FONT == "宋体"


def test_set_doc_default_font_writes_styles_xml() -> None:
    """调用后 styles.xml 的 Normal.rPr.rFonts 出现 eastAsia=宋体。"""
    doc = Document()
    set_doc_default_font(doc, DEFAULT_ASCII_FONT, DEFAULT_EA_FONT)

    normal = doc.styles["Normal"]
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert rFonts.get(qn("w:eastAsia")) == "宋体"


def test_set_doc_default_font_overrides_existing() -> None:
    """重复调用以最后一次为准。"""
    doc = Document()
    set_doc_default_font(doc, "Arial", "微软雅黑")
    set_doc_default_font(doc, DEFAULT_ASCII_FONT, DEFAULT_EA_FONT)
    normal = doc.styles["Normal"]
    rFonts = normal.element.get_or_add_rPr().find(qn("w:rFonts"))
    assert rFonts.get(qn("w:eastAsia")) == "宋体"


def test_set_doc_default_font_patches_title_style() -> None:
    """Title 样式必须也写入 eastAsia，否则日语 Word 默认 theme 会渲染
    为 ＭＳ ゴシック。这是 L-2 手测关注点的修复。"""
    doc = Document()
    set_doc_default_font(doc, DEFAULT_ASCII_FONT, DEFAULT_EA_FONT)
    title = doc.styles["Title"]
    rFonts = title.element.get_or_add_rPr().find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert rFonts.get(qn("w:eastAsia")) == "宋体"
    # theme 引用必须清掉，否则优先级压过显式 eastAsia
    assert rFonts.get(qn("w:asciiTheme")) is None
    assert rFonts.get(qn("w:eastAsiaTheme")) is None


def test_set_doc_default_font_patches_heading_styles() -> None:
    """Heading 1/2/3 样式必须也写入 eastAsia，理由同 Title。"""
    doc = Document()
    set_doc_default_font(doc, DEFAULT_ASCII_FONT, DEFAULT_EA_FONT)
    for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
        heading = doc.styles[heading_name]
        rFonts = heading.element.get_or_add_rPr().find(qn("w:rFonts"))
        assert rFonts is not None, f"{heading_name} missing rFonts"
        assert rFonts.get(qn("w:eastAsia")) == "宋体", (
            f"{heading_name}.eastAsia != 宋体"
        )
        assert rFonts.get(qn("w:eastAsiaTheme")) is None, (
            f"{heading_name} still has eastAsiaTheme ref"
        )


def test_set_doc_default_font_patches_list_styles() -> None:
    """List Bullet / List Number 也写入 eastAsia（font_family 透传时
    bullet/numbered 段落也用用户字体）。"""
    doc = Document()
    set_doc_default_font(doc, "Arial", "微软雅黑")
    for list_name in ("List Bullet", "List Number"):
        list_style = doc.styles[list_name]
        rFonts = list_style.element.get_or_add_rPr().find(qn("w:rFonts"))
        assert rFonts is not None, f"{list_name} missing rFonts"
        assert rFonts.get(qn("w:eastAsia")) == "微软雅黑"


def test_set_doc_default_font_propagates_font_family_to_all_styles() -> None:
    """font_family='微软雅黑' 时 Title / Heading 1 / List Bullet 也都是
    微软雅黑（不仅仅是 Normal）。这是用户报告 D 文件标题仍是 ＭＳ ゴシック
    的根因修复。"""
    doc = Document()
    set_doc_default_font(doc, "Arial", "微软雅黑")
    for style_name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "List Bullet",
        "List Number",
    ):
        style = doc.styles[style_name]
        rFonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        assert rFonts.get(qn("w:eastAsia")) == "微软雅黑", style_name

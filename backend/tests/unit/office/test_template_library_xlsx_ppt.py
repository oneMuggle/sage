"""Unit tests for xlsx/ppt template support in the 模板库 (round 3 — N2).

Covers:
- builtin registry: BUILTIN_TEMPLATES stays the 6 Word entries;
  ALL_BUILTIN_TEMPLATES adds the 2 excel + 2 ppt templates
- xlsx builders reopen via openpyxl: literal {{markers}}, freeze panes,
  a real =SUM() formula, header styling, number-format columns
- pptx builders reopen via python-pptx: 3 slides each, literal {{markers}}
- instantiate (excel): marker replace into office/excel/<uuid>/, formula kept,
  frozen panes kept, unfilled markers stay literal and are listed
- instantiate (ppt): slide count / marker replace / multiline values become
  real paragraphs; table-cell markers fill too
- workspace user templates: .xlsx/.pptx in office/templates/ are listed with
  the right doc_type and instantiated (by filename or bare stem); broken
  files are skipped
- response-shape contract: filename extension must match the template's
  doc_type (word templates still reject .xlsx outputs and vice versa)
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from backend.office.errors import OfficePathError
from backend.office.models import OfficeDocType
from backend.office.template_library import (
    ALL_BUILTIN_TEMPLATES,
    BUILTIN_EXCEL_TEMPLATES,
    BUILTIN_PPT_TEMPLATES,
    BUILTIN_TEMPLATES,
    instantiate_template,
    list_templates,
)

# ──────────────────────────────────────────────────────────────────────
# builtin registry
# ──────────────────────────────────────────────────────────────────────


def test_builtin_registry_splits_by_doc_type():
    # BUILTIN_TEMPLATES stays the Word-only registry (existing callers rely on it)
    assert len(BUILTIN_TEMPLATES) == 6
    assert all(spec.doc_type is OfficeDocType.WORD for spec in BUILTIN_TEMPLATES)
    assert len(BUILTIN_EXCEL_TEMPLATES) == 2
    assert all(spec.doc_type is OfficeDocType.EXCEL for spec in BUILTIN_EXCEL_TEMPLATES)
    assert len(BUILTIN_PPT_TEMPLATES) == 2
    assert all(spec.doc_type is OfficeDocType.PPT for spec in BUILTIN_PPT_TEMPLATES)
    assert len(ALL_BUILTIN_TEMPLATES) == 10
    ids = [spec.id for spec in ALL_BUILTIN_TEMPLATES]
    assert len(set(ids)) == len(ids)


def test_builtin_office_ids_names_and_doc_types():
    by_id = {spec.id: spec for spec in ALL_BUILTIN_TEMPLATES}
    assert by_id["budget_sheet"].name == "预算表"
    assert by_id["budget_sheet"].doc_type is OfficeDocType.EXCEL
    assert by_id["inventory"].name == "库存台账"
    assert by_id["inventory"].doc_type is OfficeDocType.EXCEL
    assert by_id["kickoff_deck"].name == "项目启动会"
    assert by_id["kickoff_deck"].doc_type is OfficeDocType.PPT
    assert by_id["weekly_sync"].name == "周例会"
    assert by_id["weekly_sync"].doc_type is OfficeDocType.PPT


def test_list_templates_includes_office_builtins():
    response = list_templates(None)
    assert [t.source for t in response.templates].count("builtin") == 10
    budget = next(t for t in response.templates if t.id == "budget_sheet")
    assert budget.doc_type is OfficeDocType.EXCEL
    assert budget.filename is None
    assert {p.name for p in budget.placeholders} == {
        "title",
        "owner",
        "period",
        "total",
        "item1",
        "item2",
        "item3",
    }
    deck = next(t for t in response.templates if t.id == "kickoff_deck")
    assert deck.doc_type is OfficeDocType.PPT
    assert {p.name for p in deck.placeholders} == {
        "project",
        "presenter",
        "date",
        "agenda_items",
        "roles",
    }


# ──────────────────────────────────────────────────────────────────────
# xlsx builtin builders
# ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "spec", BUILTIN_EXCEL_TEMPLATES, ids=[spec.id for spec in BUILTIN_EXCEL_TEMPLATES]
)
def test_builtin_xlsx_build_reopens_with_markers(spec):
    path = spec.build()
    try:
        workbook = load_workbook(str(path))
        all_text = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        all_text.append(cell.value)
        joined = "\n".join(all_text)
        for placeholder in spec.placeholders:
            assert "{{" + placeholder.name + "}}" in joined
        for worksheet in workbook.worksheets:
            assert worksheet.freeze_panes, f"{worksheet.title} not frozen"
    finally:
        path.unlink(missing_ok=True)


def test_budget_sheet_build_layout():
    spec = {s.id: s for s in ALL_BUILTIN_TEMPLATES}["budget_sheet"]
    path = spec.build()
    try:
        workbook = load_workbook(str(path))
        assert [sheet.title for sheet in workbook.worksheets] == ["概览", "明细"]
        detail = workbook["明细"]
        assert [detail.cell(row=1, column=col).value for col in range(1, 5)] == [
            "项目",
            "类别",
            "金额",
            "备注",
        ]
        assert detail["C5"].value == "=SUM(C2:C4)"
        assert detail["C5"].font.bold
        assert detail["A1"].fill.start_color.rgb.endswith("4472C4")
    finally:
        path.unlink(missing_ok=True)


def test_inventory_build_number_formats():
    spec = {s.id: s for s in ALL_BUILTIN_TEMPLATES}["inventory"]
    path = spec.build()
    try:
        workbook = load_workbook(str(path))
        sheet = workbook.active
        assert sheet["A1"].value == "{{warehouse}}库存台账"
        assert sheet.freeze_panes == "A3"
        assert sheet["E3"].number_format == "0"
        assert sheet["F3"].number_format == "#,##0.00"
        assert sheet["H3"].number_format == "yyyy-mm-dd"
    finally:
        path.unlink(missing_ok=True)


# ──────────────────────────────────────────────────────────────────────
# pptx builtin builders
# ──────────────────────────────────────────────────────────────────────


def _deck_frame_texts(prs):
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    return texts


@pytest.mark.parametrize(
    "spec",
    BUILTIN_PPT_TEMPLATES,
    ids=[spec.id for spec in BUILTIN_PPT_TEMPLATES],
)
def test_builtin_ppt_build_reopens_with_markers(spec):
    path = spec.build()
    try:
        prs = Presentation(str(path))
        assert len(prs.slides) == 3
        joined = "\n".join(_deck_frame_texts(prs))
        for placeholder in spec.placeholders:
            assert "{{" + placeholder.name + "}}" in joined
    finally:
        path.unlink(missing_ok=True)


# ──────────────────────────────────────────────────────────────────────
# excel 实例化（instantiate 路径）
# ──────────────────────────────────────────────────────────────────────

BUDGET_DATA = {
    "title": "启明星 Q3 预算",
    "owner": "张三",
    "period": "2026 年 Q3",
    "total": "120,000 元",
    "item1": "云资源",
    "item2": "外包测试",
    "item3": "差旅",
}


def test_instantiate_budget_sheet_end_to_end(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path), template_id="budget_sheet", filename="Q3预算", data=BUDGET_DATA
    )

    output = Path(result.output_path)
    # managed layout: <workspace>/office/excel/<uuid>/<filename>
    assert output.parent.parent == tmp_path / "office" / "excel"
    assert output.name == "Q3预算.xlsx"
    assert result.filename == "Q3预算.xlsx"
    assert result.file_size_bytes == output.stat().st_size
    assert result.filled_count == 7
    assert result.unfilled_placeholders == []

    workbook = load_workbook(str(output))
    overview = workbook["概览"]
    assert overview["A1"].value == "启明星 Q3 预算"
    assert overview["B2"].value == "张三"
    assert overview["B3"].value == "2026 年 Q3"
    assert overview["B4"].value == "120,000 元"
    # 真实公式原样保留
    detail = workbook["明细"]
    assert detail["C5"].value == "=SUM(C2:C4)"
    assert detail["A2"].value == "云资源"
    assert detail.freeze_panes == "A2"
    # 没有任何标记残留
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and "{{" in cell.value)


def test_instantiate_budget_sheet_lists_unfilled_marker(tmp_path: Path):
    partial = {key: value for key, value in BUDGET_DATA.items() if key != "total"}
    result = instantiate_template(
        str(tmp_path), template_id="budget_sheet", filename="预算.xlsx", data=partial
    )
    assert result.filled_count == 6
    assert result.unfilled_placeholders == ["total"]

    workbook = load_workbook(str(result.output_path))
    # 缺数据的标记保持字面量原样
    assert workbook["概览"]["B4"].value == "{{total}}"


def test_instantiate_inventory_end_to_end(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path),
        template_id="inventory",
        filename="华东台账.xlsx",
        data={"warehouse": "华东一号仓"},
    )
    assert result.filled_count == 1
    assert result.unfilled_placeholders == []

    sheet = load_workbook(str(result.output_path)).active
    assert sheet["A1"].value == "华东一号仓库存台账"
    assert sheet.freeze_panes == "A3"
    # 数字格式示例列在实例化后仍然保留
    assert sheet["F3"].number_format == "#,##0.00"


# ──────────────────────────────────────────────────────────────────────
# ppt 实例化（instantiate 路径）
# ──────────────────────────────────────────────────────────────────────

DECK_DATA = {
    "project": "启明星",
    "presenter": "李四",
    "date": "2026-09-09",
    "agenda_items": "1. 项目背景\n2. 里程碑\n3. 分工",
    "roles": "后端：张三\n前端：王五",
}


def test_instantiate_kickoff_deck_end_to_end(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path), template_id="kickoff_deck", filename="启动会", data=DECK_DATA
    )

    output = Path(result.output_path)
    assert output.parent.parent == tmp_path / "office" / "ppt"
    assert output.name == "启动会.pptx"
    assert result.filled_count == 5
    assert result.unfilled_placeholders == []

    prs = Presentation(str(output))
    assert len(prs.slides) == 3
    title_texts = _deck_frame_texts(prs)
    joined = "\n".join(title_texts)
    assert "启明星 项目启动会" in joined
    assert "汇报人：李四" in joined
    assert "日期：2026-09-09" in joined
    assert "{{" not in joined


def test_instantiate_kickoff_deck_multiline_becomes_paragraphs(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path),
        template_id="kickoff_deck",
        filename="deck.pptx",
        data={"agenda_items": "1. 项目背景\n2. 里程碑\n3. 分工", "roles": "待定"},
    )
    prs = Presentation(str(result.output_path))
    agenda_frames = [
        shape.text_frame
        for shape in prs.slides[1].shapes
        if shape.has_text_frame and "里程碑" in shape.text_frame.text
    ]
    assert len(agenda_frames) == 1
    frame = agenda_frames[0]
    assert [p.text for p in frame.paragraphs] == ["1. 项目背景", "2. 里程碑", "3. 分工"]

    # 未提供的 {{project}}/{{presenter}}/{{date}} 保持字面量并被列出
    assert sorted(result.unfilled_placeholders) == ["date", "presenter", "project"]
    first_slide_text = "\n".join(_deck_frame_texts(prs))
    assert "{{project}} 项目启动会" in first_slide_text


def test_instantiate_kickoff_deck_lists_unfilled_marker(tmp_path: Path):
    data = dict(DECK_DATA)
    del data["roles"]
    result = instantiate_template(
        str(tmp_path), template_id="kickoff_deck", filename="deck.pptx", data=data
    )
    assert result.filled_count == 4
    assert result.unfilled_placeholders == ["roles"]
    output_texts = "\n".join(_deck_frame_texts(Presentation(str(result.output_path))))
    assert "{{roles}}" in output_texts


def test_instantiate_weekly_sync_end_to_end(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path),
        template_id="weekly_sync",
        filename="周例会.pptx",
        data={"week": "W37", "done": "模板库上线", "next": "联调", "risks": "无"},
    )
    assert result.filled_count == 4
    prs = Presentation(str(result.output_path))
    joined = "\n".join(_deck_frame_texts(prs))
    assert "周例会（W37）" in joined
    assert "模板库上线" in joined
    assert "{{" not in joined


def test_instantiate_pptx_fills_table_cell_marker(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, 914400, 914400, 4572000, 1828800).table
    table.cell(0, 0).text = "成员"
    table.cell(1, 0).text = "{{member}}"
    prs.save(str(templates_dir / "roster.pptx"))

    result = instantiate_template(
        str(tmp_path),
        workspace_template="roster.pptx",
        filename="roster-filled.pptx",
        data={"member": "赵六"},
    )
    assert result.filled_count == 1
    prs_out = Presentation(str(result.output_path))
    table_out = next(
        shape.table
        for shape in prs_out.slides[0].shapes
        if shape.has_table
    )
    assert table_out.cell(1, 0).text == "赵六"


# ──────────────────────────────────────────────────────────────────────
# workspace user templates (.xlsx / .pptx)
# ──────────────────────────────────────────────────────────────────────


def test_workspace_xlsx_template_listed_and_instantiated(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    workbook = Workbook()
    workbook.active["A1"] = "你好，{{name}}"
    workbook.active["A2"] = "=1+1"
    workbook.save(str(templates_dir / "team.xlsx"))

    response = list_templates(str(tmp_path))
    entry = next(t for t in response.templates if t.filename == "team.xlsx")
    assert entry.source == "workspace"
    assert entry.doc_type is OfficeDocType.EXCEL
    assert entry.id.startswith("ws_")
    assert {p.name for p in entry.placeholders} == {"name"}

    result = instantiate_template(
        str(tmp_path),
        workspace_template="team.xlsx",
        filename="团队名录",
        data={"name": "张三"},
    )
    output = Path(result.output_path)
    assert output.parent.parent == tmp_path / "office" / "excel"
    assert output.name == "团队名录.xlsx"
    sheet = load_workbook(str(output)).active
    assert sheet["A1"].value == "你好，张三"
    assert sheet["A2"].value == "=1+1"


def test_workspace_template_without_extension_resolves_xlsx(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    workbook = Workbook()
    workbook.active["A1"] = "{{name}}"
    workbook.save(str(templates_dir / "team.xlsx"))

    result = instantiate_template(
        str(tmp_path), workspace_template="team", filename="out.xlsx", data={"name": "ok"}
    )
    assert result.filled_count == 1


def test_workspace_pptx_template_listed_and_instantiated(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(914400, 914400, 4572000, 914400).text_frame.text = (
        "主题：{{topic}}"
    )
    prs.save(str(templates_dir / "sync.pptx"))

    response = list_templates(str(tmp_path))
    entry = next(t for t in response.templates if t.filename == "sync.pptx")
    assert entry.doc_type is OfficeDocType.PPT
    assert {p.name for p in entry.placeholders} == {"topic"}

    result = instantiate_template(
        str(tmp_path),
        workspace_template="sync.pptx",
        filename="sync-filled.pptx",
        data={"topic": "排期"},
    )
    texts = "\n".join(
        shape.text_frame.text
        for shape in Presentation(str(result.output_path)).slides[0].shapes
        if shape.has_text_frame
    )
    assert "主题：排期" in texts


def test_broken_workspace_xlsx_is_skipped_not_fatal(tmp_path: Path):
    templates_dir = tmp_path / "office" / "templates"
    templates_dir.mkdir(parents=True)
    workbook = Workbook()
    workbook.active["A1"] = "{{name}}"
    workbook.save(str(templates_dir / "good.xlsx"))
    (templates_dir / "broken.xlsx").write_text("not a zip at all")

    response = list_templates(str(tmp_path))
    workspace_entries = [t for t in response.templates if t.source == "workspace"]
    assert [t.filename for t in workspace_entries] == ["good.xlsx"]


# ──────────────────────────────────────────────────────────────────────
# response-shape contract (doc_type-aware filename validation)
# ──────────────────────────────────────────────────────────────────────


def test_instantiate_rejects_mismatched_output_extension(tmp_path: Path):
    # excel 模板不接受 .docx 输出名
    with pytest.raises(OfficePathError):
        instantiate_template(
            str(tmp_path),
            template_id="budget_sheet",
            filename="预算.docx",
            data=BUDGET_DATA,
        )
    # word 模板也不接受 .xlsx 输出名
    with pytest.raises(OfficePathError):
        instantiate_template(
            str(tmp_path), template_id="weekly_report", filename="周报.xlsx", data={}
        )


def test_instantiate_unknown_office_template_id_raises(tmp_path: Path):
    from backend.office.errors import OfficeFileNotFoundError

    with pytest.raises(OfficeFileNotFoundError):
        instantiate_template(
            str(tmp_path), template_id="no_such_sheet", filename="x.xlsx", data={}
        )


def test_word_templates_still_fill_after_dispatch_change(tmp_path: Path):
    result = instantiate_template(
        str(tmp_path),
        template_id="weekly_report",
        filename="周报.docx",
        data={
            "title": "t",
            "author": "a",
            "week_range": "w",
            "highlights": "h",
            "next_plan": "n",
            "risks": "r",
        },
    )
    text = "\n".join(p.text for p in Document(str(result.output_path)).paragraphs)
    assert "{{" not in text

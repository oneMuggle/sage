"""Unit tests for the PDF → Word converter (round-3 N3, pdf_to_word.py).

The fixture PDFs are built with PyMuPDF itself (write-then-read testing,
same convention as the rest of the office suite): page 1 carries a 28pt
title (→ Heading 1), two 11pt body lines and a vector-ruled 2x2 table;
page 2 carries a 16pt section title (→ Heading 2) and one body line.
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest

from backend.office import pdf_to_word
from backend.office.errors import OfficeFileNotFoundError
from backend.office.pdf_to_word import PdfToWordRequest, convert_pdf_to_word

pytest.importorskip("docx")


def _write_test_pdf(path: Path) -> Path:
    """Two pages: big title + body + ruled 2x2 table; section title + body."""
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 100), "Quarterly Report", fontsize=28, fontname="helv")
    page.insert_text(
        (72, 140), "This is the first body line of the document.", fontsize=11, fontname="helv"
    )
    page.insert_text(
        (72, 160), "A second body line follows right here.", fontsize=11, fontname="helv"
    )
    # 2x2 ruled grid (find_tables detects vector rules).
    xs = (72, 272, 472)
    ys = (300, 340, 380)
    for x in xs:
        page.draw_line(pymupdf.Point(x, ys[0]), pymupdf.Point(x, ys[-1]))
    for y in ys:
        page.draw_line(pymupdf.Point(xs[0], y), pymupdf.Point(xs[-1], y))
    page.insert_text((82, 330), "Alpha", fontsize=11, fontname="helv")
    page.insert_text((282, 330), "Beta", fontsize=11, fontname="helv")
    page.insert_text((82, 370), "Gamma", fontsize=11, fontname="helv")
    page.insert_text((282, 370), "Delta", fontsize=11, fontname="helv")
    # Page 2: 16pt section title (11 * 1.3 <= 16 < 11 * 1.6 → Heading 2).
    page2 = pdf.new_page(width=595, height=842)
    page2.insert_text((72, 100), "Section Two", fontsize=16, fontname="helv")
    page2.insert_text((72, 130), "Content of the second page.", fontsize=11, fontname="helv")
    pdf.save(str(path))
    pdf.close()
    return path


@pytest.fixture()
def source_pdf(tmp_path: Path) -> Path:
    """Source PDF inside a fresh workspace directory."""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    return _write_test_pdf(workspace / "report.pdf")


# ──────────────────────────────────────────────────────────────────────
# Service function
# ──────────────────────────────────────────────────────────────────────


def test_convert_rebuilds_headings_body_and_table(source_pdf: Path, tmp_path: Path):
    from docx import Document

    workspace = source_pdf.parent
    result = convert_pdf_to_word(source_pdf, workspace, "converted.docx")

    assert result.ok is True, result.error
    assert result.error is None
    output = Path(result.output_path)
    assert output.is_file()
    assert output == workspace / "office" / "word" / output.parent.name / "converted.docx"
    assert result.filename == "converted.docx"
    assert result.file_size_bytes == output.stat().st_size > 0
    assert result.page_count == 2
    assert result.table_count == 1
    assert result.image_count == 0
    assert result.paragraph_count >= 3  # title + two body lines (page 1 alone)

    doc = Document(str(output))
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    # 标题启发式：28pt ≥ body(11) × 1.6 → Heading 1；16pt ≥ 11 × 1.3 → Heading 2。
    assert "Heading 1" in styles
    assert "Heading 2" in styles
    assert "Quarterly Report" in texts
    assert "Section Two" in texts
    # 表格单元格文本不再以正文段落重复出现（de-duplication）。
    assert "Alpha" not in texts
    # 表格本体转成 docx 表格。
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Alpha"
    assert table.cell(1, 1).text == "Delta"


def test_convert_default_filename_uses_pdf_stem(source_pdf: Path):
    result = convert_pdf_to_word(source_pdf, source_pdf.parent, "")
    assert result.ok is True
    assert result.filename == "report.docx"
    assert Path(result.output_path).name == "report.docx"


def test_convert_missing_source_returns_error(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    result = convert_pdf_to_word(workspace / "ghost.pdf", workspace, "out.docx")
    assert result.ok is False
    assert result.error  # 人类可读的失败原因
    assert result.output_path is None


def test_convert_source_outside_workspace_rejected(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    outside = tmp_path / "outside.pdf"
    _write_test_pdf(outside)
    result = convert_pdf_to_word(outside, workspace, "out.docx")
    assert result.ok is False
    assert result.output_path is None


def test_convert_oversized_pdf_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """超过 MAX_PDF_SIZE 的源 PDF 在 fitz 打开前就被拒绝（守卫常量可注入）。"""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = _write_test_pdf(workspace / "big.pdf")
    monkeypatch.setattr(pdf_to_word, "MAX_PDF_SIZE", 10)  # 10 bytes → 任何 PDF 都超限
    result = convert_pdf_to_word(source, workspace, "out.docx")
    assert result.ok is False
    assert "size limit" in result.error
    assert result.output_path is None


def test_convert_invalid_workspace_rejected(tmp_path: Path):
    source = _write_test_pdf(tmp_path / "report.pdf")
    result = convert_pdf_to_word(source, tmp_path / "no-such-workspace", "out.docx")
    assert result.ok is False


def test_convert_scanned_pdf_yields_empty_document(tmp_path: Path):
    """无文本层的扫描件：转换成功但正文为空（OCR 是非目标，明确降级）。"""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "scanned.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(str(source))
    pdf.close()
    result = convert_pdf_to_word(source, workspace, "out.docx")
    assert result.ok is True
    assert result.paragraph_count == 0
    assert result.table_count == 0
    assert result.page_count == 1


# ──────────────────────────────────────────────────────────────────────
# Route (POST /office/pdf/to-word) — persistence + lineage
# ──────────────────────────────────────────────────────────────────────


def _seed_managed_pdf(workspace: Path, doc_id: str) -> Path:
    """Create a managed source PDF + its DB row; return the file path."""
    from backend.data.database import get_database
    from backend.office.models import (
        OfficeDocStatus,
        OfficeDocType,
        OfficeDocumentMetadata,
        OfficeDocumentSummary,
    )
    from backend.office.storage import generate_document_dir, save_document

    directory = generate_document_dir(workspace, OfficeDocType.PDF, doc_id)
    managed = directory / "source.pdf"
    _write_test_pdf(managed)
    now = 1_700_000_000_000
    save_document(
        get_database().get_connection(),
        OfficeDocumentSummary(
            id=doc_id,
            workspace_path=str(workspace.resolve()),
            doc_type=OfficeDocType.PDF,
            original_filename=None,
            generated_filename="source.pdf",
            status=OfficeDocStatus.PARSED,
            created_at=now,
            updated_at=now,
            metadata=OfficeDocumentMetadata(file_size_bytes=managed.stat().st_size),
        ),
    )
    return managed


def test_route_persists_word_row_with_derived_from(tmp_path: Path):
    from backend.api.office_routes import pdf_to_word_endpoint
    from backend.data.database import get_database

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    managed = _seed_managed_pdf(workspace, "pdf-src")

    result = pdf_to_word_endpoint(
        PdfToWordRequest(
            workspace_path=str(workspace),
            file_path=str(managed),
            out_filename="converted.docx",
            source_doc_id="pdf-src",
        )
    )
    assert result.ok is True, result.error

    row = get_database().get_connection().execute(
        "SELECT doc_type, status, derived_from FROM office_documents "
        "WHERE generated_filename = 'converted.docx'",
    ).fetchone()
    assert row is not None
    assert row["doc_type"] == "word"
    assert row["status"] == "generated"
    assert row["derived_from"] == "pdf-src"


def test_route_without_source_doc_id_leaves_lineage_null(tmp_path: Path):
    from backend.api.office_routes import pdf_to_word_endpoint
    from backend.data.database import get_database

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    managed = _seed_managed_pdf(workspace, "pdf-src-2")

    result = pdf_to_word_endpoint(
        PdfToWordRequest(workspace_path=str(workspace), file_path=str(managed))
    )
    assert result.ok is True, result.error
    assert result.filename == "source.docx"  # 默认名 = 源 PDF stem + .docx
    row = get_database().get_connection().execute(
        "SELECT derived_from FROM office_documents WHERE generated_filename = 'source.docx'",
    ).fetchone()
    assert row is not None
    assert row["derived_from"] is None


def test_route_unknown_source_doc_id_raises_404_error(tmp_path: Path):
    from backend.api.office_routes import pdf_to_word_endpoint

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    managed = _seed_managed_pdf(workspace, "pdf-src-3")
    with pytest.raises(OfficeFileNotFoundError):
        pdf_to_word_endpoint(
            PdfToWordRequest(
                workspace_path=str(workspace),
                file_path=str(managed),
                source_doc_id="ghost-id",
            )
        )

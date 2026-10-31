"""Tests for Word template profile analysis and M1a fill."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from backend.office.errors import OfficeTemplateFillError
from backend.office.word_template_profile import analyze_word_template_profile
from backend.office.word_template_profile_fill import fill_word_template_profile


def _add_content_control(document, tag: str, text: str) -> None:
    paragraph = document.add_paragraph(text)
    sdt = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    tag_element = OxmlElement("w:tag")
    tag_element.set(qn("w:val"), tag)
    props.append(tag_element)
    sdt.append(props)
    content = OxmlElement("w:sdtContent")
    paragraph._p.addprevious(sdt)
    content.append(paragraph._p)
    sdt.append(content)


def _add_bookmark(paragraph, name: str, ident: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), ident)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), ident)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _build_template(path: Path) -> Path:
    document = Document()
    _add_content_control(document, "project_title", "Old Title")

    company = document.add_paragraph()
    run = company.add_run("Old Co")
    run.bold = True
    run.font.size = Pt(12)
    _add_bookmark(company, "company", "101")

    document.add_paragraph("Funding: <<funding>>")
    document.add_paragraph("Author: {{author}}")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "[[dataset]]"
    document.save(str(path))
    return path


def test_analyze_profile_detects_all_slot_kinds(tmp_path: Path) -> None:
    template = _build_template(tmp_path / "template.docx")

    profile = analyze_word_template_profile(template, workspace_path=str(tmp_path))

    assert profile["kind"] == "word-template-profile"
    assert profile["source_path"] == "template.docx"
    slots = {slot["name"]: slot for slot in profile["slots"]}
    assert slots["project_title"]["kind"] == "content_control"
    assert slots["project_title"]["current_value"] == "Old Title"
    assert slots["company"]["kind"] == "bookmark"
    assert slots["company"]["current_value"] == "Old Co"
    assert slots["funding"]["kind"] == "placeholder"
    assert slots["dataset"]["anchor"]["table_index"] == 0
    assert slots["author"]["kind"] == "jinja"
    assert profile["summary"]["has_jinja_control"] is False


def test_fill_profile_replaces_slots_and_preserves_formatting(tmp_path: Path) -> None:
    template = _build_template(tmp_path / "template.docx")

    result = fill_word_template_profile(
        template,
        {
            "project_title": "Sage",
            "company": "New Co",
            "funding": "NSF",
            "author": "Alice",
            "dataset": "data-v1",
        },
        "filled.docx",
        workspace_path=str(tmp_path),
    )

    assert result["ok"] is True
    assert result["filled_count"] == 5
    assert result["residual_placeholders"] == []
    output = Document(str(tmp_path / "filled.docx"))
    text = "\n".join(t.text or "" for t in output.element.body.iter(qn("w:t")))
    assert "Sage" in text
    assert "New Co" in text
    assert "NSF" in text
    assert "Alice" in text
    assert "Old Title" not in text
    company = next(p for p in output.paragraphs if "New Co" in p.text)
    assert company.runs[0].bold is True
    assert company.runs[0].font.size == Pt(12)
    assert output.tables[0].cell(0, 0).text == "data-v1"


def test_fill_rejects_missing_required_slot(tmp_path: Path) -> None:
    template = _build_template(tmp_path / "template.docx")

    with pytest.raises(OfficeTemplateFillError):
        fill_word_template_profile(
            template,
            {"company": "New Co"},
            "filled.docx",
            workspace_path=str(tmp_path),
        )


def test_fill_rejects_jinja_control_flow(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("{% if flag %}A{% endif %}")
    template = tmp_path / "control.docx"
    document.save(str(template))

    with pytest.raises(OfficeTemplateFillError):
        fill_word_template_profile(
            template,
            {"flag": "x"},
            "out.docx",
            workspace_path=str(tmp_path),
        )


def test_fill_refuses_overwrite(tmp_path: Path) -> None:
    template = _build_template(tmp_path / "template.docx")
    (tmp_path / "filled.docx").write_bytes(b"existing")

    with pytest.raises(OfficeTemplateFillError):
        fill_word_template_profile(
            template,
            {"project_title": "Sage", "company": "C", "funding": "F", "author": "A"},
            "filled.docx",
            workspace_path=str(tmp_path),
        )

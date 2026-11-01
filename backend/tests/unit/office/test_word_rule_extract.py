"""Unit tests for deterministic Word rule extraction (M3a)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from backend.office.errors import OfficeParseError
from backend.office.word_rule_extract import extract_word_rules


def _set_east_asia(style: object, name: str) -> None:
    rpr = style.element.get_or_add_rPr()  # type: ignore[attr-defined]
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _add_page_number(document: Document) -> None:
    footer = document.sections[0].footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)


def _spec_doc(path: Path, paragraphs: list[str]) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def _template_doc(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    normal = document.styles["Normal"]
    normal.font.name = "仿宋_GB2312"
    normal.font.size = Pt(16)
    _set_east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.0)
    for index in range(3):
        document.add_paragraph(f"正文段落 {index + 1}")
    _add_page_number(document)
    document.save(str(path))
    return path


def _rules(pack) -> dict:
    return {rule.kind: rule for rule in pack.rules}


def _spec_paragraphs() -> list[str]:
    return [
        "本规范适用于 A4 纸张，页边距上 3.7 厘米、下 3.5 厘米、左 2.8 厘米、右 2.6 厘米。",
        "正文采用仿宋_GB2312 三号字，行距固定值 28 磅，首行缩进 2 字符。",
        "页脚应包含阿拉伯数字页码。",
        "标题采用 1、1.1、1.1.1 编号。",
        "论文应包括：摘要、引言、结论、参考文献。",
        "图和表题注应连续编号。",
    ]


def test_extract_spec_rules_with_sources_and_confidences(tmp_path: Path) -> None:
    path = _spec_doc(tmp_path / "spec.docx", _spec_paragraphs())

    result = extract_word_rules(path, workspace_path=str(tmp_path))

    assert result.ok is True
    assert result.file_type == "spec"
    assert result.source_path == "spec.docx"
    rules = _rules(result.pack)
    assert set(rules) == {
        "page.size",
        "page.margins",
        "page.page_number",
        "style.font",
        "style.paragraph",
        "structure.heading_sequence",
        "structure.required_headings",
        "content.caption_sequence",
    }
    margins = rules["page.margins"].params
    assert margins["top_cm"] == pytest.approx(3.7)
    assert margins["bottom_cm"] == pytest.approx(3.5)
    assert margins["left_cm"] == pytest.approx(2.8)
    assert margins["right_cm"] == pytest.approx(2.6)
    font = rules["style.font"].params
    assert font["font_east_asia"] == "仿宋_GB2312"
    assert font["size_pt"] == pytest.approx(16.0)
    paragraph = rules["style.paragraph"].params
    assert paragraph["line_spacing_pt"] == pytest.approx(28.0)
    assert paragraph["first_line_indent_cm"] == pytest.approx(1.13)
    assert rules["structure.required_headings"].params["headings"] == [
        "摘要",
        "引言",
        "结论",
        "参考文献",
    ]
    assert all(rule.source.strip() for rule in result.pack.rules)
    assert all(candidate.confidence > 0 for candidate in result.candidates)


def test_extraction_is_reproducible(tmp_path: Path) -> None:
    path = _spec_doc(tmp_path / "spec.docx", _spec_paragraphs())

    first = extract_word_rules(path, workspace_path=str(tmp_path))
    second = extract_word_rules(path, workspace_path=str(tmp_path))

    assert first.pack.model_dump() == second.pack.model_dump()
    assert [issue.model_dump() for issue in first.conflicts] == [
        issue.model_dump() for issue in second.conflicts
    ]


def test_duplicate_clauses_are_merged(tmp_path: Path) -> None:
    path = _spec_doc(
        tmp_path / "dup.docx",
        [
            "页边距上 3.7 厘米、下 3.5 厘米、左 2.8 厘米、右 2.6 厘米。",
            "页边距上 3.7 厘米、下 3.5 厘米、左 2.8 厘米、右 2.6 厘米。",
        ],
    )

    result = extract_word_rules(path, workspace_path=str(tmp_path))

    margin_rules = [rule for rule in result.pack.rules if rule.kind == "page.margins"]
    assert len(margin_rules) == 1
    assert result.conflicts == []


def test_conflicting_margins_are_marked(tmp_path: Path) -> None:
    path = _spec_doc(
        tmp_path / "conflict.docx",
        ["上边距 3.7 厘米。", "上边距 2.5 厘米。"],
    )

    result = extract_word_rules(path, workspace_path=str(tmp_path))

    assert len(result.conflicts) == 1
    conflict = result.conflicts[0]
    assert conflict.kind == "page.margins"
    assert conflict.rule_ids
    margin_candidates = [
        candidate for candidate in result.candidates if candidate.rule.kind == "page.margins"
    ]
    assert len(margin_candidates) == 2
    assert all(candidate.conflict_group for candidate in margin_candidates)


def test_char_indent_conversion_and_missing_font_warning(tmp_path: Path) -> None:
    with_font = _spec_doc(
        tmp_path / "with-font.docx",
        ["正文采用宋体 16pt，首行缩进 2 字符。"],
    )
    converted = extract_word_rules(with_font, workspace_path=str(tmp_path))
    paragraph = next(
        rule for rule in converted.pack.rules if rule.kind == "style.paragraph"
    )
    assert paragraph.params["first_line_indent_cm"] == pytest.approx(1.13)

    without_font = _spec_doc(
        tmp_path / "without-font.docx",
        ["首行缩进 2 字符。"],
    )
    skipped = extract_word_rules(without_font, workspace_path=str(tmp_path))
    assert all("first_line_indent_cm" not in rule.params for rule in skipped.pack.rules)
    assert any(
        warning["code"] == "char_indent_unconverted" for warning in skipped.warnings
    )


def test_extract_template_dominant_values(tmp_path: Path) -> None:
    path = _template_doc(tmp_path / "template.docx")

    result = extract_word_rules(
        path, file_type="template", workspace_path=str(tmp_path)
    )

    assert result.ok is True
    assert result.file_type == "template"
    rules = _rules(result.pack)
    assert rules["page.size"].params == {"size": "A4"}
    assert rules["page.orientation"].params == {"orientation": "portrait"}
    assert rules["page.margins"].params["top_cm"] == pytest.approx(3.7)
    assert rules["page.margins"].params["right_cm"] == pytest.approx(2.6)
    assert rules["style.font"].params["font_east_asia"] == "仿宋_GB2312"
    assert rules["style.font"].params["size_pt"] == pytest.approx(16.0)
    paragraph = rules["style.paragraph"].params
    assert paragraph["line_spacing"] == pytest.approx(1.5)
    assert paragraph["first_line_indent_cm"] == pytest.approx(1.0)
    assert rules["page.page_number"].params == {"required": True}
    assert all(candidate.confidence >= 0.5 for candidate in result.candidates)


def test_template_without_body_samples_warns(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("只有标题", level=1)
    path = tmp_path / "heading-only.docx"
    document.save(str(path))

    result = extract_word_rules(
        path, file_type="template", workspace_path=str(tmp_path)
    )

    assert any(
        warning["code"] == "template_no_body_samples" for warning in result.warnings
    )


def test_extract_txt_spec(tmp_path: Path) -> None:
    path = tmp_path / "spec.txt"
    path.write_text(
        "页面使用 A4 纸张。\n正文采用宋体小四。\n",
        encoding="utf-8",
    )

    result = extract_word_rules(path, workspace_path=str(tmp_path))

    assert result.ok is True
    assert result.source_path == "spec.txt"
    assert "page.size" in _rules(result.pack)
    assert all(
        "第 1 段" in rule.source or "第 2 段" in rule.source
        for rule in result.pack.rules
    )


def test_extract_rejects_invalid_types(tmp_path: Path) -> None:
    docx_path = _spec_doc(tmp_path / "doc.docx", ["A4 纸张"])

    with pytest.raises(OfficeParseError):
        extract_word_rules(docx_path, file_type="bogus", workspace_path=str(tmp_path))

    text_path = tmp_path / "template.txt"
    text_path.write_text("A4", encoding="utf-8")
    with pytest.raises(OfficeParseError):
        extract_word_rules(
            text_path, file_type="template", workspace_path=str(tmp_path)
        )

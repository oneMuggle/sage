"""Tests for the 6-rule journal validator.

RED/GREEN:
- test_validate_good_filled_paper_against_simple_chinese_template_returns_no_errors
- test_validate_bad_filled_paper_detects_violations
- test_validate_docx_missing_section_returns_warning
"""
from __future__ import annotations

from pathlib import Path

from backend.office.journal.models import ViolationSeverity
from backend.office.journal.parser import parse_journal_spec
from backend.office.journal.validator import validate_document

FIXTURE_DIR = Path(__file__).resolve().parent.parent.parent.parent / "fixtures" / "journal"


def test_validate_good_filled_paper_against_simple_chinese_template_returns_no_errors():
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    from docx import Document

    doc = Document(str(FIXTURE_DIR / "good_filled_paper.docx"))
    violations = validate_document(doc, spec)
    error_violations = [v for v in violations if v.severity is ViolationSeverity.ERROR]
    assert error_violations == []


def test_validate_bad_filled_paper_detects_violations():
    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    from docx import Document

    doc = Document(str(FIXTURE_DIR / "bad_filled_paper.docx"))
    violations = validate_document(doc, spec)
    rule_ids = {v.rule_id for v in violations}
    # bad 文件字体/字号/行距/边距全错，至少 4 个 ERROR
    assert len(violations) >= 4
    assert any("font" in r or "spacing" in r or "margin" in r or "size" in r for r in rule_ids)


def test_validate_docx_missing_section_returns_warning():
    """缺失 '关键词' 章节 → 1 条 warning violation"""
    from docx import Document

    spec = parse_journal_spec(FIXTURE_DIR / "simple_chinese_template.docx")
    doc = Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("只有摘要没有关键词。")
    violations = validate_document(doc, spec)
    missing = [v for v in violations if "关键词" in v.message]
    assert len(missing) >= 1
    assert missing[0].severity is ViolationSeverity.WARNING

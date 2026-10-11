"""多形态项目上下文预算水位与工作区资产台账探测服务 (2026-10-10)。

对应 docs/mcp-sage-project-management-organization-plan-20261010.md Phase P2：
1. 计算五层受控上下文水位（L1 SAGE.md / L2 Metadata+Constraints / L3 ProjectProfile / L4 Materials）；
2. 扫描项目工作区生成多态资产台账：
   - 一般档案 (business): 案卷文书成品台账 (.docx / .xlsx / .pptx / .pdf) 与四级目录归属；
   - 科研项目 (research): 文献库 (.bib / .pdf)、实验记录 (.ipynb / .py / .csv) 与论文手稿 (.tex / .docx)；
   - 代码工程 (coding): 工程构建/配置指示文件、SAGE.md 与 .sage/hooks.json 状态。
"""

from __future__ import annotations

import json
import logging
import os
import shutil
import sys
from pathlib import Path
from typing import Any, Dict, List

from backend.chat.project_context import (
    PER_FILE_CHAR_CAP,
    TOTAL_CHAR_CAP,
    build_constraints_block,
    build_project_materials_block,
    build_project_metadata_block,
    discover_project_context,
)
from backend.data.database import get_database
from backend.data.project_constraint_repo import ProjectConstraintRepository
from backend.data.project_material_repo import ProjectMaterialRepository
from backend.data.project_repo import Project
from backend.memory.project_profile import ProjectProfileStore

logger = logging.getLogger(__name__)

OFFICE_EXTENSIONS = frozenset({".docx", ".xlsx", ".pptx", ".pdf", ".doc", ".xls", ".ppt"})
RESEARCH_EXTENSIONS = frozenset({".bib", ".tex", ".ipynb", ".pdf", ".csv", ".tsv", ".ris"})
CODING_INDICATOR_FILES = (
    "package.json",
    "pyproject.toml",
    "tsconfig.json",
    "Cargo.toml",
    "go.mod",
    "requirements.txt",
    "Makefile",
    "Dockerfile",
)
SKIP_DIRS = frozenset(
    {
        ".git",
        "node_modules",
        "__pycache__",
        ".venv",
        "venv",
        "dist",
        "build",
        ".next",
        ".vite",
        "target",
        ".worktrees",
    }
)
MAX_SCANNED_FILES = 200


def inspect_project_context_budget(project: Project) -> Dict[str, Any]:
    """计算项目五层受控上下文的实时字符占用与水位比例。"""
    l1_block = discover_project_context(project.path).render()
    l2_meta_block = build_project_metadata_block(project)
    l2_constraints_block = build_constraints_block(project.id)

    m_repo = ProjectMaterialRepository()
    all_materials = m_repo.list_by_project(project.id)
    active_materials = m_repo.get_active_materials_for_project(project.id)
    l4_materials_block = build_project_materials_block(active_materials)

    c_repo = ProjectConstraintRepository()
    enabled_constraints = c_repo.list_by_project(project.id, enabled_only=True)

    l3_profile_chars = 0
    try:
        profile_store = ProjectProfileStore(get_database())
        profile_store.load()
        l3_profile_chars = len(profile_store.get_snapshot(project.path))
    except Exception as exc:  # noqa: BLE001
        logger.debug("inspect_project_context_budget profile skip: %s", exc)

    l1_chars = len(l1_block)
    l2_meta_chars = len(l2_meta_block)
    l2_constraints_chars = len(l2_constraints_block)
    l4_materials_chars = len(l4_materials_block)

    total_chars = (
        l1_chars + l2_meta_chars + l2_constraints_chars + l3_profile_chars + l4_materials_chars
    )
    cap_chars = TOTAL_CHAR_CAP
    usage_ratio = round(min(total_chars / float(cap_chars), 1.0), 4) if cap_chars > 0 else 0.0

    return {
        "project_id": project.id,
        "l1_conventions_chars": l1_chars,
        "l2_metadata_chars": l2_meta_chars,
        "l2_constraints_chars": l2_constraints_chars,
        "l3_profile_chars": l3_profile_chars,
        "l4_materials_chars": l4_materials_chars,
        "total_chars": total_chars,
        "cap_chars": cap_chars,
        "per_file_cap_chars": PER_FILE_CHAR_CAP,
        "usage_ratio": usage_ratio,
        "active_materials_count": len(active_materials),
        "total_materials_count": len(all_materials),
        "enabled_constraints_count": len(enabled_constraints),
    }


def _classify_dossier_stage(rel_path: str) -> str:
    top = rel_path.replace("\\", "/").split("/")[0]
    if top.startswith("00_") or "立项" in top or "需求" in top:
        return "立项与需求"
    if top.startswith("01_") or "参考" in top or "依据" in top or "法规" in top:
        return "参考依据"
    if top.startswith("02_") or "草稿" in top or "编制" in top:
        return "过程草稿"
    if top.startswith("03_") or "定稿" in top or "归档" in top or "签发" in top:
        return "定稿归档"
    return "根目录文书"


def _classify_research_stage(rel_path: str, ext: str) -> str:
    top = rel_path.replace("\\", "/").split("/")[0].lower()
    if "literature" in top or ext in (".bib", ".ris"):
        return "文献与引文库"
    if "experiment" in top or ext in (".ipynb", ".csv", ".tsv"):
        return "实验与数据集"
    if "manuscript" in top or ext == ".tex":
        return "论文手稿"
    if "submission" in top or "rebuttal" in top:
        return "投稿与审稿回复"
    return "科研文档"


def inspect_project_workspace_overview(project: Project) -> Dict[str, Any]:
    """只读扫描工作区顶层与一级子目录，输出多形态项目工作台所需的资产清单。"""
    root = Path(project.path)
    has_sage_md = False
    has_hooks_json = False
    coding_indicators: List[str] = []
    office_deliverables: List[Dict[str, Any]] = []
    research_artifacts: List[Dict[str, Any]] = []
    directory_summary: List[Dict[str, Any]] = []

    if not root.is_dir():
        return {
            "project_id": project.id,
            "has_sage_md": False,
            "has_hooks_json": False,
            "coding_indicators": [],
            "office_deliverables": [],
            "research_artifacts": [],
            "directory_summary": [],
        }

    for candidate in ("SAGE.md", "CLAUDE.md", "AGENTS.md"):
        if (root / candidate).is_file():
            has_sage_md = True
            break

    if (root / ".sage" / "hooks.json").is_file():
        has_hooks_json = True

    if (root / ".git").exists():
        coding_indicators.append(".git")
    for indicator in CODING_INDICATOR_FILES:
        if (root / indicator).is_file():
            coding_indicators.append(indicator)

    scanned = 0
    try:
        entries = sorted(root.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
    except OSError:
        entries = []

    for entry in entries:
        if scanned >= MAX_SCANNED_FILES:
            break
        name = entry.name
        if name in SKIP_DIRS or (name.startswith(".") and name != ".sage"):
            continue
        try:
            if entry.is_dir():
                sub_count = 0
                for child in sorted(entry.iterdir(), key=lambda p: p.name.lower()):
                    if child.name.startswith("."):
                        continue
                    if child.is_file():
                        sub_count += 1
                        if scanned < MAX_SCANNED_FILES:
                            scanned += 1
                            _collect_file_artifact(
                                root, child, office_deliverables, research_artifacts
                            )
                directory_summary.append({"name": name, "file_count": sub_count})
            elif entry.is_file():
                scanned += 1
                _collect_file_artifact(root, entry, office_deliverables, research_artifacts)
        except OSError:
            continue

    office_deliverables.sort(key=lambda item: item["modified_at"], reverse=True)
    research_artifacts.sort(key=lambda item: item["modified_at"], reverse=True)

    return {
        "project_id": project.id,
        "has_sage_md": has_sage_md,
        "has_hooks_json": has_hooks_json,
        "coding_indicators": coding_indicators,
        "office_deliverables": office_deliverables[:30],
        "research_artifacts": research_artifacts[:30],
        "directory_summary": directory_summary[:20],
    }


def _collect_file_artifact(
    root: Path,
    file_path: Path,
    office_deliverables: List[Dict[str, Any]],
    research_artifacts: List[Dict[str, Any]],
) -> None:
    ext = file_path.suffix.lower()
    if ext not in OFFICE_EXTENSIONS and ext not in RESEARCH_EXTENSIONS:
        return
    try:
        stat = file_path.stat()
        rel_path = os.path.relpath(str(file_path), str(root)).replace("\\", "/")
    except OSError:
        return

    base_info = {
        "name": file_path.name,
        "relative_path": rel_path,
        "ext": ext,
        "size_bytes": int(stat.st_size),
        "modified_at": int(stat.st_mtime * 1000),
    }
    if ext in OFFICE_EXTENSIONS:
        office_deliverables.append(
            {
                **base_info,
                "category": _classify_dossier_stage(rel_path),
            }
        )
    if ext in RESEARCH_EXTENSIONS:
        research_artifacts.append(
            {
                **base_info,
                "category": _classify_research_stage(rel_path, ext),
            }
        )


__all__ = [
    "inspect_project_context_budget",
    "inspect_project_workspace_overview",
]


def extract_workspace_file_as_material_content(
    workspace_root: Path, relative_path: str
) -> str:
    """安全读取工作区内指定相对路径文件并提取为可注入受控资料池的文本。

    支持纯文本/代码/Markdown/BibTeX/LaTeX/CSV、Jupyter Notebook (.ipynb)
    以及 Office 文档 (.docx / .xlsx / .pptx) 的结构化文本提取。
    拒绝任何越界路径或符号链接逃逸。
    """
    clean_rel = (relative_path or "").strip().replace("\\", "/").lstrip("/")
    if not clean_rel or ".." in clean_rel.split("/"):
        raise ValueError("非法的相对路径")

    root_real = Path(os.path.realpath(str(workspace_root)))
    candidate = root_real / clean_rel
    resolved = Path(os.path.realpath(str(candidate)))

    try:
        resolved.relative_to(root_real)
    except ValueError as exc:
        raise ValueError("文件路径越界，仅允许读取项目工作区内部文件") from exc

    if not resolved.is_file():
        raise FileNotFoundError(f"文件不存在: {clean_rel}")

    ext = resolved.suffix.lower()
    extracted = ""

    if ext == ".docx":
        try:
            import docx  # type: ignore[import-not-found]

            doc = docx.Document(str(resolved))
            paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            extracted = "\n".join(paras)
        except Exception as exc:  # noqa: BLE001
            logger.debug("docx extract fallback for %s: %s", clean_rel, exc)
            extracted = f"[Word 文档: {clean_rel}, 大小: {resolved.stat().st_size} 字节]"
    elif ext == ".xlsx":
        try:
            import openpyxl  # type: ignore[import-not-found]

            wb = openpyxl.load_workbook(str(resolved), read_only=True, data_only=True)
            sheet_lines: List[str] = []
            for sheet_name in wb.sheetnames[:5]:
                ws = wb[sheet_name]
                sheet_lines.append(f"### 工作表: {sheet_name}")
                for row in list(ws.iter_rows(values_only=True))[:40]:
                    vals = [str(cell) if cell is not None else "" for cell in row]
                    if any(vals):
                        sheet_lines.append(" | ".join(vals))
            wb.close()
            extracted = "\n".join(sheet_lines)
        except Exception as exc:  # noqa: BLE001
            logger.debug("xlsx extract fallback for %s: %s", clean_rel, exc)
            extracted = f"[Excel 表格: {clean_rel}, 大小: {resolved.stat().st_size} 字节]"
    elif ext == ".pptx":
        try:
            import pptx  # type: ignore[import-not-found]

            prs = pptx.Presentation(str(resolved))
            slide_lines: List[str] = []
            for idx, slide in enumerate(prs.slides, start=1):
                texts = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text.strip()
                ]
                if texts:
                    slide_lines.append(f"### Slide {idx}\n" + "\n".join(texts))
            extracted = "\n".join(slide_lines)
        except Exception as exc:  # noqa: BLE001
            logger.debug("pptx extract fallback for %s: %s", clean_rel, exc)
            extracted = f"[PPT 演示文稿: {clean_rel}, 大小: {resolved.stat().st_size} 字节]"
    elif ext == ".ipynb":
        try:
            nb = json.loads(resolved.read_text(encoding="utf-8", errors="replace"))
            cell_lines: List[str] = []
            for cell in nb.get("cells", [])[:40]:
                ctype = cell.get("cell_type", "code")
                src = "".join(cell.get("source", [])).strip()
                if src:
                    cell_lines.append(f"[{ctype}]\n{src}")
            extracted = "\n\n".join(cell_lines)
        except Exception as exc:  # noqa: BLE001
            logger.debug("ipynb extract fallback for %s: %s", clean_rel, exc)
            extracted = resolved.read_text(encoding="utf-8", errors="replace")
    elif ext == ".pdf":
        extracted = f"[PDF 文献/案卷: {clean_rel}, 大小: {resolved.stat().st_size} 字节]"
    else:
        extracted = resolved.read_text(encoding="utf-8", errors="replace")

    body = (extracted or "").strip()
    if not body:
        body = f"[文件: {clean_rel}, 大小: {resolved.stat().st_size} 字节]"

    header = f"[来源文件: {clean_rel}]\n"
    max_body = max(PER_FILE_CHAR_CAP - len(header), 1000)
    if len(body) > max_body:
        body = body[:max_body] + "\n...[文件内容已按单资料预算截断]"
    return header + body


def inspect_project_runtime_and_hooks(project_id: str, workspace_root: str) -> Dict[str, Any]:
    """轻量只读诊断项目语言清单、本地运行时满足度、测试命令入口与 .sage/hooks.json 状态。"""
    root = Path(workspace_root)
    detected_languages: List[str] = []
    test_commands: List[str] = []
    recommendations: List[str] = []

    if root.exists() and root.is_dir():
        pkg_json = root / "package.json"
        if pkg_json.is_file():
            detected_languages.append("TypeScript/Node.js" if (root / "tsconfig.json").is_file() else "Node.js")
            try:
                pkg_data = json.loads(pkg_json.read_text(encoding="utf-8", errors="replace"))
                scripts = pkg_data.get("scripts") or {}
                if isinstance(scripts, dict):
                    for key in ("test", "lint", "typecheck", "build", "dev"):
                        if key in scripts:
                            test_commands.append(f"npm run {key}")
            except Exception as exc:  # noqa: BLE001
                logger.debug("package.json parse skipped: %s", exc)

        if (root / "pyproject.toml").is_file() or (root / "requirements.txt").is_file() or (root / "pytest.ini").is_file():
            detected_languages.append("Python")
            test_commands.append("pytest")

        if (root / "Cargo.toml").is_file():
            detected_languages.append("Rust")
            test_commands.append("cargo test")

        if (root / "go.mod").is_file():
            detected_languages.append("Go")
            test_commands.append("go test ./...")

    available_runtimes: List[str] = []
    if sys.executable:
        available_runtimes.append("python")
    for bin_name in ("node", "git", "cargo", "go"):
        if shutil.which(bin_name):
            available_runtimes.append(bin_name)

    missing_runtimes: List[str] = []
    for lang in detected_languages:
        if "Node.js" in lang and "node" not in available_runtimes:
            missing_runtimes.append("node")
        if lang == "Rust" and "cargo" not in available_runtimes:
            missing_runtimes.append("cargo")
        if lang == "Go" and "go" not in available_runtimes:
            missing_runtimes.append("go")

    if missing_runtimes:
        level = "unsatisfied" if len(missing_runtimes) == len(detected_languages) else "partial"
        recommendations.append("缺少本地运行时: " + ", ".join(missing_runtimes))
    else:
        level = "satisfied"

    if not (root / "SAGE.md").is_file() and not (root / "CLAUDE.md").is_file() and not (root / "AGENTS.md").is_file():
        recommendations.append("建议生成 SAGE.md 声明构建/测试命令与模块边界")

    hooks_file = root / ".sage" / "hooks.json"
    hooks_config_exists = hooks_file.is_file()
    hooks_count = 0
    if hooks_config_exists:
        try:
            hooks_raw = json.loads(hooks_file.read_text(encoding="utf-8", errors="replace"))
            if isinstance(hooks_raw, dict):
                hooks_list = hooks_raw.get("hooks", hooks_raw)
                hooks_count = len(hooks_list) if isinstance(hooks_list, (list, dict)) else 1  # noqa: UP038
            elif isinstance(hooks_raw, list):
                hooks_count = len(hooks_raw)
        except Exception as exc:  # noqa: BLE001
            logger.debug("hooks.json parse skipped: %s", exc)

    hooks_trusted = False
    try:
        from backend.hooks.project_hooks import is_project_trusted  # type: ignore[import-not-found]

        hooks_trusted = bool(is_project_trusted(str(root)))
    except Exception:  # noqa: BLE001
        hooks_trusted = False

    return {
        "project_id": project_id,
        "level": level,
        "detected_languages": detected_languages,
        "available_runtimes": available_runtimes,
        "test_commands": test_commands[:6],
        "hooks_config_exists": hooks_config_exists,
        "hooks_count": hooks_count,
        "hooks_trusted": hooks_trusted,
        "recommendations": recommendations,
    }
